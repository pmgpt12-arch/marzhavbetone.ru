"""P11 buyer regression: eight-file identity, economics, PDFs and style.
Static only; native Word and independent acceptance are separate gates.
"""
from pathlib import Path
from docx import Document
from zipfile import ZipFile
import hashlib, json, re, subprocess
R=Path(__file__).resolve().parents[1]
B=R/"tools/candidates/portfolio-owner-pdf-20261013/13-raschet-metrami"
M=R/"tools/candidates/portfolio-owner-pdf-20261013/p11-t6-dayone-20261014/release-manifest-HOLD.json"
def calc(sale,discount,tax,fees,debt,rate,months,work):
    finance=debt*rate*months/1200
    net=sale*(1-discount/100)-tax-fees-finance
    return round(net),round(net-work)
def check():
    manifest=json.loads(M.read_text())
    files=manifest["files"]
    assert len(files)==8
    assert all((B/n).is_file() for n in files)
    assert all(hashlib.sha256((B/n).read_bytes()).hexdigest()==v for n,v in files.items())
    assert all(n in (B/"00-START-HERE.txt").read_text() and n in (B/"MANIFEST.md").read_text() for n in files)
    d=Document(B/"52-chek-list-raschet-metrami.docx")
    ts=[t for t in d.tables if len(t.rows)==10]
    assert len(ts)==1
    t=ts[0]
    assert [r.cells[0].text[:1] for r in t.rows[1:]]==list("АБВГДЕЖЗИ")
    assert "А × (1 − Б / 100)" in t.rows[3].cells[1].text
    assert "/ 1200" in t.rows[6].cells[1].text
    assert "В − Г − Д − Е" in t.rows[7].cells[1].text
    assert "Ж − З" in t.rows[9].cells[1].text
    cases=[(10000000,5,600000,200000,4000000,12,10,9000000,8300000,-700000),(10000000,0,0,0,0,0,0,9000000,10000000,1000000),(5000000,20,500000,100000,2000000,24,12,4500000,2920000,-1580000)]
    assert all(calc(*v[:8])==v[8:] for v in cases)
    assert calc(*cases[0][:8]) != (9500000,-700000)
    assert calc(10000000,5,600000,200000,4000000,12,10,9000000) != calc(10000000,5,600000,200000,4000000,0,10,9000000)
    for n in ("51-algoritm-raschet-metrami.pdf","55-pamyatka-punkt-ustupki-escrow.pdf"):
        txt=subprocess.check_output(["pdftotext",str(B/n),"-"]).decode()
        assert "{{" not in txt and "}}" not in txt
    for n in ("52-chek-list-raschet-metrami.docx","54-zapros-sostoyanie-escrow.docx"):
        d=Document(B/n)
        assert ZipFile(B/n).testzip() is None
        for p in list(d.paragraphs)+[p for t in d.tables for row in t.rows for cell in row.cells for p in cell.paragraphs]:
            assert all(not r.italic for r in p.runs)
    for n in ("51-algoritm-raschet-metrami.docx","55-pamyatka-punkt-ustupki-escrow.docx"):
        d=Document(B/n)
        borders=d.styles["Title"]._element.xpath(".//w:pBdr/w:bottom")
        assert len(borders)==1 and borders[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color")=="977A48"
    import importlib.util
    src=R/"tools/candidates/portfolio-owner-pdf-20261013/p11-f02-f03-20261014/execution.py"
    spec=importlib.util.spec_from_file_location("p11_f02_f03_readonly",src)
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    trees={k:mod.read_doc((mod.C/n).read_bytes()) for k,n in mod.FILES.items()}
    positive,negative=mod.tests(trees)
    assert len(positive)==13 and all(positive.values()) and len(negative)==15 and all(negative.values())
    return {"status":"PASS_STATIC","files":8,"numeric_cases":len(cases),"negative_math_checks":2,"pdf_placeholders":0,"scope":"NOT_NATIVE_NOT_INDEPENDENT_NOT_RELEASE"}
if __name__=="__main__":
    print(json.dumps(check(),ensure_ascii=False))
