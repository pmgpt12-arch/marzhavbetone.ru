#!/usr/bin/env python3
"""Check approved editorial/layout deltas; preserve every numerical formula and feature."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
import json,hashlib,re,copy
import openpyxl
R=Path(__file__).resolve().parents[1];B=R/'tools/candidates/remaining-batch-r3-20261018'
I=Path('/tmp/mb-r3/input');O=Path('/tmp/mb-r3/work');X={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
edits=json.loads((B/'text-edits.json').read_text());fs=json.loads((B/'formula-text-edits.json').read_text());layouts=json.loads((B/'layout-edits.json').read_text())
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def xml(e):return E.tostring(e,method='c14n') if e is not None else None
def cells(root):return {c.get('r'):c for c in root.findall('.//x:c',X)}
literal=re.compile(r'"(?:[^"]|"")*"')
results=[]
pathmap=json.loads((B/'form-originals-audit/workbook-path-map.json').read_text()) if (B/'form-originals-audit/workbook-path-map.json').exists() else {}
for p in sorted(I.rglob('*.xlsx')):
 sku=p.parent.name;key=sku+'/'+p.name;q=O/pathmap.get(key,key);plans=layouts.get(key,{'sheets':{}})['sheets']
 allowed={(e['loc'].split(':')[0],e['loc'].split(':')[1]):e for e in edits if e['sku']==sku and e['file']==p.name}
 allowedf={(e['part'],e['cell']):e for e in fs if e['sku']==sku and e['file']==p.name}
 a=openpyxl.load_workbook(p,data_only=False);b=openpyxl.load_workbook(q,data_only=False)
 assert a.sheetnames==b.sheetnames,(key,'sheets')
 counts={'formula_count':0,'display_formula_changes':0,'text_changes':0}
 for i,(wa,wb) in enumerate(zip(a,b),1):
  part=f'xl/worksheets/sheet{i}.xml';plan=plans.get(part,{})
  assert wa.freeze_panes==wb.freeze_panes,(key,'freeze')
  refs={c.coordinate for row in wa for c in row if c.value is not None}|{c.coordinate for row in wb for c in row if c.value is not None}
  for ref in refs:
   ca=wa[ref];cb=wb[ref]
   if ca.data_type=='f':
    counts['formula_count']+=1;e=allowedf.get((part,ref))
    if e:
     assert ca.value==e['old'] and cb.value==e['new'],(key,part,ref,'formula edit')
     assert literal.sub('""',ca.value)==literal.sub('""',cb.value),(key,ref,'formula calculation changed')
     counts['display_formula_changes']+=1
    else:assert ca.value==cb.value,(key,ref,'unapproved formula change')
   elif ca.value!=cb.value:
    e=allowed.get((part,ref));assert e and ca.value==e['old'] and cb.value==e['new'],(key,part,ref,'unapproved value',ca.value,cb.value)
    counts['text_changes']+=1
  assert set(str(v) for v in wb.merged_cells.ranges)==set(str(v) for v in wa.merged_cells.ranges)|set(plan.get('merges',[])),(key,part,'merges')
 with ZipFile(p) as za,ZipFile(q) as zb:
  assert za.testzip() is None and zb.testzip() is None
  assert set(za.namelist())==set(zb.namelist()),(key,'part inventory')
  sa=E.fromstring(za.read('xl/styles.xml'));sb=E.fromstring(zb.read('xl/styles.xml'))
  for ca,cb in zip(sa,sb):
   assert ca.tag==cb.tag
   if E.QName(ca).localname=='cellXfs':
    assert len(cb)>=len(ca)
    for x,y in zip(ca,cb):assert xml(x)==xml(y),(key,'original style changed')
    for added in list(cb)[len(ca):]:
     aa=copy.deepcopy(added);aa.attrib.pop('applyAlignment',None);al=aa.find('x:alignment',X)
     assert al is not None and al.get('wrapText')=='1';al.attrib.pop('wrapText',None)
     def base(x):
      x=copy.deepcopy(x);x.attrib.pop('applyAlignment',None);z=x.find('x:alignment',X)
      if z is not None:z.attrib.pop('wrapText',None)
      if z is not None and not z.attrib and len(z)==0:x.remove(z)
      return xml(x)
     assert any(base(aa)==base(original) for original in ca),(key,'new style changed more than wrap')
   else:assert xml(ca)==xml(cb),(key,'unapproved style resource')
  for n in za.namelist():
   if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n):
    ra=E.fromstring(za.read(n));rb=E.fromstring(zb.read(n));plan=plans.get(n,{})
    for tag in ['dataValidations','conditionalFormatting','sheetProtection','sheetViews','autoFilter','tableParts','pageMargins','printOptions','headerFooter','colBreaks']:
     assert [xml(e) for e in ra.findall('x:'+tag,X)]==[xml(e) for e in rb.findall('x:'+tag,X)],(key,n,tag)
    for tag,option in [('rowBreaks','break_before'),('cols','columns')]:
     if option not in plan:assert xml(ra.find('x:'+tag,X))==xml(rb.find('x:'+tag,X)),(key,n,tag)
    pa=ra.find('x:pageSetup',X);pb=rb.find('x:pageSetup',X)
    if 'fit_height' in plan:
     assert pb.get('fitToHeight')==str(plan['fit_height']);pb=copy.deepcopy(pb)
     if pa is not None and pa.get('fitToHeight') is not None:pb.set('fitToHeight',pa.get('fitToHeight'))
     else:pb.attrib.pop('fitToHeight',None)
    assert xml(pa)==xml(pb),(key,n,'pageSetup')
    ac=cells(ra);bc=cells(rb);styleallowed=set(plan.get('wrap',[]))|set(plan.get('copy_style',{}))
    for ref,c in ac.items():
     if ref not in styleallowed:assert c.get('s','0')==bc[ref].get('s','0'),(key,n,ref,'style changed')
    rowsa={e.get('r'):e for e in ra.findall('x:sheetData/x:row',X)}
    rowsb={e.get('r'):e for e in rb.findall('x:sheetData/x:row',X)}
    editedrows={re.search(r'\d+',ref)[0] for part,ref in allowed if part==n}|set(plan.get('row_heights',{}))
    for row,rr in rowsa.items():
     aa=dict(rr.attrib);bb=dict(rowsb[row].attrib)
     if row in editedrows:
      for attr in ['ht','customHeight']:aa.pop(attr,None);bb.pop(attr,None)
     assert aa==bb,(key,n,row,'row metadata')
    for row,height in plan.get('row_heights',{}).items():assert float(rowsb[row].get('ht'))==height
   elif n=='xl/styles.xml':continue
   elif n=='xl/workbook.xml':
    ra=E.fromstring(za.read(n));rb=E.fromstring(zb.read(n))
    na=ra.find('x:definedNames',X);nb=rb.find('x:definedNames',X)
    allowednames=set()
    for part,pl in plans.items():
     idx=str(int(re.search(r'sheet(\d+)',part)[1])-1)
     for option,tag in [('print_area','_xlnm.Print_Area'),('repeat_rows','_xlnm.Print_Titles')]:
      if option in pl:allowednames.add((idx,tag))
    if sku=='P1' and p.name.startswith(('06-','07-','08-')):allowednames|={('0','_xlnm.Print_Area'),('0','_xlnm.Print_Titles')}
    def othernames(node):return [xml(v) for v in node if (v.get('localSheetId'),v.get('name')) not in allowednames] if node is not None else []
    assert othernames(na)==othernames(nb),(key,'unapproved definedName')
    if na is not None:ra.remove(na)
    if nb is not None:rb.remove(nb)
    assert xml(ra)==xml(rb),(key,'workbook metadata')
   else:
    expected=za.read(n).replace(b'+00:00Z',b'Z') if n=='docProps/core.xml' else za.read(n)
    assert expected==zb.read(n),(key,n,'unexpected part change')
 assert counts['text_changes']==len(allowed),(key,'missing text edits',counts,len(allowed))
 assert counts['display_formula_changes']==len(allowedf),(key,'missing formula edits')
 results.append({'sku':sku,'source_file':p.name,'file':q.name,'sha256':digest(q),**counts,'numerical_formula_skeletons':'IDENTICAL','functional_features':'PRESERVED','status':'PASS'})
for lock in json.loads((B/'original-form-lock.json').read_text()):assert digest(O/lock['sku']/lock['member'])==lock['sha256'],('published original modified',lock)
report={'status':'PASS','workbooks':results,'formula_count':sum(x['formula_count'] for x in results),'formula_display_changes':sum(x['display_formula_changes'] for x in results),'published_originals':'BYTE_IDENTICAL','native_office':'PENDING'}
(B/'evidence/workbook-preservation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print('PRESERVATION_PASS',len(results),'books',report['formula_count'],'formulas',report['formula_display_changes'],'display edits')
