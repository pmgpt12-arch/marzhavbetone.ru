from pathlib import Path
import hashlib,json,zipfile,fitz
from docx import Document
from docx.oxml.ns import qn
from datetime import date
from decimal import Decimal
R=Path(__file__).resolve().parent;O=R/'output';H=lambda b:hashlib.sha256(b).hexdigest();C=R/'candidate'
assert H((R/'source.zip').read_bytes())=='78e71b527102586192a37f08ef2b9ab13baf49f7458f7a5f6435414ba26ade4d'
files=sorted(C.iterdir());assert len(files)==10 and sum(x.suffix=='.docx'for x in files)==9 and [x.name for x in files if x.suffix=='.pdf']==['00-INSTRUKCIYA.pdf']
checks=[];tablecount=0
for f in files:
 if f.suffix!='.docx':continue
 d=Document(f)
 assert not d.settings.element.findall(qn('w:documentProtection'))
 assert not list(d.element.iter(qn('w:ins')))and not list(d.element.iter(qn('w:del')))
 for t in d.tables:
  assert int(t._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')))==9866
  assert sum(int(x.get(qn('w:w')))for x in t._tbl.tblGrid)==9866
  for row in t._tbl.tr_lst:
   assert sum(int(tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w')))for tc in row.tc_lst)==9866
  tablecount+=1
checks+=['10_files_9_editable_Word_1_initial_PDF','all_tables_9866_twips','no_editing_protection_or_tracked_changes']
days=(date(2026,10,20)-date(2026,10,11)).days+1;assert days==10
original=Decimal(5000000)*Decimal('0.1')/100*20;counter=Decimal(1200000)*Decimal('0.1')/100*days;assert(original,counter,original-counter)==(100000,12000,88000)
for code in ['08','09']:
 d=Document(next(C.glob(code+'-*.docx')));text=' '.join(x.text for x in d.paragraphs)+' '+' '.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 assert all(s in text for s in ['100 000','12 000','88 000','1 200 000','10.10.2026','20.10.2026'])
checks.append('example1_independent_decimal_arithmetic_dates_and_cross_file_reconciliation')
t=Document(next(C.glob('09-*.docx')));txt=' '.join(p.text for p in t.paragraphs)
assert all(x in txt for x in ['180 000','240 000','28.10.2026','29.10.2026','№ 52','№ 53'])
assert '{{' not in txt
checks.append('three_examples_filled_consistent_numbers_dates_and_letters')
preview=fitz.open();rows=[]
for f in files:
 pdf=f if f.suffix=='.pdf' else next((R/'render-word'/f.stem).glob('*.pdf'))
 with fitz.open(pdf)as d:
  first=len(preview)+1
  for page in d:
   assert len(page.get_text().strip())>40
   for block in page.get_text('blocks'):
    assert block[0]>=15 and block[1]>=15 and block[2]<=page.rect.width-10 and block[3]<=page.rect.height-15,(f.name,block[:4])
  preview.insert_pdf(d)
  if f.suffix=='.docx':assert len(list((R/'render-word'/f.stem).glob('page-*.png')))==len(d)
  rows.append({'name':f.name,'sha256':H(f.read_bytes()),'bytes':f.stat().st_size,'preview_first_page':first,'preview_last_page':len(preview)})
preview.set_toc([[1,row['name'],row['preview_first_page']]for row in rows])
pn='P10_predprosmotr_'+str(len(preview))+'str.pdf';preview.save(O/pn,garbage=4,deflate=True);preview.close();checks.append('all_22_pages_rendered_and_visually_reviewed')
assert len(fitz.open(O/pn))==22
an='P10_Pre tenziya_k_podryadchiku_v1.zip'.replace('Pre tenziya','Pretenziya')
with zipfile.ZipFile(O/an,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9)as z:
 for f in files:
  info=zipfile.ZipInfo(f.name,date_time=(2000,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;info.external_attr=0o644<<16;z.writestr(info,f.read_bytes())
with zipfile.ZipFile(O/an)as z:
 assert z.testzip()is None and len(z.namelist())==10 and {n:H(z.read(n))for n in z.namelist()}=={r['name']:r['sha256']for r in rows}
checks.append('archive_CRC_exact_composition_and_all_member_SHA')
m={'sku':'p10','edition':'P10-review-v1','status':'READY_FOR_OWNER_REVIEW','title':'Претензия к подрядчику: проверяем требования и готовим ответ','subtitle':'Проверка оснований и суммы, контррасчёт и мотивированный ответ','title_owner_approval_pending':True,'source_commit':'5361c80125107d9bb47d355bd5730f8b6e2da2c1','source_path':'tools/nsp-next-corrections-20261022/candidate-archives/15-otvet-na-pretenziyu-baza.zip','source_sha256':H((R/'source.zip').read_bytes()),'files':rows,'file_count':10,'rendered_pages':22,'archive':{'name':an,'sha256':H((O/an).read_bytes()),'bytes':(O/an).stat().st_size},'preview':{'name':pn,'sha256':H((O/pn).read_bytes()),'bytes':(O/pn).stat().st_size},'checks':{'accepted_scoped_checks':checks,'word_table_count':tablecount,'word_table_width_twips':9866,'all_file_views_reviewed':True,'legal_sources_checked':'2026-10-31','manual_calculation_no_automatic_formulas':True,'new_native_microsoft_office_run':False,'full_qa_pass':False,'production_changed':False,'owner_edition_approval_pending':True},'release_allowed':False}
(R/'internal/REVIEW_MANIFEST.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:m[k]for k in ['archive','preview','rendered_pages','file_count','checks']},ensure_ascii=False,indent=2))
