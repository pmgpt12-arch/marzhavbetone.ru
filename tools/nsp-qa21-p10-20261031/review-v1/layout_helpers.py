from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips,Pt
W=9866
def elem(parent,tag,attrs):
    node=parent.find(qn(tag))
    if node is None: node=OxmlElement(tag); parent.append(node)
    for k,v in attrs.items(): node.set(qn(k),str(v))
    return node

def norm(doc):
    for sec in doc.sections: sec.left_margin=Twips(1020); sec.right_margin=Twips(1020)
    for t in doc.tables:
        count=len(t.columns)
        first=t.cell(0,0).text
        if count==2: widths=[3700,6166]
        elif count==3: widths=[550,6200,3116] if first=='№' else [1450,4350,4066]
        elif count==4: widths=[1850,2150,3166,2700]
        elif count==5: widths=[750,3000,2250,1050,2816]
        else: widths=[W//count]*count; widths[-1]+=W-sum(widths)
        t.autofit=False
        pr=t._tbl.tblPr
        elem(pr,'w:tblW',{'w:type':'dxa','w:w':W})
        elem(pr,'w:tblInd',{'w:type':'dxa','w:w':0})
        elem(pr,'w:jc',{'w:val':'left'})
        borders=elem(pr,'w:tblBorders',{})
        for edge in ['top','left','bottom','right','insideH','insideV']:
            elem(borders,'w:'+edge,{'w:val':'single','w:sz':4,'w:color':'D9D9D9'})
        margins=elem(pr,'w:tblCellMar',{})
        for edge,v in [('top',65),('bottom',65),('left',90),('right',90)]:elem(margins,'w:'+edge,{'w:type':'dxa','w:w':v})
        for col,w in zip(t.columns,widths): col.width=Twips(w)
        for rowidx,tr in enumerate(t._tbl.tr_lst):
            rp=tr.get_or_add_trPr(); elem(rp,'w:cantSplit',{})
            for height in list(rp.findall(qn('w:trHeight'))):
                if height.get(qn('w:hRule'))=='exact': height.set(qn('w:hRule'),'atLeast')
            pos=0
            for tc in tr.tc_lst:
                span=tc.grid_span; tp=tc.get_or_add_tcPr()
                elem(tp,'w:tcW',{'w:type':'dxa','w:w':sum(widths[pos:pos+span])});pos+=span
                elem(tp,'w:vAlign',{'w:val':'center'})
                for tag in ['w:noWrap','w:tcFitText']:
                    bad=tp.find(qn(tag))
                    if bad is not None:tp.remove(bad)
            if rowidx==0 and count>2:elem(rp,'w:tblHeader',{})
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:
                        if run.font.size and run.font.size.pt>11:run.font.size=Pt(10)
    return doc

