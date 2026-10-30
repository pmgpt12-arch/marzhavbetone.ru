from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parent
f=next((R/'candidate').glob('02-*.docx'));d=Document(f)
next(p for p in d.paragraphs if p.text=='Правовые ориентиры по предмету требования').paragraph_format.page_break_before=True
d.save(f)
f=next((R/'candidate').glob('03-*.docx'));d=Document(f)
for p in list(d.paragraphs):
 if p.text.startswith('Повторите раздел по каждому коду'):p._element.getparent().remove(p._element)
d.save(f)
