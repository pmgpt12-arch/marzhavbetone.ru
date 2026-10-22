from pathlib import Path
from docx import Document
from docx.shared import Pt
import json,hashlib
Q=Path(__file__).resolve().parents[1];manifest=json.loads((Q/'content-manifest.json').read_text())
for rec in manifest:
 if not rec['path'].endswith('.docx'):continue
 p=Q/'candidates'/rec['path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==rec['sha256'];d=Document(p)
 if p.name.startswith('01-'):
  for para in d.paragraphs:
   if para.text.startswith('Правомерность удержания проверена'):para.text='Основание неоплаты проверено по п. [номер] Договора и закону; его нарушения: [конкретные факты]. Отсутствие отдельного согласия не исключает зачёт по ст. 410 ГК РФ.'
   if para.text.startswith('применимой ответственности:'):para.text='неустойки по Договору [сумма] руб. либо процентов по ст. 395 ГК РФ [сумма] руб.;'
   if para.text.startswith('Одновременно неустойка и проценты'):para.text='Одновременное взыскание за одно нарушение - только по исключению закона или договора (п. 4 ст. 395 ГК РФ): [основание / не применяется].'
  for para in d.paragraphs:
   if any(t in para.text for t in ['Основание неоплаты','неустойки по Договору','Одновременное взыскание']):
    for run in para.runs:run.font.name='Arial';run.font.size=Pt(10.5)
 else:
  start=next(i for i,para in enumerate(d.paragraphs)if para.text=='Контроль сторон и расчёта')
  for i,para in enumerate(d.paragraphs):
   if i>=start or para.text.startswith('Этот текст запроса'):
    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(10)
  d.paragraphs[start].paragraph_format.keep_with_next=True;d.paragraphs[start+1].paragraph_format.keep_together=True
 d.save(p);rec['previous_candidate_sha256']=rec['sha256'];rec['sha256']=hashlib.sha256(p.read_bytes()).hexdigest();b=p.read_bytes();rec['git_blob']=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest();rec['layout_refinement']='match source body font; avoid detached signature/heading'
(Q/'content-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2));print('REFINED_NEW_DOCX_ONLY')
