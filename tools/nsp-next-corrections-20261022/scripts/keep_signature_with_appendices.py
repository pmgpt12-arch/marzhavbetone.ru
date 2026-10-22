from pathlib import Path
from docx import Document
import json,hashlib
Q=Path(__file__).resolve().parents[1];m=json.loads((Q/'content-manifest.json').read_text());rec=next(r for r in m if r['path'].endswith('/01-pretenziya-na-uderzhanie.docx'));p=Q/'candidates'/rec['path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==rec['sha256'];d=Document(p);start=next(i for i,p in enumerate(d.paragraphs)if p.text.startswith('4. ПРИЛОЖЕНИЯ'))
for i,p0 in enumerate(d.paragraphs):
 if i>=start:p0.paragraph_format.keep_together=True;p0.paragraph_format.keep_with_next=i<len(d.paragraphs)-1
 d.paragraphs[-2].paragraph_format.keep_with_next=True
d.save(p);rec['previous_candidate_sha256']=rec['sha256'];b=p.read_bytes();rec['sha256']=hashlib.sha256(b).hexdigest();rec['git_blob']=hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest();rec['layout_refinement']='Keep appendix list with signature; no detached signature page';(Q/'content-manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2));print(rec['sha256'])
