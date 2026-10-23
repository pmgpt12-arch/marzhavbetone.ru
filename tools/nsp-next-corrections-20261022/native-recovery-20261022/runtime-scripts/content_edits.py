from pathlib import Path
import json,hashlib,re,html,io
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import fitz
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
from reportlab.lib.styles import ParagraphStyle
Q=Path(__file__).resolve().parents[1];R=Q.parent/'next-corrections';OUT=Q/'candidates'; records=[]
def record(path,source):
 b=(OUT/path).read_bytes();records.append({'path':path,'base_sha256':hashlib.sha256(source).hexdigest(),'sha256':hashlib.sha256(b).hexdigest(),'git_blob':hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest(),'status':'CANDIDATE_BUILT','office':'PENDING'if path.endswith('.docx')else'NOT_APPLICABLE'})
def save_doc(path,d,raw):
 p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);d.save(p);record(path,raw)
def link(d,p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
norm395='https://www.consultant.ru/document/cons_doc_LAW_5142/93383099dc4f06165d8d4aac4eeb159db4f6da3d/'
escrow='https://www.consultant.ru/document/cons_doc_LAW_51038/57da6efc7ca337d428cf526d01e70925ce5bdcb0/'
price='https://www.consultant.ru/document/cons_doc_LAW_51038/7e20edcc51ba599c70fb328204e3ac1226e7d912/'
off='https://www.consultant.ru/document/cons_doc_LAW_5142/8226847de7f2eeb92a5dc53e6289e35d1c74e4c2/'
offset='https://www.consultant.ru/document/cons_doc_LAW_5142/ce09e59d11f3a60c30a1601c0f6c7cceb9835dfa/'
support='Поддержка: nastoronepodryadchika@yandex.ru. В теме укажите SKU и файл, в сообщении - страницу либо лист/ячейку, пример ввода и ожидаемый результат. Сообщите об ошибке, приложив копию без конфиденциальных данных.'
path='products-storage/07-uderzhaniya-shtrafy-zachety/01-pretenziya-na-uderzhanie.docx';raw=(R/path).read_bytes();d=Document(io.BytesIO(raw))
for p in d.paragraphs:
 if 'Удержание сумм без письменного согласования' in p.text:p.text='Правомерность удержания проверена по пункту [номер] Договора и применимому закону. Само отсутствие отдельного согласия Подрядчика не исключает предусмотренный законом зачёт. В данном случае основание для неоплаты отсутствует по следующим причинам: [конкретные нарушения основания, суммы и срока].'
 if p.text.strip()=='неустойки по договору - [сумма] руб.,':p.text='применимой ответственности: [неустойка по пункту Договора, сумма] либо [проценты по ст. 395 ГК РФ, сумма];'
 if 'процентов по ст. 395 ГК РФ - [сумма]' in p.text:p.text='Одновременно неустойка и проценты за одно денежное нарушение заявляются только при прямо применимом исключении закона или договора (п. 4 ст. 395 ГК РФ): [основание исключения / исключить одновременное требование].'
save_doc(path,d,raw)
path='products-storage/13-raschet-metrami/54-zapros-sostoyanie-escrow.docx';raw=(R/path).read_bytes();d=Document(io.BytesIO(raw))
for p in d.paragraphs:
 if p.text.startswith('До подписания уступки или ДДУ'):p.text='Этот текст запроса применяется к уступке существующего ДДУ. При новом ДДУ до его регистрации проверяются проект, цена, банк и будущий плательщик; требование уже наполненного эскроу к этой стадии не применяется. После регистрации контролируются срок и фактическое внесение денег. Для готового объекта проверяются ЕГРН и документы передачи, а не эскроу.'
 p.paragraph_format.widow_control=True
p=d.add_paragraph('Контроль сторон и расчёта');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
d.add_paragraph('До подписания запишите отдельно: должника за работы; цедента или застройщика; будущего участника ДДУ; лицо, вносящее остаток цены. Если это разные лица, их участие и расчёты должны иметь самостоятельное основание. Запрос банковских сведений сам по себе не прекращает долг за работы.')
d.add_paragraph('Для нового ДДУ: подписанный проект и регистрация, затем оплата на эскроу по согласованному графику. Для уступки: зарегистрированный ДДУ, подтверждённая оплата либо одновременный перевод долга с необходимым согласием кредитора, регистрация уступки. Цена уступки и цена ДДУ учитываются раздельно. Для готового объекта: право продавца, обременения, основание передачи, регистрация перехода права и акт.')
d.add_paragraph('В соглашении о расчётах зафиксируйте, какую часть одного долга за работы, по какому основанию и при каком подтверждённом событии погашают. Отступное по ст. 409 ГК РФ и зачёт по ст. 410 ГК РФ не взаимозаменяемы. При участии третьего лица требуется согласованная схема всех сторон.')
p=d.add_paragraph('Нормы для сверки: ')
for label,url in [('214-ФЗ ст. 15.4 ч. 2–3',escrow),('214-ФЗ ст. 15.5 ч. 4',price),('ГК РФ ст. 409',off),('ГК РФ ст. 410',offset)]:link(d,p,label,url);p.add_run('; ')
d.add_paragraph(support);p=d.add_paragraph('Связанный маршрут при риске банкротства: ');link(d,p,'P9 - защита требований','https://nastoronepodryadchika.ru/products/p9-bankrotstvo-genpodryadchika.html')
save_doc(path,d,raw)
pdfmetrics.registerFont(TTFont('NSP','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('NSPB','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
body=ParagraphStyle('body',fontName='NSP',fontSize=9.3,leading=13,spaceAfter=5)
head=ParagraphStyle('head',fontName='NSPB',fontSize=12,leading=16,spaceBefore=9,spaceAfter=7,keepWithNext=True)
title=ParagraphStyle('title',parent=head,fontSize=16,leading=20,spaceAfter=15)
def para(t,style=body):return Paragraph(html.escape(t).replace('\n',' '),style)
def pdf_link(label,url):return Paragraph(f'<link href="{html.escape(url)}" color="#355d7a">{html.escape(label)}</link>',body)
def make_pdf(path,extra,replace):
 raw=(R/path).read_bytes();doc=fitz.open(stream=raw,filetype='pdf');story=[]
 for pi,page in enumerate(doc):
  for bi,b in enumerate(page.get_text('blocks')):
   t=b[4].strip();t=' '.join(t.split())
   if not t or t.isdigit()or t.startswith('СИСТЕМА ЗАЩИТЫ ПОДРЯДЧИКА /'):continue
   for before,after in replace:t=t.replace(before,after)
   if not t:continue
   heading=bool(re.match(r'^(Шаг \d|ШАГ \d|Назначение$|Когда применяется$|Граница применения$|Важно, прочитайте|Другие документы|Документы для|Почему один|ФОРМУЛИРОВКА)',t))
   story.append(para(t,title if pi==0 and bi==0 else head if heading else body))
 story.append(Spacer(1,12))
 for kind,t in extra:
  if kind=='heading':story.append(para(t,head))
  elif kind=='link':story.append(pdf_link(*t))
  else:story.append(para(t))
 out=OUT/path;out.parent.mkdir(parents=True,exist_ok=True)
 SimpleDocTemplate(str(out),pagesize=(595.28,841.89),leftMargin=48,rightMargin=48,topMargin=44,bottomMargin=44,title=Path(path).stem,author='На стороне подрядчика').build(story)
 record(path,raw)
 for i,page in enumerate(fitz.open(out)):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(Q/'proof'/(Path(path).stem+f'-page-{i+1}.png'))
routes=[('heading','Три маршрута и проверка одного долга за работы'),('text','Новый ДДУ: определите застройщика, будущего участника и плательщика цены. До регистрации проверяйте проект и график, а не требуйте справку об уже внесённой цене. Оплата - после регистрации в установленные сроки, для эскроу денежное исполнение подтверждает банк. Регистрация ДДУ сама по себе не равна погашению долга за работы.'),('text','Уступка существующего ДДУ: отдельно проверьте цедента, зарегистрированный ДДУ, оплаченную цену либо одновременный перевод остатка долга с необходимым согласием кредитора. Не смешивайте цену уступки и пополнение эскроу. До подписания передаточного акта проверьте допустимость уступки по ст. 11 Закона №214-ФЗ.'),('text','Готовая недвижимость: установите правообладателя, обременения, зарегистрированный переход права и передачу объекта. Эскроу по ДДУ не является обязательным универсальным доказательством для этой ветви.'),('text','Карта сделки: должник за работы / передающее объект или право лицо / получатель / плательщик остатка цены / кредитор по ДДУ. Для каждого укажите договор, сумму, действие, срок и подтверждение. При разных лицах согласуйте их участие; двусторонний зачёт между иными лицами не возникает автоматически.'),('text','Погашение работ: назовите одно правовое основание, сумму и событие, подтверждающие прекращение долга. При отступном учитывайте его предоставление (ст. 409 ГК РФ). При зачёте проверьте встречность, однородность, сроки и заявление (ст. 410 ГК РФ). Для регистрации прав и действий третьего лица отдельно согласуйте условия и доказательства. До подтверждённого события не отмечайте работы оплаченными.'),('heading','Проверка экономики без двойного вычитания'),('text','Сопоставляйте один и тот же погашаемый долг. Чистая стоимость недвижимости после налогов, дисконта, продажи и финансирования минус дополнительные доплаты подрядчика плюс денежная часть расчёта. Не вычитайте доплату второй раз, если она уже включена в чистую стоимость.'),('text','Учебный пример 1: 8,3 млн до дополнительной доплаты - 3 млн = 5,3 млн; при долге за работы 9 млн разница составляет -3,7 млн. Пример 2: 6,5 млн чистой стоимости после всех доплат + 3 млн денежной части - 9 млн долга = +0,5 млн. Это не значения вашей сделки.'),('text','Сверьте результат с уже включёнными файлами 52 и 56. Если плательщик, событие погашения либо подтверждение оплаты не установлены, отложите подписание и запросите недостающие документы.'),('heading','Нормы и помощь'),('link',('214-ФЗ ст. 15.4 ч. 2–3 - режим эскроу',escrow)),('link',('214-ФЗ ст. 15.5 ч. 4 - внесение денег после регистрации',price)),('link',('ГК РФ ст. 409 - предоставление отступного',off)),('link',('ГК РФ ст. 410 - условия зачёта',offset)),('text',support),('link',('Связанный маршрут: P9 - банкротство генподрядчика','https://nastoronepodryadchika.ru/products/p9-bankrotstvo-genpodryadchika.html'))]
make_pdf('products-storage/13-raschet-metrami/51-algoritm-raschet-metrami.pdf',routes,[('это не оплата работ: подрядчик получает не квартиру, а долг за неё.','это передача права с дополнительной обязанностью: определите её размер, плательщика и влияние на чистую стоимость. До проверки не считайте долг за работы погашенным.'),('Эскроу вообще не открыт → стоп, → юрист.','При уступке отсутствие необходимых сведений об эскроу - основание остановить согласование. При новом ДДУ сначала установите порядок регистрации, открытия счёта и последующей оплаты.'),('Передаваемое право проверено на оплату:','При уступке передаваемое право проверено на оплату:')])
make_pdf('products-storage/13-raschet-metrami/55-pamyatka-punkt-ustupki-escrow.pdf',routes,[])
make_pdf('products-storage/07-uderzhaniya-shtrafy-zachety/07-algoritm-proverki-uderzhaniy.pdf',[('heading','Выберите применимое основание ответственности'),('text','За одно нарушение денежного обязательства при предусмотренной договором неустойке проценты по статье 395 по общему правилу не добавляются. Исключение должно следовать из применимого закона или договора (п. 4 ст. 395 ГК РФ). Не переносите одновременно обе строки в претензию без такого основания.'),('text','Пример: основной долг 100 000 руб. и договорная неустойка 2 000 руб. дают 102 000 до судебных расходов. Дополнительные 1 500 руб. процентов не прибавляются автоматически. При отсутствии договорной неустойки отдельно проверьте основание, срок и расчёт процентов.'),('text','Разделите зачёт по ст. 410 ГК РФ, договорный порядок удержания и убытки. Для каждого нужны собственные основания; отдельное согласие на зачёт не является универсальным требованием закона. Сохраните договор, уведомление, контррасчёт и доказательство доставки возражений.'),('link',('ГК РФ ст. 395 п. 4',norm395)),('link',('ГК РФ ст. 410',offset)),('text',support),('link',('Связанный маршрут: взыскание оплаты P1','https://nastoronepodryadchika.ru/products/p1-oplata-po-ks2.html'))],[('+ Пени по договору','Применимая неустойка по договору либо проценты по ст. 395 ГК РФ'),('+ Проценты по ст. 395 ГК РФ','Одновременное взыскание - только при применимом исключении закона или договора')])
additions={
'products-storage/10-avans-usloviya-i-pribyl/00-START-HERE.txt':'''ПРОВЕРКА ПО ДОГОВОРУ\nТребование 07 применяйте, когда право на аванс предусмотрено договором, согласованное условие выплаты наступило, срок истёк. Предложение об авансе в переговорах не создаёт уже возникший долг. Сверьте п. 2 ст. 711 ГК РФ и условия договора; изменение согласовывается сторонами по ст. 450 ГК РФ. При требовании возврата отдельно определите основание прекращения договора, стоимость подтверждённых работ и уже проведённые платежи; расчёт 09 не заменяет правовую квалификацию.\nПРИМЕР\nПолучен аванс 3 млн, подтверждены работы на 4 млн, дополнительно оплачено 0,5 млн. Простой пример сальдо: 4 - 3 - 0,5 = 0,5 млн в пользу подрядчика. Не вычитайте аванс повторно, если он уже учтён в показателе оплаты. При спорных работах отделите подтверждённый и спорный объём. Ожидаемый результат перенесите в 10 с приложениями и доказательством доставки.\nНормы: ГК РФ ст. 711 п. 2 и ст. 450 - проверяются по применимому договору; ссылки на нормы требования и возврата содержатся также в 07/08/10.\n''',
'products-storage/11-bankrotstvo-genpodryadchika/00-START-HERE.txt':'''РЕЕСТР 04 НЕ СЧИТАЕТ РАЗМЕР ТРЕБОВАНИЯ\nКС-2 и КС-3 по одним работам на 100 000 руб. дают сумму значений документов 200 000 руб. Долг в этом примере до оплат - 100 000 руб. Отделите повторные подтверждения одной работы, основной долг и санкции; сопоставьте расчёт с описью 05 и заявлением 08. Статус неполноты 04 требует устранения до подачи, но отсутствие замечаний модели не означает готовность заявления.\nТекущие платежи: ст. 5 п. 1–2 Закона №127-ФЗ, возникновение обязательства, а не одна дата КС. https://www.consultant.ru/document/cons_doc_LAW_39331/b98b0fa41fd3af9378e035703cf9c55508b074fe/\nПроцедуру, сроки, адресатов, пошлину, льготу либо ходатайство и подтверждение направления сверяйте по 02/05/08/09. Решение принимает суд.\n''',
'products-storage/02-dopraboty-bez-poter/00-START-HERE.txt':'''ДОПОЛНИТЕЛЬНАЯ ПРОВЕРКА\nВ расчёте 07 сумма по заполненным строкам предварительна, пока есть незавершённые позиции. Заполните название, количество, цену и применимый коэффициент; число неполных строк должно стать нулём. Пустая цена не равна нулевой согласованной стоимости.\nГК РФ ст. 743 п. 3–4: сообщите о необходимости неучтённых работ и увеличения сметы. Если ответа нет, после установленного срока приостановите соответствующие работы. Без исполнения обязанности уведомления право на оплату и убытки может быть утрачено; исключение требует доказать необходимость немедленных действий в интересах заказчика, например для предотвращения гибели или повреждения объекта. Выполненный без уведомления объём не легализуется одним новым письмом. Соберите переписку, фактическую хронологию и доказательства необходимости и передайте на правовую оценку.\nhttps://www.consultant.ru/document/cons_doc_LAW_9027/b48c040fb20b324efb339e89ad832027d4b5d97a/\n'''}
for path,addition in additions.items():
 if '/10-avans-'in path:addition+='ГК РФ ст. 711 п. 2: https://www.consultant.ru/document/cons_doc_LAW_9027/46561f5271883fcf47f3e9859d3a59eb56bfcd96/\nГК РФ ст. 450 п. 1: https://legalacts.ru/kodeks/GK-RF-chast-1/razdel-iii/podrazdel-2/glava-29/statja-450/\n'
 raw=(R/path).read_bytes();p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(raw.decode().rstrip()+'\n\n'+addition+'\n'+support+'\nСвязанные решения: https://nastoronepodryadchika.ru/katalog.html#group-do-podpisaniya\n');record(path,raw)
(Q/'content-manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
print('CONTENT_CANDIDATES',len(records))
