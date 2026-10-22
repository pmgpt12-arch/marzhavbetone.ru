from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import io,json,hashlib
q=Path(__file__).resolve().parents[1];r=q.parent/'next-corrections/downloads';records=[]
names=['01-proverochnyy-list-komplekta-ks.docx','03-tipovye-prichiny-vozvrata.docx']
with ZipFile(r/'vozvrat-ks.zip')as z:
 for name in names:
  raw=z.read(name);d=Document(io.BytesIO(raw));changes=[]
  for node in d.element.body.iter(qn('w:p')):
   p=Paragraph(node,d);t=p.text;new=t.replace('КС-2 (форма по ОКУД 0322001)','КС-2 (форма по ОКУД 0322005)').replace('КС-3 (форма по ОКУД 0322002)','КС-3 (форма по ОКУД 0322001)').replace('Подписи подрядчика (руководитель, бухгалтер)','Подписи уполномоченных представителей; проверить полномочия').replace('Журнал учета выполненных работ (КС-6)','Журнал учёта выполненных работ (КС-6а)').replace('Ведомость ресурсов (КС-4)','Ведомость ресурсов, если предусмотрена договором').replace('Раздел 4. Обязательные приложения','Раздел 4. Приложения по договору и применимым требованиям').replace('Отсутствие подписи бухгалтера на КС-3','Отсутствие подписи уполномоченного представителя стороны').replace('Причина 6. Нет подписи бухгалтера на КС-3','Причина 6. Не подтверждены подпись и полномочия').replace('На КС-3 подписаны только руководители, а бухгалтер не подписал.Решение: по форме КС-3 подпись бухгалтера обязательна. Проверяйте перед передачей.','Проверьте подписи представителей заказчика и подрядчика и их полномочия. Унифицированная форма КС-3 не устанавливает отдельную обязательную подпись бухгалтера. Внутреннее визирование бухгалтерией не подменяет подпись стороны и проверяется по принятому порядку.').replace('Сумма по КС-3 = сумме по КС-2','Сверены период, состав работ, накопительные суммы и НДС КС-2/КС-3')
   if t.startswith('КС-2 - 1 500 000'):new='Сверьте одинаковый период, состав предъявленных работ, накопительные графы и НДС. Не сравнивайте механически разные показатели КС-2 и КС-3. При расхождении найдите основание и исправьте расчёт, сохранив согласованный порядок округления.'
   if t.startswith('Сертификат на бетон действует'):new='Сопоставьте документ о качестве с видом продукции, партией, датой выпуска и применимыми требованиями. Истечение срока сертификата к дате работ само по себе не доказывает непригодность ранее выпущенной продукции. Запросите нужные подтверждения и проверьте паспорт партии; не присваивайте всем паспортам единый срок действия.'
   new=new.replace('На КС-3 подписаны только руководители, а бухгалтер не подписал.','Проверьте подписи уполномоченных представителей заказчика и подрядчика и основания их полномочий.').replace('Решение: по форме КС-3 подпись бухгалтера обязательна. Проверяйте перед передачей.','Унифицированная форма КС-3 не устанавливает отдельную обязательную подпись бухгалтера. Внутреннее согласование бухгалтерией определяется принятым порядком и не заменяет подписи сторон.').replace('Просроченные сертификаты на материалы','Документы о качестве не сопоставлены с продукцией, партией и применимыми требованиями')
   if new!=t:p.text=new;changes.append({'before':t,'after':new})
  d.add_paragraph('Состав приложений зависит от работ, проекта, договора и обязательных требований. КС-6а - журнал учёта выполненных работ; он не заменяет общий журнал работ. Унифицированные формы применяйте с учётом установленной формы первичного документа и требований конкретной сделки (ст. 9 ч. 2, 4 Закона №402-ФЗ).')
  for label,url in [('402-ФЗ ст. 9: реквизиты и формы','https://www.consultant.ru/document/cons_doc_LAW_122855/b24121ba7152f33673608559f2ca844ef5b6a74c/'),('КС-6а: форма и назначение','https://www.consultant.ru/document/cons_doc_LAW_26303/dc9b2aa6e521cdef7e924813c565fbec7c69ca6b/')]:
   p=d.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True));rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text=label;rr.append(tt);h.append(rr);p._p.append(h)
  d.add_paragraph('Поддержка: nastoronepodryadchika@yandex.ru. Тема: FREE-vozvrat-ks, имя файла и страница. Опишите расхождение и ожидаемый результат. Приложите пример без конфиденциальных данных.')
  assert 'подпись бухгалтера обязательна'not in '\n'.join(Paragraph(n,d).text for n in d.element.body.iter(qn('w:p')))
  out=q/'candidates/free/vozvrat-ks'/name;out.parent.mkdir(parents=True,exist_ok=True);d.save(out)
  records.append({'archive':'downloads/vozvrat-ks.zip','member':name,'path':str(out.relative_to(q/'candidates')),'base_sha256':hashlib.sha256(raw).hexdigest(),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'changes':changes,'status':'CANDIDATE_BUILT_NATIVE_WORD_BLOCKED'})
(q/'free-docx-manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2));print('FREE_DOCX',len(records),'TARGETED_REPLACEMENTS',sum(len(x['changes'])for x in records))
