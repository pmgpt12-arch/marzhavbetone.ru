#!/usr/bin/env python3
"""Книги P5 открываются без ошибок и считают итог.

Issue pmgpt12-arch/marzhavbetone.ru#294 (MB001-R1-P5). В двух покупательских
книгах P5 итоговая ячейка суммировала диапазон, в который входила сама:
`05-reestr-uderzhaniy.xlsx` E5 `=SUM(E2:E20)`, `06-grafik-vozmeshcheniya.xlsx`
D4 `=SUM(D2:D10)`. LibreOffice Calc показывал `Err:522` (циклическая ссылка),
после сохранения — `#VALUE!`. Имена листов были 37 и 33 символа при пределе
Excel 31. Существующий `test_delivery_artifacts.py` этого не ловил.

Два слоя:
1. Статический — только стандартная библиотека, по XML книги: ни одна
   формула не ссылается на собственную ячейку, имена листов ≤ 31 символа,
   итоговые формулы на месте.
2. Пересчёт в LibreOffice Calc headless: книга открывается, пересчитывается
   и сохраняется; в результате нет ошибок формул (`#VALUE!`, `#REF!`,
   `#NAME?`, `Err:5xx` …), формулы и имена листов пережили сохранение, а
   итог действительно суммирует введённые суммы. Нет Calc или openpyxl —
   слой печатается строкой «ПРОПУСК» с причиной, а не проходит молча;
   `P5_REQUIRE_CALC=1` превращает пропуск в провал.

Проверяются собранные книги, а не намерения генератора. Каталог можно
подменить — так тест прогоняется на версии из `main` до исправления:

    python3 tools/test_p5_spreadsheets.py
    python3 tools/test_p5_spreadsheets.py --dir /путь/к/книгам
    P5_BOOKS_DIR=/путь/к/книгам python3 -m pytest tools/test_p5_spreadsheets.py -q
"""
from __future__ import annotations

import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
P5 = ROOT / "products-storage" / "07-uderzhaniya-shtrafy-zachety"

ПРЕДЕЛ_ИМЕНИ_ЛИСТА = 31

# Итоги, ради которых книги существуют: без них «исправление» могло бы
# просто удалить формулу.
ИТОГИ = {
    "05-reestr-uderzhaniy.xlsx": "SUM(E2:E20)",
    "06-grafik-vozmeshcheniya.xlsx": "SUM(D2:D10)",
}

ЯЧЕЙКА_XML = re.compile(r'<c r="([A-Z]{1,3})(\d+)"[^>]*?(?:/>|>(.*?)</c>)', re.S)
ФОРМУЛА_XML = re.compile(r"<f[^>]*>([^<]*)</f>")
ЛИСТ_XML = re.compile(r'<sheet\b[^>]*\bname="([^"]*)"')
СТРОКОВЫЙ_ЛИТЕРАЛ = re.compile(r'"(?:[^"]|"")*"')
# Ссылка или диапазон, возможно с именем листа: A1, $A$1, A1:B20, 'Лист'!A1
ССЫЛКА = re.compile(
    r"(?:(?:'((?:[^']|'')+)'|([^\s'!(),:;=+\-*/&^<>]+))!)?"
    r"(?<![A-Za-z0-9_.])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![A-Za-z0-9_(])")
ОШИБКА_ФОРМУЛЫ = re.compile(
    r"#(?:VALUE|REF|NAME|DIV/0|N/A|NUM|NULL)[!?]?|Err:\d{3}|#ЗНАЧ|#ССЫЛ|#ИМЯ|#ДЕЛ/0")

ПРОПУСКИ: list[str] = []


def каталог() -> Path:
    for i, a in enumerate(sys.argv):
        if a == "--dir" and i + 1 < len(sys.argv):
            return Path(sys.argv[i + 1])
    return Path(os.environ.get("P5_BOOKS_DIR", P5))


def книги() -> list[Path]:
    найдено = sorted(каталог().glob("*.xlsx"))
    assert найдено, f"в {каталог()} нет книг — проверка ничего не проверяет"
    return найдено


def _col(буквы: str) -> int:
    n = 0
    for ch in буквы:
        n = n * 26 + ord(ch) - 64
    return n


def листы(книга: Path) -> list[tuple[str, str]]:
    """(имя листа, XML листа) в порядке книги."""
    with zipfile.ZipFile(книга) as z:
        имена = [html.unescape(n) for n in ЛИСТ_XML.findall(z.read("xl/workbook.xml").decode("utf-8"))]
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        wb = z.read("xl/workbook.xml").decode("utf-8")
        результат = []
        for m in re.finditer(r'<sheet\b[^>]*>', wb):
            тег = m.group(0)
            имя = html.unescape(re.search(r'\bname="([^"]*)"', тег).group(1))
            rid = re.search(r'r:id="([^"]+)"', тег).group(1)
            цель = re.search(rf'<Relationship\b[^>]*Id="{rid}"[^>]*Target="([^"]+)"', rels) \
                or re.search(rf'<Relationship\b[^>]*Target="([^"]+)"[^>]*Id="{rid}"', rels)
            путь = цель.group(1).lstrip("/")
            путь = путь if путь.startswith("xl/") else "xl/" + путь
            результат.append((имя, z.read(путь).decode("utf-8")))
        assert [и for и, _ in результат] == имена
        return результат


def формулы(xml: str) -> list[tuple[str, str]]:
    """(адрес, формула) для каждой ячейки листа с формулой."""
    out = []
    for m in ЯЧЕЙКА_XML.finditer(xml):
        f = ФОРМУЛА_XML.search(m.group(3) or "")
        if f and f.group(1):
            out.append((m.group(1) + m.group(2), html.unescape(f.group(1))))
    return out


def самоссылка(лист: str, адрес: str, формула: str) -> str | None:
    """Ссылка формулы, которая накрывает её собственную ячейку."""
    m = re.fullmatch(r"([A-Z]{1,3})(\d+)", адрес)
    col, row = _col(m.group(1)), int(m.group(2))
    текст = СТРОКОВЫЙ_ЛИТЕРАЛ.sub('""', формула)
    for s in ССЫЛКА.finditer(текст):
        чужой = s.group(1) or s.group(2)
        if чужой and чужой.replace("''", "'") != лист:
            continue
        c1, r1 = _col(s.group(3)), int(s.group(4))
        c2, r2 = (_col(s.group(5)), int(s.group(6))) if s.group(5) else (c1, r1)
        if min(c1, c2) <= col <= max(c1, c2) and min(r1, r2) <= row <= max(r1, r2):
            return s.group(0)
    return None


# ─────────────────────────── 1. статический слой ───────────────────────────

def test_формула_не_ссылается_на_свою_ячейку() -> None:
    плохие = []
    for книга in книги():
        for лист, xml in листы(книга):
            for адрес, формула in формулы(xml):
                ссылка = самоссылка(лист, адрес, формула)
                if ссылка:
                    плохие.append(f"{книга.name} [{лист}] {адрес} ={формула} — {ссылка} включает {адрес}")
    assert not плохие, "циклическая ссылка на собственную ячейку (Calc: Err:522):\n  " + "\n  ".join(плохие)


def test_имя_листа_не_длиннее_31_символа() -> None:
    длинные = [f"{книга.name}: «{лист}» — {len(лист)} символов"
               for книга in книги() for лист, _ in листы(книга)
               if len(лист) > ПРЕДЕЛ_ИМЕНИ_ЛИСТА]
    assert not длинные, f"имя листа длиннее {ПРЕДЕЛ_ИМЕНИ_ЛИСТА} (предел Excel):\n  " + "\n  ".join(длинные)


def test_итоговые_формулы_на_месте() -> None:
    по_имени = {к.name: к for к in книги()}
    нет = []
    for имя, итог in ИТОГИ.items():
        книга = по_имени.get(имя)
        if not книга:
            нет.append(f"{имя}: книги нет")
            continue
        все = [ф.replace(" ", "").lstrip("=") for _, xml in листы(книга) for _, ф in формулы(xml)]
        if итог not in все:
            нет.append(f"{имя}: нет формулы ={итог}, есть {все}")
    assert not нет, "итоговая формула пропала:\n  " + "\n  ".join(нет)


def test_детектор_самоссылки_различает_случаи() -> None:
    """Сторож самой проверки: иначе зелёный результат ничего не значит."""
    assert самоссылка("Л", "E5", "SUM(E2:E20)") == "E2:E20"
    assert самоссылка("Л", "D4", "SUM($D$2:$D$10)")
    assert самоссылка("Л", "A1", "A1+1")
    assert самоссылка("Л", "E21", "SUM(E2:E20)") is None
    assert самоссылка("Л", "E5", "SUM(Другой!E2:E20)") is None
    assert самоссылка("Л", "E5", "SUM('Л'!E2:E20)")
    assert самоссылка("Л", "E5", 'IF(A1="E5",1,0)') is None
    assert самоссылка("Л", "B2", "LOG10(100)") is None


# ─────────────────────────── 2. пересчёт в Calc ───────────────────────────

def _calc_ready() -> str | None:
    soffice = shutil.which("soffice")
    if not soffice:
        return "soffice не найден"
    # libreoffice-core ставит soffice без Calc: xlsx он не открывает
    # («source file could not be loaded»). Это неполная среда, не провал книг.
    программа = Path(soffice).resolve().parent
    if not (программа / "libsclo.so").exists():
        return f"LibreOffice без модуля Calc ({программа}), установлен только core"
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return "нет openpyxl"
    return None


def пересчитать(src: list[Path], tmp: Path) -> dict[str, Path]:
    """Открывает книги в Calc headless, пересчитывает и сохраняет в xlsx."""
    out = tmp / "calc"
    cmd = ["soffice", f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--norestore",
           "--calc", "--convert-to", "xlsx", "--outdir", str(out), *map(str, src)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    результат = {p.name: out / p.name for p in src}
    нет = [n for n, p in результат.items() if not p.exists()]
    assert not нет, f"Calc не открыл/не сохранил {нет}: {r.stdout[-500:]} {r.stderr[-500:]}"
    return результат


def ошибки_после_calc(книга: Path) -> list[str]:
    """Ячейки-ошибки в книге, сохранённой Calc (t="e" или текст ошибки)."""
    найдено = []
    for лист, xml in листы(книга):
        for m in ЯЧЕЙКА_XML.finditer(xml):
            тело = m.group(3) or ""
            v = re.search(r"<v>([^<]*)</v>", тело)
            значение = html.unescape(v.group(1)) if v else ""
            if 't="e"' in m.group(0)[:m.group(0).find(">")] or ОШИБКА_ФОРМУЛЫ.fullmatch(значение):
                найдено.append(f"[{лист}] {m.group(1)}{m.group(2)} = {значение or '(ошибка)'}")
    return найдено


def test_calc_пересчитывает_без_ошибок() -> None:
    причина = _calc_ready()
    if причина:
        ПРОПУСКИ.append(f"пересчёт Calc: {причина}")
        if os.environ.get("P5_REQUIRE_CALC") == "1":
            raise AssertionError(f"P5_REQUIRE_CALC=1, а Calc недоступен: {причина}")
        return
    import openpyxl
    tmp = Path(tempfile.mkdtemp(prefix="p5-calc-"))
    try:
        исходные = книги()
        # Функциональные копии: в строки ввода вписаны суммы, итог обязан их сложить.
        заполненные = []
        ожидание: dict[str, tuple[str, float]] = {}
        for книга in исходные:
            итог = ИТОГИ.get(книга.name)
            if not итог:
                continue
            wb = openpyxl.load_workbook(книга)
            ws = wb.worksheets[0]
            столбец = re.match(r"SUM\(([A-Z]+)", итог).group(1)
            адрес = next(c.coordinate for row in ws.iter_rows() for c in row
                         if isinstance(c.value, str) and c.value.replace(" ", "").lstrip("=") == итог)
            ws[f"{столбец}2"], ws[f"{столбец}3"] = 1500.5, 2500
            копия = tmp / f"filled-{книга.name}"
            wb.save(копия)
            заполненные.append(копия)
            ожидание[копия.name] = (адрес, 4000.5)
        calc = пересчитать(исходные + заполненные, tmp)

        провалы = []
        for книга in исходные:
            после = calc[книга.name]
            for ош in ошибки_после_calc(после):
                провалы.append(f"{книга.name}: {ош}")
            до_имена = [л for л, _ in листы(книга)]
            после_имена = [л for л, _ in листы(после)]
            if до_имена != после_имена:
                провалы.append(f"{книга.name}: Calc изменил имена листов {до_имена} → {после_имена}")
            до_ф = sorted(ф.replace(" ", "").lstrip("=") for _, x in листы(книга) for _, ф in формулы(x))
            после_ф = sorted(ф.replace(" ", "").lstrip("=").removeprefix("of:=")
                             for _, x in листы(после) for _, ф in формулы(x))
            if до_ф != после_ф:
                провалы.append(f"{книга.name}: формулы после сохранения Calc {после_ф}, было {до_ф}")
        for имя, (адрес, сумма) in ожидание.items():
            wb = openpyxl.load_workbook(calc[имя], data_only=True)
            значение = wb.worksheets[0][адрес].value
            if not isinstance(значение, (int, float)) or abs(значение - сумма) > 1e-9:
                провалы.append(f"{имя}: итог {адрес} = {значение!r}, ожидалось {сумма}")
        assert not провалы, "после пересчёта Calc:\n  " + "\n  ".join(провалы)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main() -> int:
    провал = 0
    print(f"книги: {каталог()}")
    for имя, проверка in sorted(globals().items()):
        if not имя.startswith("test_"):
            continue
        try:
            проверка()
        except AssertionError as ошибка:
            провал += 1
            print(f"ПРОВАЛ  {имя}\n{ошибка}")
        else:
            print(f"ок  {имя}")
    for п in ПРОПУСКИ:
        print(f"ПРОПУСК  {п}")
    return 1 if провал else 0


if __name__ == "__main__":
    sys.exit(main())
