#!/usr/bin/env python3
"""P1 delivery list and mirrored buyer assets stay aligned."""
from hashlib import sha256
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "products-storage/01-zakrytie-rabot"
MIRROR = ROOT / "products-storage/04-polnyy-komplekt-pto/01-30-bazovye-pakety/01-zakrytie-rabot"
deliverables = [p for p in BASE.iterdir() if p.is_file() and p.suffix.lower() in (".docx", ".pdf", ".xlsx")]
for path in deliverables:
    partner = MIRROR / path.name
    assert partner.is_file(), (path.name, "mirror missing")
    assert sha256(path.read_bytes()).digest() == sha256(partner.read_bytes()).digest(), (path.name, "mirror differs")
instruction = Document(BASE / "00-INSTRUKCIYA.docx")
listed = [row.cells[0].text for row in instruction.tables[0].rows[1:]]
expected = sorted(p.name for p in deliverables if p.name not in ("00-INSTRUKCIYA.docx", "00-INSTRUKCIYA.pdf"))
assert sorted(listed) == expected, ("instruction differs", sorted(set(expected) - set(listed)), sorted(set(listed) - set(expected)))
assert len(listed) == 16, len(listed)
assert any("16.08.2026" in p.text for p in instruction.paragraphs), "source date not disclosed"
print("P1 delivery alignment: 16 listed files, all mirrored buyer assets PASS")
