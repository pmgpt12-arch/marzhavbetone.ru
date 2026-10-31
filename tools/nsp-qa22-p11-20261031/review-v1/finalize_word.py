from pathlib import Path
from docx import Document
from docx.shared import Pt
R=Path(__file__).resolve().parent
f=next((R/'candidate').glob('53-*.docx'));d=Document(f)
for p in d.paragraphs:
 if p.text.startswith('5. При предлагаемых'):p.text='5. Сведения об известных ограничениях передачи или финансирования в доступной для раскрытия части и документы о необходимых согласиях.'
d.save(f)
f=next((R/'candidate').glob('56-*.docx'));d=Document(f)
d.styles['Heading 1'].paragraph_format.space_before=Pt(6)
d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.save(f)
