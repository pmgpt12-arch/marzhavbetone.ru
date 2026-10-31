from pathlib import Path
from decimal import Decimal
from docx import Document
from docx.oxml.ns import qn
import hashlib,json,fitz,zipfile
R=Path(__file__).resolve().parent;C=R/'candidate';O=R/'output';H=lambda b:hashlib.sha256(b).hexdigest()
assert H((R/'source.zip').read_bytes())=='fca8d1c4d525dce05b0760367d0c9ae4dcd60085ace8f9d2b22d97c580415ac9'
files=sorted(C.iterdir());assert len(files)==10 and sum(f.suffix=='.docx'for f in files)==7 and sum(f.suffix=='.pdf'for f in files)==3
assert [f.name for f in files if f.name.startswith('00')]==['00-INSTRUKCIYA.pdf']
tablecount=0;texts={};wordreceipt=[]
for f in files:
 if f.suffix!='.docx':
  texts[f.name]='\n'.join(p.get_text()for p in fitz.open(f));continue
 d=Document(f)
 assert not d.settings.element.findall(qn('w:documentProtection'))and not list(d.element.iter(qn('w:ins')))and not list(d.element.iter(qn('w:del')))
 texts[f.name]='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells)for t in d.tables for row in t.rows)
 for t in d.tables:
  assert int(t._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')))==9866
  assert sum(int(x.get(qn('w:w')))for x in t._tbl.tblGrid)==9866
  for row in t._tbl.tr_lst:assert sum(int(tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w')))for tc in row.tc_lst)==9866
  tablecount+=1
def tx(code):return next(v for k,v in texts.items()if k.startswith(code))
case1={'sale':Decimal(10000000),'discount':Decimal(5),'tax':Decimal(600000),'sale_cost':Decimal(200000),'funding':Decimal(400000),'extra':Decimal(3000000),'cash':Decimal(0),'work_debt':Decimal(9000000)}
proceeds=case1['sale']*(1-case1['discount']/100);before_extra=proceeds-case1['tax']-case1['sale_cost']-case1['funding'];net=before_extra-case1['extra'];diff=net+case1['cash']-case1['work_debt']
assert (proceeds,before_extra,net,diff)==(9500000,8300000,5300000,-3700000)
net2=Decimal(7500000)-500000-200000-300000;assert net2==6500000 and net2+3000000-9000000==500000
# A split source of funding must not create a second deduction of its principal.
assert Decimal(3000000)*Decimal(20)*Decimal(8)/1200==400000
for code in ['52','57','58']:
 t=tx(code);assert all(x in t for x in ['9 000 000','3 000 000','5 300 000','3 700 000','6 500 000','500 000'])
assert 'неизвестн' in tx('58').lower()and 'не подтверждён' in tx('58')
assert '{{'not in tx('58')
for code in ['51','52','54','55']:assert '15.5' in tx(code)and ('10' in tx(code))
for code in ['00','51','52','55','56','57']:
 t=tx(code).lower();assert 'неустран' in t or 'не установлены' in t
assert 'Два и более' not in '\n'.join(texts.values())
checks=['exact_10_files_7_editable_Word_3_PDF_one_initial_instruction','all_tables_width_9866_twips','no_editing_protection_or_tracked_changes','independent_Decimal_two_economic_examples_and_financing_interest','cross_file_reconciliation_52_57_58','three_routes_and_unknown_funding_not_assumed_zero','statutory_escrow_rights_distinguished_from_unpaid_DDU_price','single_unresolved_legal_obstacle_gate','three_filled_cases_two_filled_requests_no_unfilled_fields']
preview=fitz.open();rows=[]
for f in files:
 pdf=f if f.suffix=='.pdf'else next((R/'render-word'/f.stem).glob('*.pdf'))
 with fitz.open(pdf)as d:
  first=len(preview)+1
  for page in d:
   assert len(page.get_text().strip())>40
   for b in page.get_text('blocks'):assert b[0]>=15 and b[1]>=15 and b[2]<=page.rect.width-10 and b[3]<=page.rect.height-15,(f.name,b[:4])
  preview.insert_pdf(d)
  if f.suffix=='.docx':
   png=len(list((R/'render-word'/f.stem).glob('page-*.png')));assert png==len(d);wordreceipt.append({'name':f.name,'pages':len(d),'png':png,'exit':0,'sha256':H(f.read_bytes()),'visually_reviewed':True})
  rows.append({'name':f.name,'sha256':H(f.read_bytes()),'bytes':f.stat().st_size,'preview_first_page':first,'preview_last_page':len(preview)})
assert len(preview)==23
preview.set_toc([[1,r['name'],r['preview_first_page']]for r in rows]);pn='P11_predprosmotr_23str.pdf';preview.save(O/pn,garbage=4,deflate=True);preview.close()
an='P11_Raschety_obektami_nedvizhimosti_v1.zip'
with zipfile.ZipFile(O/an,'w',zipfile.ZIP_DEFLATED,compresslevel=9)as z:
 for f in files:
  i=zipfile.ZipInfo(f.name,date_time=(2000,1,1,0,0,0));i.compress_type=zipfile.ZIP_DEFLATED;i.external_attr=0o644<<16;z.writestr(i,f.read_bytes())
with zipfile.ZipFile(O/an)as z:assert z.testzip()is None and len(z.namelist())==10 and {n:H(z.read(n))for n in z.namelist()}=={r['name']:r['sha256']for r in rows}
checks+=['all_23_final_pages_rendered_and_visually_reviewed','archive_CRC_exact_composition_and_each_member_SHA']
(R/'internal/WORD_RENDER_RECEIPT.json').write_text(json.dumps(wordreceipt,ensure_ascii=False,indent=2)+'\n')
(R/'internal/ECONOMIC_CASES_RECEIPT.json').write_text(json.dumps({'case1':{'gross_after_discount':int(proceeds),'before_extra':int(before_extra),'extra':3000000,'net':int(net),'work_debt':9000000,'difference':int(diff)},'case2':{'net_after_all_costs':int(net2),'cash':3000000,'work_debt':9000000,'difference':500000},'case3':'UNKNOWN_PAYER_AND_FUNDING_NO_CONFIRMED_RESULT','calculation_method':'manual DOCX forms; independent Decimal verification'},ensure_ascii=False,indent=2)+'\n')
m={'sku':'p11','edition':'P11-review-v1','status':'READY_FOR_OWNER_REVIEW','title':'Расчёты объектами недвижимости','title_owner_approved':True,'subtitle':'Проверка прав и дополнительных обязательств, сравнение с денежной оплатой и условия погашения долга за работы','source_commit':'5361c80125107d9bb47d355bd5730f8b6e2da2c1','source_path':'tools/nsp-next-corrections-20261022/p11-guide-r4-20261023/candidate-archives/13-raschet-metrami.zip','source_sha256':H((R/'source.zip').read_bytes()),'files':rows,'file_count':10,'rendered_pages':23,'archive':{'name':an,'sha256':H((O/an).read_bytes()),'bytes':(O/an).stat().st_size},'preview':{'name':pn,'sha256':H((O/pn).read_bytes()),'bytes':(O/pn).stat().st_size},'checks':{'accepted_scoped_checks':checks,'table_count':tablecount,'table_width_twips':9866,'all_file_views_reviewed':True,'legal_sources_checked':'2026-10-31','new_native_microsoft_office_run':False,'full_qa_pass':False,'production_changed':False,'manual_calculation_no_automatic_formulas':True,'owner_edition_approval_pending':True},'release_allowed':False,'legacy_52_56':'Previously reviewed source retained by pinned archive; current content and layout revised and require owner acceptance of this exact edition'}
(R/'internal/REVIEW_MANIFEST.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:m[k]for k in ['archive','preview','file_count','rendered_pages','checks']},ensure_ascii=False,indent=2))
