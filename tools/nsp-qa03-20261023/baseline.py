"""Freeze actual public-source tree, accepted QA01 letter and protected candidates."""
from pathlib import Path
import json,hashlib,zipfile,shutil,subprocess
from lxml import etree as E
from docx import Document
O=Path(__file__).resolve().parent;R=O.parent.parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip()
assert head=='1ee5055b63611cd91bb09bb61a8b33bf1efba1b8'
source=O/'source';source.mkdir(exist_ok=False)
archive=Path('/tmp/qa03-source-uderzhaniya.zip');assert sha(archive)=='07f372b8cd43b3a6a31f20834290e99ac25a2012ed91472ec01cd12004281122'
with zipfile.ZipFile(archive)as z:
 assert z.testzip()is None and len(z.infolist())==5
 for i in z.infolist():
  assert '/'not in i.filename and '\\'not in i.filename
  (source/i.filename).write_bytes(z.read(i))
accepted=R/'tools/nsp-qa01-20261023/retention/candidate/04-pismo-o-vozvrate-uderzhaniya.docx'
assert sha(accepted)=='5647f33e362ed59338ea3fc46783f2d61572093bd0378353e6b3e459b1337e60'
(O/'accepted-qa01').mkdir();shutil.copyfile(accepted,O/'accepted-qa01'/accepted.name)
candidate=O/'candidate';candidate.mkdir()
for p in source.iterdir():shutil.copyfile(p,candidate/p.name)
shutil.copyfile(accepted,candidate/accepted.name)
q2=R/'tools/nsp-qa02-20261023';q2manifest=json.loads((q2/'MANIFEST.json').read_text());assert sha(q2/q2manifest['archive'])==q2manifest['archive_sha256']
for f in q2manifest['entries']:assert sha(q2/'candidate'/f['name'])==f['candidate_sha256']
matrix=json.loads((R/'tools/nsp-qa01-20261023/READINESS_MATRIX.json').read_text());row=next(x for x in matrix['free']if x['id']=='FREE/uderzhaniya')
prior=R/'tools/nsp-qa01-20261023/retention/uderzhaniya-QA-only.zip';assert sha(prior)==row['candidate_sha256']
with zipfile.ZipFile(prior)as z:
 assert z.testzip()is None
 for p in candidate.iterdir():assert z.read(p.name)==p.read_bytes()
protected=json.loads((q2/'BASELINE.json').read_text())['preserved_13']
for f in protected:
 assert sha(R.parent/'qa01-repo'/f['path'])==f['sha256']
 blob=subprocess.check_output(['git','rev-parse',head+':'+f['path']],cwd=R,text=True).strip();assert blob==f['git_blob']
baseline=dict(source_head=head,branch='qa/nsp-qa03-uderzhaniya-20261023',source_zip_sha256=sha(archive),qa01_zip_sha256=sha(prior),accepted_qa01_letter_sha256=sha(accepted),source_files=[dict(name=p.name,bytes=p.stat().st_size,sha256=sha(p),starting_candidate_sha256=sha(candidate/p.name))for p in sorted(source.iterdir())],qa01_matrix=row,protected_13=protected,qa02_manifest_sha256=sha(q2/'MANIFEST.json'),qa02_archive_sha256=q2manifest['archive_sha256'])
(O/'BASELINE.json').write_text(json.dumps(baseline,ensure_ascii=False,indent=2)+'\n')
texts={}
for p in list(source.glob('*.docx'))+[accepted]:
 d=Document(p);texts[('accepted-qa01/'if p==accepted else'source/')+p.name]=dict(paragraphs=[dict(index=i,style=x.style.name,text=''.join(x._p.itertext())if False else ''.join(x._p.xpath('.//w:t/text()')))for i,x in enumerate(d.paragraphs)],tables=[[[c.text for c in row.cells]for row in t.rows]for t in d.tables])
(O/'SOURCE_TEXT.json').write_text(json.dumps(texts,ensure_ascii=False,indent=2)+'\n')
print('BASELINE PASS: public-source ZIP5, latest QA01 ZIP5, accepted letter, QA02 and protected13')
