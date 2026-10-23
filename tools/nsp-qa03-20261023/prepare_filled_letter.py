"""Synthetic worked proof, copied from the candidate; never a sixth ZIP attachment."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy
import json,hashlib,xml.etree.ElementTree as E
O=Path(__file__).resolve().parent;src=O/'candidate/04-pismo-o-vozvrate-uderzhaniya.docx';d=Document(src)
replacements={'[№ и дата]':'УЧ-01 от 15.10.2025','[наименование второй стороны]':'ООО «Учебный заказчик»','[наименование вашей организации]':'ООО «Учебный подрядчик»','[наименование объекта]':'«Учебный объект»','[номер пункта]':'8.1','[условие возврата дословно из договора]':'в течение пяти рабочих дней после окончания гарантийного периода','[дата наступления]':'15.10.2026','[наименование, номер и дата документа]':'акт №1 от 15.10.2025; п. 8.2 договора','[пункт, число и тип дней; крайний день либо срок после доставки этого первичного требования]':'п. 8.2, пять рабочих дней; крайний день 22.10.2026','[дата]':'23.10.2026','[сумма]':'80 000 рублей','[сумма прописью]':'Восемьдесят тысяч рублей','[Укажите фактические сведения о замечаниях к вашим работам и документы, подтверждающие выполнение относящихся к возврату условий договора.]':'По учебному примеру относящиеся к возврату замечания не предъявлены; подтверждение — акт сверки от 23.10.2026.','[банковские реквизиты]':'учебные реквизиты — не для перевода денег','[адрес]':'условный адрес для учебной проверки','[электронная почта]':'qa@example.invalid','[документа, подтверждающего наступление события]':'акта №1 и выписки из п. 8.2 договора','[Копии актов об устранении замечаний / иные документы]':'Акт сверки от 23.10.2026; ПП №2 от 20.10.2026','[__]':'1','[Должность]':'Уполномоченное лицо','[ФИО]':'УЧЕБНЫЙ ПРИМЕР','[подпись]':'подпись не ставится'}
for p in d.paragraphs:
 if any(k in p.text for k in replacements):
  text=p.text
  for a,b in replacements.items():text=text.replace(a,b)
  p.text=text
# Retain only the actual sendable text plus its calculation, preserving their XML.
body=d._element.body;start=next(i for i,e in enumerate(body)if e.tag==qn('w:p')and ''.join(e.itertext()).startswith('Текст письма'))
end=next(i for i,e in enumerate(body)if e.tag==qn('w:p')and 'Как отправлять'in ''.join(e.itertext()))
keep=list(body)[start:end];sect=deepcopy(body.sectPr)
for e in list(body):body.remove(e)
for e in keep:body.append(e)
body.append(sect)
t=d.tables[0]
for j,x in enumerate(['1','КС-2 №1 от 15.10.2025','1 000 000','100 000','20 000','0','80 000']):t.cell(1,j).text=x
last=t.rows[-1].cells
for j,x in enumerate(['Итого','Итого','1 000 000','100 000','20 000','0','80 000']):last[j].text=x
folder=O/'native-extra/word/nsp-qa03-word-filled-v1-20261023';folder.mkdir(parents=True,exist_ok=False);p=folder/'04-filled-example.docx';d.save(p)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
r=E.Element('job',id=folder.name,qa_only='true');E.SubElement(r,'file',name=p.name,tag='q3df',sha256=sha(p));E.indent(r);E.ElementTree(r).write(folder/'job.xml',encoding='utf-8',xml_declaration=True)
(O/'proof/FILLED_LETTER.json').write_text(json.dumps(dict(source_sha256=sha(src),proof_sha256=sha(p),synthetic_only=True,amount=80000,held=100000,cash_already_returned=20000,other_discharge=0,event='2026-10-15',contract_deadline='2026-10-22',as_of='2026-10-23',not_in_zip=True),indent=2)+'\n')
print('Synthetic two-part letter proof prepared; candidate unchanged')
