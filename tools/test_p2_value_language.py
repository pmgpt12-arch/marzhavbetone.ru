"""Buyer P2 files speak about use, contents and next actions without disclaimer rhetoric."""
from pathlib import Path
from docx import Document
import fitz
import zipfile
import re

BASE = Path(__file__).resolve().parents[1] / 'products-storage/02-dopraboty-bez-poter'
FORBIDDEN = [
    r'шаблон\w*',
    r'юрист\w*',
    r'типов\w*',
    r'юридическ\w*\s+консультац\w*',
    r'важно, прочитайте до использования',
    r'это типовой шаблон, а не',
    r'не юридическая консультация',
    r'не заменяет (?:юридическую|консультацию|проверку|бухгалтерию)',
    r'не (?:является )?гаранти[ея]',
    r'не гарантирует (?:оплату|взыскание)',
    r'чего (?:этот документ|этот комплект) не делает',
    r'зачем этот документ',
    r'когда шаблон откладывается',
    r'откладывайте шаблон',
    r'главный риск файла',
    r'ложную уверенность',
    r'шаблон требует адаптации',
    r'этот комплект односторонний акт не покрывает',
]

def text_of(path):
    if path.suffix == '.docx':
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] +
                         [cell.text for table in doc.tables for row in table.rows for cell in row.cells] +
                         [p.text for section in doc.sections for p in section.header.paragraphs] +
                         [p.text for section in doc.sections for p in section.footer.paragraphs])
    if path.suffix == '.pdf':
        return '\n'.join(page.get_text() for page in fitz.open(path))
    if path.suffix == '.xlsx':
        with zipfile.ZipFile(path) as archive:
            return '\n'.join(archive.read(n).decode() for n in archive.namelist()
                             if n.startswith('xl/') and n.endswith('.xml'))
    return path.read_text()

files = [p for p in BASE.iterdir() if p.suffix.lower() in {'.docx', '.pdf', '.xlsx', '.txt'}]
assert len(files) == 16, len(files)
failures = [(p.name, phrase) for p in files for phrase in FORBIDDEN
            if re.search(phrase, text_of(p), re.I)]
page = (BASE.parents[1] / 'products/p2-dopolnitelnye-raboty.html').read_text()
failures += [('products/p2-dopolnitelnye-raboty.html', phrase) for phrase in FORBIDDEN
             if re.search(phrase, page, re.I)]
CANONICAL = 'Дополнительные работы: как получить оплату'
email_paths = [
    BASE / '00-PISMO-POSLE-POKUPKI.txt',
    BASE.parent / '04-polnyy-komplekt-pto/01-30-bazovye-pakety/02-dopraboty-bez-poter/00-PISMO-POSLE-POKUPKI.txt',
]
for email_path in email_paths:
    email = text_of(email_path)
    if CANONICAL not in email or 'Допработы без потерь' in email:
        failures.append((str(email_path.relative_to(BASE.parents[1])), 'canonical product name'))
    if '00-INSTRUKCIYA.pdf' not in email or '00_Инструкция.pdf' in email:
        failures.append((str(email_path.relative_to(BASE.parents[1])), 'delivered instruction filename'))
    failures += [(str(email_path.relative_to(BASE.parents[1])), phrase) for phrase in FORBIDDEN
                 if re.search(phrase, email, re.I)]

config = (BASE.parents[1] / 'products-config.php').read_text()
assert '$body .= "- документы редактируются в Microsoft Word и Excel;\\n";' in config
assert '$body .= "- шаблоны редактируются в Microsoft Word и Excel;\\n";' not in config

assert not failures, failures
print(f'P2 value language PASS: {len(files)} buyer-facing files, mirrored email, runtime email and product page')
