from pathlib import Path
from copy import deepcopy
from collections import Counter
import hashlib, io, json, re, shutil, zipfile
from docx import Document
from docx.shared import Twips, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_LEFT
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parent
SRC=ROOT/'source'; OUT=ROOT/'candidate'; INTERNAL=ROOT/'internal'
assert not OUT.exists(), 'Inspect previous candidate before rebuilding'
OUT.mkdir(); INTERNAL.mkdir(exist_ok=True)
H=lambda b:hashlib.sha256(b).hexdigest()
assert H((ROOT/'source.zip').read_bytes())=='e3dfe51b4f6d5c61c235037d7e737cb2dee3e688ecde844c1eb120a450c5e2b2'
W=9866
def elem(parent,tag,attrs):
    node=parent.find(qn(tag))
    if node is None: node=OxmlElement(tag); parent.append(node)
    for k,v in attrs.items(): node.set(qn(k),str(v))
    return node

def norm(doc):
    for sec in doc.sections: sec.left_margin=Twips(1020); sec.right_margin=Twips(1020)
    for t in doc.tables:
        count=len(t.columns)
        first=t.cell(0,0).text
        if count==2: widths=[3700,6166]
        elif count==3: widths=[550,6200,3116] if first=='№' else [1450,4350,4066]
        elif count==4: widths=[1850,2150,3166,2700]
        elif count==5: widths=[750,3000,2250,1050,2816]
        else: widths=[W//count]*count; widths[-1]+=W-sum(widths)
        t.autofit=False
        pr=t._tbl.tblPr
        elem(pr,'w:tblW',{'w:type':'dxa','w:w':W})
        elem(pr,'w:tblInd',{'w:type':'dxa','w:w':0})
        elem(pr,'w:jc',{'w:val':'left'})
        borders=elem(pr,'w:tblBorders',{})
        for edge in ['top','left','bottom','right','insideH','insideV']:
            elem(borders,'w:'+edge,{'w:val':'single','w:sz':4,'w:color':'D9D9D9'})
        margins=elem(pr,'w:tblCellMar',{})
        for edge,v in [('top',65),('bottom',65),('left',90),('right',90)]:elem(margins,'w:'+edge,{'w:type':'dxa','w:w':v})
        for col,w in zip(t.columns,widths): col.width=Twips(w)
        for rowidx,tr in enumerate(t._tbl.tr_lst):
            rp=tr.get_or_add_trPr(); elem(rp,'w:cantSplit',{})
            for height in list(rp.findall(qn('w:trHeight'))):
                if height.get(qn('w:hRule'))=='exact': height.set(qn('w:hRule'),'atLeast')
            pos=0
            for tc in tr.tc_lst:
                span=tc.grid_span; tp=tc.get_or_add_tcPr()
                elem(tp,'w:tcW',{'w:type':'dxa','w:w':sum(widths[pos:pos+span])});pos+=span
                elem(tp,'w:vAlign',{'w:val':'center'})
                for tag in ['w:noWrap','w:tcFitText']:
                    bad=tp.find(qn(tag))
                    if bad is not None:tp.remove(bad)
            if rowidx==0 and count>2:elem(rp,'w:tblHeader',{})
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:
                        if run.font.size and run.font.size.pt>11:run.font.size=Pt(10)
    return doc

def replace(doc,old,new):
    found=0
    for p in doc.paragraphs:
        if old in p.text:p.text=p.text.replace(old,new);found+=1
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if old in cell.text:cell.text=cell.text.replace(old,new);found+=1
    assert found,old

changes=[]
for f in sorted(SRC.iterdir()):
    if f.suffix!='.docx':continue
    d=Document(f)
    if f.name=='10-sravnitelnaya-tablica.docx':
        replace(d,'Состав: 7 DOCX, 2 XLSX, 1 PDF и инструкция TXT — 11 покупательских файлов.','Состав: 7 DOCX, 2 XLSX и 2 PDF — 11 покупательских файлов. Начните с 00-INSTRUKCIYA.pdf.')
        replace(d,'Пять заполняемых форм для применимых скрытых работ. Выбирайте по виду работ и проектным основаниям.','Одна редактируемая рабочая форма для применимых скрытых работ. Реквизиты и состав участников сверяйте с требованиями объекта; порядок применения — в PDF-инструкции.')
    if f.name=='09-tipovye-oshibki.docx':
        replace(d,'Выберите применимую форму №03; состав участников определите по фактическим условиям.','Используйте рабочую форму №03 после сверки реквизитов и состава участников с применимыми требованиями объекта.')
    if f.name=='07-sroki-hraneniya.docx':
        d.paragraphs[0].text='Определение сроков хранения документов комплекта КС'
        p=next(p for p in d.paragraphs if p.text.startswith('Ст. 29 Федерального закона'))
        p.text='Первичные учетные документы хранят по правилам архивного дела, но не менее пяти лет после отчетного года: часть 1 статьи 29 Федерального закона от 06.12.2011 № 402-ФЗ. При более длительном применимом сроке или незавершенном споре уничтожение не допускайте. Источник: https://www.consultant.ru/document/cons_doc_LAW_122855/eb7eaf7415d317fa9daf978473137c9209c1e1d7/'
        d.tables[0].cell(1,3).text='Не менее пяти лет после отчетного года; проверьте более длительные архивные, налоговые, договорные и спорные основания.'
        for index in [2,3]:
            d.tables[0].cell(index,3).text='Для исполнительной документации — весь срок эксплуатации объекта по пункту 8 Порядка №344/пр. Держателя и передачу определяйте отдельно.'
        p=next(p for p in d.paragraphs if p.text.startswith('Приказ Минстроя России от 16.05.2023'))
        p.text='Исполнительная документация хранится весь срок эксплуатации объекта: пункт 8 Порядка, утвержденного приказом Минстроя России от 16.05.2023 №344/пр. В процессе строительства ее хранит лицо, осуществляющее строительство (пункт 6). Передачу на постоянное хранение оформляют по пункту 7 и условиям договора. Не считайте передачу заказчику разрешением уничтожить требуемые экземпляры и доказательства у подрядчика. Источник: https://www.consultant.ru/document/cons_doc_LAW_448832/'
        p=next(p for p in d.paragraphs if p.text.startswith('Приказ Росархива'))
        p.text='Приказ Росархива от 20.12.2019 №236: срок выбирают по точной категории и функции документа. Укажите применимую статью перечня и событие начала срока в номенклатуре дел. Источник: https://publication.pravo.gov.ru/Document/View/0001202002070036'
        d.tables[0].cell(7,2).text='Предмет переписки: цена, приемка, недостатки, претензия, электронная подпись. Определите точную категорию по содержанию вопроса.'
    if f.name=='05-pismo-o-nedostatkah.docx':
        block=next(p for p in d.paragraphs if '\nКому:' in p.text and '\nОт:' in p.text)
        lines=block.text.splitlines();assert len(lines)==3 and lines[0].startswith('Исх.')
        t=d.add_table(rows=2, cols=2)
        for row,line in zip(t.rows,lines[1:]):
            label,body=line.split(':',1);row.cells[0].text=label;row.cells[1].text=body.strip()
        block.text=lines[0];block._p.addnext(t._tbl)
        block=next(p for p in d.paragraphs if p.text.startswith('[Должность уполномоченного лица]'))
        lines=block.text.splitlines();assert len(lines)==3
        t=d.add_table(rows=3, cols=2)
        values=[('Подпись уполномоченного лица',lines[0]),('Основание полномочий',lines[1]),('Печать',lines[2])]
        for row,values in zip(t.rows,values):row.cells[0].text,row.cells[1].text=values
        block._p.addprevious(t._tbl);block._p.getparent().remove(block._p)
    if f.name=='04-protokol-vnutrenney-proverki.docx':
        paras=[p for p in d.paragraphs if p.text.startswith(('Проверяющий __________________ /','Повторно принято __________________ /'))]
        assert len(paras)==2
        t=d.add_table(rows=2,cols=2)
        for row,p in zip(t.rows,paras):
            row.cells[0].text='Проверяющий' if p.text.startswith('Проверяющий') else 'Повторная проверка'
            row.cells[1].text='__________________ / __________________ / Дата __________'
        paras[0]._p.addprevious(t._tbl)
        for p in paras:p._p.getparent().remove(p._p)
    norm(d)
    if f.name=='10-sravnitelnaya-tablica.docx':
        widths=[550,2800,6516];t=d.tables[0]
        for col,w in zip(t.columns,widths):col.width=Twips(w)
        for row in t.rows:
            for c,w in zip(row.cells,widths):elem(c._tc.get_or_add_tcPr(),'w:tcW',{'w:type':'dxa','w:w':w})
    if f.name=='07-sroki-hraneniya.docx':
        for i,p in enumerate(d.paragraphs):
            if p.text.startswith('2. Алгоритм'):p.paragraph_format.page_break_before=True
            for run in p.runs:
                run.font.name='Arial';run.font.size=Pt(16 if i==0 else 11 if p.text.startswith(('1.','2.','3.')) else 9.5)
        for ri,row in enumerate(d.tables[0].rows):
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:run.font.name='Arial';run.font.size=Pt(9.5)
    d.save(OUT/f.name)
    changes.append({'name':f.name,'source_sha256':H(f.read_bytes()),'sha256':H((OUT/f.name).read_bytes()),'scope':'table geometry and listed content corrections'})

# The owner's accepted editable form remains byte-identical to the published P2 form.
ao=(ROOT/'accepted-aosr.docx').read_bytes()
assert H(ao)=='a0b0e8f48f7f0b8310a96358c226038a2066a17c159021d6f34f742b7969d7d7'
(OUT/'03-akt-skrytyh-rabot.docx').write_bytes(ao)
for name in ['02-reestr-prilozheniy.xlsx','06-zhurnal-peredachi.xlsx','08-algoritm-proverki.pdf']:
    shutil.copyfile(SRC/name,OUT/name)

pdfmetrics.registerFont(TTFont('Nsp','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('NspBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
styles={
 'title':ParagraphStyle('Title',fontName='NspBold',fontSize=19,leading=24,spaceAfter=15),
 'h':ParagraphStyle('Heading',fontName='NspBold',fontSize=12,leading=16,spaceBefore=12,spaceAfter=6,keepWithNext=True),
 'body':ParagraphStyle('Body',fontName='Nsp',fontSize=9.7,leading=14,spaceAfter=6),
 'small':ParagraphStyle('Small',fontName='Nsp',fontSize=8.3,leading=11.3,spaceAfter=5),
 'cell':ParagraphStyle('Cell',fontName='Nsp',fontSize=8.5,leading=11.5),
}
story=[]
def p(s,style='body'):story.append(Paragraph(escape(s).replace('\n','<br/>'),styles[style]))
def h(s):p(s,'h')
def step(num,title,body):h(str(num)+' '+title);p(body)
p('Проверка и передача КС-2 и КС-3','title')
p('P3  КС без возврата  |  Инструкция к комплекту  |  Редакция от 30 октября 2026 года','small')
p('Комплект помогает проверить документы до передачи заказчику, устранить замечания и сохранить подтверждение вручения. КС-2, КС-3, договор и смету берите из документов вашего объекта.')
step(1,'Соберите исходные документы','Зафиксируйте объект, договор, период, действующую смету и изменения. Срок рассмотрения, состав приложений и канал передачи найдите в договоре. Откройте 08-algoritm-proverki.pdf и проверочный лист №01.')
step(2,'Определите полный перечень приложений','В №02 внесите документы, требуемые для вашего вида работ по применимым нормам, проекту и договору. Учитывайте весь требуемый перечень, а не только имеющиеся файлы. Неприменимость пункта обоснуйте.')
step(3,'Проверьте объемы и суммы','Пройдите 40 пунктов №01. Разделите текущий период, итог с начала года и с начала работ. НДС и стоимость без налога сверяйте отдельно. Стоимость КС-3 не подменяет остаток долга.')
step(4,'Оформите скрытые работы','Используйте редактируемый файл №03 для применимых скрытых работ. До подписания проверьте состав реквизитов, участников, приложений и полномочий по требованиям объекта. Порядок применения формы — на странице 3 инструкции.')
step(5,'Закройте замечания','В №04 запишите замечание, действие, ответственного и срок. После исправления укажите документ, дату и результат повторной проверки. Положительное решение по комплекту принимают после закрытия замечаний.')
step(6,'Зафиксируйте передачу и ответ','В №06 разделите отправку и получение; укажите состав, адресата и подтверждение. Для ответа на замечания используйте №05. Повторную передачу регистрируйте отдельно, сохраняя историю.')
step(7,'Сохраните рабочее дело','Примените №07 к конкретным документам. Условия хранения, держателя и подтверждения передачи определите до завершения работы по объекту. Типичные ошибки и способы исправления — №09.')
story.append(PageBreak())
p('Состав комплекта','title')
p('11 файлов: 7 DOCX, 2 XLSX и 2 PDF. Начальная инструкция одна — 00-INSTRUKCIYA.pdf. Редактируемые формы Word заполняйте по фактическим документам объекта.')
purposes={
'00-INSTRUKCIYA.pdf':'Начало работы, состав, применение форм и таблиц',
'01-proverochnyy-list-komplekta-ks.docx':'40 пунктов проверки КС-2, КС-3 и приложений',
'02-reestr-prilozheniy.xlsx':'Полный перечень приложений и проверка реквизитов',
'03-akt-skrytyh-rabot.docx':'Редактируемая рабочая форма акта скрытых работ',
'04-protokol-vnutrenney-proverki.docx':'Замечания, исправления и повторная проверка',
'05-pismo-o-nedostatkah.docx':'Ответ заказчику по замечаниям и повторной передаче',
'06-zhurnal-peredachi.xlsx':'Отправка, получение, сроки, возврат и повторная передача',
'07-sroki-hraneniya.docx':'Классификация документов и основания хранения',
'08-algoritm-proverki.pdf':'Краткий порядок проверки и передачи комплекта',
'09-tipovye-oshibki.docx':'15 ошибок и адресные способы исправления',
'10-sravnitelnaya-tablica.docx':'Инструменты комплекта и проверяемые результаты'}
data=[[Paragraph('Файл',styles['cell']),Paragraph('Назначение',styles['cell'])]]
for name,desc in purposes.items():
    soft=escape(name).replace('-','-<wbr/>').replace('_','_<wbr/>')
    data.append([Paragraph(soft,styles['cell']),Paragraph(escape(desc),styles['cell'])])
t=Table(data,colWidths=[245,248],repeatRows=1,hAlign='LEFT')
t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E9EEF2')),('GRID',(0,0),(-1,-1),.45,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
h('Границы результата')
p('Проверенный перечень приложений подтверждает состав внесенных данных. Отправка комплекта, получение адресатом, приемка работ и оплата — отдельные события. Для каждого сохраняйте свое подтверждение и применимый срок.')
story.append(PageBreak())
p('Работа с формами и Excel','title')
h('Акт скрытых работ №03')
p('№03 — редактируемая рабочая форма акта скрытых работ. Она отличается от рекомендуемого образца приложения №3 к составу исполнительной документации по приказу Минстроя №344/пр. Перед применением сопоставьте форму с утвержденным перечнем ИД объекта и требуемым набором реквизитов.')
p('Укажите работы, участок, оси и отметки, проектные основания, даты, материалы, документы качества и связанные приложения. Проверьте роли участников и полномочия на дату подписания. Дополните рабочую форму необходимыми реквизитами; если для объекта принят рекомендуемый образец №344/пр, используйте согласованную на объекте форму. Освидетельствование проводят до закрытия работ последующими работами.')
p('Один и тот же рабочий шаблон может применяться к разным видам скрытых работ после проверки условий объекта. Он не заменяет специальные испытания, схемы и документы, необходимые для конкретного результата работ.')
h('Реестр приложений №02')
p('Сначала прочитайте пояснения внутри книги. Заполните полный применимый перечень, затем вид, номер, дату, количество листов и статус каждого документа. Дату вводите как дату Excel. Пустая строка или автоматический статус не доказывают полноту обязательного перечня; ее подтверждает проверяющий.')
h('Журнал передачи №06')
p('Указывайте отправку и подтверждение получения раздельно. Сохраните подписанную опись, отметку адресата или электронную квитанцию с привязкой к конкретным документам и редакциям. Договорной срок ответа задайте явно; перед проверкой просрочки проверьте контрольную дату.')
h('Общие правила заполнения')
p('Демонстрационные записи замените фактическими данными. Формулы не затирайте. Работайте в пределах настроенного диапазона; для следующей партии используйте отдельную копию книги. В Word замените поля в квадратных скобках и служебные пояснения рабочим текстом. Подписывает уполномоченное лицо.')
story.append(PageBreak())
p('Нормы и контроль перед отправкой','title')
h('Применимость документов')
p('КС-2 и КС-3 происходят из постановления Госкомстата России от 11.11.1999 №100. Применимость конкретных унифицированных форм определяйте по договору, специальным требованиям и правилам первичных учетных документов по статье 9 Закона №402-ФЗ. Само наличие КС не заменяет проверку фактической сдачи и приемки работ.')
p('Состав и порядок ведения исполнительной документации определены приказом Минстроя России от 16.05.2023 №344/пр в редакции от 23.06.2025. Приложение №1 к приказу содержит состав ИД; рекомендуемый образец АОСР находится в приложении №3 именно к этому составу. Не считайте все приложения обязательными для каждого акта.')
h('Хранение')
p('Для первичных учетных документов действует минимум пять лет после отчетного года по части 1 статьи 29 Закона №402-ФЗ. Исполнительную документацию хранят весь срок эксплуатации объекта по пункту 8 Порядка №344/пр; держателя и передачу определяют отдельно по пунктам 6–7 и договору. Для других документов проверьте фактическую функцию, архивный перечень и номенклатуру дел. При незавершенном споре документы сохраняйте.')
h('Проверьте результат')
p('Объект, договор и период однозначны. Объемы и суммы сверены. Перечень приложений подтвержден. Замечания закрыты и повторно проверены. Получатель, канал передачи и срок определены. Доказательства отправки и получения сохранены отдельно от решения о приемке и факта оплаты.')
h('Источники')
for text,url in [
('Постановление Госкомстата №100','https://www.consultant.ru/document/cons_doc_LAW_26273/'),
('Закон №402-ФЗ','https://www.consultant.ru/document/cons_doc_LAW_122855/'),
('Приказ Минстроя №344/пр','https://www.consultant.ru/document/cons_doc_LAW_448832/'),
('Официальная публикация приказа №344/пр','https://publication.pravo.gov.ru/document/0001202306010019'),
('Изменение №369/пр','https://publication.pravo.gov.ru/document/0001202507220009')]:
    story.append(Paragraph(escape(text)+': <link href="'+url+'" color="#244A60">'+escape(url)+'</link>',styles['small']))
p('Полный текст в некоммерческой версии КонсультантПлюс может быть доступен по расписанию. Нормативные ссылки проверены 30 октября 2026 года по доступным текстам; применимую редакцию проверяйте на дату использования.','small')
h('Помощь')
p('nastoronepodryadchika@yandex.ru  |  Тема: P3 — комплект КС. Укажите файл и пункт; для Excel — лист и ячейку. При необходимости приложите обезличенный пример. Если документы приняты, а срок оплаты прошел, используйте отдельный маршрут P1 на сайте nastoronepodryadchika.ru.')
def footer(canvas,doc):
    canvas.setFont('Nsp',8);canvas.setFillColor(colors.HexColor('#55606A'))
    canvas.drawString(51,810,'НА СТОРОНЕ ПОДРЯДЧИКА  /  P3')
    canvas.drawString(51,29,'Проверка и передача КС-2 и КС-3')
    canvas.drawRightString(544,29,str(doc.page))
pdf=OUT/'00-INSTRUKCIYA.pdf'
SimpleDocTemplate(str(pdf),pagesize=(595.276,841.89),leftMargin=51,rightMargin=51,topMargin=60,bottomMargin=48).build(story,onFirstPage=footer,onLaterPages=footer)
names=sorted(p.name for p in OUT.iterdir())
assert set(names)==set(purposes)
assert Counter(p.suffix for p in OUT.iterdir())=={'.docx':7,'.xlsx':2,'.pdf':2}
report={'status':'CANDIDATE_BUILT_VISUAL_PENDING','sku':'P3','edition':'P3-layout-v1','source_commit':'5361c80125107d9bb47d355bd5730f8b6e2da2c1','source_archive_sha256':H((ROOT/'source.zip').read_bytes()),'source_candidate_path':'tools/nsp-next-corrections-20261022/candidate-archives/05-ks-bez-vozvrata.zip','file_count':11,'single_instruction':'00-INSTRUKCIYA.pdf','editable_aosr_sha256':H(ao),'editable_aosr_reused_from_accepted_p2':True,'aosr_not_represented_as_exact_344_template':True,'excel_unchanged':all((OUT/n).read_bytes()==(SRC/n).read_bytes()for n in ['02-reestr-prilozheniy.xlsx','06-zhurnal-peredachi.xlsx']),'changes':changes,'files':[{'name':n,'bytes':(OUT/n).stat().st_size,'sha256':H((OUT/n).read_bytes())}for n in names],'owner_approval':'PENDING','release_allowed':False,'full_qa_pass':False,'native_tests_repeated':False}
(INTERNAL/'BUILD_RECEIPT.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items()if k not in ['changes','files']},ensure_ascii=False,indent=2))
