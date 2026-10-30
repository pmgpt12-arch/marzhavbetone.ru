from pathlib import Path
from collections import Counter
from lxml import etree
import hashlib,json,re,zipfile
from pypdf import PdfReader,PdfWriter
from docx import Document
ROOT=Path(__file__).resolve().parent;B=ROOT/'candidate';I=ROOT/'internal';O=ROOT/'output';O.mkdir(exist_ok=True)
H=lambda b:hashlib.sha256(b).hexdigest()
build=json.loads((I/'BUILD_RECEIPT.json').read_text());checks=[]
def check(name,condition):
    assert condition,name
    checks.append({'name':name,'status':'PASS'})
files=sorted(B.iterdir())
check('exact_11_members',len(files)==11 and all(p.is_file()for p in files))
check('one_initial_PDF_instruction',[p.name for p in files if p.name.startswith('00-')]==['00-INSTRUKCIYA.pdf'])
check('exact_format_counts',Counter(p.suffix for p in files)=={'.docx':7,'.xlsx':2,'.pdf':2})
for n in ['02-reestr-prilozheniy.xlsx','06-zhurnal-peredachi.xlsx','08-algoritm-proverki.pdf']:
    check('unchanged_'+n,(B/n).read_bytes()==(ROOT/'source'/n).read_bytes())
check('owner_accepted_editable_AOSR_exact',H((B/'03-akt-skrytyh-rabot.docx').read_bytes())=='a0b0e8f48f7f0b8310a96358c226038a2066a17c159021d6f34f742b7969d7d7')
render=json.loads((I/'RENDER_RECEIPT.json').read_text())
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};q=lambda n:'{'+ns['w']+'}'+n
tables=[];pagecounts={}
for p in files:
    if p.suffix!='.docx':continue
    with zipfile.ZipFile(p)as z:
        check('docx_crc_'+p.name,z.testzip()is None)
        xml=etree.fromstring(z.read('word/document.xml'))
        settings=etree.fromstring(z.read('word/settings.xml'))
        prot=settings.find('w:documentProtection',ns)
        check('editable_'+p.name,prot is None or prot.get(q('enforcement'))in [None,'0','false'])
        for ti,t in enumerate(xml.findall('.//w:tbl',ns),1):
            widths=[int(c.get(q('w')))for c in t.findall('w:tblGrid/w:gridCol',ns)]
            check('table_width_'+p.name+'_'+str(ti),sum(widths)==9866)
            layout=t.find('w:tblPr/w:tblLayout',ns)
            check('table_fixed_'+p.name+'_'+str(ti),layout is not None and layout.get(q('type'))=='fixed')
            tables.append({'file':p.name,'table':ti,'width_twips':sum(widths)})
        text=' '.join(xml.xpath('//w:t/text()',namespaces=ns))
        check('no_stale_guide_reference_'+p.name,not any(x in text for x in ['00-START-HERE.txt','03-akt-skrytyh-rabot.pdf','Пять заполняемых форм']))
    row=next(r for r in render if r['file']==p.name)
    check('exact_rendered_input_'+p.name,row['exit_code']==0 and row['sha256']==H(p.read_bytes()))
    pdf=ROOT/'render-final'/p.stem/(p.stem+'.pdf')
    pagecounts[p.name]=len(PdfReader(pdf).pages)
    check('render_has_all_page_images_'+p.name,all((pdf.parent/('page-'+str(k)+'.png')).is_file()for k in range(1,pagecounts[p.name]+1)))
check('all_15_tables_verified',len(tables)==15)
check('checklist_main_text_unchanged',Document(B/'01-proverochnyy-list-komplekta-ks.docx').element.body.xpath('.//w:t/text()')==Document(ROOT/'source/01-proverochnyy-list-komplekta-ks.docx').element.body.xpath('.//w:t/text()'))
for n in ['05-pismo-o-nedostatkah.docx','04-protokol-vnutrenney-proverki.docx']:
    source=Document(ROOT/'source'/n);out=Document(B/n)
    text=' '.join(out.element.body.xpath('.//w:t/text()'))
    required=[p.text for p in source.paragraphs if p.text.startswith(('По договору','Передаём документы','По разногласиям','Приложения:','Состав приложений','Для каждого замечания','Ответственный за передачу'))]
    check('required_outgoing_text_preserved_'+n,all(s in text for s in required))
guide=re.sub(r'\s+',' ',' '.join(p.extract_text()for p in PdfReader(B/'00-INSTRUKCIYA.pdf').pages))
check('guide_4_pages',len(PdfReader(B/'00-INSTRUKCIYA.pdf').pages)==4)
check('guide_lists_exact_files',all(p.name in guide for p in files))
check('guide_explains_working_AOSR','отличается от рекомендуемого образца' in guide)
check('guide_retention_separates_functions','весь срок эксплуатации объекта' in guide and 'пять лет после отчетного года' in guide)
check('no_obsolete_TXT_in_guide','00-START-HERE.txt'not in guide)
zipout=O/'P3_KS_bez_vozvrata_v1.zip';assert not zipout.exists()
with zipfile.ZipFile(zipout,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6)as z:
    for p in files:
        info=zipfile.ZipInfo(p.name,(2026,10,30,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,p.read_bytes())
with zipfile.ZipFile(zipout)as z:
    check('final_zip_crc_and_member_hashes',z.testzip()is None and z.namelist()==[p.name for p in files]and all(H(z.read(p.name))==H(p.read_bytes())for p in files))
preview=O/'P3_predprosmotr_19str.pdf';assert not preview.exists()
writer=PdfWriter();writer.append(B/'00-INSTRUKCIYA.pdf')
for p in files:
    if p.suffix=='.docx':writer.append(ROOT/'render-final'/p.stem/(p.stem+'.pdf'))
writer.append(B/'08-algoritm-proverki.pdf')
writer.add_metadata({'/Title':'P3 КС без возврата Предпросмотр Word и PDF','/Author':'На стороне подрядчика'})
with preview.open('wb')as f:writer.write(f)
check('preview_19_pages',len(PdfReader(preview).pages)==19)
receipt={**build,'status':'LAYOUT_COMPOSITION_AND_CRC_PASS_OWNER_REVIEW_PENDING','table_count':len(tables),'table_geometry':tables,'word_pdf_preview_pages':19,'native_office_retest_performed':False,'new_full_product_qa_claimed':False,'render_engine':'LibreOffice via packaged render_docx.py; not a native Microsoft Office receipt','checks':checks,'check_count':len(checks),'files':[{'name':p.name,'bytes':p.stat().st_size,'sha256':H(p.read_bytes())}for p in files],'archive':{'name':zipout.name,'sha256':H(zipout.read_bytes()),'bytes':zipout.stat().st_size},'preview':{'name':preview.name,'sha256':H(preview.read_bytes()),'bytes':preview.stat().st_size},'word_pages':pagecounts,'publication_permission':'PENDING_OWNER_APPROVAL'}
(I/'REVIEW_MANIFEST.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:receipt[k]for k in ['status','check_count','file_count','table_count','word_pdf_preview_pages','archive','preview']},ensure_ascii=False,indent=2))
