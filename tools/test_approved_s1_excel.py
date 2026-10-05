from pathlib import Path
import sys
import warnings
import pytest
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent))
import approved_s1_excel as layout

@pytest.mark.parametrize("filename", sorted(layout.APPROVED))
def test_native_owner_bytes_survive(filename, tmp_path):
    template = layout.TEMPLATES / filename
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        generated = openpyxl.load_workbook(template)
    output = tmp_path / filename
    import build_s1_candidate as builder
    {"02-proverka-i-kontrol-otveta.xlsx": builder.f02, "04-uchet-raschetov-i-otpravok.xlsx": builder.f04, "08-raschet-procentov-395.xlsx": builder.f08}[filename](tmp_path)
    assert output.read_bytes() == template.read_bytes()

def test_formula_divergence_blocks_before_existing_output_is_touched(tmp_path):
    filename = "04-uchet-raschetov-i-otpravok.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        generated = openpyxl.load_workbook(layout.TEMPLATES / filename)
    generated["Взаиморасчёты"]["I5"] = "=1"
    output = tmp_path / filename
    output.write_bytes(b"previous approved output")
    with pytest.raises(ValueError, match="semantic mismatch"):
        layout.write_approved_layout(generated, output)
    assert output.read_bytes() == b"previous approved output"
    assert not list(tmp_path.glob("*.tmp"))

def test_corrupt_template_fails_before_output(tmp_path, monkeypatch):
    filename = "04-uchet-raschetov-i-otpravok.xlsx"
    (tmp_path / filename).write_bytes(b"invalid")
    monkeypatch.setattr(layout, "TEMPLATES", tmp_path)
    with pytest.raises(ValueError, match="hash mismatch"):
        layout.write_approved_layout(None, tmp_path / "out" / filename)

def test_validation_divergence_blocks_before_output(tmp_path):
    filename = "04-uchet-raschetov-i-otpravok.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        generated = openpyxl.load_workbook(layout.TEMPLATES / filename)
    # Native x14 is read through XML by the gate; canonical generator owns normal rules.
    import build_s1_candidate as builder
    captured = {}
    original = layout.write_approved_layout
    try:
        layout.write_approved_layout = lambda wb, path: captured.update(workbook=wb)
        builder.f04(tmp_path)
    finally:
        layout.write_approved_layout = original
    generated = captured["workbook"]
    generated["Взаиморасчёты"].data_validations.dataValidation[0].formula1 = '"unexpected option"'
    output = tmp_path / filename
    with pytest.raises(ValueError, match="validation mismatch"):
        layout.write_approved_layout(generated, output)
    assert not output.exists()
