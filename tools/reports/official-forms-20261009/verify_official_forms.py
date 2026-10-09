from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).parent
ARCHIVES={
 'P1':ROOT.parent/'downloads/P1_Akty_KS2_KS3_19_files_OFFICIAL_FORMS_2026-10-09.zip',
 'P2':ROOT.parent/'downloads/P2_Dopolnitelnye_Raboty_15_files_OFFICIAL_FORMS_2026-10-09.zip'
}
counts={'P1':19,'P2':15}
for product,archive in ARCHIVES.items():
    with ZipFile(archive) as z:
        names=z.namelist()
        assert len(names)==counts[product] and len(names)==len(set(names))
        assert set(names)=={p.name for p in (ROOT/product).iterdir() if p.is_file()}
        for name in names:
            assert z.read(name)==(ROOT/product/name).read_bytes(),(product,name)
        instruction=Document(BytesIO(z.read('00-INSTRUKCIYA.docx')))
        text='\n'.join(p.text for p in instruction.paragraphs)
        pdf='\n'.join(p.extract_text() or '' for p in PdfReader(BytesIO(z.read('00-INSTRUKCIYA.pdf'))).pages)
        assert 'Происхождение и применение форм' in text and 'Происхождение и применение форм' in pdf
        assert 'КС-6а' in text if product=='P1' else '344/пр' in text
        if product=='P1':
            for name,code,ncols in [('01-ks-2.docx','0322005',8),('02-ks-3.docx','0322001',6)]:
                d=Document(BytesIO(z.read(name)))
                alltext=' '.join(p.text for p in d.paragraphs)+' '.join(c.text for t in d.tables for row in t.rows for c in row.cells)
                assert code in alltext and 'Госкомстата' in alltext
                assert any(len(t.columns)==ncols for t in d.tables)
            ks2=Document(BytesIO(z.read('01-ks-2.docx')))
            ks3=Document(BytesIO(z.read('02-ks-3.docx')))
            assert all(t in ' '.join(c.text for c in ks2.tables[1].rows[0].cells) for t in ['позиции по смете','единичной расценки'])
            assert all(t in ' '.join(c.text for c in ks3.tables[1].rows[0].cells) for t in ['Код','отчетный период'])
        else:
            d=Document(BytesIO(z.read('04-akt-skrytyh-rabot.docx')))
            body='\n'.join(p.text for p in d.paragraphs)
            assert '344/пр' in body and 'Рекомендуемый образец' in body
            assert all(f'{i}. ' in body for i in range(1,8))
            assert body.count('Представитель лица, осуществляющего строительство')>=4
            assert 'Приложения' in body and 'Подписи представителей' in body
    print(product,counts[product],'files and official-form gates passed')
