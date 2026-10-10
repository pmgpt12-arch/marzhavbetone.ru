"""Read-only comparison and bounded XML structural checks for native controls.
Usage: python verify_preservation.py SOURCE OUTPUT
Writes only OUTPUT/independent-review.json. Never modifies XLSX.
"""
import argparse, copy, hashlib, json, zipfile
from pathlib import Path
import openpyxl
from lxml import etree as E
SHEET_ORDER='sheetPr dimension sheetViews sheetFormatPr cols sheetData sheetCalcPr sheetProtection protectedRanges scenarios autoFilter sortState dataConsolidate customSheetViews mergeCells phoneticPr conditionalFormatting dataValidations hyperlinks printOptions pageMargins pageSetup headerFooter rowBreaks colBreaks customProperties cellWatches ignoredErrors smartTags drawing legacyDrawing legacyDrawingHF picture oleObjects controls webPublishItems tableParts extLst'.split()
BOOK_ORDER='fileVersion fileSharing workbookPr workbookProtection bookViews sheets functionGroups externalReferences definedNames calcPr oleSize customWorkbookViews pivotCaches smartTagPr smartTagTypes webPublishing fileRecoveryPr webPublishObjects extLst'.split()
def run(src,out):
    report_path=out/'independent-review.json'
    old=json.loads(report_path.read_text()) if report_path.exists() else None
    history=[] if old is None else old.get('history',[])+[{k:v for k,v in old.items() if k!='history'}]
    result={'scope':'Independent mechanical XLSX comparison; no normative, UI, print-render or full G3 acceptance','history':history,'files':[],'defects':[]}
    for name in ('07-raschet-stoimosti.xlsx','08-zhurnal-doprabot.xlsx'):
        a=openpyxl.load_workbook(src/name);b=openpyxl.load_workbook(out/name)
        vals=[];styles=[];locks=[];extra=[]
        for sa,sb in zip(a,b):
            for row in sa:
                for ca in row:
                    cb=sb[ca.coordinate]
                    if ca.value!=cb.value: vals.append([sa.title,ca.coordinate])
                    for k in ('font','fill','border','alignment','number_format'):
                        if copy.copy(getattr(ca,k))!=copy.copy(getattr(cb,k)): styles.append([sa.title,ca.coordinate,k])
                    expected=not(sa.title=='Реестр' and 4<=ca.row<=103 and not(name.startswith('07') and ca.column_letter in ('F','I')))
                    if cb.protection.locked!=expected: locks.append([sa.title,ca.coordinate])
            for row in sb:
                for cb in row:
                    if cb.value is not None and sa[cb.coordinate].value is None: extra.append([sb.title,cb.coordinate])
        f={'file':name,'sha256':hashlib.sha256((out/name).read_bytes()).hexdigest(),'sheet_names_equal':a.sheetnames==b.sheetnames,'values_formulas_differences':vals,'unexpected_new_values':extra,'visual_style_differences':styles,'lock_mismatches':locks,'freeze_merge_dimensions_preserved':all(sa.freeze_panes==sb.freeze_panes and str(sa.merged_cells)==str(sb.merged_cells) and sa.max_row==sb.max_row and sa.max_column==sb.max_column for sa,sb in zip(a,b)),'row_column_dimensions_preserved':all({k:dict(v) for k,v in sa.row_dimensions.items()}=={k:dict(v) for k,v in sb.row_dimensions.items()} and {k:dict(v) for k,v in sa.column_dimensions.items()}=={k:dict(v) for k,v in sb.column_dimensions.items()} for sa,sb in zip(a,b)),'validations_preserved':all(sa.data_validations==sb.data_validations for sa,sb in zip(a,b)),'conditional_formatting_preserved':all(sa.conditional_formatting._cf_rules==sb.conditional_formatting._cf_rules for sa,sb in zip(a,b)),'sheet_protection_enabled':all(s.protection.sheet and s.protection.selectLockedCells and not s.protection.selectUnlockedCells for s in b),'print_areas':{s.title:str(s.print_area) for s in b},'print_repeat_rows':b['Реестр'].print_title_rows,'print_setup':{s.title:{'paperSize':s.page_setup.paperSize,'orientation':s.page_setup.orientation,'fitToWidth':s.page_setup.fitToWidth,'fitToHeight':s.page_setup.fitToHeight} for s in b},'schema_issues':[]}
        with zipfile.ZipFile(out/name) as z:
            for p,order in [('xl/workbook.xml',BOOK_ORDER),('xl/worksheets/sheet1.xml',SHEET_ORDER),('xl/worksheets/sheet2.xml',SHEET_ORDER)]:
                root=E.fromstring(z.read(p));tags=[E.QName(c).localname for c in root]
                ranks=[order.index(t) if t in order else -1 for t in tags]
                if -1 in ranks or ranks!=sorted(ranks):f['schema_issues'].append({'part':p,'defect':'Schema child order incorrect','tags':tags})
                for t in ('printOptions','pageMargins','pageSetup','headerFooter','sheetProtection','definedNames'):
                    if tags.count(t)>1:f['schema_issues'].append({'part':p,'defect':'Duplicate singleton '+t})
                if p.startswith('xl/worksheets'):
                    for t in ('pageMargins','printOptions','pageSetup','sheetProtection'):
                        if tags.count(t)!=1:f['schema_issues'].append({'part':p,'defect':'Expected exactly one '+t})
        result['files'].append(f)
        for issue in f['schema_issues']:result['defects'].append(dict(file=name,**issue))
        for k in ('values_formulas_differences','unexpected_new_values','visual_style_differences','lock_mismatches'):
            if f[k]:result['defects'].append({'file':name,'defect':k,'detail':f[k]})
        for k in ('sheet_names_equal','freeze_merge_dimensions_preserved','row_column_dimensions_preserved','validations_preserved','conditional_formatting_preserved','sheet_protection_enabled'):
            if not f[k]:result['defects'].append({'file':name,'defect':k})
    result['result']='PASS_MECHANICAL' if not result['defects'] else 'FAIL_MECHANICAL'
    report_path.write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps({'result':result['result'],'files':[{k:f[k] for k in ('file','sha256','print_setup')} for f in result['files']],'defects':result['defects']},ensure_ascii=False,indent=2))
    return 0 if not result['defects'] else 1
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source',type=Path);parser.add_argument('output',type=Path);args=parser.parse_args()
    raise SystemExit(run(args.source,args.output))
