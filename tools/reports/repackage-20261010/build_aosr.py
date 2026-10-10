from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
lines=json.loads((ROOT/'aosr-source.json').read_text())
starts=re.compile(r'^(Объект капитального|Застройщик,|Лицо,|АКТ$|освидетельствования скрытых|N |Представитель |произвели |и составили |[1-7]\. |Дополнительные сведения|Акт составлен|Приложения)')
blocks=[];buf=[]
for line in lines:
 line=line.strip()
 if not line:continue
 if starts.match(line) or re.match(r'^_+',line):
  if buf:blocks.append(' '.join(buf));buf=[]
  if re.match(r'^_+',line):blocks.append(line);continue
 buf.append(line)
if buf:blocks.append(' '.join(buf))

variants={
 'p3':('03-akt-skrytyh-rabot.docx',['Бетонирование','Гидроизоляция','Электромонтажные работы до закрытия','Трубопроводы до закрытия','Иные скрытые работы']),
 'p4':('31-pyat-aktov-skrytyh-rabot.docx',['Основания и подготовка','Армирование','Гидроизоляция','Инженерные сети до закрытия','Утепление и огнезащита'])}
hints={
 'Бетонирование':'Укажите конструкцию, участок и отметки; материалы бетонной смеси и документы о качестве; результаты контроля, предусмотренные проектом. Определите, какие выполненные работы будут скрыты последующей операцией.',
 'Гидроизоляция':'Укажите основание, участок, материал, число слоёв и проектную толщину; приложите документы о качестве и результаты предусмотренного проектом контроля.',
 'Электромонтажные работы до закрытия':'Укажите трассу, помещение, марки и сечения кабелей, способ прокладки и крепления; приложите исполнительную схему и результаты предусмотренных испытаний.',
 'Трубопроводы до закрытия':'Укажите участок и назначение трубопровода, материалы, диаметры, соединения и проектные отметки; назовите документы о проведённых испытаниях, если они требуются для этих работ.',
 'Иные скрытые работы':'Определите работы, которые будут скрыты следующей операцией, и точное место. Укажите проектные решения, материалы, способы проверки и разрешённые последующие работы.',
 'Основания и подготовка':'Укажите оси, отметки, состав основания, толщины и проектные решения; приложите результаты измерений и испытаний, предусмотренных проектом.',
 'Армирование':'Укажите конструкцию и участок, класс и диаметры арматуры, шаг, соединения и защитный слой; приложите документы о качестве и исполнительную схему.',
 'Инженерные сети до закрытия':'Укажите участок сети, трассу, проектные отметки, материалы и способ монтажа; приложите исполнительную схему и документы о предусмотренных испытаниях.',
 'Утепление и огнезащита':'Выберите вид работ и укажите участок, материал, проектную толщину и способ нанесения или крепления; приложите документы о качестве и результаты предусмотренного контроля.'}
for sku,(name,kinds) in variants.items():
 d=Document();sec=d.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
 n=d.styles['Normal'];n.font.name='Times New Roman';n.font.size=Pt(10);n.font.color.rgb=RGBColor.from_string('262F37');n.paragraph_format.space_after=Pt(1)
 n.paragraph_format.widow_control=True
 for style in d.styles:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for sn,size in [('Title',17),('Heading 1',12)]:
  s=d.styles[sn];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
  s.paragraph_format.keep_with_next=True
 foot=sec.footer.paragraphs[0];foot.alignment=2
 foot.add_run('Маржа в бетоне  ·  ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
 for k,kind in enumerate(kinds,1):
  if k>1:d.add_page_break()
  d.add_paragraph(f'Вариант {k} {kind}',style='Title')
  p=d.add_paragraph('Рекомендуемый образец приложения № 3 к составу исполнительной документации, утверждённому приказом Минстроя России от 16.05.2023 № 344/пр.')
  for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('8E6F32')
  p=d.add_paragraph('Укажите фактические работы, место и проектные решения в пунктах 1 и 2. Материалы и результаты контроля внесите в пункты 3 и 4. Подписи оформите по составу участников освидетельствования, указанному в форме.')
  d.add_paragraph(hints[kind])
  p.paragraph_format.space_after=Pt(8)
  signature_start=[i for i,b in enumerate(blocks) if b.startswith('Представитель ')][-5]
  for bi,b in enumerate(blocks):
   p=d.add_paragraph(b)
   if bi==signature_start:p.paragraph_format.page_break_before=True
   if bi>=signature_start:p.paragraph_format.space_after=Pt(7)
   if starts.match(b):p.paragraph_format.keep_with_next=True
   if b.startswith('_') and bi+1<len(blocks) and blocks[bi+1]=='(фамилия, инициалы) (подпись)':p.paragraph_format.keep_with_next=True
   if b.startswith('('):
    for r in p.runs:r.font.size=Pt(9)
   if b=='АКТ' or b=='освидетельствования скрытых работ':
    p.alignment=1
    for r in p.runs:r.bold=True;r.font.size=Pt(12)
   if re.match(r'^[1-7]\. ',b):
    for r in p.runs:r.bold=True
 d.core_properties.title='Пять форм актов освидетельствования скрытых работ'
 d.core_properties.subject='Приказ Минстроя России 344/пр приложение 3'
 d.save(ROOT/sku/name)
 text=' '.join(p.text for p in d.paragraphs)
 assert text.count('1. К освидетельствованию предъявлены следующие работы:')==5
 assert text.count('7. Разрешается производство последующих работ')==5
 assert text.count('(фамилия, инициалы) (подпись)')==25
 assert 'Подписи всех трех сторон обязательны' not in text
 print(sku,name,len(d.paragraphs))
