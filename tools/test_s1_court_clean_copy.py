"""Native Save-As/delete workflow, not a twelfth delivered buyer file."""
import re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from test_s1_court_appendices import (N10, ROOT, blocks, section, text_of,
    rows, attachment_table, TOKENS, CLAIM_BEFORE, SIGNATURE,
    make_register, fill_claim, selected, FILL)


def clean_copy(src, dst):
    """Emulate buyer Save As then remove everything before court header/after signature.
    Delete native block elements; retain styles, numbering, section settings and table.
    """
    doc = Document(src)
    body = list(doc.element.body)
    start = next(i for i, e in enumerate(body)
                 if e.tag.endswith('}p') and Paragraph(e, doc).text.startswith('В Арбитражный суд '))
    end = next(i for i, e in enumerate(body[start:], start)
               if e.tag.endswith('}p') and '_______________' in Paragraph(e, doc).text
               and ('ДАТА_ПОДПИСАНИЯ_ИСКА' in Paragraph(e, doc).text or '05.10.2026' in Paragraph(e, doc).text))
    for i, e in enumerate(body):
        if not start <= i <= end and not e.tag.endswith('}sectPr'):
            doc.element.body.remove(e)
    doc.save(dst)
    return dst


def test_clean_native_copy_retains_tokens_table_and_legal_text(tmp_path):
    out = clean_copy(N10, tmp_path/'clean-template.docx')
    doc = Document(out)
    pars = [p.text for p in doc.paragraphs if p.text]
    assert pars == CLAIM_BEFORE[1:] + ['Приложения:', SIGNATURE]
    assert Counter(re.findall(r'\{\{[^}]*\}\}', text_of(blocks(doc)))) == Counter(TOKENS)
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 19
    assert rows(doc.tables[0])[0] == ['№','Документ (наименование, № и дата)','Листов']
    assert all(r._tr.xpath('./w:trPr/w:cantSplit') for r in doc.tables[0].rows)
    assert doc.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert doc.sections[0].page_width == Document(N10).sections[0].page_width


def test_buyer_can_find_copy_boundaries_in_03_and_10a():
    a = text_of(section(Document(N10),'Раздел А.'))
    algorithm = Document(ROOT/'03-algoritm-dejstviy.docx')
    pars = [p.text for p in algorithm.paragraphs]
    i = next(i for i,t in enumerate(pars) if t.startswith('Шаг 12.'))
    assert 'чистую копию' in '\n'.join(pars[i:])
    for need in ('Сохраните','В Арбитражный суд','подписи с датой','PDF','разделы В и Г','Исходный файл 10'):
        assert need in a
    b = text_of(section(Document(N10),'Раздел Б.'))
    assert 'Реквизиты искового заявления определены' not in b
    assert 'Как сохранить' not in b


def test_actual_filled_copy_has_11_attachments_and_zero_service(tmp_path):
    book = make_register(tmp_path/'04-fixture.xlsx')
    filled = fill_claim(N10,book,tmp_path/'filled-workbook.docx')
    out = clean_copy(filled,tmp_path/'court-copy.docx')
    doc = Document(out)
    got = rows(doc.tables[0])[1:]
    entries = selected(book)
    assert len(entries)==11
    assert got == [[str(i),f'{n}: {nd}',str(s)] for i,(n,nd,s) in enumerate(entries,1)]
    text = text_of(blocks(doc))
    assert '{{' not in text
    for service in ('Раздел А','Раздел Б','Раздел В','Раздел Г','Файл 10','файл 04','файла 04','проверочный лист','Сверьте шаблон','Важно, прочитайте','Редакция-кандидат','Как сохранить'):
        assert service not in text
    assert doc.paragraphs[0].text == 'В Арбитражный суд города Москвы'
    assert doc.paragraphs[-1].text.endswith('05.10.2026')


if __name__=='__main__':
    import sys
    out = Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
    FILL['{{ГОСПОШЛИНА}}']='30 000,00' # Synthetic print fixture, not a legal fee determination.
    book=make_register(out/'04-fixture.xlsx')
    full=fill_claim(N10,book,out/'filled-file10.docx')
    clean_copy(full,out/'court-copy.docx')
    print(out/'court-copy.docx')
