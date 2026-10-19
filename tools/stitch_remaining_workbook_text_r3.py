#!/usr/bin/env python3
"""Transplant approved Artifact edits into original OOXML without feature loss."""
from pathlib import Path
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
import json,re,argparse
from normalize_xlsx_element_order import normalize_parts
parser=argparse.ArgumentParser();parser.add_argument('--only');args=parser.parse_args()
BASE=Path(__file__).resolve().parents[1]/'tools/candidates/remaining-batch-r3-20261018'
IN=Path('/tmp/mb-r3/input');OUT=Path('/tmp/mb-r3/work');AUTH=Path('/tmp/mb-r3/authored')
X={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};Q='{'+X['x']+'}'
layouts=json.loads((BASE/'layout-edits.json').read_text());formulas=json.loads((BASE/'formula-text-edits.json').read_text())
def dump(r):return E.tostring(r,encoding='utf-8',xml_declaration=True)
def colnum(s):
 n=0
 for c in s:n=n*26+ord(c)-64
 return n
def cell(root,ref,create=False):
 c=root.find(f'.//x:c[@r="{ref}"]',X)
 if c is not None or not create:return c
 data=root.find('x:sheetData',X);num=int(re.search(r'\d+',ref)[0]);row=data.find(f'x:row[@r="{num}"]',X)
 if row is None:
  row=E.Element(Q+'row',r=str(num));data.insert(next((i for i,e in enumerate(data) if int(e.get('r'))>num),len(data)),row)
 c=E.Element(Q+'c',r=ref);row.insert(next((i for i,e in enumerate(row) if colnum(re.match('[A-Z]+',e.get('r'))[0])>colnum(re.match('[A-Z]+',ref)[0])),len(row)),c);return c
def text(z,c):
 if c.get('t')=='s':
  ss=E.fromstring(z.read('xl/sharedStrings.xml'));return ''.join(list(ss)[int(c.find('x:v',X).text)].itertext())
 if c.get('t')=='str':return c.find('x:v',X).text or ''
 return ''.join(c.xpath('.//x:t/text()',namespaces=X))
order='sheetPr dimension sheetViews sheetFormatPr cols sheetData sheetCalcPr sheetProtection protectedRanges scenarios autoFilter sortState dataConsolidate customSheetViews mergeCells phoneticPr conditionalFormatting dataValidations hyperlinks printOptions pageMargins pageSetup headerFooter rowBreaks colBreaks customProperties cellWatches ignoredErrors smartTags drawing legacyDrawing legacyDrawingHF picture oleObjects controls webPublishItems tableParts extLst'.split()
def ensure(root,tag):
 e=root.find('x:'+tag,X)
 if e is None:
  e=E.Element(Q+tag);rank=order.index(tag);root.insert(next((i for i,c in enumerate(root) if E.QName(c).localname in order and order.index(E.QName(c).localname)>rank),len(root)),e)
 return e
changes=[]
for patch in sorted(Path('/tmp/mb-r3/patches').rglob('*.json')):
 edits=json.loads(patch.read_text());sku=patch.parent.name;name=patch.stem+'.xlsx';key=sku+'/'+name
 if args.only and key!=args.only:continue
 layout=layouts.get(key,{'sheets':{}});fs=[f for f in formulas if f['sku']==sku and f['file']==name]
 with ZipFile(IN/sku/name) as z:parts={n:z.read(n) for n in z.namelist()}
 authored=ZipFile(AUTH/sku/name) if edits or fs or layout['sheets'] else None
 styles=E.fromstring(parts['xl/styles.xml']);xfs=styles.find('x:cellXfs',X);style_changed=False
 for part in sorted({e['part'] for e in edits+fs}|set(layout['sheets'])):
  old=E.fromstring(parts[part]);new=E.fromstring(authored.read(part));plan=layout['sheets'].get(part,{})
  for target,source in plan.get('copy_style',{}).items():cell(old,target,True).set('s',cell(old,source).get('s','0'))
  for edit in [e for e in edits if e['part']==part]:
   a=cell(old,edit['cell'],True);b=cell(new,edit['cell']);assert b is not None
   assert a.find('x:f',X) is None,(key,edit['cell'],'formula overwrite')
   s=text(authored,b);assert s==edit['text'],(key,edit['cell'],'authored text mismatch')
   for child in list(a):a.remove(child)
   a.set('t','inlineStr');t=E.SubElement(E.SubElement(a,Q+'is'),Q+'t');t.text=s;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
   row=a.getparent();nr=new.find('x:sheetData/x:row[@r="'+row.get('r')+'"]',X)
   if nr.get('ht') and float(nr.get('ht'))>float(row.get('ht','15')):row.set('ht',str(round(float(nr.get('ht')),2)));row.set('customHeight','1')
  for f in [f for f in fs if f['part']==part]:
   a=cell(old,f['cell']);b=cell(new,f['cell']);assert '='+a.find('x:f',X).text==f['old']
   assert b is not None and '='+b.find('x:f',X).text==f['new'],(key,f['cell'],'authored formula mismatch')
   a.find('x:f',X).text=f['new'][1:]
   for v in a.findall('x:v',X):a.remove(v)
   a.attrib.pop('t',None)
  for ref in plan.get('wrap',[]):
   a=cell(old,ref);b=cell(new,ref);ast=E.fromstring(authored.read('xl/styles.xml'));axf=ast.find('x:cellXfs',X)[int(b.get('s','0'))];alignment=axf.find('x:alignment',X)
   assert alignment is not None and alignment.get('wrapText') in {'1','true'},(key,ref,'wrap not authored')
   xf=deepcopy(xfs[int(a.get('s','0'))]);al=xf.find('x:alignment',X)
   if al is None:al=E.SubElement(xf,Q+'alignment')
   al.set('wrapText','1');xf.set('applyAlignment','1');existing=next((i for i,v in enumerate(xfs) if E.tostring(v)==E.tostring(xf)),None)
   if existing is None:existing=len(xfs);xfs.append(xf);xfs.set('count',str(len(xfs)));style_changed=True
   a.set('s',str(existing))
  for ref in plan.get('merges',[]):
   from openpyxl.utils.cell import range_boundaries,get_column_letter
   x1,y1,x2,y2=range_boundaries(ref)
   for y in range(y1,y2+1):
    for x in range(x1,x2+1):
     rr=get_column_letter(x)+str(y);c=cell(old,rr)
     assert rr==get_column_letter(x1)+str(y1) or c is None or not any(ch.tag in {Q+'f',Q+'v',Q+'is'} for ch in c),(key,ref,'merge would hide content')
   mc=ensure(old,'mergeCells')
   if not any(e.get('ref')==ref for e in mc):E.SubElement(mc,Q+'mergeCell',ref=ref)
   mc.set('count',str(len(mc)))
  for row,height in plan.get('row_heights',{}).items():
   rr=cell(old,'A'+row,True).getparent();rr.set('ht',str(height));rr.set('customHeight','1')
  if plan.get('columns'):
   cols=ensure(old,'cols');source=list(cols);cols.clear()
   for c in source:
    for i in range(int(c.get('min')),int(c.get('max'))+1):
     cc=deepcopy(c);cc.set('min',str(i));cc.set('max',str(i));cols.append(cc)
   for letter,width in plan['columns'].items():
    idx=colnum(letter);cc=next((c for c in cols if int(c.get('min'))==idx),None)
    if cc is None:cc=E.SubElement(cols,Q+'col',min=str(idx),max=str(idx))
    ac=new.find('x:cols',X);found=next(c for c in ac if int(c.get('min'))<=idx<=int(c.get('max')))
    cc.set('width',found.get('width'));cc.set('customWidth','1')
  if 'break_before' in plan:
   rb=ensure(old,'rowBreaks');rb.clear()
   for row in plan['break_before']:E.SubElement(rb,Q+'brk',id=str(row-1),min='0',max='16383',man='1')
   rb.set('count',str(len(rb)));rb.set('manualBreakCount',str(len(rb)))
  if 'fit_height' in plan:ensure(old,'pageSetup').set('fitToHeight',str(plan['fit_height']))
  if plan.get('copy_style'):
   from openpyxl.utils.cell import range_boundaries,get_column_letter
   dim=ensure(old,'dimension');x1,y1,x2,y2=range_boundaries(dim.get('ref','A1'))
   for ref in plan['copy_style']:
    x,y,_,_=range_boundaries(ref);x2=max(x2,x);y2=max(y2,y)
   dim.set('ref',f'{get_column_letter(x1)}{y1}:{get_column_letter(x2)}{y2}')
  parts[part]=dump(old)
 if style_changed:parts['xl/styles.xml']=dump(styles)
 repairs=parts['docProps/core.xml'].count(b'+00:00Z');parts['docProps/core.xml']=parts['docProps/core.xml'].replace(b'+00:00Z',b'Z')
 wb=E.fromstring(parts['xl/workbook.xml']);names=wb.find('x:definedNames',X);wb_changed=False
 def name_set(idx,tag,ref):
  global names,wb_changed
  if names is None:
   names=E.Element(Q+'definedNames');wb.insert(list(wb).index(wb.find('x:sheets',X))+1,names)
  for n in list(names):
   if n.get('name')==tag and n.get('localSheetId')==str(idx):names.remove(n)
  if ref is not None:E.SubElement(names,Q+'definedName',name=tag,localSheetId=str(idx)).text=ref
  wb_changed=True
 for part,plan in layout['sheets'].items():
  idx=int(re.search(r'sheet(\d+)',part)[1])-1;sheet=wb.find('x:sheets',X)[idx].get('name').replace("'","''")
  if 'print_area' in plan:name_set(idx,'_xlnm.Print_Area',"'"+sheet+"'!"+re.sub(r'([A-Z]+)(\d+)',r'$\1$\2',plan['print_area']))
  if 'repeat_rows' in plan:name_set(idx,'_xlnm.Print_Titles',"'"+sheet+"'!"+re.sub(r'(\d+)',r'$\1',plan['repeat_rows']) if plan['repeat_rows'] else None)
 if sku=='P1' and name.startswith(('06-','07-','08-')):
  name_set(0,'_xlnm.Print_Area',"'Реестр'!$A$1:$"+('K' if name.startswith('07-') else 'J')+'$103');name_set(0,'_xlnm.Print_Titles',"'Реестр'!$1:$3")
 if wb_changed:parts['xl/workbook.xml']=dump(wb)
 order_repairs=normalize_parts(parts)
 with ZipFile(OUT/sku/name,'w',ZIP_DEFLATED) as z:
  for n,b in sorted(parts.items()):
   info=ZipInfo(n,(2026,10,18,0,0,0));info.compress_type=ZIP_DEFLATED;z.writestr(info,b)
 changes.append({'sku':sku,'file':name,'authored_text_cells':len(edits),'formula_display_changes':len(fs),'double_timezone_repairs':repairs,'element_order_repairs':order_repairs})
 if authored:authored.close()
 print(changes[-1])
Path('/tmp/mb-r3/workbook-stitch-report.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
