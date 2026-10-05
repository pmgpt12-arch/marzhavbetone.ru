"""F-2 (06): письмо заказчику печатается без служебного текста, а служебный
текст не потерян — он в файле 03, шаг 8."""
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent / "candidates" / "s1-oplata-za-raboty"
N06 = ROOT / "06-uvedomlenie-o-prosrochke.docx"
A03 = ROOT / "03-algoritm-dejstviy.docx"

# Видимый текст письма — без изменений относительно 0ae2bf0, кроме убранных
# служебных блоков.
LETTER = [
    "Исх. № {{ИСХОДЯЩИЙ_НОМЕР}} от {{ДАТА_ДОКУМЕНТА}}",
    "Кому: {{НАИМЕНОВАНИЕ_ЗАКАЗЧИКА}}, ИНН {{ИНН_ЗАКАЗЧИКА}}, ОГРН {{ОГРН_ЗАКАЗЧИКА}}",
    "Адрес (по выписке ЕГРЮЛ): {{АДРЕС_ЗАКАЗЧИКА}}; адрес для корреспонденции по договору: {{АДРЕС_ПО_ДОГОВОРУ}}",
    "От: {{НАИМЕНОВАНИЕ_СУБПОДРЯДЧИКА}}, ИНН {{ИНН_СУБПОДРЯДЧИКА}}, ОГРН {{ОГРН_СУБПОДРЯДЧИКА}}, адрес: {{АДРЕС_СУБПОДРЯДЧИКА}}",
    "УВЕДОМЛЕНИЕ О ПРОСРОЧКЕ ОПЛАТЫ",
    "По договору № {{НОМЕР_ДОГОВОРА}} от {{ДАТА_ДОГОВОРА}} Субподрядчик выполнил, а Заказчик принял работы по актам:",
    "Срок оплаты по п. {{ПУНКТ_ДОГОВОРА_ОБ_ОПЛАТЕ}} договора истёк по каждому акту в дату, указанную в таблице. Задолженность на {{ДАТА_ДОКУМЕНТА}} составляет {{СУММА_ДОЛГА}} руб. (реестр взаиморасчётов прилагается).",
    "Просим оплатить задолженность в срок до {{ДАТА_ОПЛАТЫ_ПО_ТРЕБОВАНИЮ}} по реквизитам: {{БАНКОВСКИЕ_РЕКВИЗИТЫ}}. Если у Заказчика есть возражения по сумме, просим сообщить их письменно в тот же срок.",
    "При неоплате Субподрядчик направит претензию с требованием уплаты задолженности и процентов по ст. 395 ГК РФ и обратится в арбитражный суд.",
    "Приложение: реестр взаиморасчётов на {{ДАТА_РЕЕСТРА_ВЗАИМОРАСЧЁТОВ}}.",
    "{{ДОЛЖНОСТЬ_ПОДПИСАНТА}}, {{ОСНОВАНИЕ_ПОЛНОМОЧИЙ}}  _______________  / {{ФИО}} /",
]
TABLE_HEAD = ["№", "Акт (№, дата)", "Сумма, руб.", "Срок оплаты по договору", "Оплачено, руб.", "Остаток, руб."]

# Служебный текст, который не должен печататься в письме заказчику.
SERVICE = [
    "Файл 06", "Важно, прочитайте", "рабочий шаблон", "юридической консультац", "Редакция-кандидат",
    "юридическую приёмку", "Как подготовить приложение", "После отправки", "Первое письменное требование",
    "надёжнее пройти шаг", "файл 01", "файла 04", "файл 07", "Calc", "Excel", "Звонок контактному лицу",
    "Уведомление о просрочке оплаты",  # заголовок продукта; заголовок письма — прописными
]

# Указания, перенесённые из 06 в 03, шаг 8 — дословно.
MOVED = [
    "первое письменное требование после просрочки. Короче претензии; может направляться до неё. Если в письме есть основание, сумма, срок и последствия — суд может оценить его и как претензию, но надёжнее пройти шаг 10.",
    "Как подготовить приложение",
    "Приложение — лист «Взаиморасчёты» файла 04, только заполненные строки.",
    "в Calc — «Файл → Экспорт в PDF → Выделение», в Excel — «Печать → Напечатать выделенный фрагмент». Свод справа (колонки L–M) прикладывать не нужно.",
    "колонки A–J",
    "После отправки",
    "Направьте способом из файла 07; строка в файле 04, лист «Реестр передачи», с доказательством.",
    "Строка в листе «Контроль ответа» (файл 02) со сроком ответа.",
    "Звонок контактному лицу; итог — письмом «по итогам разговора фиксируем…».",
    "«Важно, прочитайте до использования»",
    "Таблица актов (строки 1–12 файла 06)",
    "«Долг по актам» файла 04 (строки 5–34)",
    "шапку таблицы (строка 4)",
]


def printed(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            parts += [p.text for p in hf.paragraphs]
    return "\n".join(parts)


def step8(doc) -> str:
    texts = [p.text for p in doc.paragraphs]
    start = next(i for i, t in enumerate(texts) if t.startswith("Шаг 8. "))
    end = next(i for i, t in enumerate(texts) if t.startswith("Шаг 9. "))
    return "\n".join(texts[start:end])


def test_06_prints_only_the_letter():
    doc = Document(N06)
    body = [p.text for p in doc.paragraphs if p.text.strip()]
    assert body == LETTER
    assert not [p for p in doc.paragraphs if p.style.name in ("Title", "Subtitle") or p.style.name.startswith("Heading")]
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == TABLE_HEAD


def test_06_has_no_service_leakage():
    text = printed(Document(N06))
    leaked = [s for s in SERVICE if s in text]
    assert not leaked, f"служебный текст в письме заказчику: {leaked}"


def test_06_fields_are_unique_and_fill_completely():
    doc = Document(N06)
    text = printed(doc)
    assert "{{ДАТА}}" not in text
    filled = text
    fields = set(re.findall(r"\{\{[^{}]+\}\}", text))
    for f in fields:
        filled = filled.replace(f, "X")
    assert "{{" not in filled and "}}" not in filled


def test_03_step8_keeps_moved_guidance():
    s8 = step8(Document(A03))
    lost = [m for m in MOVED if m not in s8]
    assert not lost, f"указания из 06 потеряны в 03, шаг 8: {lost}"


def test_03_still_carries_disclaimer_and_candidate_note():
    import build_s1_candidate as B
    text = printed(Document(A03))
    assert B.DISCLAIMER in text and B.CANDIDATE_NOTE in text
