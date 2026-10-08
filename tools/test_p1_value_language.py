"""Buyer P1 files speak about use, contents and next actions without disclaimer rhetoric."""
from pathlib import Path
from docx import Document
import fitz
import zipfile
import re

BASE = Path(__file__).resolve().parents[1] / 'products-storage/01-zakrytie-rabot'
FORBIDDEN = [
    r'шаблон\w*',
    r'юрист\w*',
    r'типов\w*',
    r'юридическ\w*\s+консультац\w*',
    r'важно, прочитайте до использования',
    r'это типовой шаблон, а не',
    r'не юридическая консультация',
    r'не заменяет (?:юридическую|консультацию|проверку|бухгалтерию)',
    r'не (?:является )?гаранти[ея]',
    r'не гарантирует (?:оплату|взыскание)',
    r'чего (?:этот документ|этот комплект) не делает',
    r'зачем этот документ',
    r'когда шаблон откладывается',
    r'откладывайте шаблон',
    r'главный риск файла',
    r'ложную уверенность',
    r'шаблон требует адаптации',
    r'этот комплект односторонний акт не покрывает',
]

def text_of(path):
    if path.suffix == '.docx':
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] +
                         [cell.text for table in doc.tables for row in table.rows for cell in row.cells] +
                         [p.text for section in doc.sections for p in section.header.paragraphs] +
                         [p.text for section in doc.sections for p in section.footer.paragraphs])
    if path.suffix == '.pdf':
        return '\n'.join(page.get_text() for page in fitz.open(path))
    if path.suffix == '.xlsx':
        with zipfile.ZipFile(path) as archive:
            return '\n'.join(archive.read(n).decode() for n in archive.namelist()
                             if n.startswith('xl/') and n.endswith('.xml'))
    return path.read_text()

files = [p for p in BASE.iterdir() if p.suffix.lower() in {'.docx', '.pdf', '.xlsx', '.txt'}]
assert len(files) == 20, len(files)
failures = [(p.name, phrase) for p in files for phrase in FORBIDDEN
            if re.search(phrase, text_of(p), re.I)]
page = (BASE.parents[1] / 'products/p1-oplata-po-ks2.html').read_text()
failures += [('products/p1-oplata-po-ks2.html', phrase) for phrase in FORBIDDEN
             if re.search(phrase, page, re.I)]
assert not failures, failures
print(f'P1 value language PASS: {len(files)} buyer-facing files and product page')
