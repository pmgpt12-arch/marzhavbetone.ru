#!/usr/bin/env python3
"""Итоговые формулы двух бесплатных книг считаются в LibreOffice Calc.

ЗАЧЕМ. Draft PR #293 открыл обе книги в Calc и увидел `Err:522` —
циклическую ссылку — вместо сумм (Issue #296):

    00-free-dopy-ne-v-podarok/04-reestr-doprabot.xlsx
        G2 = G2*D2         сумма строки умножала саму себя, а цена в F
        G5 = SUM(G2:G10)   итог в G5 суммировал и G5
    00-free-ks-bez-vozvrata/02-reestr-prilozheniy.xlsx
        E7 = SUM(E2:E20)   итог в E7 суммировал и E7

Читатель получал таблицу, которая на первой же строке показывает ошибку.
Формулы пишут генераторы `build_free_02.py` и `build_free_03.py`, поэтому
тест смотрит и на книги, и на генераторы, и на архивы в downloads/.

Слои:
1. Статика (openpyxl): ни одна формула не ссылается на свою ячейку;
   итог суммирует ровно строки данных над собой.
2. Генераторы: два прогона во временную папку дают побайтно одинаковые
   книги, и они совпадают с лежащими в репозитории.
3. Архивы: в downloads/<ключ>.zip лежат байт в байт те же книги.
4. Calc: `soffice --headless` пересчитывает книги; в выводе нет `Err:`,
   `#VALUE!`, `#REF!`, `#NAME?` и прочих ошибок, суммы равны ожидаемым.
   Нет `soffice` или модуля Calc — это ПРОВАЛ, а не пропуск: без пересчёта
   дефект, ради которого тест написан, не виден.

    python3 tools/test_free_calc_formulas.py

Код возврата 1 при любом провале.
"""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
STORAGE = ROOT / "products-storage"

# книга → генератор, функция сборки, архив, ожидаемые значения после Calc
BOOKS = {
    "00-free-dopy-ne-v-podarok/04-reestr-doprabot.xlsx": {
        "builder": "build_free_02.py",
        "build": "build_free_02",
        "zip": "downloads/dop-raboty.zip",
        # 150 м2 × 850 ₽ = 127 500 ₽ в строке и в итоге
        "expect": {"G2": 127500, "G5": 127500},
        "total": "G5",
    },
    "00-free-ks-bez-vozvrata/02-reestr-prilozheniy.xlsx": {
        "builder": "build_free_03.py",
        "build": "build_free_03",
        "zip": "downloads/vozvrat-ks.zip",
        # 3 + 2 + 5 + 1 листов
        "expect": {"E7": 11},
        "total": "E7",
    },
}

ERRORS = re.compile(r"Err:\d+|#VALUE!|#REF!|#NAME\?|#DIV/0!|#N/A|#NUM!|#NULL!")
REF = re.compile(r"(?<![A-Za-z_\"])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")

PROBLEMS: list[str] = []


def fail(msg: str) -> None:
    PROBLEMS.append(msg)
    print(f"ПРОВАЛ  {msg}")


def col_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def refs(formula: str) -> list[tuple[int, int, int, int]]:
    """Прямоугольники (кол1, стр1, кол2, стр2), на которые ссылается формула."""
    bare = re.sub(r'"[^"]*"', '""', formula)          # строки внутри формулы
    out = []
    for c1, r1, c2, r2 in REF.findall(bare):
        c2, r2 = c2 or c1, r2 or r1
        out.append((col_num(c1), int(r1), col_num(c2), int(r2)))
    return out


# ───────────────────────────── 1. статика ─────────────────────────────

def check_static(rel: str, spec: dict) -> None:
    import openpyxl
    before = len(PROBLEMS)
    ws = openpyxl.load_workbook(STORAGE / rel).active
    formulas = {c.coordinate: c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")}
    if not formulas:
        fail(f"{rel}: в книге нет ни одной формулы")
    for coord, f in formulas.items():
        cell = ws[coord]
        for c1, r1, c2, r2 in refs(f):
            if c1 <= cell.column <= c2 and r1 <= cell.row <= r2:
                fail(f"{rel}: {coord} = {f} ссылается на собственную ячейку")

    total = ws[spec["total"]]
    f = total.value if isinstance(total.value, str) else ""
    m = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", f)
    if not m:
        fail(f"{rel}: {spec['total']} не итог SUM, а {total.value!r}")
        return
    if m.group(1) != total.column_letter or m.group(3) != total.column_letter:
        fail(f"{rel}: {spec['total']} = {f} суммирует чужой столбец")
    if int(m.group(2)) != 2 or int(m.group(4)) != total.row - 1:
        fail(f"{rel}: {spec['total']} = {f} — ожидался диапазон со 2-й "
             f"строки до {total.row - 1}-й, прямо над итогом")
    if len(PROBLEMS) == before:
        print(f"ok      {rel}: формулы {formulas}")


# ───────────────────────────── 2. генераторы ─────────────────────────────

def run_builder(spec: dict, base: Path) -> None:
    path = STORAGE / spec["builder"]
    mod_spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(mod_spec)
    mod_spec.loader.exec_module(module)
    module.BASE_DIR = str(base)
    getattr(module, spec["build"])()


def check_builders(tmp: Path) -> None:
    runs = []
    for i in (1, 2):
        base = tmp / f"run{i}"
        for rel in BOOKS:
            (base / rel).parent.mkdir(parents=True, exist_ok=True)
        for spec in {s["builder"]: s for s in BOOKS.values()}.values():
            run_builder(spec, base)
        runs.append(base)
    for rel in BOOKS:
        a, b = (r / rel for r in runs)
        committed = STORAGE / rel
        if a.read_bytes() != b.read_bytes():
            fail(f"{rel}: два прогона генератора дали разные байты")
        elif a.read_bytes() != committed.read_bytes():
            fail(f"{rel}: книга в репозитории не та, что собирает генератор — "
                 f"пересоберите {BOOKS[rel]['builder']}")
        else:
            digest = hashlib.sha256(a.read_bytes()).hexdigest()[:16]
            print(f"ok      {rel}: генератор воспроизводим, sha256 {digest}")


# ───────────────────────────── 3. архивы ─────────────────────────────

def check_zips() -> None:
    for rel, spec in BOOKS.items():
        archive = ROOT / spec["zip"]
        name = Path(rel).name
        with zipfile.ZipFile(archive) as z:
            if name not in z.namelist():
                fail(f"{spec['zip']}: нет {name}")
                continue
            if z.read(name) != (STORAGE / rel).read_bytes():
                fail(f"{spec['zip']}: {name} в архиве не совпадает с книгой — "
                     f"запустите tools/build_free_zips.py")
                continue
        print(f"ok      {spec['zip']}: {name} совпадает с книгой")


# ───────────────────────────── 4. Calc ─────────────────────────────

def calc_missing() -> str | None:
    soffice = shutil.which("soffice")
    if not soffice:
        return "soffice не найден"
    program = Path(soffice).resolve().parent
    if not (program / "libsclo.so").exists():
        return f"LibreOffice без модуля Calc ({program}): нужен libreoffice-calc"
    return None


def check_calc(tmp: Path) -> None:
    import openpyxl
    missing = calc_missing()
    if missing:
        fail(f"пересчёт в Calc невозможен: {missing}")
        return
    src = tmp / "calc-in"
    src.mkdir()
    books = {rel: shutil.copy(STORAGE / rel, src / Path(rel).name) for rel in BOOKS}
    for fmt in ("csv", "xlsx"):
        out = tmp / f"calc-{fmt}"
        cmd = ["soffice", f"-env:UserInstallation=file://{tmp}/profile",
               "--headless", "--norestore", "--calc", "--convert-to", fmt,
               "--outdir", str(out), *map(str, books.values())]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        for rel in BOOKS:
            if not (out / Path(rel).with_suffix(f".{fmt}").name).exists():
                fail(f"{rel}: Calc не пересчитал книгу в {fmt}: "
                     f"{r.stdout[-300:]} {r.stderr[-300:]}")

    for rel, spec in BOOKS.items():
        name = Path(rel)
        csv_path = tmp / "calc-csv" / name.with_suffix(".csv").name
        xlsx_path = tmp / "calc-xlsx" / name.name
        if not csv_path.exists() or not xlsx_path.exists():
            continue
        shown = csv_path.read_text(encoding="utf-8", errors="replace")
        errors = sorted(set(ERRORS.findall(shown)))
        if errors:
            fail(f"{rel}: после пересчёта в Calc ошибки {errors}")
        ws = openpyxl.load_workbook(xlsx_path, data_only=True).active
        for coord, want in spec["expect"].items():
            got = ws[coord].value
            if got != want:
                fail(f"{rel}: {coord} после Calc = {got!r}, ожидалось {want}")
        if not errors:
            rows = list(csv.reader(shown.splitlines()))
            values = {c: ws[c].value for c in spec["expect"]}
            print(f"ok      {rel}: Calc без ошибок, {values}, строк {len(rows)}")


def main() -> int:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("ПРОВАЛ  нет openpyxl — книги не прочитать")
        return 1
    tmp = Path(tempfile.mkdtemp(prefix="free-calc-"))
    try:
        for rel, spec in BOOKS.items():
            check_static(rel, spec)
        check_builders(tmp)
        check_zips()
        check_calc(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"\nпровалов: {len(PROBLEMS)}")
    return 1 if PROBLEMS else 0


if __name__ == "__main__":
    sys.exit(main())
