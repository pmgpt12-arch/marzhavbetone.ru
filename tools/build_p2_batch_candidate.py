#!/usr/bin/env python3
"""P2 scoped buyer edition, retaining approved calculation structures and source forms."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
import shutil,subprocess
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'products-storage/02-dopraboty-bez-poter'
OUT=ROOT/'tools/candidates/remaining-batch-20261018/P2'
if OUT.exists():shutil.rmtree(OUT)
OUT.mkdir(parents=True,exist_ok=True)
for f in SRC.iterdir():
 if f.suffix.lower() in ('.docx','.xlsx','.pdf','.txt') and not f.name.startswith('00-PISMO'):
  shutil.copy2(f,OUT/f.name)
names=sorted(x.name for x in OUT.iterdir() if x.is_file())
intro="""С ЧЕГО НАЧАТЬ
Система оформления и получения оплаты за дополнительные работы

Результат: каждая дополнительная позиция связана с исходным объёмом, поручением или уведомлением, согласованными условиями цены и срока, фактом выполнения, приёмкой и оплатой.

До начала работ откройте 10-algoritm-doprabot.pdf, маршрут А. Сверьте договор, смету и проект; получите поручение 01 или направьте уведомление 03, согласуйте объём в форме 02. Расчёт стоимости — 07; если меняются цена и срок, оформите соглашения 11 и 12 до выполнения соответствующего объёма.

Если работы уже выполнены, откройте маршрут Б файла 10. Запишите фактические даты, инициатора и документы, подтверждающие объём. Направьте уведомление 03 текущей датой, оформите решение по позициям 02 и расчёт 07. Ведите журнал 08, фотофиксацию 09 и акты освидетельствования 04 до закрытия скрываемого участка.

При отсутствии решения или необходимости остановить затронутые работы используйте 06 и сохраните доказательство доставки. Отдельно согласуйте смещение графика с обоснованием 05. После выполнения передайте приёмочные документы по договору и сверьте включение дополнительного объёма в оплату.

Перед подписанием каждого документа проверьте полномочия сторон, договорные пункты, перечень приложений, суммы, даты и границы. Сохраните подписанные версии и подтверждения направления и получения.

СОСТАВ АРХИВА
"""
(OUT/'00-START-HERE.txt').write_text(intro+'\n'.join(names)+'\n',encoding='utf-8')

d=Document(OUT/'00-INSTRUKCIYA.docx')
for p in d.paragraphs:
 if p.text.startswith('Инструкция к комплекту'):p.text='Система оформления и получения оплаты за дополнительные работы'
for sec in d.sections:
 for p in list(sec.header.paragraphs)+list(sec.footer.paragraphs):
  if 'МАРЖА В БЕТОНЕ' in p.text:p.text=p.text.replace('МАРЖА В БЕТОНЕ','НА СТОРОНЕ ПОДРЯДЧИКА')
d.save(OUT/'00-INSTRUKCIYA.docx')
res=subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo_p2_instruction_batch','--headless','--convert-to','pdf','--outdir',str(OUT),str(OUT/'00-INSTRUKCIYA.docx')],capture_output=True,text=True)
if res.returncode:raise RuntimeError(res.stderr)
# Existing form introductions duplicate the route and precede the sendable body.
for prefix in ('01','02','03','05','06','11','12'):
 f=next(OUT.glob(prefix+'-*.docx'))
 d=Document(f);p=d.paragraphs[1];p._element.getparent().remove(p._element)
 d.save(f)
def make_pdf(name,title,sections):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
 normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9.3);normal.font.color.rgb=RGBColor(43,50,57)
 normal.paragraph_format.space_after=Pt(6)
 p=d.add_paragraph('НА СТОРОНЕ ПОДРЯДЧИКА  /  P2');p.runs[0].font.size=Pt(8);p.runs[0].font.color.rgb=RGBColor(174,128,49)
 p=d.add_paragraph(title);p.runs[0].font.size=Pt(17);p.runs[0].bold=True
 for head,items in sections:
  h=d.add_paragraph(head);h.runs[0].bold=True;h.runs[0].font.size=Pt(11)
  for n,item in enumerate(items,1):
   p=d.add_paragraph();r=p.add_run(f'{n:02d}  ');r.bold=True;r.font.color.rgb=RGBColor(174,128,49);p.add_run(item)
 temp=OUT/(name+'.docx');d.save(temp)
 res=subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo_p2_cards_batch','--headless','--convert-to','pdf','--outdir',str(OUT),str(temp)],capture_output=True,text=True)
 if res.returncode:raise RuntimeError(res.stderr)
 temp.unlink()

make_pdf('09-checklist-fotofiksacii','Фотофиксация дополнительных работ',[
('До начала и во время работ',[
'Общий план: объект, участок и устойчивый ориентир.',
'Границы: начало и конец объёма, смежные конструкции, привязка к чертежу.',
'Причина: обнаруженное обстоятельство до изменения или закрытия конструкции.',
'Масштаб: измерительный инструмент, единицы, обмерная ведомость или схема.',
'Ход работ: фотографии этапов, материала, операций и фактического объёма.',
'Скрываемое: снимки до закрытия участка, связь с актом 04 и датой осмотра.']),
('После выполнения и перед передачей',[
'Итоговый вид: результат, границы и детали, значимые для приёмки.',
'Идентификация: файлы названы по объекту, участку, работе и дате.',
'Оригиналы с метаданными сохранены отдельно от копий для переписки.',
'Позиции формы 02 сопоставлены с фото, журналом 08, актами и чертежами.',
'Перечень файлов направлен уполномоченному адресату; подтверждение доставки сохранено.'])
])
make_pdf('10-algoritm-doprabot','Порядок оформления дополнительных работ',[
('Маршрут А  Работы ещё не начаты',[
'Сверьте договор, проект, смету и границы исходных работ.',
'Получите поручение 01 или направьте уведомление 03 с доказательствами.',
'Согласуйте объём в форме 02, схему и полномочия подписанта.',
'Рассчитайте цену в 07 и влияние на график в 05; подпишите соглашения 11 и 12 по изменяемым условиям.',
'Ведите журнал 08, сохраняйте фото 09, освидетельствуйте скрываемое актом 04.',
'Передайте приёмочные документы и сверьте включение в оплату.']),
('Маршрут Б  Работы уже выполнены',[
'Запишите фактические даты, инициатора, его полномочия, переписку, журнал и фото.',
'Отделите дополнительные позиции от исходного объёма; восстановите границы и обмеры.',
'Направьте уведомление 03 текущей датой и запросите решение по форме 02.',
'Представьте расчёт 07 и оформите соглашения 11 и 12 по изменяемым условиям.',
'Свяжите согласованные позиции с приёмкой, КС-2, журналом 08 и подтверждением оплаты.']),
('Если ответ отсутствует или работы остановлены',[
'Сохраните уведомление и доказательство доставки. Зафиксируйте затронутые работы, меры сохранности и фактические даты; используйте письмо 06 по применимому основанию.'])
])
# Print registers across all columns with repeated headings.
for name,last in (('07-raschet-stoimosti.xlsx','J'),('08-zhurnal-doprabot.xlsx','L')):
 w=load_workbook(OUT/name);sh=w['Реестр']
 sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3
 sh.sheet_properties.pageSetUpPr.fitToPage=True
 sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
 sh.print_area=f'A1:{last}104';sh.print_title_rows='1:3'
 w.save(OUT/name)
for name in ('07-raschet-stoimosti.xlsx','08-zhurnal-doprabot.xlsx'):
 a=load_workbook(SRC/name);b=load_workbook(OUT/name)
 assert {(s.title,c.coordinate):c.value for s in a for row in s for c in row if c.data_type=='f'}=={(s.title,c.coordinate):c.value for s in b for row in s for c in row if c.data_type=='f'}
print('P2_PREPARED',len(list(OUT.iterdir())))
