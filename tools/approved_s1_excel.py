from __future__ import annotations
import hashlib
import json
import warnings
import io
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.sax.saxutils import escape
import openpyxl

TEMPLATES = Path(__file__).resolve().parent / "templates" / "s1-owner-excel"
APPROVED = json.loads((TEMPLATES / "approved-sha256.json").read_text())

def _targets(square_reference):
    for area in str(square_reference).split():
        left, top, right, bottom = openpyxl.utils.range_boundaries(area)
        for row in range(top, bottom + 1):
            for column in range(left, right + 1):
                yield openpyxl.utils.get_column_letter(column) + str(row)

def _native_validations(native):
    result = {}
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(io.BytesIO(native)) as archive:
        relationships = {node.attrib["Id"]: node.attrib["Target"] for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook.findall("{" + namespace + "}sheets/{" + namespace + "}sheet"):
            target = relationships[sheet.attrib["{" + relationship_namespace + "}id"]]
            filename = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            tree = ET.fromstring(archive.read(filename))
            for node in tree.iter():
                if node.tag.rsplit("}", 1)[-1] != "dataValidation":
                    continue
                fields = {child.tag.rsplit("}", 1)[-1]: "".join(child.itertext()).strip() for child in node}
                rule = (node.attrib.get("type", ""), node.attrib.get("operator", ""), node.attrib.get("allowBlank", "0") in ("1", "true"), fields.get("formula1", ""), fields.get("formula2", ""))
                for coordinate in _targets(node.attrib.get("sqref", fields.get("sqref", ""))):
                    result[(sheet.attrib["name"], coordinate)] = rule
    return result

def _generated_validations(generated):
    result = {}
    for sheet in generated.worksheets:
        for validation in sheet.data_validations.dataValidation:
            rule = (validation.type or "", validation.operator or "", bool(validation.allowBlank), str(validation.formula1 or ""), str(validation.formula2 or ""))
            for coordinate in _targets(validation.sqref):
                result[(sheet.title, coordinate)] = rule
    return result

CORRECTIONS = Path(__file__).resolve().parent / "templates" / "s1-reviewed-corrections"

def _scoped_names(workbook):
    result = {}
    for scope, names in [(None, workbook.defined_names)] + [(s.title, s.defined_names) for s in workbook.worksheets]:
        for key, value in names.items():
            node = value.to_tree()
            result[(scope, key)] = (tuple(sorted(node.attrib.items())), node.text)
    return result

def _check_semantics(actual, expected, filename):
    if actual.sheetnames != expected.sheetnames:
        raise ValueError("Approved Excel sheet mismatch: " + filename)
    for name in actual.sheetnames:
        a_sheet, e_sheet = actual[name], expected[name]
        for coordinate in set(a_sheet._cells) | set(e_sheet._cells):
            a, e = a_sheet.cell(*coordinate), e_sheet.cell(*coordinate)
            if (None if a.value == "" else a.value) != (None if e.value == "" else e.value):
                raise ValueError(f"Approved Excel semantic mismatch: {filename}/{name}!{a.coordinate}")
    if _scoped_names(actual) != _scoped_names(expected):
        raise ValueError("Approved Excel defined-name mismatch: " + filename)

def _load_native(native):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*extension is not supported.*")
        return openpyxl.load_workbook(io.BytesIO(native))

def _reviewed_native(native, filename):
    records = json.loads((CORRECTIONS / "corrections.json").read_text())
    if filename not in records:
        return native
    record = records[filename]
    if record["independent_review"] != "ACCEPT" or not record["review_report"] or not record["review_report_sha256"]:
        raise ValueError("Excel correction independent review pending: " + filename)
    report = Path(__file__).resolve().parent / record["review_report"]
    if hashlib.sha256(report.read_bytes()).hexdigest() != record["review_report_sha256"]:
        raise ValueError("Excel correction review-report hash mismatch: " + filename)
    if record["owner_sha256"] != APPROVED[filename] or record["owner_approval_of_new_text"] is not False:
        raise ValueError("Excel correction owner lineage mismatch: " + filename)
    corrected = (CORRECTIONS / filename).read_bytes()
    if hashlib.sha256(corrected).hexdigest() != record["correction_sha256"]:
        raise ValueError("Excel correction artifact hash mismatch: " + filename)
    with zipfile.ZipFile(io.BytesIO(native)) as a, zipfile.ZipFile(io.BytesIO(corrected)) as b:
        if a.namelist() != b.namelist() or len(set(a.namelist())) != len(a.namelist()):
            raise ValueError("Excel correction ZIP members mismatch: " + filename)
        for part in a.namelist():
            old, new = a.read(part), b.read(part)
            if part == "xl/sharedStrings.xml":
                old_text, new_text = escape(record["old_text"]).encode(), escape(record["new_text"]).encode()
                if old.count(old_text) != 1 or old.replace(old_text, new_text) != new:
                    raise ValueError("Excel correction shared-string scope mismatch: " + filename)
            elif old != new:
                raise ValueError("Excel correction unexpected ZIP part: " + filename + "/" + part)
    original, expected = _load_native(native), _load_native(corrected)
    cell = original[record["sheet"]][record["cell"]]
    if cell.value != record["old_text"] or expected[record["sheet"]][record["cell"]].value != record["new_text"]:
        raise ValueError("Excel correction target mismatch: " + filename)
    cell.value = record["new_text"]  # in-memory comparison only; never save owner/native with openpyxl
    _check_semantics(original, expected, filename)
    if _native_validations(native) != _native_validations(corrected):
        raise ValueError("Excel correction validation mismatch: " + filename)
    return corrected

def write_approved_layout(generated, output: Path) -> None:
    """Write owner-native bytes or a separately reviewed correction after semantic checks."""
    native = (TEMPLATES / output.name).read_bytes()
    if hashlib.sha256(native).hexdigest() != APPROVED[output.name]:
        raise ValueError("Approved Excel template hash mismatch: " + output.name)
    native = _reviewed_native(native, output.name)
    approved = _load_native(native)
    _check_semantics(generated, approved, output.name)
    if _generated_validations(generated) != _native_validations(native):
        raise ValueError("Approved Excel validation mismatch: " + output.name)
    temporary = output.with_name(output.name + ".approved.tmp")
    try:
        temporary.write_bytes(native)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()
