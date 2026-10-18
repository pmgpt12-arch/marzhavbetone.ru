from pathlib import Path
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
from io import BytesIO
from lxml import etree as E
import subprocess,hashlib,json,re,shutil
import openpyxl
P=Path(__file__).resolve().parent;R=P/'R1';QA=Path('/tmp/p13-docs-qa')
names=['00-START-HERE.pdf','60-slovar-poley-smeta.docx','61-algoritm-proverki-smety.docx','62-chek-list-vypadayushchih-zatrat.docx','63-reestr-zatrat-po-obektu.xlsx','64-zapros-o-vklyuchenii-zatrat.docx','65-protokol-raznoglasiy-po-smete.docx','66-pamyatka-nakladnye-i-smetnaya-pribyl.docx','67-pamyatka-zimnee-udorozhanie-i-vzis.docx','68-pamyatka-musor-transport-nepredvidennye.docx']
assert sorted(x.name for x in R.iterdir())==sorted(names)
sha=lambda b:hashlib.sha256(b).hexdigest()
def doc_text(p):
 with ZipFile(p) as z:
  root=E.fromstring(z.read('word/document.xml'))
  return '\n'.join(root.xpath('//*[local-name()="p"]//*[local-name()="t"]/text()'))
texts={n:doc_text(R/n) for n in names if n.endswith('.docx')}
dictionary=texts[names[1]]
fields=set()
for name in [names[5],names[6]]:fields.update(re.findall(r'\{\{([^{}]+)\}\}',texts[name]))
assert all(f in dictionary for f in fields),fields
assert texts[names[3]].count('☐')==62
assert texts[names[5]].startswith('Кому:')
assert texts[names[6]].startswith('Протокол разногласий по смете')
pdftext=subprocess.check_output(['pdftotext',str(R/names[0]),'-'],text=True)
assert all(n in pdftext for n in names[1:]),'PDF file references'
source=openpyxl.load_workbook(P/'workbook-build/print-source.xlsx',data_only=False)
final=openpyxl.load_workbook(R/names[4],data_only=False)
a=source['Реестр затрат'];b=final['Реестр затрат']
assert [[a.cell(i,j).value for j in range(1,4)] for i in range(1,53)]==[[b.cell(i,j).value for j in range(1,4)] for i in range(1,53)]
assert all(b.cell(i,j).value is None for i in range(2,53) for j in range(4,11))
assert b['H55'].value is None
assert len([c for row in b for c in row if c.data_type=='f'])==3
for d in b.data_validations.dataValidation:assert d.showErrorMessage and d.errorStyle=='stop'
for s in source.sheetnames:
 assert str(source[s].print_area)==str(final[s].print_area)
 assert source[s].print_title_rows==final[s].print_title_rows
 assert source[s].page_setup==final[s].page_setup
native=json.loads((P/'native-report.json').read_text());assert native['status']=='PASS'
assert len(native['scenarios'])==10
qadir=P/'print-qa';qadir.mkdir(exist_ok=True)
pages={};expected=[1,1,1,3,3,2,1,1,1,1]
for n,want in zip(names,expected):
 pdf=R/n if n.endswith('.pdf') else QA/(Path(n).stem+'.pdf')
 info=subprocess.check_output(['pdfinfo',str(pdf)],text=True)
 count=int(re.search(r'Pages:\s*(\d+)',info).group(1));assert count==want,(n,count,want)
 content=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True).split('\f')
 assert all(p.strip() for p in content[:count]),(n,'blank page')
 pages[n]=count
 if not n.endswith('.pdf'):shutil.copyfile(pdf,qadir/pdf.name)
def build_zip():
 out=BytesIO()
 with ZipFile(out,'w',ZIP_DEFLATED,compresslevel=9) as z:
  for n in names:
   info=ZipInfo(n,date_time=(2026,10,18,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o100644<<16
   z.writestr(info,(R/n).read_bytes(),compress_type=ZIP_DEFLATED,compresslevel=9)
 return out.getvalue()
blob=build_zip();assert blob==build_zip()
archive=P/'P13-smeta-10-files-business-R1.zip';archive.write_bytes(blob)
with ZipFile(archive) as z:assert z.testzip() is None and z.namelist()==names
gate=subprocess.run(['python3','tools/editorial_business_style_gate.py',str(archive)],text=True,capture_output=True)
(P/'editorial-gate.txt').write_text(gate.stdout+gate.stderr);assert gate.returncode==0,gate.stdout+gate.stderr
files=[{'name':n,'size':(R/n).stat().st_size,'sha256':sha((R/n).read_bytes()),'print_pages':pages[n]} for n in names]
manifest={'product':'P13','title':'Система проверки сметной документации и выявления неоплачиваемых затрат','status':'READY_FOR_OWNER_BUYER_REVIEW_NOT_PUBLISHED','target_price_rub':17900,'branch':'candidate/mb001-p13-owner-review-20261018','base_commit':'6323b4fc74f6a92bc19588bd03ba18ef5a7857c9','archive':{'name':archive.name,'size':len(blob),'sha256':sha(blob),'files':len(names),'crc':'PASS','reproducible_from_exact_files':True},'files':files,'changed_files':names,'checks':{'editorial_gate':gate.stdout.strip(),'manual_semantic_review':'Completed for every buyer document and workbook instruction','dictionary_fields':len(fields),'checklist_points':62,'all_print_pages_visually_reviewed':True,'xlsx_print':'2 x A3 landscape register with repeating header, 1 x A4 landscape guide','xlsx_native_scenarios':'10 PASS in LibreOffice 24.2.7.2; final exact SHA in native-report.json','xlsx_formulas':'H54 and H56 corrected; C53 completeness warning added; no other original formulas','xlsx_preservation_report':'workbook-build/preservation-report.json','xlsx_blank_form_no_fixture_data':'PASS','microsoft_excel_native_test':False},'sources':{'checklist_only':'portfolio-owner-editorial-legal-v4-20261013/16-chto-v-smete-ne-oplatyat: preserved 62 checklist points with local business wording edits','print_workbook':'PR #496 at 467b942c632cb00c431ae1c3d1394ce2b5249dfb, SHA 2262235a4f2d3513bed99c58a261053181bf5510af684e9a22453ac529c63e52; print acceptance does not imply acceptance of old content or formulas'},'owner_acceptance':'Pending; review forms in Word and register in Excel; publication requires owner authorization','production_changes':False}
(P/'r1-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(manifest['archive'],ensure_ascii=False));print(gate.stdout.strip())
