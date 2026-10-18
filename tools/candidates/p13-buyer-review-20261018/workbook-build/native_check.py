"""Disposable OOXML input fixtures; LibreOffice recalculates actual buyer formulas."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
import subprocess,json,sys,math
import openpyxl
source=Path(sys.argv[1]).resolve();work=Path('/tmp/p13-native');work.mkdir(exist_ok=True)
inputs=work/'inputs';outputs=work/'outputs';inputs.mkdir(exist_ok=True);outputs.mkdir(exist_ok=True)
NS={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};Q='{'+NS['x']+'}'
with ZipFile(source) as z:parts={n:z.read(n) for n in z.namelist()}
def fixture(name,values):
 root=E.fromstring(parts['xl/worksheets/sheet1.xml']);data=root.find('x:sheetData',NS)
 for addr,value in values.items():
  c=root.find(f'.//x:c[@r="{addr}"]',NS)
  if c is None:
   row=data.find(f'x:row[@r="{"".join(filter(str.isdigit,addr))}"]',NS);c=E.SubElement(row,Q+'c',r=addr)
  for x in list(c):c.remove(x)
  c.attrib.pop('t',None)
  if value is None:continue
  if isinstance(value,str):c.set('t','inlineStr');E.SubElement(E.SubElement(c,Q+'is'),Q+'t').text=value
  else:E.SubElement(c,Q+'v').text=str(value)
 for cell in root.findall('.//x:c',NS):
  if cell.find('x:f',NS) is not None:
   v=cell.find('x:v',NS)
   if v is not None:cell.remove(v)
 for row in data:
  def col(c):
   n=0
   for a in ''.join(filter(str.isalpha,c.get('r'))):n=n*26+ord(a)-64
   return n
  row[:]=sorted(row,key=col)
 pp=parts.copy();pp['xl/worksheets/sheet1.xml']=E.tostring(root,encoding='utf-8',xml_declaration=True)
 with ZipFile(inputs/(name+'.xlsx'),'w',ZIP_DEFLATED) as z:
  for n,b in pp.items():z.writestr(n,b)
base={**{f'D{i}':'не применяется' for i in range(2,53)},**{f'I{i}':'не применяется' for i in range(2,53)},'H55':1000000}
commercial=base.copy()
for row,status,amount,doc in [(2,'в запросе',200000,''),(3,'в протоколе разногласий',80000,''),(4,'покрыто, проверено',100000,'Смета п.4'),(5,'принято',50000,'Соглашение'),(6,'риск принят',30000,'Решение директора')]:
 commercial.update({f'D{row}':'нет',f'I{row}':status,f'H{row}':amount,f'G{row}':doc})
cases=[('01-blank',{},0,None,'Заполните покрытие'),('02-nonapplicable',base,0,0,'Проверка заполнена'),
 ('03-commercial',commercial,310000,.31,'Проверка заполнена'),
 ('04-missing',{**commercial,'H2':None},110000,.11,'без числовой оценки'),
 ('05-text',{**commercial,'H2':'200 тыс'},110000,.11,'без числовой оценки'),
 ('06-proof',{**commercial,'G4':None},310000,.31,'документы покрытия'),
 ('07-zero-profit',{**commercial,'H55':0},310000,None,'Проверка заполнена'),
 ('08-last-row',{**commercial,'D52':'нет','I52':'в запросе','H52':25000},335000,.335,'Проверка заполнена'),
 ('09-excluded-stale',{**commercial,'H52':25000},310000,.31,'Проверка заполнена'),
 ('10-missing-status',{**base,'I52':None},0,0,'Заполните статус')]
for name,values,*_ in cases:fixture(name,values)
subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo-p13-native','--headless','--convert-to','xlsx','--outdir',str(outputs)]+[str(p) for p in sorted(inputs.glob('*.xlsx'))],check=True,stdout=subprocess.DEVNULL)
results=[]
for name,_,amount,ratio,message in cases:
 w=openpyxl.load_workbook(outputs/(name+'.xlsx'),data_only=True);s=w['Реестр затрат'];a=s['H54'].value;r=s['H56'].value;m=s['C53'].value
 assert a==amount,(name,'amount',a,amount)
 assert (r is None or r=='') if ratio is None else isinstance(r,(int,float)) and math.isclose(r,ratio), (name,'ratio',r,ratio)
 assert message in str(m),(name,'message',m,message)
 errors=[c.coordinate for sh in w for row in sh for c in row if c.data_type=='e'];assert not errors,(name,errors)
 results.append({'case':name,'H54':a,'H56':r,'message':m,'status':'PASS'})
native=openpyxl.load_workbook(outputs/'03-commercial.xlsx',data_only=False)
dv=list(native['Реестр затрат'].data_validations.dataValidation)
assert {str(r) for d in dv for r in d.sqref.ranges}=={'D2:D52','F2:F52','I2:I52','H2:H52','H55'},[(str(d.sqref),d.type,d.formula1) for d in dv]
statuses=[d for d in dv if str(d.sqref)=='I2:I52'];assert statuses and 'D$2' in statuses[0].formula1,statuses
assert all(d.showErrorMessage and d.errorStyle=='stop' for d in dv)
raw=openpyxl.load_workbook(source,data_only=False)
assert len(raw['Реестр затрат'].data_validations.dataValidation)==5
assert raw['Реестр затрат'].print_title_rows=='$1:$1'
assert len([c for row in raw['Реестр затрат'] for c in row if c.data_type=='f'])==3
import hashlib
report={'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'status':'PASS','engine':subprocess.check_output(['libreoffice','--version'],text=True).strip(),'scenarios':results,'native_roundtrip_validation_ranges':5,'native_roundtrip_validation_rules':len(dv),'native_excel_test':False,'scope':'LibreOffice recalculation, OOXML structure and print QA; Microsoft Excel remains buyer acceptance environment'}
Path(sys.argv[2]).write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False))
