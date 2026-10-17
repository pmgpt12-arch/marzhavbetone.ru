#!/usr/bin/env python3
"""T3 buyer archive structure, field dictionary and first-page form gate."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from docx import Document
import re, subprocess, sys
archive=Path(sys.argv[1]); expected=['00-START-HERE.pdf','01-poryadok-deystviy.docx','02-vozrazheniya-na-akt.docx','03-vozrazheniya-na-zachet.docx','04-slovar-poley.docx']
with ZipFile(archive) as z:
 assert z.testzip() is None and len(z.namelist())==5 and set(z.namelist())==set(expected), 'ZIP CRC or buyer inventory'
 pdf=subprocess.run(['pdftotext','-','-'],input=z.read(expected[0]),capture_output=True,check=True).stdout.decode()
 assert all(name in pdf for name in expected[1:]), 'Instruction references'
 docs={name:Document(BytesIO(z.read(name))) for name in expected[1:]}
 words=' '.join([p.text for p in docs[expected[4]].paragraphs]+[c.text for t in docs[expected[4]].tables for row in t.rows for c in row.cells])
 for name in expected[2:4]:
  d=docs[name]; paras=[p.text.strip() for p in d.paragraphs if p.text.strip()]
  assert paras[0].startswith('Кому:') and paras[1].startswith('Адрес:') and any(p.startswith('Исх. №') for p in paras[:5]), 'Form first '+name
  fields=set(re.findall(r'\{\{([А-ЯЁ_]+)\}\}', '\n'.join(paras)))
  assert fields and all(field in words for field in fields), 'Dictionary gaps '+name+' '+str(fields-set(words.split()))
 print('T3_BUYER_GATE_PASS files=5 CRC=ok PDF_links=ok forms_first=ok fields_in_dictionary=ok')
