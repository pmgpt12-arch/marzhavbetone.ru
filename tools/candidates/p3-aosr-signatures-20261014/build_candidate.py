"""P3: one reusable existing AOSR source -> five complete working acts.
Only isolated outputs. Other P3 issues and independent/native acceptance remain open.
"""
import copy
import hashlib
import io
import json
import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT/'candidate'
OUT.mkdir(exist_ok=True)
EXPECTED = {
 '01-proverochnyy-list-komplekta-ks.docx':'194c55b064ee7a86a115bbe3db5f51d141bd542c3d3fecbcfb9a61e04c0cd10d',
 '03-akt-skrytyh-rabot.docx':'82e02be86bb370aba267f9ec2c698e01f05bc5f523bd13485362c699853fc854',
 '09-tipovye-oshibki.docx':'8d91dff418c122555d8b524cf7b262c6c2a1f374bfd6dfe3678b0e38f148a675',
 '04-akt-skrytyh-rabot.docx':'cf769ace49643dfedfe156f229ebb7c5aed82426e82e6fc262899d9f3699fe80'}
def sha(b): return hashlib.sha256(b).hexdigest()
inputs = {n:(ROOT/'inputs'/n).read_bytes() for n in EXPECTED}
for n,b in inputs.items():
    if sha(b) != EXPECTED[n]: raise ValueError('SOURCE_DRIFT: '+n)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

DELTAS = {
 '01-proverochnyy-list-komplekta-ks.docx': {
  'Подписи: руководитель + главный инженер + бухгалтер':'КС-2 подписывают уполномоченные представители сторон; документы о полномочиях проверены. Внутреннее согласование проверено отдельно.',
  'КС-11: номера актов скрытых работ совпадают':'АОСР: наименования, участки и номера актов сопоставлены с соответствующими скрытыми работами в КС-2.',
  'Все работы из КС-2 покрыты актами':'АОСР оформлены на подлежащие освидетельствованию скрытые работы; применимость подтверждена проектом и перечнем работ объекта.',
  'Акты подписаны представителем заказчика':'У АОСР есть подпись уполномоченного представителя заказчика по строительному контролю; полномочия и НРС проверены по применимости.',
  'Акты подписаны представителем автора проекта':'Представитель проектировщика участвует и подписывает АОСР при его привлечении для проверки соответствия работ проекту.',
  'Акты подписаны руководителем работ':'Подписи представителей лица, осуществляющего строительство, и его строительного контроля проверены; исполнитель по отдельному договору участвует по применимости.',
  'Даты актов меньше или равны дате окончания периода КС-2':'Указаны фактические даты работ и освидетельствования до их закрытия последующими работами; связь с периодом КС-2 проверена без изменения фактических дат.'},
 '09-tipovye-oshibki.docx': {
  'Ошибка 4. Отсутствие актов скрытых работ':'Ошибка 4. Нет АОСР на работы, подлежащие освидетельствованию',
  'ПРИЧИНА: Работы выполнены, акты составлены, но не приложены к КС-2.\nУСТРАНЕНИЕ: Реестр приложений с автопроверкой (входит в полный комплект).':'ПРИЧИНА: Не определены скрытые работы либо не подтверждено их освидетельствование.\nУСТРАНЕНИЕ: Сверьте проект и перечень работ объекта. Проверьте соответствующие АОСР, участки, даты, полномочия и применимость участников. Сопоставьте акты с приложениями к комплекту; замечания внесите в протокол №04.',
  'Ошибка 7. Отсутствие подписи бухгалтера':'Ошибка 7. Нет подписи или полномочий представителя стороны',
  'ПРИЧИНА: Бухгалтер не участвует в процессе, считает что его подпись не нужна.\nУСТРАНЕНИЕ: Включите бухгалтера в процесс внутренней проверки. Обучите требованиям.':'ПРИЧИНА: КС-2 не подписан уполномоченным представителем стороны либо его полномочия не подтверждены.\nУСТРАНЕНИЕ: Проверьте подписантов и основания полномочий. Внутреннюю проверку бухгалтерией оформляйте отдельно: должность бухгалтера сама по себе не делает его подпись обязательной в КС-2.'}}

receipts = []
for name, mapping in DELTAS.items():
    original = inputs[name]
    with ZipFile(io.BytesIO(original)) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        found = set()
        for p in xml.iter(W+'p'):
            current = ''.join((x.text or '') if x.tag==W+'t' else '\n' for x in p.iter() if x.tag in [W+'t', W+'br'])
            if current not in mapping: continue
            if current in found: raise ValueError('NON_UNIQUE_ANCHOR: '+current)
            found.add(current)
            first = p.find(W+'r')
            rp = copy.deepcopy(first.find(W+'rPr')) if first is not None and first.find(W+'rPr') is not None else None
            for child in list(p):
                if child.tag != W+'pPr': p.remove(child)
            r = etree.SubElement(p,W+'r')
            if rp is not None: r.append(rp)
            for i,line in enumerate(mapping[current].split('\n')):
                if i: etree.SubElement(r,W+'br')
                t = etree.SubElement(r,W+'t');t.text=line
        if found != set(mapping): raise ValueError('MISSING_ANCHORS: '+str(set(mapping)-found))
        patched = etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
        buffer = io.BytesIO()
        with ZipFile(buffer,'w',ZIP_DEFLATED) as dst:
            for info in z.infolist(): dst.writestr(info,patched if info.filename=='word/document.xml' else z.read(info.filename))
    data=buffer.getvalue();(OUT/name).write_bytes(data)
    with ZipFile(io.BytesIO(original)) as a,ZipFile(io.BytesIO(data)) as b:
        changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
        assert changed==['word/document.xml']
    receipts.append({'file':name,'source_sha256':sha(original),'sha256':sha(data),'exact_text_replacements':len(mapping),'changed_parts':changed})

master = Document(io.BytesIO(inputs['04-akt-skrytyh-rabot.docx']))
master_text = [p.text for p in master.paragraphs]
assert len(master_text)==53 and all(any(s.startswith(str(i)+'. ') for s in master_text) for i in range(1,8))
# Supplement the source's condensed prompts with the applicable source requisites.
COMMON = dict(enumerate(master_text))
COMMON[1]='На основе рекомендуемого образца: приложение №3 к составу исполнительной документации, утверждённому приказом Минстроя России от 16.05.2023 №344/пр.'
COMMON[5]='Для юрлица: наименование, ОГРН, ИНН, адрес, телефон/факс. Для ИП: ФИО, адрес, ОГРНИП, ИНН. Для физлица, не ИП: ФИО, паспортные данные, адрес, телефон/факс. СРО: наименование, ОГРН, ИНН — когда членство требуется.'
COMMON[7]='Для юрлица: наименование, ОГРН, ИНН, адрес, телефон/факс. Для ИП: ФИО, адрес, ОГРНИП, ИНН. СРО: наименование, ОГРН, ИНН — когда членство требуется.'
COMMON[9]=COMMON[7]
COMMON[13]='Должность, фамилия, инициалы; документ о полномочиях; НРС — кроме случаев, когда членство в СРО не требуется. При контроле по договору: реквизиты организации (наименование, ОГРН, ИНН, адрес) либо ИП (ФИО, адрес, ОГРНИП, ИНН).'
COMMON[15]='Должность (при наличии), фамилия, инициалы, реквизиты документа о полномочиях.'
COMMON[17]='Должность, фамилия, инициалы, документ о полномочиях; НРС — кроме случаев, когда членство в СРО не требуется.'
COMMON[19]='Должность, фамилия, инициалы, документ о полномочиях; наименование, ОГРН, ИНН и адрес организации либо ФИО, адрес, ОГРНИП и ИНН индивидуального предпринимателя.'
COMMON[21]=COMMON[19]
TYPES=['Бетонирование','Гидроизоляция','Электромонтаж','Отопление и водоснабжение','Иные скрытые работы']
FILLABLE={2:'ОБЪЕКТ',4:'ЗАКАЗЧИК',6:'СТРОИТЕЛЬ',8:'ПРОЕКТИРОВЩИК',12:'ПРЕДСТАВИТЕЛЬ_ЗАКАЗЧИКА',14:'ПРЕДСТАВИТЕЛЬ_СТРОИТЕЛЯ',16:'СТРОЙКОНТРОЛЬ_СТРОИТЕЛЯ',18:'ПРЕДСТАВИТЕЛЬ_ПРОЕКТИРОВЩИКА',20:'ПРЕДСТАВИТЕЛЬ_ИСПОЛНИТЕЛЯ',23:'РАБОТЫ_И_МЕСТО',25:'ПРОЕКТ_И_ЧЕРТЕЖИ',27:'МАТЕРИАЛЫ_И_ДОКУМЕНТЫ',29:'СХЕМЫ_И_ИСПЫТАНИЯ',31:'ДАТЫ_РАБОТ',33:'ТРЕБОВАНИЯ_И_ПРОЕКТ',35:'ПОСЛЕДУЮЩИЕ_РАБОТЫ',37:'ДОПОЛНИТЕЛЬНЫЕ_СВЕДЕНИЯ',38:'ЭКЗЕМПЛЯРЫ',40:'ПРИЛОЖЕНИЯ'}
HELPERS={3,5,7,9,13,15,17,19,21,24,26,28,30,32,34,36,39,41}

def document_base():
    d=Document(io.BytesIO(inputs['04-akt-skrytyh-rabot.docx']))
    for node in list(d._element.body):
        if node.tag!=qn('w:sectPr'): d._element.body.remove(node)
    for section in d.sections:
        section.page_width=Cm(21);section.page_height=Cm(29.7)
        section.top_margin=Cm(1.55);section.bottom_margin=Cm(1.55)
        section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
        section.header_distance=Cm(0.55);section.footer_distance=Cm(0.55)
        for x in [section.header,section.footer]:
            for p in x.paragraphs:p.text=''
        p=section.header.paragraphs[0];p.text='МАРЖА В БЕТОНЕ  /  КОНТРОЛЬ СКРЫТЫХ РАБОТ'
        p.runs[0].font.size=Pt(8);p.runs[0].font.color.rgb=RGBColor.from_string('756A50')
        p=section.footer.paragraphs[0];p.text='АОСР  /  '
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    for style in d.styles:
        if style.type in (1,2):
            style.font.name='Arial';style.font.italic=False;style.font.color.rgb=RGBColor.from_string('262626')
    normal=d.styles['Normal'];normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(2);normal.paragraph_format.line_spacing=1
    return d

def add(d,text,size=10,bold=False,helper=False,new_page=False):
    p=d.add_paragraph();p.paragraph_format.page_break_before=new_page
    p.paragraph_format.space_after=Pt(2 if helper else 3)
    p.paragraph_format.keep_together=True
    r=p.add_run(text);r.font.name='Arial';r.font.size=Pt(size);r.bold=bold;r.italic=False
    r.font.color.rgb=RGBColor.from_string('5A5A5A' if helper else '262626')
    return p

def append_act(d,index,kind,values=None):
    prefix='A'+str(index)+'_'
    for i in range(53):
        text=COMMON[i]
        if i==0:text='АКТ ОСВИДЕТЕЛЬСТВОВАНИЯ СКРЫТЫХ РАБОТ — '+kind
        if i in FILLABLE:
            text=re.sub(r'_{10,}','{{'+prefix+FILLABLE[i]+'}}',text)
        if i==10:text='АКТ № {{'+prefix+'НОМЕР_АКТА}} от {{'+prefix+'ДАТА_АКТА}}'
        if i==22:text='Произвели осмотр работ, выполненных {{'+prefix+'ИСПОЛНИТЕЛЬ_РАБОТ}}, и составили настоящий акт о нижеследующем.'
        if i==32:text='Укажите фактические даты начала и окончания работ.'
        if i in [44,46,48,50,52]:text='{{'+prefix+'ФИО_'+str((i-42)//2)+'}}    Подпись __________________'
        if values:
            for key,value in values.items():text=text.replace('{{'+prefix+key+'}}',value)
        p=add(d,text,size=13 if i==0 else (8.5 if i in HELPERS else 10),bold=i in [0,10,11,42],helper=i in HELPERS,new_page=i in [0,23])
        if i in FILLABLE or i in [11,42,43,45,47,49,51]:p.paragraph_format.keep_with_next=True

full=document_base()
add(full,'ПЯТЬ РАБОЧИХ ФОРМ АОСР',size=18,bold=True)
add(full,'Выберите форму по виду скрытых работ',size=12,bold=True)
add(full,'Каждая форма ниже является отдельным заполняемым актом: реквизиты объекта и участников, пункты 1–7, приложения и подписи. Общая структура едина; наименование вида работ служит навигацией, а не самостоятельным нормативным основанием.')
for i,kind in enumerate(TYPES,1):add(full,str(i)+'. '+kind)
add(full,'Порядок заполнения',size=12,bold=True)
for text in [
 'Скопируйте выбранный акт в отдельный документ. Заполните поля в двойных фигурных скобках фактическими сведениями и удалите подсказки. В пятой форме замените общее название конкретным видом работ.',
 'Применимость АОСР определяйте по проекту и перечню работ объекта: это не акт на каждую строку КС-2. Впишите точные участки, оси и отметки; укажите документы качества и результаты контроля, относящиеся к этим работам.',
 'Участников определяйте по их функциям и условиям привлечения, указанным в форме. Подпись проектировщика не является безусловным требованием для каждого акта. Поля неприменимого участника и соответствующий блок подписи исключайте согласованно с условиями объекта.',
 'Указывайте фактические даты. Освидетельствование и разрешение последующих работ должны быть оформлены до закрытия проверяемых работ; не подгоняйте дату акта под период КС-2.',
 'При подготовке КС используйте проверочный лист №01. Недостатки фиксируйте в протоколе №04; сверку приложений и дат в реестре №02 выполняйте по документам, а не только по отметке «ОК».']:
    add(full,text)
add(full,'Основание: приложение №3 к составу исполнительной документации по приказу Минстроя России от 16.05.2023 №344/пр. Варианты используют одну структуру формы; числовые критерии для конкретных работ берутся из проекта и применимых требований.',size=9,helper=True)
for i,kind in enumerate(TYPES,1):append_act(full,i,kind)
full.save(OUT/'03-akt-skrytyh-rabot.docx')

# Synthetic filling is QA evidence, not a signed construction record.
fixture=document_base()
vals={key:'ТЕСТ: заполненное значение '+key.lower().replace('_',' ') for key in FILLABLE.values()}
vals.update({'НОМЕР_АКТА':'УЧ-003','ДАТА_АКТА':'14.10.2026','ИСПОЛНИТЕЛЬ_РАБОТ':'ООО «Подрядчик» (учебный пример)',
 'ОБЪЕКТ':'Учебный объект; секция 1, оси А–Б/1–2; адрес условный',
 'РАБОТЫ_И_МЕСТО':'Гидроизоляция участка фундамента в осях А–Б/1–2; отметка −1,200',
 'ПРОЕКТ_И_ЧЕРТЕЖИ':'Учебная РД КР-01, лист 12; ООО «Проектировщик»',
 'МАТЕРИАЛЫ_И_ДОКУМЕНТЫ':'Учебный материал М-01; документ качества №ДК-01 от 12.10.2026',
 'СХЕМЫ_И_ИСПЫТАНИЯ':'Учебная исполнительная схема №ИС-03 от 14.10.2026; протокол контроля №ПК-03',
 'ДАТЫ_РАБОТ':'Начало 12.10.2026; окончание 14.10.2026',
 'ПОСЛЕДУЮЩИЕ_РАБОТЫ':'Защитный слой на указанном участке после освидетельствования',
 'ЭКЗЕМПЛЯРЫ':'2','ДОПОЛНИТЕЛЬНЫЕ_СВЕДЕНИЯ':'Учебный пример заполнения. Не является актом выполненных работ.',
 'ПРИЛОЖЕНИЯ':'Учебная ИС-03, 1 лист; ДК-01, 1 лист; ПК-03, 1 лист'})
for i in range(1,6):vals['ФИО_'+str(i)]='УЧЕБНЫЙ ПОДПИСАНТ '+str(i)
append_act(fixture,1,'Гидроизоляция / пример заполнения',vals)
fixture.save(OUT/'03-filled-example-QA.docx')

new_text='\n'.join(p.text for p in full.paragraphs)
checks={
 'source_identity_4_of_4':all(sha(inputs[n])==s for n,s in EXPECTED.items()),
 'five_complete_acts':sum(p.text.startswith('АКТ №') for p in full.paragraphs)==5,
 'all_seven_sections_per_act':all(sum(p.text.startswith(str(i)+'. ') and 'форма' not in p.text.lower() for p in full.paragraphs)>=5 for i in range(1,8)),
 'five_signature_sections':new_text.count('Подписи представителей')==5,
 'conditional_designer_per_act':new_text.count('если привлечен для проверки соответствия работ проекту')==10,
 'conditional_subcontractor_per_act':new_text.count('если работы выполнены иным лицом по договору')==10,
 'five_appendices':sum(p.text.startswith('Приложения  ') for p in full.paragraphs)==5,
 'old_description_only_stubs_removed':not any(x in new_text for x in ['[Форма с','[Форма акта','[Универсальный','Подписи всех трех сторон обязательны']),
 'no_banned_language':not any(x in new_text.lower() for x in ['шаблон','юридическая консультация','не заменяет юриста']),
 'filled_fixture_no_tokens':not any('{{' in p.text for p in fixture.paragraphs),
 'original_inputs_unchanged':all((ROOT/'inputs'/n).read_bytes()==b for n,b in inputs.items())}
for n in ['03-akt-skrytyh-rabot.docx','03-filled-example-QA.docx']:
    b=(OUT/n).read_bytes();receipts.append({'file':n,'sha256':sha(b),'bytes':len(b),'source_master_sha256':EXPECTED['04-akt-skrytyh-rabot.docx']})
    with ZipFile(io.BytesIO(b)) as z:checks['crc_'+n]=z.testzip() is None
report={'id':'P3-AOSR-SIGNATURES-REMEDIATION-20261014','claim':6167364456,'base':'df29cf76472e7d8f83b759e69f60b46b25581fb7','status':'AUTHOR_STRUCTURAL_PASS' if all(checks.values()) else 'FAIL_NOT_VERIFIED','checks':checks,'passed':sum(checks.values()),'total':len(checks),'artifacts':receipts,'source_deltas':DELTAS,'native_word':'NOT_RUN','visual':'PENDING','independent_subject_review':'PENDING','release':'HOLD','G1_G8_accepted_products':0,'open_other_p3_findings':['07 retention sources','02/06 workbooks','08 route','remaining 01/09 claims outside this delta','10 marketing/price assertions','final nomination and brand acceptance']}
(ROOT/'author-structural.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'status':report['status'],'checks':checks,'artifacts':receipts},ensure_ascii=False,indent=2))
if not all(checks.values()):raise SystemExit(1)
