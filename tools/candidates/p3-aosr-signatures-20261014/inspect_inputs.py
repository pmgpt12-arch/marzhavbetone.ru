"""Read four pinned GitHub documents; write only isolated copies and metadata."""
import base64
import hashlib
import io
import json
import subprocess
from pathlib import Path
from docx import Document
ROOT = Path(__file__).resolve().parent
REF = 'df29cf76472e7d8f83b759e69f60b46b25581fb7'
REPO = 'repos/pmgpt12-arch/marzhavbetone.ru'
PREFIX = 'tools/candidates/portfolio-owner-editorial-v3-20261013/05-ks-bez-vozvrata/'
PATHS = [PREFIX+n for n in ['01-proverochnyy-list-komplekta-ks.docx','03-akt-skrytyh-rabot.docx','09-tipovye-oshibki.docx']]
PATHS += ['products-storage/02-dopraboty-bez-poter/04-akt-skrytyh-rabot.docx']
(ROOT/'inputs').mkdir(exist_ok=True)
records = []
for path in PATHS:
    raw = subprocess.check_output(['gh','api',REPO+'/contents/'+path+'?ref='+REF], timeout=12)
    obj = json.loads(raw)
    data = base64.b64decode(obj['content'])
    digest = hashlib.sha256(data).hexdigest()
    blob = hashlib.sha1(b'blob '+str(len(data)).encode()+b'\0'+data).hexdigest()
    assert blob == obj['sha'], path
    dest = ROOT/'inputs'/Path(path).name
    if dest.exists():
        assert dest.read_bytes() == data, 'Input collision'
    else:
        dest.write_bytes(data)
    doc = Document(io.BytesIO(data))
    paras = [{'index':i,'text':p.text} for i,p in enumerate(doc.paragraphs)]
    rows = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    entry = {'path':path,'sha256':digest,'git_blob':blob,'bytes':len(data),'paragraphs':paras,'tables':rows}
    records.append(entry)
    print('FILE',path,'SHA',digest,'BLOB',blob)
    if Path(path).name[:2] in ['03','04']:
        print(json.dumps({'paragraphs':paras,'tables':rows},ensure_ascii=False))
    else:
        print(json.dumps([p for p in paras if any(s in p['text'].lower() for s in ['подпис','аоср','скрыт','кс-11','кс-4'])],ensure_ascii=False))
(ROOT/'input-evidence.json').write_text(json.dumps({'ref':REF,'files':records},ensure_ascii=False,indent=2)+'\n')
print('INPUTS_VERIFIED',len(records))
