"""Вспомогательный дамп: ячейки листов xlsx и текст docx (для чтения проверяющим)."""
import sys
from pathlib import Path

import docx
import openpyxl

p = Path(sys.argv[1])
if p.suffix == ".xlsx":
    wb = openpyxl.load_workbook(p)
    for name in (sys.argv[2:] or wb.sheetnames):
        ws = wb[name]
        print(f"=== {name} {ws.dimensions}")
        for row in ws.iter_rows(max_row=min(ws.max_row, 60)):
            for c in row:
                if c.value is not None:
                    print(c.coordinate, repr(c.value)[:400])
else:
    d = docx.Document(p)
    body = d.element.body
    for el in body.iter():
        if el.tag.endswith("}p"):
            t = "".join(x.text or "" for x in el.iter() if x.tag.endswith("}t"))
            if t.strip():
                print(t)
