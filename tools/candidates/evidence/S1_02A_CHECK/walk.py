"""Независимый прогон С-1, С-3, С-4 по файлам, распакованным из штатного ZIP
(<scratch>/kit, см. zip_fingerprint.py). Данные — вымышленные данные #325 §2.
Книги заполняются openpyxl и пересчитываются в Calc; документы заполняются и
конвертируются в PDF в Writer. Каждый запуск LibreOffice — с отдельным
профилем <scratch>/lo-profile. Продуктовые файлы не меняются.
Запуск: python3 walk.py <scratch>"""
import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path

import docx
import openpyxl

S = Path(sys.argv[1]).resolve()
KIT = S / "kit"
PROFILE = f"file://{S}/lo-profile"
D = lambda s: dt.datetime.strptime(s, "%d.%m.%Y")  # noqa: E731
ACTS = [("А-1", "КС-2 № 1 от 14.04.2026", 400000, "14.04.2026", "14.05.2026"),
        ("А-2", "КС-2 № 2 от 19.05.2026", 300000, "19.05.2026", "18.06.2026"),
        ("А-3", "КС-2 № 3 от 16.06.2026", 200000, "16.06.2026", "16.07.2026")]
# Ответы на 16 вопросов «Маршрута» (C5:C20) для чистого случая А.
BASE = ["Да", "Да", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет", "Нет", "Да", "Да", "Нет", "Нет", "Да"]
ROUTE = {  # вариант → {номер вопроса: ответ}
    "С-1 (чистый А)": {},
    "С-3 (только письмо)": {15: "Да"},
    "письмо + соглашение": {14: "Да", 15: "Да"},
    "письмо, сумма не сверена": {15: "Да", 16: "Нет"},
    "С-4 (B4 по акту 2 + B5 в акте 3)": {6: "Да", 9: "Да"},
    "B4 + B2 + B5 + B1": {6: "Да", 8: "Да", 9: "Да", 11: "Да"},
    "B5 + письмо": {9: "Да", 15: "Да"},
}
SC = {
    "С-1": dict(ctrl=["Не вышел на переговоры", "Молчание после срока", "Молчание после срока"], acts08=[0, 1, 2], calc="10.08.2026"),
    "С-3": dict(ctrl=["Прислал гарантийное письмо без соглашения", "Обещал оплатить в названную дату", "Молчание после срока"],
                acts08=[0, 1, 2], calc="17.08.2026"),
    "С-4": dict(ctrl=["Не вышел на переговоры", "Молчание после срока", "Возражает по сумме или вернул акт сверки с расхождением",
                      "Отказ с мотивами: объём, качество, подпись", "Молчание после срока"], acts08=[0], calc="10.08.2026"),
}


def calc(files: list[Path], out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation={PROFILE}", "--headless", "--norestore", "--calc",
                    "--convert-to", "xlsx", "--outdir", str(out), *map(str, files)], capture_output=True, timeout=600, check=True)


def writer_pdf(files: list[Path], out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation={PROFILE}", "--headless", "--norestore",
                    "--convert-to", "pdf", "--outdir", str(out), *map(str, files)], capture_output=True, timeout=600, check=True)


def errors(wb) -> int:
    return sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#")
               and c.value[:4] in ("#REF", "#VAL", "#NAM", "#DIV", "#N/A", "#NUM", "#NUL"))


def main() -> None:
    work = S / "walk"
    shutil.rmtree(work, ignore_errors=True)
    # 1. «Маршрут» 02: варианты ответов → код, итог, следующий шаг, предупреждение, «Ещё границы».
    src = work / "route-src"
    src.mkdir(parents=True)
    names = {}
    for i, (name, ch) in enumerate(ROUTE.items()):
        wb = openpyxl.load_workbook(KIT / "02-proverka-i-kontrol-otveta.xlsx")
        ws = wb["Маршрут"]
        for q, a in enumerate(BASE, start=1):
            ws.cell(row=4 + q, column=3, value=ch.get(q, a))
        ws["C27"] = D("14.05.2026")
        p = src / f"route{i}.xlsx"
        wb.save(p)
        names[p.name] = name
    calc(sorted(src.iterdir()), work / "route-calc")
    print("## 02 «Маршрут» после пересчёта в Calc")
    for fn, name in names.items():
        ws = openpyxl.load_workbook(work / "route-calc" / fn, data_only=True)["Маршрут"]
        print(f"- {name}: C22={ws['C22'].value!r}; C23={ws['C23'].value!r}")
        print(f"    C24={ws['C24'].value!r}")
        print(f"    C25={ws['C25'].value!r}; C26={ws['C26'].value!r}; C28={ws['C28'].value}")

    # 2. Сценарии: 02 «Контроль ответа», 04, затем 08 по значениям из 04.
    for sc, cfg in SC.items():
        d = work / sc
        (d / "fill").mkdir(parents=True)
        wb = openpyxl.load_workbook(KIT / "02-proverka-i-kontrol-otveta.xlsx")
        ws = wb["Контроль ответа"]
        for r, react in enumerate(cfg["ctrl"], start=5):
            ws.cell(row=r, column=2, value=D("20.07.2026"))
            ws.cell(row=r, column=3, value=f"событие {r - 4}")
            ws.cell(row=r, column=4, value=D("01.09.2026"))
            ws.cell(row=r, column=6, value=react)
        wb.save(d / "fill" / "02.xlsx")
        wb = openpyxl.load_workbook(KIT / "04-uchet-raschetov-i-otpravok.xlsx")
        v, g = wb["Взаиморасчёты"], wb["Долг по актам"]
        for i, (aid, doc, amt, signed, last) in enumerate(ACTS):
            r = 5 + i
            v.cell(row=r, column=2, value=D(signed)); v.cell(row=r, column=3, value="Начисление по акту")
            v.cell(row=r, column=4, value=doc); v.cell(row=r, column=5, value=amt); v.cell(row=r, column=6, value=aid)
            g.cell(row=r, column=2, value=aid); g.cell(row=r, column=3, value=doc); g.cell(row=r, column=4, value=D(last))
        wb.save(d / "fill" / "04.xlsx")
        calc([d / "fill" / "02.xlsx", d / "fill" / "04.xlsx"], d / "calc")
        k = openpyxl.load_workbook(d / "calc" / "02.xlsx", data_only=True)
        b = openpyxl.load_workbook(d / "calc" / "04.xlsx", data_only=True)
        print(f"\n## {sc}")
        print("02 «Контроль ответа», G:")
        for r, react in enumerate(cfg["ctrl"], start=5):
            print(f"  {react} → {k['Контроль ответа'].cell(row=r, column=7).value}")
        v, g = b["Взаиморасчёты"], b["Долг по актам"]
        j = [v.cell(row=r, column=10).value for r in range(5, 505)]
        print(f"04 «Взаиморасчёты»: I5:I7={[v.cell(row=r, column=9).value for r in (5, 6, 7)]}; J5:J7={j[:3]}; "
              f"непустых J в J8:J504 = {sum(1 for x in j[3:] if x not in (None, ''))}; M10={v['M10'].value}; M11={v['M11'].value}")
        rows = [[g.cell(row=r, column=c).value for c in (2, 4, 5, 8, 10, 11, 12)] for r in (5, 6, 7)]
        print(f"04 «Долг по актам» B,D,E,H,J,K,L: {rows}")
        print(f"ошибок «#…» в 02/04: {errors(k)}/{errors(b)}")
        wb = openpyxl.load_workbook(KIT / "08-raschet-procentov-395.xlsx")
        wb["Ввод"]["B4"] = D(cfg["calc"]); wb["Ввод"]["A12"] = D("01.01.2026"); wb["Ввод"]["B12"] = 20
        for n, i in enumerate(cfg["acts08"]):
            r = 5 + n
            aid, doc, _, _, _ = ACTS[i]
            a = wb["Акты"]
            a.cell(row=r, column=2, value=aid); a.cell(row=r, column=3, value=doc)
            a.cell(row=r, column=4, value=rows[i][3]); a.cell(row=r, column=5, value=rows[i][2]); a.cell(row=r, column=6, value=rows[i][4])
        wb.save(d / "fill" / "08.xlsx")
        calc([d / "fill" / "08.xlsx"], d / "calc8")
        p = openpyxl.load_workbook(d / "calc8" / "08.xlsx", data_only=True)
        akt = p["Акты"]
        print(f"08 «Акты» H: {[akt.cell(row=5 + n, column=8).value for n in range(len(cfg['acts08']))]}; "
              f"«Расчёт» D2={p['Расчёт']['D2'].value!r}; D3={p['Расчёт']['D3'].value!r}; ошибок «#…»: {errors(p)}")
        cfg["pct"] = p["Расчёт"]["D3"].value
        cfg["rows"] = rows

    # 3. Writer: 06 и 09 по С-1/С-3/С-4, 10 (иск) по С-1, 05 по С-3.
    print("\n## Writer: заполнение и PDF")
    for sc, cfg in SC.items():
        d = work / sc
        acts = [ACTS[i] for i in cfg["acts08"]]
        debt = sum(a[2] for a in acts)
        vals = {"СУММА_ДОЛГА": f"{debt:,}".replace(",", " "), "СУММА_ПРОЦЕНТОВ": f"{cfg['pct']:.2f}".replace(".", ","),
                "ДАТА_РАСЧЁТА": cfg["calc"], "ДАТА_ПО": cfg["calc"]}
        for f in ("06-uvedomlenie-o-prosrochke", "09-pretenziya") + (("10-obrashchenie-v-sud",) if sc == "С-1" else ()) \
                + (("05-peregovory-i-perenos-sroka",) if sc == "С-3" else ()):
            doc_ = docx.Document(KIT / f"{f}.docx")
            left = set()

            def fill(par):
                t = par.text
                if "{{" not in t:
                    return
                t2 = re.sub(r"\{\{([^}]+)\}\}", lambda m: vals.get(m.group(1), f"[ВЫМЫШЛ. {m.group(1)}]"), t)
                for r_ in par.runs[1:]:
                    r_.text = ""
                if par.runs:
                    par.runs[0].text = t2
            for par in doc_.paragraphs:
                fill(par)
            for t in doc_.tables:
                hdr = [c.text for c in t.rows[0].cells]
                if any("Последний день срока оплаты" in h or "Срок оплаты по договору" in h for h in hdr) and len(t.rows) > len(acts):
                    for n, a in enumerate(acts, start=1):
                        cells = t.rows[n].cells
                        for ci, h in enumerate(hdr):
                            if "Акт" in h or "Документ" in h:
                                cells[ci].text = a[1]
                            elif "Сумма" in h:
                                cells[ci].text = f"{a[2]:,}".replace(",", " ")
                            elif "Последний день" in h or "Срок оплаты" in h:
                                cells[ci].text = a[4]
                            elif "Подписан" in h or h.strip() == "Дата":
                                cells[ci].text = a[3]
                for row in t.rows:
                    for c in row.cells:
                        for par in c.paragraphs:
                            fill(par)
            for par in doc_.paragraphs:
                left |= set(re.findall(r"\{\{[^}]+\}\}", par.text))
            out = d / "fill" / f"{f}.docx"
            doc_.save(out)
            writer_pdf([out], d / "pdf")
            pdf = d / "pdf" / f"{f}.pdf"
            pages = re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1)
            txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
            print(f"- {sc} {f}: PDF {pages} стр.; «{{{{» в PDF: {txt.count('{{')}; "
                  f"даты актов в PDF: {[a[4] for a in acts if a[4] in txt]}; сумма процентов в PDF: {vals['СУММА_ПРОЦЕНТОВ'] in txt}")


main()
