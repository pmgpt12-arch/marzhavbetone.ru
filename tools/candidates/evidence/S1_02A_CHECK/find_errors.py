"""Где в пересчитанной книге стоят значения-ошибки и что там в исходнике."""
import sys
import warnings

import openpyxl

warnings.simplefilter("ignore")
calc = openpyxl.load_workbook(sys.argv[1], data_only=True)
src = openpyxl.load_workbook(sys.argv[1])
for ws in calc:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value[:4] in ("#REF", "#VAL", "#NAM", "#DIV", "#N/A", "#NUM", "#NUL"):
                print(ws.title, c.coordinate, repr(c.value), "| формула:", repr(src[ws.title][c.coordinate].value)[:200],
                      "| формат:", src[ws.title][c.coordinate].number_format)
