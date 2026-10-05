"""P5 factual literal inventory, exact buyers, no legal verdict or source mutation."""
import pathlib,subprocess,json,hashlib,re,zipfile,xml.etree.ElementTree as ET,datetime,collections,sys,shutil
S=pathlib.Path('/home/denis/projects/marzhavbetone.ru')
R=pathlib.Path('/home/denis/.local/state/claude-dispatcher/recovery-20261004/p5-current-citation-inventory-20261005')
PIN='29e137c2d8822e0a3b0b2cf80a6443d33b3726c3'
PACK='products-storage/07-uderzhaniya-shtrafy-zachety'
PROOF=pathlib.Path('/home/denis/.local/state/claude-dispatcher/recovery-20261004/p5-actual-php-zip-20261005/receipt.json')
FROZEN=pathlib.Path('/home/denis/.local/state/claude-dispatcher/recovery-20261004/s1-current-citation-inventory-20261005/extract.py')
def h(b):return hashlib.sha256(b).hexdigest()
def git(*a):return subprocess.check_output(['git','-C',str(S),*a],timeout=15)
def norm(t):return re.sub(r'\s+',' ',t).strip()
R.mkdir(exist_ok=False)
required={'normative_contract':('/home/denis/projects/ai-business-os/docs/Normative_Contract.md','cc7daf5d6515035576adb4854ce4523a4cd9b4b2ab95cebef63c4104d1fb99bb'),'task_protocol':('/home/denis/projects/ai-business-os/docs/Task_Setting_Protocol.md','337fe8864579a95dd4af53e44ad036016145a4be82aac5560e535a09fa56e686')}
assert h(FROZEN.read_bytes())=='e202c8502ae73d60b1ad07219770d440f76a665fa60d54de1ae20d75680415c0'
deps={}
for key,(p,digest) in required.items():
 raw=pathlib.Path(p).read_bytes();assert h(raw)==digest;deps[key]={'path':p,'sha256':digest,'bytes':len(raw)}
 shutil.copyfile(p,R/(key+'.md'))
proof=json.loads(PROOF.read_text()); assert proof['status']=='PASS_ACTUAL_PHP_BUYER_ZIP_BYTES'
members=proof['buyer_members'];assert len(members)==11
assert git('rev-parse',PIN).decode().strip()==PIN
inputs=[]
for m in members:
 rel=PACK+'/'+m['name'];raw=git('show',PIN+':'+rel);assert h(raw)==m['sha256']
 p=R/'exact-buyers'/m['name'];p.parent.mkdir(exist_ok=True);p.write_bytes(raw)
 inputs.append({'path':rel,'name':m['name'],'source_sha':PIN,'git_blob':git('rev-parse',PIN+':'+rel).decode().strip(),'sha256':h(raw),'size':len(raw),'local_path':str(p)})
passport={'execution_pattern':'one_shot','primary_result':'P5 exact-current factual citation inventory for lawyer handoff','feedback_loop_required':False,'checkpoint_policy':'verified_only','parent':'MB001 autonomous parallel preparation; no release/legal gate waived','source_sha':PIN,'scope':'Read-only all 11 actual buyer files; own isolated outputs only','inputs':inputs,'required_sources':deps,'reference_parser_sha256':h(FROZEN.read_bytes()),'reference_parser_reuse':'article/marker regex and locator schema approach only; no S1 old maps imported','dependencies':{'actual_php_receipt':str(PROOF),'receipt_sha256':h(PROOF.read_bytes()),'native04_05_exact':True},'executor':'Python3/openpyxl read_only/OOXML/pdftotext, API USD0, no LLM','route':['pin inputs and passport','extract exact paragraphs/cells/TXT lines/PDF physical pages','record raw literal spans and unresolved legal markers','selfcheck exact all locators and hashes','freeze READY for independent root coverage review'],'acceptance':'11/11 exact members; raw quote spans replay; no inferred whole-page code; ranges retained; no legal verdict or invented norm IDs','stop':'missing/mismatched exact input or tool; preserve logs and no source writes','outputs':['citation-inventory.json','source-text-units.json (server only)','input-hashes.json','verification-receipt.json','citation-inventory-report.md','SHA256SUMS'],'source_currentness':'NOT_VERIFIED','human_legal':'NOT_PERFORMED','api_usd':0}
(R/'passport.json').write_text(json.dumps(passport,ensure_ascii=False,indent=2))
sys.path.append('/home/denis/projects/ai-business-os/.venv/lib/python3.12/site-packages')
import openpyxl
# Retain exact approved grammar; unresolved markers expose scope beyond it.
codes=r'ГК|АПК|НК|ГПК|ТК|КоАП|БК|СК|ЗК|ЖК'
num=r'\d+(?:\.\d+)?'
prefix=r'(?:(?:пп?|пункт(?:а|ы|ов|у|ом)?|ч|част(?:ь|и|ью)|абз|абзац(?:а|ы|ев)?)\.?\s*'+num+r'\s*){0,4}'
artword=r'(?:ст\.|стать(?:я|и|ю|е|ей|ями))'
refpat=re.compile(r'(?<![А-Яа-яA-Za-z])'+prefix+artword+r'\s*'+num+r'(?:\s*(?:,|и|–|-)\s*'+num+r')*(?:\s*(?:'+codes+r')(?:\s*РФ)?)?',re.I)
lawpat=re.compile(r'\b(?:'+codes+r')\b|\bст\.\s*\d|\bстать(?:я|и|ю|е|ей|ями)\b|\d+\s*[-–]ФЗ\b|Пленум|Верховн\w*\s+[Сс]уд\w*|[Фф]едеральн\w*\s+закон\w*|[Пп]риказ\w*|[Пп]остановлен\w*|\bзакон\w*|\bГОСТ\b|\bСП\s*\d|\bСНиП\b|\bФСБУ\b',re.I)
codepat=re.compile(r'\b('+codes+r')\b',re.I)
canonical={c.lower():c for c in ['ГК','АПК','НК','ГПК','ТК','КоАП','БК','СК','ЗК','ЖК']}
WX='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
units=[];docs=[]
for item in inputs:
 p=pathlib.Path(item['local_path']);start=len(units);scope=[]
 if p.suffix=='.docx':
  with zipfile.ZipFile(p) as z:
   parts=[n for n in z.namelist() if n=='word/document.xml' or re.fullmatch(r'word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml',n)]
   for part in parts:
    root=ET.fromstring(z.read(part));parent={ch:node for node in root.iter() for ch in node};nonempty=0
    for ordinal,par in enumerate(root.iter(WX+'p'),1):
     text=''.join(n.text or '' for n in par.iter(WX+'t'))
     if not text.strip():continue
     nonempty+=1;lineage=[];node=par
     while node in parent:
      up=parent[node];siblings=[ch for ch in up if ch.tag==node.tag];lineage.append(node.tag.split('}')[-1]+'['+str(siblings.index(node)+1)+']');node=up
     units.append({'document':item['path'],'location':{'kind':'docx_paragraph','part':part,'global_paragraph_ordinal_including_empty':ordinal,'nonempty_paragraph_ordinal':nonempty,'xml_path':'/'+'/'.join(reversed(lineage))},'text':text})
   scope={'xml_parts':parts,'limitations':'w:t text only; drawing text/field codes/images are outside extraction'}
 elif p.suffix=='.xlsx':
  book=openpyxl.load_workbook(p,read_only=True,data_only=False)
  sheets=[]
  for sheet in book.worksheets:
   sheets.append({'name':sheet.title,'max_row':sheet.max_row,'max_column':sheet.max_column})
   for row in sheet.iter_rows():
    for cell in row:
     if cell.value is not None:units.append({'document':item['path'],'location':{'kind':'xlsx_cell','sheet':sheet.title,'cell':cell.coordinate,'data_type':cell.data_type,'value_mode':'formula' if cell.data_type=='f' else 'stored_value'},'text':str(cell.value)})
  book.close();scope={'sheets':sheets,'limitations':'stored cells/formulas only; drawings/comments/printed display outside extraction; no workbook saved'}
 elif p.suffix=='.pdf':
  info=subprocess.check_output(['pdfinfo',str(p)],text=True); count=int(re.search(r'^Pages:\s+(\d+)',info,re.M).group(1))
  rawtext=subprocess.check_output(['pdftotext','-layout','-enc','UTF-8',str(p),'-']);text=rawtext.decode('utf-8')
  textpath=R/(p.stem+'.pdftotext.txt');textpath.write_bytes(rawtext);pages=text.split('\f')
  if not pages[-1].strip():pages.pop()
  assert len(pages)==count,(len(pages),count)
  pagestats=[]
  for ordinal,page in enumerate(pages,1):
   units.append({'document':item['path'],'location':{'kind':'pdf_page_text','physical_page':ordinal,'extraction':'pdftotext -layout -enc UTF-8','offset_basis':'Unicode characters in exact page text, excludes form-feed separator'},'text':page})
   pagestats.append({'physical_page':ordinal,'chars':len(page),'nonempty':bool(page.strip()),'empty_or_image_only_text_scope':not bool(page.strip()),'lines':len(page.splitlines())})
  scope={'pages':pagestats,'pdftotext_path':str(textpath),'pdftotext_sha256':h(rawtext),'limitations':'text layer only; image content/layout/reading order not legal-reviewed; empty pages explicitly flagged'}
 elif p.suffix=='.txt':
  nonempty=0
  for line_no,text in enumerate(p.read_text().splitlines(),1):
   if not text.strip():continue
   nonempty+=1;units.append({'document':item['path'],'location':{'kind':'txt_line','physical_line':line_no,'nonempty_line':nonempty},'text':text})
  scope={'limitations':'UTF-8 physical lines, no display interpretation'}
 else:raise AssertionError(p.suffix)
 docs.append(dict(item,text_units=len(units)-start,extraction_scope=scope))
occurrences=[];markers=[]
def locate(unit,a,b):
 loc=dict(unit['location'])
 if loc['kind']=='pdf_page_text':
  text=unit['text'];ls=text.rfind('\n',0,a)+1;end=text.find('\n',b);end=len(text) if end<0 else end
  loc.update({'page_char_span':[a,b],'physical_line_start':text.count('\n',0,a)+1,'physical_line_end':text.count('\n',0,max(a,b-1))+1,'line_start_char_offset':ls,'column_start_0based':a-ls,'line_context_char_span':[ls,end]})
 return loc
for unit in units:
 text=unit['text'];matches=list(refpat.finditer(text))
 for m in matches:
  literal=m.group();explicit=codepat.findall(literal)
  if unit['location']['kind']=='pdf_page_text':
   ls=text.rfind('\n',0,m.start())+1;le=text.find('\n',m.end());le=len(text) if le<0 else le
   context=text[ls:le] if '\n' not in text[m.start():m.end()] else literal
   context_rule='one_explicit_code_in_same_physical_pdf_line'
  else:context=text;context_rule='one_explicit_code_in_same_text_unit'
  contextcodes=sorted(set(canonical[c.lower()] for c in codepat.findall(context)))
  code=canonical[explicit[-1].lower()] if explicit else contextcodes[0] if len(contextcodes)==1 else None
  binding='explicit_literal' if explicit else context_rule if code else 'UNRESOLVED_NO_OR_MULTIPLE_EXPLICIT_CODE'
  tail=re.split(artword,literal,maxsplit=1,flags=re.I)[-1];articles=re.findall(num,re.split(codepat,tail,maxsplit=1)[0])
  a=max(0,m.start()-70);b=min(len(text),m.end()+100)
  occurrences.append({'occurrence_id':'ref-'+str(len(occurrences)+1).zfill(4),'document':unit['document'],'file_sha256':next(d['sha256'] for d in docs if d['path']==unit['document']),'location':locate(unit,m.start(),m.end()),'unit_location':unit['location'],'literal':literal,'span':[m.start(),m.end()],'short_quote':text[a:b],'quote_span':[a,b],'code':code,'code_binding':binding,'articles_as_written':articles,'range_or_list_syntax':bool(re.search(r'\d\s*(?:[–-]|,|и)\s*\d',tail)),'registry_norm_id':None,'normative_verdict':'NOT_VERIFIED'})
 for m in lawpat.finditer(text):
  if any(x.start()<=m.start() and m.end()<=x.end() for x in matches):continue
  a=max(0,m.start()-70);b=min(len(text),m.end()+100)
  markers.append({'marker_id':'marker-'+str(len(markers)+1).zfill(4),'document':unit['document'],'location':locate(unit,m.start(),m.end()),'unit_location':unit['location'],'marker':m.group(),'span':[m.start(),m.end()],'short_quote':text[a:b],'quote_span':[a,b],'classification':'UNRESOLVED_OR_NON_ARTICLE_LEGAL_MARKER','reason':'law token outside article grammar; code-only/general language is not a defect finding','registry_norm_id':None,'normative_verdict':'NOT_VERIFIED'})
groups=collections.OrderedDict()
for o in occurrences:
 key=(norm(o['literal']).lower(),o['code'],tuple(o['articles_as_written']))
 groups.setdefault(key,{'literal_normalized':norm(o['literal']),'code':o['code'],'articles_as_written':o['articles_as_written'],'occurrence_ids':[]})['occurrence_ids'].append(o['occurrence_id'])
dedup=[dict(group_id='group-'+str(i).zfill(3),**g) for i,g in enumerate(groups.values(),1)]
idx={(u['document'],json.dumps(u['location'],sort_keys=True)):u['text'] for u in units}
for row in occurrences+markers:
 text=idx[(row['document'],json.dumps(row['unit_location'],sort_keys=True))];a,b=row['span'];assert text[a:b]==row.get('literal',row.get('marker'));a,b=row['quote_span'];assert text[a:b]==row['short_quote']
assert sorted(x for g in dedup for x in g['occurrence_ids'])==sorted(o['occurrence_id'] for o in occurrences)
for d in docs:assert h(pathlib.Path(d['local_path']).read_bytes())==d['sha256']==h(git('show',PIN+':'+d['path']))
inventory={'observed':datetime.datetime.now(datetime.timezone.utc).isoformat(),'source_sha':PIN,'documents':docs,'parser':{'ref_regex':refpat.pattern,'law_marker_regex':lawpat.pattern,'code_binding':'explicit literal or unique same text unit; PDF never whole-page code inference, only single physical line; ambiguity remains unresolved','article_list_rule':'numbers as written; ranges never expanded','dedup_key':'normalized literal + code + numbers as written; occurrence IDs retained','scope':'factual literal citation inventory, not legal-mechanism completeness or a normative registry'},'occurrences':occurrences,'deduplicated_groups':dedup,'unresolved_legal_markers':markers,'source_currentness':'NOT_VERIFIED','human_legal':'NOT_PERFORMED','registry_norm_ids_assigned':False}
def write(name,data): (R/name).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
write('input-hashes.json',{'source_sha':PIN,'buyer_files':inputs,'required_inputs':deps,'actual_php_receipt_sha256':h(PROOF.read_bytes()),'reference_parser_sha256':h(FROZEN.read_bytes())})
write('citation-inventory.json',inventory)
write('source-text-units.json',{'source_sha':PIN,'units':units})
citationunits=[u for u in units if list(refpat.finditer(u['text'])) or list(lawpat.finditer(u['text']))]
write('citation-text-units.json',{'source_sha':PIN,'units':citationunits})
receipt={'status':'VERIFIED_FACTUAL_INVENTORY_PENDING_INDEPENDENT_ROOT_REVIEW','source_sha':PIN,'buyer_files':len(docs),'text_units':len(units),'formula_units':sum(u['location'].get('value_mode')=='formula' for u in units),'citation_text_units':len(citationunits),'literal_occurrences':len(occurrences),'dedup_groups':len(dedup),'unresolved_markers':len(markers),'all_span_quote_locator_checks':True,'all11_current_gitbytes_unchanged':True,'native04_05_exact_inherited_byte_binding':True,'source_currentness':'NOT_VERIFIED','human_legal':'NOT_PERFORMED','model_calls':0,'api_usd':0,'executor':{'python':sys.version,'openpyxl':openpyxl.__version__},'server_only_large_unit_manifest':{'path':str(R/'source-text-units.json'),'sha256':h((R/'source-text-units.json').read_bytes()),'size':(R/'source-text-units.json').stat().st_size}}
write('verification-receipt.json',receipt)
counts=collections.Counter(o['document'] for o in occurrences);mc=collections.Counter(o['document'] for o in markers)
lines=['# P5 current factual citation inventory','','Exact accepted source: '+PIN+'. Preparation for lawyer handoff only. Currentness NOT_VERIFIED; human legal NOT_PERFORMED.','','All 11 inputs are exact actual-PHP buyer membership; 04/05 retain approved native bytes. Formula-bearing books read only; no saves.','','| Buyer file | SHA256 | Literal occurrences | Unresolved markers |','|---|---|---|---|']
for d in docs:lines.append('| '+d['name']+' | '+d['sha256']+' | '+str(counts[d['path']])+' | '+str(mc[d['path']])+' |')
lines.extend(['','Groups represent literal candidates, not verified norms. Legal markers include general words and are not defect findings. Lists and ranges remain as written. Code binding is explicit or unique within paragraph/cell/TXT line; PDF only explicit literal or same physical line, never whole-page inference. PDF references retain exact physical page, character spans and physical line mapping; text-layer absence and extraction limitations are explicit. All registry IDs are null.','','Exact occurrences and raw quotes are in citation-inventory.json, compact original text in citation-text-units.json; full cell/formula corpus stays server-side, pinned by verification-receipt.json. No S1 old maps, official normative sources, legal registry, financial verdict, publication or live price changes.','','Root independent coverage review required before any commit/push.'])
(R/'citation-inventory-report.md').write_text('\n'.join(lines)+'\n')
shutil.copyfile(__file__,R/'extract.py')
write('READY.json',{'status':'READY_FOR_INDEPENDENT_ROOT_COVERAGE_REVIEW','source_sha':PIN,'script_sha256':h((R/'extract.py').read_bytes()),'map_sha256':h((R/'citation-inventory.json').read_bytes()),'compact_units_sha256':h((R/'citation-text-units.json').read_bytes()),'receipt_sha256':h((R/'verification-receipt.json').read_bytes()),'currentness':'NOT_VERIFIED','human_legal':'NOT_PERFORMED'})
(R/'SHA256SUMS').write_text('\n'.join(h(p.read_bytes())+'  '+p.name for p in sorted(R.iterdir()) if p.is_file() and p.name!='SHA256SUMS')+'\n')
print(json.dumps(receipt,ensure_ascii=False));print((R/'READY.json').read_text())
