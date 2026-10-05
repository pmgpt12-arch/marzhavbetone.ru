"""Покупательские фикстуры 03 по новой инструкции. Стороны и данные — из
предыдущей приёмки (previous-buyer-fill.py), вымышленные. Абзацы НЕ
переписываются целиком: заполняются только поля [..], ненужные абзацы
удаляются (раздел 2 шаг 1; раздел 3 шаг 6; раздел 6)."""
import re
from pathlib import Path

import docx

RUN = Path(__file__).parent
SRC = Path("/home/denis/projects/marzhavbetone.ru/.worktrees/codex-correction-388-20261005"
           "/products-storage/02-dopraboty-bez-poter/03-uvedomlenie-o-doprabotah.docx")
OUT = RUN / "filled"
OUT.mkdir(exist_ok=True)

SUB = "ООО «СтройМонолит-Тест» (вымышленная организация), ИНН 7700000001"
CUS = "АО «ГенСтрой Пример» (вымышленная организация), ИНН 7700000002"
DOG = "договор подряда № ТП-17/2026 от 12.03.2026, объект «ЖК Тестовый квартал, корпус 2»"
SUB_SIGN = "Генеральный директор ООО «СтройМонолит-Тест» И. И. Иванов"


def find(doc, start):
    ps = [p for p in doc.paragraphs if p.text.startswith(start)]
    assert len(ps) == 1, (start, len(ps))
    return ps[0]


def fill(doc, start, *values):
    """Заменяет поля [..] абзаца по порядку; остальной текст — как в форме."""
    p = find(doc, start)
    text = p.text
    for v in values:
        text, n = re.subn(r"\[[^\]]*\]", lambda m: v, text, count=1)
        assert n == 1, (start, v)
    assert not re.search(r"\[[^\]]*\]", text), text
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""
    return text


def drop(doc, start):
    p = find(doc, start)
    p._element.getparent().remove(p._element)


def head(doc, lines):
    for line in lines:
        doc.paragraphs[0].insert_paragraph_before(line)


def sign(doc, date):
    fill(doc, "Дата:", date)
    fill(doc, "Подпись:", "(личная подпись)")
    fill(doc, "Расшифровка подписи:", SUB_SIGN)


def save(doc, name):
    doc.save(OUT / name)
    rest = [p.text for p in doc.paragraphs if re.search(r"\[[^\]]*\]|/", p.text)]
    print(name, "остатки:", rest or 0)


# ---- A: раздел 2, шаг 1
d = docx.Document(SRC)
head(d, ["Исх. № 112 от 14.09.2026", f"Кому: {CUS}", f"От: {SUB}", f"Основание: {DOG}"])
fill(d, "Дополнительная работа выявлена", "11.09.2026",
     "вскрытии перекрытия на отм. +6,300 в осях 3–5/В–Г: требуется усиление узла "
     "опирания плиты, не предусмотренное рабочей документацией шифр ТК-2-КЖ")
drop(d, "Дополнительная работа выполнена")
fill(d, "До [дата]", "18.09.2026")
# ровно один вариант режима: второй вариант и косая черта удалены
p = find(d, "Работа не начинается")
p.runs[0].text = p.text.split(" / ")[0] + "."
for r in p.runs[1:]:
    r.text = ""
drop(d, "Приложения:")  # на шаге 1 ничего не прикладывается
sign(d, "14.09.2026")
save(d, "A-03-uvedomlenie-o-doprabotah.docx")

# ---- B: раздел 3, шаг 6
d = docx.Document(SRC)
head(d, ["Исх. № 131 от 05.10.2026", f"Кому: {CUS}", f"От: {SUB}", f"Основание: {DOG}"])
drop(d, "Дополнительная работа выявлена")
fill(d, "Дополнительная работа выполнена", "08.08.2026", "14.08.2026",
     "начальника участка АО «ГенСтрой Пример» С. С. Сидорова",
     "устройство приямка 1,2×1,2×0,9 м в осях 7/Д",
     "протокол совещания № 31 от 07.08.2026, общий журнал работ (записи "
     "08.08–14.08.2026), фото от 13.08.2026")
fill(d, "До [дата]", "15.10.2026")
drop(d, "Режим до решения")
drop(d, "Работа не начинается")
fill(d, "Приложения:", "описание и согласование объёма (на 2 л.); расчёт стоимости "
     "58 900,00 руб. без НДС (на 1 л.)")
sign(d, "05.10.2026")
save(d, "B-03-uvedomlenie-o-doprabotah.docx")
