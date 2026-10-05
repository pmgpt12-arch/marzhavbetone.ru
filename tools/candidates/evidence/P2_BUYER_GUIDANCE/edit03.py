"""Правка формы 03 (N1): строка случая Б и строка «Приложения».

Существующий текст не меняется. Новые абзацы — клоны соседних (тот же стиль).
"""
import copy
import sys
from pathlib import Path

import docx

SRC = Path(sys.argv[1])
d = docx.Document(SRC)
P = d.paragraphs
assert P[2].text == "Дополнительная работа выявлена [дата] при [обстоятельства]."
assert P[7].text == "Дата: [___]" and P[6].text == ""

b = copy.deepcopy(P[2]._element)
P[2]._element.addnext(b)
bp = docx.text.paragraph.Paragraph(b, P[2]._parent)
bp.runs[0].text = ("Дополнительная работа выполнена с [дата] по [дата] по просьбе "
                   "[кто просил]: [что выполнено]. Подтверждения: [чем подтверждено].")
for r in bp.runs[1:]:
    r.text = ""

a = copy.deepcopy(P[7]._element)
P[6]._element.addprevious(a)
ap = docx.text.paragraph.Paragraph(a, P[7]._parent)
ap.runs[0].text = "Приложения: [___]"
for r in ap.runs[1:]:
    r.text = ""

d.save(SRC)
for i, p in enumerate(docx.Document(SRC).paragraphs):
    print(i, p.style.name, repr(p.text))
