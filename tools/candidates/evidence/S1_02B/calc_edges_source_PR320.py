"""Независимая проверка книги 14 S1 (#286) в LibreOffice Calc: крайние значения,
ошибки в ячейках, исходный дефект T1 на том же вводе."""
import sys, shutil, tempfile, subprocess
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import openpyxl

WT = Path(sys.argv[1]); SITE = Path(sys.argv[2])
C = WT / "tools/candidates/s1-oplata-za-raboty"
sys.path.insert(0, str(WT / "tools"))
import test_s1_candidate as T  # заполнитель и пересчёт из самой ветки

def yd(d): return 366 if (d.year % 4 == 0 and (d.year % 100 or d.year % 400 == 0)) else 365

def ref(acts, end, rates, payments, excluded=()):
    """Эталон, написанный заново: по каждому акту подневно, периоды — непрерывные
    дни с одинаковыми (долг, ставка, дней в году, исключение), округление периода."""
    tot = Decimal(0)
    for aid, debt, start in acts:
        pays = [(d, Decimal(str(s))) for d, s, a in payments if a == aid]
        per = []  # [key, days]
        d = start
        while d <= end:
            db = Decimal(str(debt)) - sum((s for pd, s in pays if pd < d), Decimal(0))
            r = Decimal(str([r for rd, r in sorted(rates) if rd <= d][-1]))
            ex = any(a <= d <= b for a, b in excluded)
            key = (db, r, yd(d), ex)
            if per and per[-1][0] == key: per[-1][1] += 1
            else: per.append([key, 1])
            d += timedelta(days=1)
        for (db, r, y, ex), n in per:
            if not ex:
                tot += (db * r / 100 * n / y).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return tot

tmp = Path(tempfile.mkdtemp(prefix="edge-"))
B = C / "14-raschet-procentov-395.xlsx"
R20 = [(date(2014, 1, 1), 20)]
cases = {}   # имя -> (книга, ожидание: Decimal | "БЛОК")
def add(name, acts, end, rates, payments=(), excluded=(), expect=None, extra=None):
    p = T.заполнить_395(B, tmp / f"{name}.xlsx", acts, end, rates, payments, excluded, extra=extra)
    cases[name] = (p, expect if expect is not None else ref(acts, end, rates, payments, excluded))

# 0. исходный дефект T1: одна строка периода с оплатой внутри
add("T1_ввод", [("A", 1_000_000, date(2025, 1, 10))], date(2025, 3, 31), [(date(2024, 10, 28), 21)],
    [(date(2025, 2, 14), 300_000, "A")], expect=Decimal("38835.62"))
# 1. полная оплата в первый день просрочки
add("полная_в_1й_день", [("A", 100_000, date(2026, 1, 1))], date(2026, 3, 31), R20, [(date(2026, 1, 1), 100_000, "A")])
# 2. оплата в последний день расчёта
add("оплата_в_последний_день", [("A", 100_000, date(2026, 1, 1))], date(2026, 3, 31), R20, [(date(2026, 3, 31), 40_000, "A")])
# 3. високосный год + смена ставки в день оплаты
add("високос_смена_ставки_в_день_оплаты", [("A", 777_777.77, date(2023, 12, 15))], date(2024, 3, 10),
    [(date(2023, 10, 30), 15), (date(2023, 12, 18), 16)], [(date(2023, 12, 18), 50_000, "A")])
# 4. один день
add("один_день", [("A", 1_000_000, date(2026, 5, 5))], date(2026, 5, 5), R20)
# 5. конец раньше начала
add("конец_раньше_начала", [("A", 1_000_000, date(2026, 5, 5))], date(2026, 5, 4), R20, expect="БЛОК")
# 6. оплаты больше долга
add("оплаты_больше_долга", [("A", 100_000, date(2026, 1, 1))], date(2026, 3, 31), R20,
    [(date(2026, 1, 10), 60_000, "A"), (date(2026, 2, 10), 60_000, "A")], expect="БЛОК")
# 7. длинный срок: предел «По дням» и за пределом
s = date(2015, 1, 1)
add("предел_дней_3999", [("A", 1_000_000, s)], s + timedelta(days=3998), R20)
add("за_пределом_4000", [("A", 1_000_000, s)], s + timedelta(days=3999), R20, expect="БЛОК")
# 8. 12 актов, разные даты, у каждого оплата, несколько смен ставки, мораторий
acts = [(f"A{i}", 100_000 + 12_345.67 * i, date(2024, 1, 1) + timedelta(days=17 * i)) for i in range(12)]
pays = [(a[2] + timedelta(days=30 + i), 10_000 + i, a[0]) for i, a in enumerate(acts)]
rates = [(date(2023, 12, 18), 16), (date(2024, 7, 29), 18), (date(2024, 9, 16), 19), (date(2024, 10, 28), 21), (date(2025, 6, 9), 20)]
add("12_актов", acts, date(2025, 9, 30), rates, pays, excluded=[(date(2024, 4, 1), date(2024, 4, 30))])
# 9. ставка не покрывает начало
add("ставка_позже_начала", [("A", 100_000, date(2026, 1, 1))], date(2026, 3, 31), [(date(2026, 2, 1), 20)], expect="БЛОК")
# 10. пустая книга
cases["пустая"] = (shutil.copy(B, tmp / "пустая.xlsx"), "НЕ_ГОТОВ")
# 11. 500 оплат по одному акту
p500 = [(date(2024, 1, 1) + timedelta(days=i), 100, "A") for i in range(500)]
add("500_оплат", [("A", 1_000_000, date(2024, 1, 1))], date(2025, 6, 30), R20, p500)
# 12. оплата после последнего дня
add("оплата_после_конца", [("A", 100_000, date(2026, 1, 1))], date(2026, 3, 31), R20, [(date(2026, 4, 2), 1_000, "A")], expect="БЛОК")
# 13. копеечный долг
add("копейки", [("A", 0.01, date(2026, 1, 1))], date(2026, 12, 31), R20)

# исходная книга T1 (действующая выдача) на вводе T1_ввод, как её заполняет покупатель: один период
wb = openpyxl.load_workbook(SITE / "products-storage/09-pervyy-shag-pri-neoplate/04-raschet-procentov-395-gk.xlsx")
v, r = wb["Входные данные"], wb["Расчёт"]
v["B4"], v["B5"], v["B6"] = 1_000_000, date(2025, 1, 10), date(2025, 3, 31)
v["A12"], v["B12"] = date(2024, 10, 28), 21
v["A35"], v["B35"] = date(2025, 2, 14), 300_000
r["A4"], r["B4"] = date(2025, 1, 10), date(2025, 3, 31)
wb.save(tmp / "t1_old.xlsx")

books = {k: v[0] for k, v in cases.items()}; books["t1_old"] = tmp / "t1_old.xlsx"
out = T.пересчитать(books, tmp)
bad = 0
for k, (p, exp) in cases.items():
    wb = openpyxl.load_workbook(out[k], data_only=True)
    st, val = wb["Расчёт"]["D2"].value, wb["Расчёт"]["D3"].value
    errs = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
    if exp == "БЛОК": ok = st == "ЗАБЛОКИРОВАН"
    elif exp == "НЕ_ГОТОВ": ok = st != "ГОТОВ"
    else: ok = st == "ГОТОВ" and Decimal(str(val)).quantize(Decimal("0.01")) == exp
    ok = ok and not errs
    bad += not ok
    print(f"{'OK ' if ok else 'BAD'} {k:36} статус={st!s:14} итог={val!s:>14}  ожидалось={exp}  ошибок={len(errs)} {errs[:3]}")
wb = openpyxl.load_workbook(out["t1_old"], data_only=True)
print("T1 действующая книга, тот же ввод: проценты =", wb["Итог для документа"]["B10"].value,
      " проверка строки:", wb["Расчёт"]["H4"].value)
shutil.rmtree(tmp, ignore_errors=True)
sys.exit(1 if bad else 0)
