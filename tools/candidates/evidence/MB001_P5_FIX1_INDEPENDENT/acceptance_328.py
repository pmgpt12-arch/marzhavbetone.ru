#!/usr/bin/env python3
"""MB001-P5-CHECK-2 (#328): независимая приёмка исправлений #323.

Воспроизводимое доказательство к отчёту
tools/candidates/MB001_P5_FIX1_INDEPENDENT_ACCEPTANCE.md.

Что делает:
1. Собирает штатный ZIP P5 функцией выдачи mvb_build_product_zip('p5')
   (через харнесс tools/test_p5_documents.архив_выдачи) и сверяет хеши всех
   11 файлов с таблицами #322 §2 и #323 §6.
2. Пересобирает генератор во временной копии и сверяет 04 и 05 с ZIP.
3. Вписывает в копии 04 и 05 собственное вымышленное дело (числа не из
   тестов #323), пересчитывает Calc (отдельный профиль LibreOffice),
   читает значения из сохранённого Calc файла и сравнивает с ручным расчётом.
4. Печатает заполненные 04, 05 и прежнюю 06 в PDF, считает страницы.
5. Заполняет 01 и 11 в Writer, печатает в PDF; выписывает места текста
   о границах (гарантийное удержание, недостатки, штраф, зачёт, встречные
   требования).

Продукт и генератор не меняются. Все копии — во временном каталоге,
путь печатается. Запуск: python3 tools/candidates/evidence/MB001_P5_FIX1_INDEPENDENT/acceptance_328.py [каталог]
"""
from __future__ import annotations

import datetime as dt
import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

КОРЕНЬ = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(КОРЕНЬ / "tools"))

import openpyxl  # noqa: E402
import test_p5_documents as харнесс  # noqa: E402

# #322 §2 (голова #317 bde8b79) и #323 §6 (голова 5fed59c = 98165e5 по продукту)
ХЕШИ_322 = {
    "00-START-HERE.txt": "faa353564bd4b15a64757c7a4ae1ce596000abd60a5214917dddf2eb3161a64a",
    "01-pretenziya-na-uderzhanie.docx": "8011d569caabfb13ee0fcdd89ed53f84e935bf5f60f7b81422d668f1afd199fd",
    "02-pretenziya-na-shtraf.docx": "932c9863b9e842c9324c1e597953ea1113d92c61e23ec91b55695c85cd994596",
    "03-vozrazhenie-na-zachet.docx": "b9739bf9910ebaff28b638cc3d75b1d5ab45d6820b0605970788ac4cd520e866",
    "04-raschet-ubytkov.xlsx": "7c636244e52f5e2a24407295963edde8e1ede8a9d596de0916584839d3bd2ca8",
    "05-reestr-uderzhaniy.xlsx": "1434fecf40e9bd2788990b49177f823760a651cd518b2b3ea15c6644e3f4183e",
    "06-grafik-vozmeshcheniya.xlsx": "f28f7502a5e852045012949f292ea4583f5026945730b02b08fcd32e2c33f4bb",
    "07-algoritm-proverki-uderzhaniy.pdf": "59aa358f436f4e975cd02c5e6432b3ea93f14f3dee09bd97ce6987a66fa91ee7",
    "08-tipovye-osnovaniya.docx": "a2281f6cc825846152271f79e2734b485f9295a1d5fd0a7f2bd3a730ce45f179",
    "10-konsultaciya-po-delu.docx": "b7be831e7080428a9c6d1aced10a2439a976976a6737647baa17916fcce3852b",
    "11-trebovanie-o-vozvrate-uderzhaniya.docx": "ff6902613525ebdea12147ca1efb9f9b5446f7dc32244d0c6d2d7af86a2f494f",
}
ХЕШИ_323 = dict(ХЕШИ_322)
ХЕШИ_323["04-raschet-ubytkov.xlsx"] = "8ae9b70c48bd839986b0f38d467a0355eda3904a07089b11260c1ef53ca4b151"
ХЕШИ_323["05-reestr-uderzhaniy.xlsx"] = "1583baacc26dacc664db9d7bbb0a8ce8ea00a763341f1d93590e13c3332f12df"

ИТОГ: list[tuple[str, bool, str]] = []


def отметить(имя: str, ок: bool, подробно: str = "") -> None:
    ИТОГ.append((имя, ок, подробно))
    print(f"{'ок    ' if ок else 'ПРОВАЛ'}  {имя}" + (f" — {подробно}" if подробно else ""))


def soffice(tmp: Path, *args: str) -> None:
    профиль = tmp / "lo-profile-328"
    cmd = ["soffice", f"-env:UserInstallation=file://{профиль}", "--headless",
           "--norestore", *args]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=240)
    if r.returncode != 0:
        raise RuntimeError(f"СРЕДА: soffice {args}: {r.returncode} {r.stderr}")


def calc_пересчитать(src: Path, out: Path, tmp: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    soffice(tmp, "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src))
    return out / src.name


def в_pdf(src: Path, out: Path, tmp: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    soffice(tmp, "--convert-to", "pdf", "--outdir", str(out), str(src))
    return out / (src.stem + ".pdf")


def pdf_сведения(p: Path) -> tuple[int, str]:
    r = subprocess.run(["pdfinfo", str(p)], capture_output=True, text=True, check=True)
    стр = int(re.search(r"Pages:\s+(\d+)", r.stdout).group(1))
    размер = re.search(r"Page size:\s+(.*)", r.stdout).group(1).strip()
    return стр, размер


def pdf_текст(p: Path) -> str:
    """Текст PDF; разделители разрядов приведены к виду «14 647,50».

    Профиль LibreOffice без русской локали печатает «14,647.50»."""
    т = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True,
                       text=True, check=True).stdout
    return re.sub(r"(\d),(\d{3})", r"\1 \2", re.sub(r"(\d),(\d{3})", r"\1 \2", т)).replace(".", ",")


def sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def близко(a, b, eps=0.005) -> bool:
    return isinstance(a, (int, float)) and abs(a - b) <= eps


# ---------- 1. ZIP ----------
def шаг_zip(tmp: Path) -> dict[str, bytes]:
    архив = харнесс.архив_выдачи()
    (tmp / "zip").mkdir()
    print("\n## ZIP mvb_build_product_zip('p5')")
    for имя in sorted(архив):
        (tmp / "zip" / имя).write_bytes(архив[имя])
        h = sha(архив[имя])
        print(f"{h}  {len(архив[имя]):>7}  {имя}  "
              f"#323:{'=' if ХЕШИ_323.get(имя) == h else '≠'}  #322:{'=' if ХЕШИ_322.get(имя) == h else '≠'}")
    отметить("ZIP: 11 файлов, имена как в #322/#323", sorted(архив) == sorted(ХЕШИ_323),
             f"{len(архив)} файлов")
    отметить("ZIP: все хеши = #323 §6 (5fed59c)",
             all(sha(архив.get(n, b"")) == h for n, h in ХЕШИ_323.items()))
    изменены = sorted(n for n, h in ХЕШИ_322.items() if sha(архив.get(n, b"")) != h)
    отметить("ZIP: против #322 §2 изменены только 04 и 05",
             изменены == ["04-raschet-ubytkov.xlsx", "05-reestr-uderzhaniy.xlsx"], ", ".join(изменены))
    return архив


# ---------- 2. Воспроизводимость сборки ----------
def шаг_пересборка(tmp: Path, архив: dict[str, bytes]) -> None:
    копия = tmp / "gen" / "products-storage"
    (копия / "07-uderzhaniya-shtrafy-zachety").mkdir(parents=True)
    shutil.copy2(КОРЕНЬ / "products-storage" / "build_paid_07.py", копия)
    # reportlab нужен генератору только для 07.pdf; если его нет, подставляется
    # заглушка: 04 и 05 строит openpyxl, на их байты она не влияет. 07 при
    # этом не сверяется.
    заглушка = (
        "import sys, types\n"
        "try:\n    import reportlab\nexcept ImportError:\n"
        "    for n in ('reportlab','reportlab.pdfgen','reportlab.pdfgen.canvas','reportlab.lib',\n"
        "              'reportlab.lib.pagesizes','reportlab.lib.units','reportlab.pdfbase',\n"
        "              'reportlab.pdfbase.pdfmetrics','reportlab.pdfbase.ttfonts'):\n"
        "        sys.modules[n] = types.ModuleType(n)\n"
        "    class _C:\n"
        "        def __init__(s, *a, **k): pass\n"
        "        def __getattr__(s, n): return lambda *a, **k: None\n"
        "    m = sys.modules\n"
        "    m['reportlab.pdfgen'].canvas = m['reportlab.pdfgen.canvas']; m['reportlab.pdfgen.canvas'].Canvas = _C\n"
        "    m['reportlab.lib.pagesizes'].A4 = (595.27, 841.89); m['reportlab.lib.units'].cm = 28.35\n"
        "    m['reportlab.pdfbase.pdfmetrics'].getRegisteredFontNames = lambda: []\n"
        "    m['reportlab.pdfbase.pdfmetrics'].registerFont = lambda *a: None\n"
        "    m['reportlab.pdfbase.pdfmetrics'].stringWidth = lambda s, f, k: len(s) * k * 0.5\n"
        "    m['reportlab.pdfbase.ttfonts'].TTFont = lambda *a: None\n"
        "    print('ЗАГЛУШКА reportlab: 07.pdf не строится')\n"
        "import runpy; runpy.run_path('build_paid_07.py', run_name='__main__')\n")
    (копия / "_run.py").write_text(заглушка, encoding="utf-8")
    r = subprocess.run([sys.executable, "_run.py"], cwd=копия,
                       capture_output=True, text=True, timeout=300)
    print("        " + r.stdout.splitlines()[0] if r.stdout else "")
    if r.returncode != 0:
        отметить("пересборка генератора во временной копии", False, r.stderr[-300:])
        return
    import io
    import zipfile
    for имя in ("04-raschet-ubytkov.xlsx", "05-reestr-uderzhaniy.xlsx"):
        новый = (копия / "07-uderzhaniya-shtrafy-zachety" / имя).read_bytes()
        if sha(новый) == sha(архив[имя]):
            отметить(f"пересборка: {имя} побайтно = файлу в ZIP", True)
            continue
        a, b = zipfile.ZipFile(io.BytesIO(архив[имя])), zipfile.ZipFile(io.BytesIO(новый))
        разные = sorted(i for i in set(a.namelist()) | set(b.namelist())
                        if i not in a.namelist() or i not in b.namelist() or a.read(i) != b.read(i))
        # docProps/app.xml содержит версию openpyxl сборщика — не содержимое книги
        отметить(f"пересборка: {имя} — части книги = ZIP (кроме версии openpyxl в app.xml)",
                 разные == ["docProps/app.xml"], ", ".join(разные))


# ---------- 3. Calc: собственное дело ----------
D = dt.datetime


def ввод_04(ws, значения: dict) -> None:
    for адрес, v in значения.items():
        ws[адрес] = v


# Независимое дело «Н» (вымышленное): КС-2 3 150 000, оплачено 2 520 000,
# гарантийное 5 % = 157 500 с ненаступившим сроком; штраф 25 000; штраф
# 2,5 % от 3 150 000; пени 0,1 %/день двумя периодами; 395 двумя ставками.
ДЕЛО_Н = {
    "G5": 3_150_000, "G6": 2_520_000, "G8": 157_500,
    "G12": 25_000, "G13": 3_150_000, "G14": 2.5,
    "B18": D(2026, 3, 1), "C18": D(2026, 3, 31), "E18": 472_500, "F18": 0.1,
    "B19": D(2026, 4, 1), "C19": D(2026, 4, 15), "E19": 300_000, "F19": 0.1,
    "B27": D(2026, 3, 1), "C27": D(2026, 3, 31), "E27": 472_500, "F27": 21,
    "B28": D(2026, 4, 1), "C28": D(2026, 4, 15), "E28": 300_000, "F28": 20,
}
# Ручной расчёт (записан числами):
#   неоплачено 3 150 000 − 2 520 000 = 630 000; без гарантийного 472 500
#   штраф 2,5 %: 3 150 000 × 2,5 / 100 = 78 750
#   пени: 472 500 × 0,1 % × 31 = 14 647,50; 300 000 × 0,1 % × 15 = 4 500,00; итого 19 147,50
#   395: 472 500 × 21 % / 365 × 31 = 8 427,33; 300 000 × 20 % / 365 × 15 = 2 465,75; итого 10 893,08
ОЖИД_Н = {
    "G7": 630_000, "G9": 472_500, "G15": 78_750,
    "D18": 31, "G18": 14_647.50, "D19": 15, "G19": 4_500.00, "G23": 19_147.50,
    "D27": 31, "G27": 8_427.33, "D28": 15, "G28": 2_465.75, "G32": 10_893.08,
}


def найти_сводку(ws) -> dict[str, str]:
    """Подпись сводки → адрес G (по тексту, а не по номеру строки)."""
    out = {}
    в_сводке = False
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and a.startswith("4. СВОДКА"):
            в_сводке = True
            continue
        if в_сводке and isinstance(a, str) and ws.cell(row=r, column=7).value is not None:
            out[a] = f"G{r}"
    return out


def шаг_calc_04(tmp: Path, архив: dict[str, bytes]) -> None:
    print("\n## 04: собственное дело «Н» и граничные случаи (Calc)")
    исход = tmp / "zip" / "04-raschet-ubytkov.xlsx"
    wb0 = openpyxl.load_workbook(исход)
    ws0 = wb0.active
    # Раскладка: подтверждаем адреса по подписям, а не верим генератору
    подписи = {
        "A5": "Сумма по КС-2", "A6": "Фактически оплачено", "A7": "Неоплачено по КС-2",
        "A8": "из них удержание", "A9": "Неоплачено без", "A12": "2.1. Штраф",
        "A13": "2.2. Штраф в процентах", "A15": "сумма штрафа", "A23": "Итого пени",
        "A32": "Итого проценты",
    }
    плохо = [a for a, t in подписи.items() if not str(ws0[a].value or "").startswith(t)]
    отметить("04: раскладка по подписям совпадает с ожидаемой", not плохо, ", ".join(плохо))
    шапка = [ws0.cell(row=17, column=c).value for c in range(1, 9)]
    шапка395 = [ws0.cell(row=26, column=c).value for c in range(1, 9)]
    отметить("04: таблицы периодов в строках 17 и 26", шапка[1] == "Начало" and шапка395[5].startswith("Ключевая"),
             f"{шапка[5]} / {шапка395[5]}")

    случаи = {
        "дело_Н": (ДЕЛО_Н, ОЖИД_Н, {}),
        # переход через год: 28.12.2026–05.01.2027 = 9 дн.; 100 000 × 0,05 % × 9 = 450
        "через_год": ({"B18": D(2026, 12, 28), "C18": D(2027, 1, 5), "E18": 100_000, "F18": 0.05},
                      {"D18": 9, "G18": 450.0, "G23": 450.0}, {}),
        # високосный 2028: 01.02–29.02.2028 = 29 дн.; 395: 365 000 × 10 % / 365 × 29 = 2 900 (делитель 365)
        "високосный_2028": ({"B27": D(2028, 2, 1), "C27": D(2028, 2, 29), "E27": 365_000, "F27": 10},
                            {"D27": 29, "G27": 2_900.0}, {}),
        # ставка введена «0,05%» (Calc хранит 0,0005): ошибка единиц. «ок» здесь
        # означает, что книга молча дала 46,80 руб. вместо 4 680 — это наблюдение Н-1
        "ставка_как_процент_молча_46_80_руб": (
            {"B18": D(2026, 7, 29), "C18": D(2026, 9, 5), "E18": 240_000, "F18": 0.0005},
            {"G18": 46.80}, {"H18": ""}),
        # ставка текстом «0,05» (вставка из письма) → должно быть сообщение и пустая сумма
        "ставка_текстом": ({"B18": D(2026, 7, 29), "C18": D(2026, 9, 5), "E18": 240_000, "F18": "0,05"},
                           {}, {"G18": None, "H18": "числом"}),
        # удержание больше неоплаченного
        "удержание_больше_долга": ({"G5": 1_000_000, "G6": 900_000, "G8": 150_000},
                                   {"G7": 100_000}, {"H8": "больше", "G9": None}),
        # дата текстом «01.03.2026» (строка, не дата)
        "дата_текстом": ({"B18": "01.03.2026", "C18": D(2026, 3, 31), "E18": 1000, "F18": 0.1},
                         {}, {"G18": None, "H18": "числом"}),
        # одна строка ошибочна, вторая верна: итог пени не считается
        "одна_ошибка_из_двух": ({"B18": D(2026, 3, 1), "C18": D(2026, 3, 31), "E18": 472_500, "F18": 0.1,
                                 "B19": D(2026, 4, 15), "C19": D(2026, 4, 1), "E19": 300_000, "F19": 0.1},
                                {"G18": 14_647.50}, {"G19": None, "G23": None, "H19": "раньше",
                                                     "H23": "не считается"}),
    }
    src_dir = tmp / "04-in"
    src_dir.mkdir()
    for имя, (ввод, _, _) in случаи.items():
        wb = openpyxl.load_workbook(исход)
        ввод_04(wb.active, ввод)
        wb.save(src_dir / f"04-{имя}.xlsx")
    out = tmp / "04-calc"
    for имя in случаи:
        calc_пересчитать(src_dir / f"04-{имя}.xlsx", out, tmp)
    for имя, (_, числа, прочее) in случаи.items():
        ws = openpyxl.load_workbook(out / f"04-{имя}.xlsx", data_only=True).active
        плохо = []
        for адр, ож in числа.items():
            v = ws[адр].value
            if not близко(v, ож):
                плохо.append(f"{адр}={v!r}≠{ож}")
        for адр, ож in прочее.items():
            v = ws[адр].value
            if ож is None:
                if v not in (None, ""):
                    плохо.append(f"{адр}={v!r}, ожидалось пусто")
            elif ож == "":
                if v not in (None, ""):
                    плохо.append(f"{адр}={v!r}, ожидалось без сообщения")
            elif ож not in str(v or ""):
                плохо.append(f"{адр}={v!r}, нет «{ож}»")
        отметить(f"04 Calc: {имя}", not плохо, "; ".join(плохо) or
                 ", ".join(f"{a}={ws[a].value!r}" for a in list(числа)[:6]))
        if имя == "дело_Н":
            сводка = найти_сводку(ws)
            знач = {k: ws[a].value for k, a in сводка.items()}
            for k, v in знач.items():
                print(f"        сводка: {k} → {v!r}")
            ожид = [630_000, 157_500, 472_500, 25_000, 78_750, 19_147.50, 10_893.08]
            отметить("04 Calc: сводка дела «Н» — 7 строк раздельно, без общего итога",
                     len(знач) == 7 and all(близко(v, o) for v, o in zip(знач.values(), ожид)))
            pdf = в_pdf(out / "04-дело_Н.xlsx", tmp / "pdf", tmp)
            стр, размер = pdf_сведения(pdf)
            текст = pdf_текст(pdf)
            видно = all(s in текст for s in ("Проверка", "14 647,50", "10 893,08", "472 500,00"))
            отметить("04 печать заполненной: 1 лист A4, суммы и столбец «Проверка» видны",
                     стр == 1 and "595" in размер and видно, f"{стр} стр., {размер}")


# ---------- 4. 05 и 06 ----------
def шаг_05_06(tmp: Path) -> None:
    print("\n## 05: реестр дела «Н», печать; 06 без изменений")
    wb = openpyxl.load_workbook(tmp / "zip" / "05-reestr-uderzhaniy.xlsx")
    ws = wb.active
    шапка = [ws.cell(row=1, column=c).value for c in range(1, 14)]
    print("        шапка 05:", шапка)
    записи = [
        [1, D(2026, 4, 20), "Удержание", "КС-2 № 3, гарантийное 5 % п. 4.6", 157_500, "Срок не наступил",
         None, None, None, None, None, "", ""],
        [2, D(2026, 4, 20), "Штраф", "письмо исх. 77, п. 9.1", 25_000, "Оспаривается",
         D(2026, 4, 24), "заказное с описью", "РПО 12345678901234", D(2026, 5, 24), "п. 13.2 договора", "", ""],
        [3, D(2026, 5, 5), "Зачет", "уведомление исх. 81 — тот же штраф 25 000", 25_000, "Оспаривается",
         D(2026, 5, 8), "ЭДО", "квитанция оператора", D(2026, 6, 7), "п. 13.2 договора", "", ""],
        [4, D(2026, 5, 5), "Удержание", "без основания", 315_000, "Оспаривается",
         D(2026, 5, 8), "нарочно", "отметка о вручении", D(2026, 6, 7), "п. 13.2 договора",
         "Ответа нет", "К юристу (10)"],
    ]
    for i, запись in enumerate(записи, 2):
        for c, v in enumerate(запись, 1):
            ws.cell(row=i, column=c, value=v)
    src = tmp / "05-in"
    src.mkdir()
    wb.save(src / "05-дело_Н.xlsx")
    out = calc_пересчитать(src / "05-дело_Н.xlsx", tmp / "05-calc", tmp)
    ws2 = openpyxl.load_workbook(out, data_only=True).active
    итог, подпись, пояснение = ws2["E21"].value, ws2["D21"].value, ws2["A22"].value
    отметить("05 Calc: E21 = сумма записей 522 500,00 (штраф и зачёт того же штрафа — дважды)",
             близко(итог, 522_500), f"E21={итог!r}, D21={подпись!r}")
    отметить("05: подпись и пояснение о повторах есть", "не размер требования" in str(подпись)
             and "несколько раз" in str(пояснение))
    отметить("05: 5 полей отправки после «Статус»", шапка[6:11] == [
        "Дата отправки", "Способ отправки", "Подтверждение отправки", "Срок ответа", "Основание срока"])
    отметить("05 Calc: дата отправки сохранилась датой", isinstance(ws2["G3"].value, dt.datetime),
             repr(ws2["G3"].value))
    pdf = в_pdf(out, tmp / "pdf", tmp)
    стр, размер = pdf_сведения(pdf)
    текст = pdf_текст(pdf)
    видно = all(s in текст for s in ("Результат", "Основание", "522 500,00", "К юристу"))
    отметить("05 печать заполненной: 1 лист A4 альбомно, все 13 граф видны",
             стр == 1 and размер.startswith("841") and видно, f"{стр} стр., {размер}")
    # «Срок ответа» (J) без формата даты: дата, записанная в файл, печатается «###»
    отметить("05 печать: «Срок ответа» печатается датой, а не «###»", "###" not in текст,
             "формат в мастере 05: J3 {!r}, G3 {!r}".format(
                 *(openpyxl.load_workbook(tmp / "zip" / "05-reestr-uderzhaniy.xlsx").active[a].number_format
                   for a in ("J3", "G3"))))
    pdf6 = в_pdf(tmp / "zip" / "06-grafik-vozmeshcheniya.xlsx", tmp / "pdf", tmp)
    стр6, размер6 = pdf_сведения(pdf6)
    print(f"        06 (не менялась): {стр6} стр., {размер6}")


# ---------- 5. Word и границы ----------
ГРАНИЦЫ = {
    "гарантийное удержание / наступивший срок возврата": r"гарант|срок возврата|наступ",
    "недостатки / дефекты": r"недостат|дефект|замечан",
    "штраф / неустойка": r"штраф|неустойк|пен[иья]",
    "зачёт": r"зач[её]т",
    "встречные требования": r"встречн",
}


def шаг_word(tmp: Path) -> None:
    print("\n## Word: заполнение 01 и 11, печать; границы по тексту ZIP")
    import docx
    замены = {
        "_____": "Н-17",
        "[Наименование заказчика]": "ООО «Заказчик-Н»",
        "[Наименование подрядчика]": "ООО «Подрядчик-Н»",
    }
    for имя in ("01-pretenziya-na-uderzhanie.docx", "11-trebovanie-o-vozvrate-uderzhaniya.docx"):
        d = docx.Document(tmp / "zip" / имя)
        n = 0
        for p in d.paragraphs:
            for r in p.runs:
                for k, v in замены.items():
                    if k in r.text:
                        r.text = r.text.replace(k, v)
                        n += 1
        копия = tmp / "word-in" / имя
        копия.parent.mkdir(exist_ok=True)
        d.save(копия)
        pdf = в_pdf(копия, tmp / "pdf", tmp)
        стр, размер = pdf_сведения(pdf)
        отметить(f"Writer: {имя} заполнена ({n} замен) и напечатана", стр >= 1 and "Н-17" in pdf_текст(pdf),
                 f"{стр} стр., {размер}")

    print("\n## Границы: упоминания в тексте файлов ZIP (абзацы с совпадениями)")
    for имя in sorted((tmp / "zip").iterdir()):
        b = имя.read_bytes()
        if имя.suffix == ".docx":
            текст = харнесс.текст_docx(b)
        elif имя.suffix == ".pdf":
            текст = харнесс.текст_pdf(b)
        elif имя.suffix == ".txt":
            текст = b.decode("utf-8")
        else:
            continue
        for тема, рег in ГРАНИЦЫ.items():
            абзацы = [a.strip() for a in текст.splitlines() if re.search(рег, a, re.I)]
            print(f"   {имя.name} | {тема}: {len(абзацы)}")
            for a in абзацы[:4]:
                print(f"        · {a[:180]}")


def main() -> int:
    if len(sys.argv) > 1:
        tmp = Path(sys.argv[1]).resolve()
        tmp.mkdir(parents=True, exist_ok=False)
    else:
        tmp = Path(tempfile.mkdtemp(prefix="mvb-328-"))
    print("каталог копий:", tmp)
    print("HEAD:", subprocess.run(["git", "-C", str(КОРЕНЬ), "rev-parse", "HEAD"],
                                  capture_output=True, text=True).stdout.strip())
    print(subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip())
    архив = шаг_zip(tmp)
    шаг_пересборка(tmp, архив)
    шаг_calc_04(tmp, архив)
    шаг_05_06(tmp)
    шаг_word(tmp)
    провалы = [n for n, ок, _ in ИТОГ if not ок]
    print(f"\nИТОГ: {len(ИТОГ) - len(провалы)} ок, {len(провалы)} провал")
    for n in провалы:
        print("  ПРОВАЛ:", n)
    return 1 if провалы else 0


if __name__ == "__main__":
    sys.exit(main())
