#!/usr/bin/env python3
"""Пакет материалов P7 для независимого нормативного перечита (Р-026).

Только подготовка: собирает реальный покупательский ZIP p7 штатной PHP-функцией
выдачи, снимает хеши, извлекает полный текст каждого выдаваемого файла, текст
страницы товара, START-HERE, MANIFEST, отчётов #307 и #315 и строит карту
происхождения разделов договора. Оценок «верно / неверно» скрипт не выносит.

    python3 tools/candidates/evidence/MB001_R026_P7_PACKET/_tools/build_packet.py \\
        --pr-head-sha <sha> --pr-head-ref <ветка> --pr-base-sha <sha>

Пишет:
    tools/candidates/MB001_R026_P7_REVIEW_PACKET.md           — основной отчёт
    tools/candidates/evidence/MB001_R026_P7_PACKET/…          — полные извлечения

Проверка готовности пакета — verify_packet.py рядом.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
from lxml import etree

sys.dont_write_bytecode = True  # импорт генератора не оставляет __pycache__ в products-storage/

TOOLS = Path(__file__).resolve().parent
PACKET = TOOLS.parent
ROOT = PACKET.parents[3]
REPORT = ROOT / "tools/candidates/MB001_R026_P7_REVIEW_PACKET.md"

KIT_REL = "products-storage/03-dogovor-podryada"
KIT = ROOT / KIT_REL
PAGE_REL = "products/p7-dogovor-podryada.html"
GEN_REL = "products-storage/build_paid_03.py"
CONFIG_REL = "products-config.php"
IMPL_REL = "tools/candidates/MB001_R2_P7_IMPLEMENTATION_REPORT.md"
TEST_REL = "tools/test_p7_contract_kit.py"
PREVIEW_REL = "tools/build_preview.py"
GAP_REF = "claude/gifted-wright-otl19h"
GAP_SHA = "a1633c2bc0c4261a7d9653d1c3876f40dea0ed18"
GAP_REL = "tools/candidates/MB001_R2_P7_GAP_SPEC.md"
SERVICE = [".htaccess", "00-PISMO-POSLE-POKUPKI.txt", "MANIFEST.md"]
PACKET_REL = "tools/candidates/evidence/MB001_R026_P7_PACKET"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
FIELD = re.compile(r"\{\{[^}]+\}\}")
XL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!",
             "#NULL!", "#SPILL!", "#CALC!"}


def q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sh(*cmd, cwd=ROOT, check=True) -> str:
    return subprocess.run(cmd, cwd=cwd, check=check, capture_output=True,
                          text=True).stdout


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT))


# ─────────────────────────────── ZIP ───────────────────────────────

def build_zip(out: Path) -> Path:
    stdout = sh("php", str(TOOLS / "build_p7.php"), str(ROOT), str(out))
    return Path(stdout.strip().splitlines()[-1])


def zip_entries(path: Path) -> list[dict]:
    entries = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            data = z.read(info.filename)
            entries.append({
                "name": info.filename,
                "size": info.file_size,
                "compressed": info.compress_size,
                "sha256": sha256(data),
                "is_dir": info.is_dir(),
                "data": data,
            })
    return entries


KIND = {".docx": "DOCX", ".xlsx": "XLSX", ".pdf": "PDF", ".txt": "TXT"}


# ─────────────────────────────── DOCX ───────────────────────────────

def run_text(p) -> str:
    out = []
    for r in p.iter(q("r")):
        for node in r:
            tag = node.tag
            if tag == q("t"):
                out.append(node.text or "")
            elif tag == q("delText"):
                out.append(f"[удалено: {node.text or ''}]")
            elif tag == q("tab"):
                out.append("\t")
            elif tag in (q("br"), q("cr")):
                out.append("\n")
            elif tag == q("instrText"):
                out.append(f"[поле: {(node.text or '').strip()}]")
            elif tag == q("footnoteReference"):
                out.append(f"[сноска {node.get(q('id'))}]")
            elif tag == q("endnoteReference"):
                out.append(f"[концевая сноска {node.get(q('id'))}]")
            elif tag == q("commentReference"):
                out.append(f"[примечание {node.get(q('id'))}]")
    for fld in p.iter(q("fldSimple")):
        out.append(f"[поле: {fld.get(q('instr'))}]")
    return "".join(out)


def walk(parent, prefix: str, lines: list, counter: dict, top: bool) -> None:
    for child in parent:
        if child.tag == q("p"):
            if top:
                counter["b"] += 1
                lines.append((f"B{counter['b']:04d} P", run_text(child)))
            else:
                counter["p"] += 1
                lines.append((f"{prefix} p{counter['p']}", run_text(child)))
        elif child.tag == q("tbl"):
            if top:
                counter["b"] += 1
                counter["t"] += 1
                tprefix = f"B{counter['b']:04d} T{counter['t']}"
            else:
                counter["nt"] = counter.get("nt", 0) + 1
                tprefix = f"{prefix} вложенная-T{counter['nt']}"
            for ri, tr in enumerate(child.findall(q("tr")), 1):
                for ci, tc in enumerate(tr.findall(q("tc")), 1):
                    span = tc.find(f"{q('tcPr')}/{q('gridSpan')}")
                    vm = tc.find(f"{q('tcPr')}/{q('vMerge')}")
                    mark = ""
                    if span is not None:
                        mark += f" объединение-по-горизонтали={span.get(q('val'))}"
                    if vm is not None:
                        mark += f" объединение-по-вертикали={vm.get(q('val')) or 'continue'}"
                    sub = {"p": 0}
                    walk(tc, f"{tprefix} R{ri} C{ci}{mark}", lines, sub, False)
        elif child.tag in (q("sdt"),):
            content = child.find(q("sdtContent"))
            if content is not None:
                walk(content, prefix, lines, counter, top)


def docx_extract(data: bytes) -> dict:
    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        tmp.write(data)
        tmp.flush()
        z = zipfile.ZipFile(tmp.name)
        parts = z.namelist()
        result = {"parts": parts, "sections": [], "meta": {}}
        body = etree.fromstring(z.read("word/document.xml")).find(q("body"))
        lines: list = []
        walk(body, "", lines, {"b": 0, "t": 0}, True)
        result["sections"].append(("word/document.xml — тело", lines))
        for part in sorted(p for p in parts
                           if re.match(r"word/(header|footer)\d*\.xml$", p)):
            root = etree.fromstring(z.read(part))
            hl: list = []
            walk(root, part, hl, {"b": 0, "t": 0}, True)
            result["sections"].append((f"{part} — колонтитул", hl))
        for part, tag, label in (("word/footnotes.xml", "footnote", "сноски"),
                                 ("word/endnotes.xml", "endnote", "концевые сноски")):
            if part in parts:
                root = etree.fromstring(z.read(part))
                nl = []
                for note in root.findall(q(tag)):
                    nid = note.get(q("id"))
                    if note.get(q("type")) in ("separator", "continuationSeparator",
                                               "continuationNotice"):
                        continue
                    for i, p in enumerate(note.iter(q("p")), 1):
                        nl.append((f"{tag} id={nid} p{i}", run_text(p)))
                result["sections"].append((f"{part} — {label}", nl))
        if "word/comments.xml" in parts:
            root = etree.fromstring(z.read("word/comments.xml"))
            cl = []
            for c in root.findall(q("comment")):
                for i, p in enumerate(c.iter(q("p")), 1):
                    cl.append((f"comment id={c.get(q('id'))} автор={c.get(q('author'))} p{i}",
                               run_text(p)))
            result["sections"].append(("word/comments.xml — примечания", cl))
        for part in ("docProps/core.xml", "docProps/app.xml"):
            if part in parts:
                root = etree.fromstring(z.read(part))
                for el in root:
                    name = etree.QName(el).localname
                    if el.text and el.text.strip():
                        result["meta"][f"{part}:{name}"] = el.text.strip()
        doc_xml = z.read("word/document.xml").decode("utf-8")
        result["tracked"] = {
            "w:ins": doc_xml.count("<w:ins "), "w:del": doc_xml.count("<w:del "),
        }
    return result


def docx_text(ex: dict) -> str:
    return "\n".join(t for _, lines in ex["sections"] for _, t in lines)


def docx_render(name: str, entry: dict, ex: dict) -> str:
    body_lines = ex["sections"][0][1]
    text = docx_text(ex)
    fields = FIELD.findall(text)
    words = len(re.findall(r"\S+", "\n".join(t for _, t in body_lines)))
    out = [f"# {name}", "",
           f"sha256: {entry['sha256']}",
           f"размер: {entry['size']} байт",
           f"части архива DOCX ({len(ex['parts'])}): " + ", ".join(ex["parts"]),
           f"колонтитулы: " + (", ".join(s for s, _ in ex["sections"] if "колонтитул" in s) or "нет"),
           f"сноски: {'есть word/footnotes.xml' if 'word/footnotes.xml' in ex['parts'] else 'нет word/footnotes.xml'}",
           f"концевые сноски: {'есть word/endnotes.xml' if 'word/endnotes.xml' in ex['parts'] else 'нет word/endnotes.xml'}",
           f"примечания (comments): {'есть word/comments.xml' if 'word/comments.xml' in ex['parts'] else 'нет word/comments.xml'}",
           f"исправления (track changes) в теле: w:ins={ex['tracked']['w:ins']}, w:del={ex['tracked']['w:del']}",
           f"блоков тела: {len({l.split()[0] for l, _ in body_lines})}; строк извлечения тела: {len(body_lines)}; "
           f"слов в теле (последовательности непробельных символов): {words}",
           f"полей {{{{…}}}}: {len(fields)} вхождений, {len(set(fields))} различных",
           "метаданные: " + "; ".join(f"{k}={v}" for k, v in sorted(ex["meta"].items())),
           "",
           "Метки мест: `B0001 P` — абзац тела №1 по порядку блоков; "
           "`B0005 T1 R2 C1 p1` — блок №5, таблица №1, строка 2, колонка 1, абзац 1 "
           "ячейки. Пустой абзац показан как ∅.", ""]
    for title, lines in ex["sections"]:
        out.append(f"## {title}")
        out.append("")
        if not lines:
            out.append("(нет текста)")
        for loc, t in lines:
            out.append(f"[{loc}] {t if t else '∅'}")
        out.append("")
    if fields:
        out.append("## Поля {{…}} — перечень различных, по алфавиту")
        out.append("")
        for f in sorted(set(fields)):
            out.append(f"{f} — {fields.count(f)}")
        out.append("")
    return "\n".join(out)


# ─────────────────────────────── XLSX ───────────────────────────────

def xlsx_extract(data: bytes) -> dict:
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
        tmp.write(data)
        tmp.flush()
        wf = openpyxl.load_workbook(tmp.name, data_only=False)
        wv = openpyxl.load_workbook(tmp.name, data_only=True)
        parts = zipfile.ZipFile(tmp.name).namelist()
        sheets = []
        totals = {"cells": 0, "text": 0, "formulas": 0, "errors": 0}
        for ws in wf.worksheets:
            vs = wv[ws.title]
            cells = []
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    cached = vs[c.coordinate].value
                    is_formula = c.data_type == "f"
                    is_error = (c.data_type == "e" or
                                (isinstance(cached, str) and cached in XL_ERRORS) or
                                (isinstance(c.value, str) and c.value in XL_ERRORS))
                    totals["cells"] += 1
                    totals["formulas"] += is_formula
                    totals["errors"] += is_error
                    totals["text"] += isinstance(c.value, str) and not is_formula
                    cells.append({"coord": c.coordinate, "type": c.data_type,
                                  "value": c.value, "cached": cached,
                                  "formula": is_formula, "error": is_error,
                                  "comment": c.comment.text if c.comment else None})
            sheets.append({
                "title": ws.title, "state": ws.sheet_state,
                "dims": ws.dimensions, "merged": [str(m) for m in ws.merged_cells.ranges],
                "validations": [(str(dv.sqref), dv.type, dv.formula1)
                                for dv in ws.data_validations.dataValidation],
                "freeze": ws.freeze_panes, "cells": cells,
            })
        names = list(wf.defined_names.keys()) if hasattr(wf.defined_names, "keys") else []
        return {"sheets": sheets, "totals": totals, "parts": parts, "names": names,
                "creator": wf.properties.creator, "created": str(wf.properties.created),
                "modified": str(wf.properties.modified)}


def xlsx_render(name: str, entry: dict, ex: dict) -> str:
    t = ex["totals"]
    out = [f"# {name}", "", f"sha256: {entry['sha256']}", f"размер: {entry['size']} байт",
           f"листов: {len(ex['sheets'])}; непустых ячеек: {t['cells']}; с текстом: {t['text']}; "
           f"формул: {t['formulas']}; ошибок формул/значений: {t['errors']}",
           f"определённые имена: {', '.join(ex['names']) or 'нет'}",
           f"свойства: creator={ex['creator']}; created={ex['created']}; modified={ex['modified']}",
           f"части архива ({len(ex['parts'])}): " + ", ".join(ex["parts"]), ""]
    for s in ex["sheets"]:
        out += [f"## Лист «{s['title']}»", "",
                f"состояние: {s['state']}; диапазон: {s['dims']}; закрепление: {s['freeze']}",
                f"объединённые диапазоны: {', '.join(s['merged']) or 'нет'}",
                "проверки данных: " + ("; ".join(f"{a} {b} {c}" for a, b, c in s["validations"]) or "нет"),
                ""]
        for c in s["cells"]:
            line = f"[{s['title']}!{c['coord']}] "
            if c["formula"]:
                line += f"формула: {c['value']} | кэш: {c['cached']!r}"
            else:
                v = c["value"]
                line += v.replace("\n", "⏎") if isinstance(v, str) else repr(v)
            if c["error"]:
                line += " | ОШИБКА"
            if c["comment"]:
                line += f" | примечание: {c['comment']}"
            out.append(line)
        out.append("")
    return "\n".join(out)


# ─────────────────────────────── PDF / TXT ───────────────────────────────

def pdf_extract(data: bytes) -> dict:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
        tmp.write(data)
        tmp.flush()
        info = sh("pdfinfo", tmp.name)
        text = sh("pdftotext", "-layout", "-enc", "UTF-8", tmp.name, "-")
    meta = dict(line.split(":", 1) for line in info.splitlines() if ":" in line)
    meta = {k.strip(): v.strip() for k, v in meta.items()}
    pages = text.split("\f")
    if pages and not pages[-1].strip():
        pages = pages[:-1]
    return {"meta": meta, "pages": pages}


def pdf_render(name: str, entry: dict, ex: dict) -> str:
    m = ex["meta"]
    out = [f"# {name}", "", f"sha256: {entry['sha256']}", f"размер: {entry['size']} байт",
           f"страниц (pdfinfo Pages): {m.get('Pages')}; страниц текста (pdftotext): {len(ex['pages'])}",
           "pdfinfo: " + "; ".join(f"{k}={m[k]}" for k in
                                   ("Title", "Author", "Creator", "Producer",
                                    "CreationDate", "ModDate", "Page size") if k in m),
           "символов текста по страницам: " + ", ".join(
               f"с.{i}={len(p.strip())}" for i, p in enumerate(ex["pages"], 1)),
           "", "Извлечено `pdftotext -layout -enc UTF-8`.", ""]
    for i, p in enumerate(ex["pages"], 1):
        out += [f"## Страница {i}", "", "```text", p.rstrip(), "```", ""]
    return "\n".join(out)


def txt_render(name: str, entry: dict, text: str, note: str = "") -> str:
    out = [f"# {name}", "", f"sha256: {entry['sha256']}", f"размер: {entry['size']} байт",
           f"строк: {len(text.splitlines())}"]
    if note:
        out.append(note)
    out += ["", "Строки пронумерованы: `L0001` — первая строка файла.", ""]
    out += [f"L{i:04d} {line}" for i, line in enumerate(text.splitlines(), 1)]
    return "\n".join(out) + "\n"


# ─────────────────────────────── HTML страницы ───────────────────────────────

BLOCK = {"html", "head", "body", "header", "footer", "main", "section", "article",
         "aside", "nav", "div", "p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol",
         "li", "dl", "dt", "dd", "table", "thead", "tbody", "tr", "td", "th",
         "details", "summary", "form", "label", "button", "figure", "figcaption",
         "blockquote", "pre", "title", "option", "select"}
SKIP = {"script", "style", "noscript", "template", "svg"}
VOID = {"meta", "link", "img", "br", "hr", "input", "source", "area", "base",
        "col", "embed", "param", "track", "wbr"}


class PageParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack: list[dict] = []
        self.blocks: list[dict] = []
        self.skip = 0
        self.jsonld: list[tuple[int, str]] = []
        self.meta: list[tuple[int, str, str]] = []
        self.images: list[tuple[int, str, str]] = []
        self._in_jsonld = None
        self._buf = []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        line = self.getpos()[0]
        if tag == "script" and a.get("type") == "application/ld+json":
            self._in_jsonld = line
            self._buf = []
        if tag == "meta" and ("name" in a or "property" in a):
            self.meta.append((line, a.get("name") or a.get("property"), a.get("content", "")))
        if tag == "img":
            self.images.append((line, a.get("src", ""), a.get("alt", "")))
        if tag in VOID:
            return
        if tag in SKIP:
            self.skip += 1
        node = {"tag": tag, "cls": a.get("class", ""), "id": a.get("id", ""),
                "line": line, "hidden": "hidden" in a or a.get("aria-hidden") == "true",
                "block": tag in BLOCK, "text": [], "faq": False}
        parent_faq = any(n["faq"] for n in self.stack)
        node["faq"] = parent_faq or ("faq" in node["cls"].split() and tag == "section")
        node["hidden"] = node["hidden"] or any(n["hidden"] for n in self.stack)
        self.stack.append(node)
        if node["block"]:
            self.blocks.append(node)

    def handle_endtag(self, tag):
        if tag == "script" and self._in_jsonld is not None:
            self.jsonld.append((self._in_jsonld, "".join(self._buf)))
            self._in_jsonld = None
        if tag in VOID:
            return
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i]["tag"] == tag:
                for n in self.stack[i:]:
                    if n["tag"] in SKIP:
                        self.skip -= 1
                del self.stack[i:]
                break

    def handle_data(self, data):
        if self._in_jsonld is not None:
            self._buf.append(data)
            return
        if self.skip or not data.strip():
            return
        for n in reversed(self.stack):
            if n["block"]:
                n["text"].append(data)
                return


def page_extract(html_text: str) -> dict:
    p = PageParser()
    p.feed(html_text)
    blocks = []
    for b in p.blocks:
        t = re.sub(r"\s+", " ", "".join(b["text"])).strip()
        if t:
            label = b["tag"] + (f".{b['cls'].split()[0]}" if b["cls"] else "")
            blocks.append({"line": b["line"], "label": label, "text": t,
                           "hidden": b["hidden"], "faq": b["faq"]})
    ld = []
    for line, raw in p.jsonld:
        try:
            ld.append((line, json.loads(raw)))
        except json.JSONDecodeError as exc:
            ld.append((line, {"_ошибка_разбора": str(exc), "_сырой_текст": raw}))
    return {"blocks": blocks, "jsonld": ld, "meta": p.meta, "images": p.images}


def ld_faq(ld) -> list[tuple[str, str]]:
    out = []
    for _, obj in ld:
        for node in obj.get("@graph", [obj]) if isinstance(obj, dict) else []:
            if node.get("@type") == "FAQPage":
                for qn in node.get("mainEntity", []):
                    out.append((qn.get("name", ""), qn.get("acceptedAnswer", {}).get("text", "")))
    return out


# ─────────────────────────────── генератор и карта ───────────────────────────────

def load_generator():
    spec = importlib.util.spec_from_file_location("build_paid_03", ROOT / GEN_REL)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def generator_lines(mod) -> dict:
    src = (ROOT / GEN_REL).read_text(encoding="utf-8").splitlines()

    def span(start_pat, end_pat):
        s = next(i for i, l in enumerate(src) if re.match(start_pat, l))
        e = next(i for i in range(s + 1, len(src)) if re.match(end_pat, src[i]))
        return s, e

    entry = re.compile(r'^\s{4}\("((?:h)|(?:[\d.]+[a-z]?)|(?:table-[a-z]+))",')
    cs, ce = span(r"^CONTRACT: list", r"^\]")
    contract = [(m.group(1), i + 1) for i in range(cs, ce + 1)
                if (m := entry.match(src[i]))]
    ps, pe = span(r"^PROTOCOL: list", r"^\]")
    # Позиции бывают записаны генератором списка: *[(f"9.2.{n}", …) for n in range(a, b)]
    spread = re.compile(r'^\s{4}\*\[\(f"([\d.]+)\{n\}",.*range\((\d+), (\d+)\)')
    protocol = []
    for i in range(ps, pe + 1):
        if (m := entry.match(src[i])):
            protocol.append((m.group(1), i + 1, src[i]))
        elif (m := spread.match(src[i])):
            for n in range(int(m.group(2)), int(m.group(3))):
                protocol.append((f"{m.group(1)}{n}", i + 1, src[i]))
    keys = [k for k, _ in mod.CONTRACT]
    assert keys == [k for k, _ in contract], "порядок CONTRACT не совпал с разбором исходника"
    assert [k for k, _ in mod.PROTOCOL] == [k for k, _, _ in protocol]
    # Конец блока пункта — строка перед следующим пунктом
    contract_spans = []
    for i, (k, line) in enumerate(contract):
        end = (contract[i + 1][1] - 1) if i + 1 < len(contract) else ce + 1
        while end > line and not src[end - 1].strip():
            end -= 1
        contract_spans.append((k, line, end))
    protocol_spans = []
    for i, (k, line, text) in enumerate(protocol):
        end = (protocol[i + 1][1] - 1) if i + 1 < len(protocol) else pe
        end = max(end, line)
        while end > line and not src[end - 1].strip():
            end -= 1
        body = "\n".join(src[line - 1:end])
        protocol_spans.append((k, line, end, body))
    defs = {}
    for i, l in enumerate(src, 1):
        m = re.match(r"^(def (\w+)|([A-Z_]+)(?::[^=]+)? = )", l)
        if m:
            defs[m.group(2) or m.group(3)] = i
    return {"contract": contract_spans, "protocol": protocol_spans, "defs": defs,
            "len": len(src)}


def contract_blocks(mod) -> list[dict]:
    """Моделирует порядок блоков build_contract: какой блок тела — какой пункт."""
    b = 0
    seq = []

    def add(kind, key, n=1):
        nonlocal b
        start = b + 1
        b += n
        seq.append({"kind": kind, "key": key, "first": start, "last": b})

    add("служебные строки", "service", 2)
    add("заголовок", "title")
    add("место и дата (таблица T1)", "place")
    add("пустой абзац", "empty")
    add("преамбула", "preamble", 2)
    for key, paras in mod.CONTRACT:
        if key == "h":
            add("заголовок раздела", paras)
        elif key == "table-retention":
            add("таблица возврата гарантийного удержания", key, 1)
            add("пустой абзац", "empty")
        elif key == "table-parties":
            add("таблица реквизитов и подписей", key, 1)
        else:
            add("пункт", key, len(paras))
    return seq


REF_CLAUSE = re.compile(
    r"(?:\bп\.|\bпп\.|\bпункт(?:а|е|ы|ов|ам|ами|ах)?)\s*"
    r"((?:\d+(?:\.\d+)+)(?:\s*(?:,|и)\s*\d+(?:\.\d+)+)*)")
REF_SECTION = re.compile(
    r"\bраздел(?:а|е|ы|ов|ам|ами|ах)?\s*(?:№\s*)?"
    r"((?:\d+(?:\s*[–-]\s*\d+)?)(?:\s*(?:,|и)\s*\d+(?:\s*[–-]\s*\d+)?)*)")


def find_refs(text_lines: list[tuple[str, str]], source: str) -> list[dict]:
    refs = []
    for loc, t in text_lines:
        for m in REF_CLAUSE.finditer(t):
            for num in re.findall(r"\d+(?:\.\d+)+", m.group(1)):
                refs.append({"source": source, "loc": loc, "clause": num.rstrip(".")})
        for m in REF_SECTION.finditer(t):
            for part in re.split(r"\s*(?:,|и)\s*", m.group(1)):
                r = re.match(r"(\d+)(?:\s*[–-]\s*(\d+))?", part)
                if not r:
                    continue
                a, bb = int(r.group(1)), int(r.group(2) or r.group(1))
                for n in range(a, bb + 1):
                    refs.append({"source": source, "loc": loc, "section": str(n)})
    return refs


# ─────────────────────────────── сборка пакета ───────────────────────────────

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pr", type=int, default=315)
    ap.add_argument("--pr-head-sha", required=True)
    ap.add_argument("--pr-head-ref", required=True)
    ap.add_argument("--pr-base-sha", required=True)
    args = ap.parse_args()

    head = sh("git", "rev-parse", "HEAD").strip()
    is_anc = subprocess.run(["git", "merge-base", "--is-ancestor", args.pr_head_sha, "HEAD"],
                            cwd=ROOT).returncode == 0
    kit_same = subprocess.run(["git", "diff", "--quiet", args.pr_head_sha, "--",
                               "products-storage", "products", CONFIG_REL, "tools/candidates/MB001_R2_P7_IMPLEMENTATION_REPORT.md"],
                              cwd=ROOT).returncode == 0
    php_version = sh("php", "-r", "echo PHP_VERSION;")
    zip_class = sh("php", "-r", "echo class_exists('ZipArchive') ? 'да' : 'нет';")

    # Каталог пакета пересобирается целиком, кроме _tools
    for child in PACKET.iterdir():
        if child.name != "_tools":
            shutil.rmtree(child) if child.is_dir() else child.unlink()

    tmp = Path(tempfile.mkdtemp(prefix="p7zip-"))
    zpath = build_zip(tmp / "a")
    zpath2 = build_zip(tmp / "b")
    zbytes = zpath.read_bytes()
    zsha = sha256(zbytes)
    zsha2 = sha256(zpath2.read_bytes())
    entries = zip_entries(zpath)
    entries2 = zip_entries(zpath2)
    same_entries = [(e["name"], e["sha256"]) for e in entries] == \
                   [(e["name"], e["sha256"]) for e in entries2]

    disk = sorted(p.name for p in KIT.iterdir() if p.is_file())
    disk_payload = sorted(n for n in disk if n not in SERVICE)
    zip_names = sorted(e["name"] for e in entries)
    disk_hash = {n: sha256((KIT / n).read_bytes()) for n in disk}

    # ── извлечения
    outdir = PACKET / "zip"
    outdir.mkdir(parents=True)
    index = []
    extracted: dict[str, dict] = {}
    for e in sorted(entries, key=lambda e: e["name"]):
        name = e["name"]
        kind = KIND.get(Path(name).suffix.lower(), "другое")
        target = outdir / f"{name}.txt"
        if kind == "DOCX":
            ex = docx_extract(e["data"])
            target.write_text(docx_render(name, e, ex), encoding="utf-8")
            body = ex["sections"][0][1]
            tables = sorted({l.split()[1] for l, _ in body if len(l.split()) > 1 and l.split()[1].startswith("T")})
            words = len(re.findall(r"\S+", "\n".join(t for _, t in body)))
            hf = [s for s, _ in ex["sections"] if "колонтитул" in s]
            hf_text = [t for s, lines in ex["sections"] if "колонтитул" in s for _, t in lines if t]
            summary = (f"абзацев/ячеек тела: {len(body)}; таблиц: {len(tables)}; слов тела: {words}; "
                       f"колонтитулов: {len(hf)} ({'; '.join(hf_text) or 'без текста'}); "
                       f"сносок: {'нет' if 'word/footnotes.xml' not in ex['parts'] else 'часть есть'}; "
                       f"примечаний: {'нет' if 'word/comments.xml' not in ex['parts'] else 'часть есть'}; "
                       f"полей {{{{…}}}}: {len(FIELD.findall(docx_text(ex)))}")
            extracted[name] = {"kind": kind, "lines": [(l, t) for s, ls in ex["sections"] for l, t in ls], "ex": ex}
        elif kind == "XLSX":
            ex = xlsx_extract(e["data"])
            target.write_text(xlsx_render(name, e, ex), encoding="utf-8")
            t = ex["totals"]
            summary = (f"листов: {len(ex['sheets'])} ({', '.join(s['title'] + ' ' + s['dims'] for s in ex['sheets'])}); "
                       f"непустых ячеек: {t['cells']}; текстовых: {t['text']}; формул: {t['formulas']}; "
                       f"ошибок: {t['errors']}")
            lines = [(f"{s['title']}!{c['coord']}", str(c["value"])) for s in ex["sheets"] for c in s["cells"]]
            extracted[name] = {"kind": kind, "lines": lines, "ex": ex}
        elif kind == "PDF":
            ex = pdf_extract(e["data"])
            target.write_text(pdf_render(name, e, ex), encoding="utf-8")
            summary = (f"страниц: {ex['meta'].get('Pages')}; Producer: {ex['meta'].get('Producer', '—')}; "
                       f"символов текста: {sum(len(p.strip()) for p in ex['pages'])}")
            lines = [(f"с.{i} L{j}", l) for i, p in enumerate(ex["pages"], 1)
                     for j, l in enumerate(p.splitlines(), 1) if l.strip()]
            extracted[name] = {"kind": kind, "lines": lines, "ex": ex}
        elif kind == "TXT":
            text = e["data"].decode("utf-8")
            target.write_text(txt_render(name, e, text), encoding="utf-8")
            summary = f"строк: {len(text.splitlines())}; кодировка UTF-8"
            lines = [(f"L{i:04d}", l) for i, l in enumerate(text.splitlines(), 1)]
            extracted[name] = {"kind": kind, "lines": lines, "text": text}
        else:
            summary = "тип не извлекается"
        index.append({"name": name, "kind": kind, "summary": summary,
                      "evidence": f"{PACKET_REL}/zip/{name}.txt"})

    # ── страница, служебные файлы, отчёты
    site = PACKET / "site"
    site.mkdir()
    html_text = (ROOT / PAGE_REL).read_text(encoding="utf-8")
    page = page_extract(html_text)
    page_sha = sha256((ROOT / PAGE_REL).read_bytes())
    vis = [f"# Видимый текст {PAGE_REL}", "", f"sha256 файла: {page_sha}",
           "Каждый блок: `L<строка исходника> <тег.класс>` и его текст. "
           "Скрипты, стили, svg и noscript исключены. Пометка [скрыт] — атрибут "
           "hidden или aria-hidden у блока либо его предка (сам CSS не вычислялся).", "",
           "## meta", ""]
    vis += [f"L{l} meta {n}: {c}" for l, n, c in page["meta"]]
    vis += ["", "## img alt", ""] + [f"L{l} img src={s} alt={a!r}" for l, s, a in page["images"]]
    vis += ["", "## Блоки", ""]
    vis += [f"L{b['line']} <{b['label']}>{' [скрыт]' if b['hidden'] else ''} {b['text']}"
            for b in page["blocks"]]
    (site / "p7-dogovor-podryada.visible.txt").write_text("\n".join(vis) + "\n", encoding="utf-8")
    faq_blocks = [b for b in page["blocks"] if b["faq"]]
    faq = [f"# FAQ страницы {PAGE_REL}", "", "## Видимый блок (section.faq)", ""]
    faq += [f"L{b['line']} <{b['label']}> {b['text']}" for b in faq_blocks]
    faq += ["", "## FAQPage в JSON-LD", ""]
    ldq = ld_faq(page["jsonld"])
    for i, (qq, aa) in enumerate(ldq, 1):
        faq += [f"В{i}. {qq}", f"О{i}. {aa}", ""]
    (site / "p7-dogovor-podryada.faq.txt").write_text("\n".join(faq) + "\n", encoding="utf-8")
    (site / "p7-dogovor-podryada.jsonld.json").write_text(
        json.dumps([{"строка_исходника": l, "json": o} for l, o in page["jsonld"]],
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    visible_faq_q = [b["text"] for b in faq_blocks if b["label"].startswith(("summary", "h3", "dt", "button"))]

    service_dir = PACKET / "kit-service"
    service_dir.mkdir()
    service_info = []
    for n in SERVICE:
        p = KIT / n
        if p.is_file():
            text = p.read_text(encoding="utf-8", errors="replace")
            ent = {"sha256": sha256(p.read_bytes()), "size": p.stat().st_size}
            (service_dir / f"{n}.txt").write_text(
                txt_render(n, ent, text, "служебный файл: исключается из ZIP списком $service в mvb_build_product_zip"),
                encoding="utf-8")
            service_info.append((n, ent["sha256"], ent["size"]))

    reports = PACKET / "source-reports"
    reports.mkdir()
    impl_bytes = (ROOT / IMPL_REL).read_bytes()
    (reports / "MB001_R2_P7_IMPLEMENTATION_REPORT.md").write_bytes(impl_bytes)
    gap_present = subprocess.run(["git", "cat-file", "-e", f"{GAP_SHA}:{GAP_REL}"],
                                 cwd=ROOT).returncode == 0
    gap_bytes = b""
    if gap_present:
        gap_bytes = subprocess.run(["git", "show", f"{GAP_SHA}:{GAP_REL}"], cwd=ROOT,
                                   check=True, capture_output=True).stdout
        (reports / "MB001_R2_P7_GAP_SPEC.md").write_bytes(gap_bytes)

    # ── генератор: воспроизводимость и карта
    mod = load_generator()
    gl = generator_lines(mod)
    regen = tmp / "regen"
    mod.build(regen, pdf=False)
    regen_cmp = {}
    for n in ("01-dogovor-subpodryada.docx", "10-protokol-raznoglasiy.docx",
              "00-INSTRUKCIYA.docx", "00-START-HERE.txt"):
        regen_cmp[n] = (sha256((regen / n).read_bytes()), disk_hash[n])

    contract_ex = extracted["01-dogovor-subpodryada.docx"]["ex"]
    body01 = contract_ex["sections"][0][1]
    bid = {}
    for loc, t in body01:
        b = loc.split()[0]
        bid.setdefault(b, []).append((loc, t))
    seq = contract_blocks(mod)
    total_blocks = len(bid)
    sim_ok = seq[-1]["last"] == total_blocks
    for s in seq:
        if s["kind"] == "пункт":
            first = bid.get(f"B{s['first']:04d}", [("", "")])[0][1]
            s["match"] = first == mod.CLAUSES[s["key"]][0]
        elif s["kind"] == "заголовок раздела":
            s["match"] = bid.get(f"B{s['first']:04d}", [("", "")])[0][1] == s["key"]
        else:
            s["match"] = None
    sim_ok = sim_ok and all(s["match"] is not False for s in seq)

    proto_ex = extracted["10-protokol-raznoglasiy.docx"]["ex"]
    body10 = proto_ex["sections"][0][1]
    ptable = sorted({l.split()[0] + " " + l.split()[1] for l, _ in body10
                     if len(l.split()) > 1 and l.split()[1] == "T2"})
    ptable_id = ptable[0] if ptable else "—"
    proto_rows = []
    for i, ((key, value), (_, line, end, body)) in enumerate(zip(mod.PROTOCOL, gl["protocol"]), 2):
        if value == mod.EXCLUDE:
            kind = "Исключить"
        elif value == "without_last":
            kind = "редакция: пункт без последнего абзаца (without_last)"
        elif key not in mod.CLAUSES:
            kind = "новый пункт (в договоре нет)"
        elif "CLAUSES[" in body:
            kind = "новая редакция, производная от текста пункта (CLAUSES[…])"
        else:
            kind = "новая редакция"
        right = [t for l, t in body10 if l.startswith(f"{ptable_id} R{i} C2")]
        left = [t for l, t in body10 if l.startswith(f"{ptable_id} R{i} C1")]
        exp_right = mod.contractor_edition(key, value)
        exp_left = mod.customer_edition(key)
        proto_rows.append({"row": i, "key": key, "kind": kind, "line": line, "end": end,
                           "left_ok": left == exp_left, "right_ok": right == exp_right})

    kinds = {}
    for r in proto_rows:
        k = r["kind"].split(" (")[0].split(",")[0]
        kinds[k] = kinds.get(k, 0) + 1

    # ── ссылки на пункты и разделы из связанных текстов
    ref_sources = []
    for n in ("00-INSTRUKCIYA.docx", "00-INSTRUKCIYA.pdf", "00-START-HERE.txt", "02-perechen-rabot.xlsx",
              "03-kalendarnyy-plan.xlsx", "04-poryadok-priemki.docx", "05-grafik-platezhey.xlsx",
              "06-dopsoglashenie-obem.docx", "07-dopsoglashenie-sroki.docx", "08-krasnye-flagi.pdf",
              "09-checklist-dogovora.pdf"):
        ref_sources += find_refs(extracted[n]["lines"], n)
    ref_sources += find_refs([(f"L{b['line']}", b["text"]) for b in page["blocks"]], PAGE_REL)
    ref_sources += find_refs([(f"L{l}", c) for l, _, c in page["meta"]], PAGE_REL + " meta")

    def refs_for_section(n: str) -> list[str]:
        hits = {}
        for r in ref_sources:
            if r.get("section") == n or (r.get("clause", "").split(".")[0] == n):
                what = r.get("clause") and f"п. {r['clause']}" or f"разд. {n}"
                hits.setdefault(r["source"], []).append(f"{r['loc']} ({what})")
        return [f"{src}: {', '.join(dict.fromkeys(v))}" for src, v in hits.items()]

    def refs_for_clause(key: str) -> list[str]:
        hits = {}
        for r in ref_sources:
            if r.get("clause") == key:
                hits.setdefault(r["source"], []).append(r["loc"])
        return [f"{src}: {', '.join(dict.fromkeys(v))}" for src, v in hits.items()]

    maps = PACKET / "maps"
    maps.mkdir()
    impl_rows = {}
    for line in impl_bytes.decode("utf-8").splitlines():
        m = re.match(r"\|\s*(\d+)\.\s", line)
        if m and m.group(1) not in impl_rows:
            impl_rows[m.group(1)] = line
    impl_line = {}
    for i, line in enumerate(impl_bytes.decode("utf-8").splitlines(), 1):
        m = re.match(r"\|\s*(\d+)\.\s", line)
        if m and m.group(1) not in impl_line:
            impl_line[m.group(1)] = i
        if line.startswith("| Преамбула") and "pre" not in impl_line:
            impl_line["pre"] = i

    clause_proto = {r["key"]: r for r in proto_rows}
    cl = ["# Карта происхождения — пункты договора (01-dogovor-subpodryada.docx)", "",
          "Источник каждой строки — список `CONTRACT` генератора "
          f"`{GEN_REL}` (функция `build_contract`, строка {gl['defs']['build_contract']}). "
          "Место в DOCX — метки блоков из "
          f"`{PACKET_REL}/zip/01-dogovor-subpodryada.docx.txt`. "
          "Позиция протокола — строка таблицы "
          f"`{ptable_id}` в `{PACKET_REL}/zip/10-protokol-raznoglasiy.docx.txt`. "
          "«Ссылки» — явные упоминания номера пункта в других выдаваемых файлах и на странице "
          "(поиск по шаблону «п./пункт N.N»); тематические совпадения без номера не сопоставлялись.",
          "", "| Пункт | Абзацев | Блоки в 01 | Первый абзац совпал с генератором | Генератор (строки) | "
          "Позиция протокола 10 | Ссылки на пункт в других файлах и на странице |",
          "|---|---|---|---|---|---|---|"]
    gsrc = (ROOT / GEN_REL).read_text(encoding="utf-8").splitlines()

    def const_end(name):
        start = gl["defs"][name]
        return next(i for i in range(start, len(gsrc) + 1) if gsrc[i - 1].startswith("]"))

    gspan_list = gl["contract"]
    gi = 0
    for s in seq:
        if s["kind"] in ("пункт",):
            while gspan_list[gi][0] != s["key"]:
                gi += 1
            _, a, b = gspan_list[gi]
            gi += 1
            pr = clause_proto.get(s["key"])
            prs = f"R{pr['row']}: {pr['kind']}" if pr else "нет позиции"
            refs = "; ".join(refs_for_clause(s["key"])) or "—"
            consts = sorted(set(re.findall(r"\b(OVERHEADS|GUARANTEES)\b", "\n".join(gsrc[a - 1:b]))))
            extra = "".join(f"; текст — константа `{c}` `{GEN_REL}:{gl['defs'][c]}-{const_end(c)}`" for c in consts)
            label = s["key"] if re.fullmatch(r"\d+(?:\.\d+)+", s["key"]) and not s["key"].endswith(".0") \
                else f"{s['key']} (ключ генератора; вводный абзац без номера)"
            cl.append(f"| {label} | {s['last'] - s['first'] + 1} | B{s['first']:04d}–B{s['last']:04d} | "
                      f"{'да' if s['match'] else 'нет'} | `{GEN_REL}:{a}-{b}`{extra} | {prs} | {refs} |")
    new_points = [r for r in proto_rows if r["kind"].startswith("новый пункт")]
    cl += ["", "Позиции протокола к пунктам, которых в шаблоне договора нет: " +
           (", ".join(f"{r['key']} (R{r['row']})" for r in new_points) or "нет"), ""]
    (maps / "contract_clauses.md").write_text("\n".join(cl) + "\n", encoding="utf-8")

    pl = ["# Карта протокола разногласий (10-protokol-raznoglasiy.docx)", "",
          f"Таблица `{ptable_id}`: строка R1 — шапка «Редакция Заказчика / Редакция Подрядчика». "
          "Левая колонка строится `customer_edition()` из текста пункта договора (`CLAUSES`), "
          "правая — `contractor_edition()` из списка `PROTOCOL` генератора. "
          "Столбцы «совпал» сравнивают извлечённый из DOCX текст с тем, что генератор выдаёт "
          "для этой позиции сейчас.", "",
          "| Строка | Пункт | Вид позиции | Генератор (правая колонка) | Левая колонка совпала | "
          "Правая колонка совпала | Пункт в 01 |", "|---|---|---|---|---|---|---|"]
    blk = {s["key"]: s for s in seq if s["kind"] == "пункт"}
    for r in proto_rows:
        b = blk.get(r["key"])
        where = "B%04d" % b["first"] if b else "в 01 нет"
        pl.append(f"| R{r['row']} | {r['key']} | {r['kind']} | `{GEN_REL}:{r['line']}-{r['end']}` | "
                  f"{'да' if r['left_ok'] else 'нет'} | {'да' if r['right_ok'] else 'нет'} | {where} |")
    (maps / "protocol_rows.md").write_text("\n".join(pl) + "\n", encoding="utf-8")

    # ── файлы 02–09: история в git
    history = {}
    for n in zip_names:
        path = f"{KIT_REL}/{n}"
        first = sh("git", "log", "--diff-filter=A", "--format=%h %ad", "--date=short", "--", path).strip().splitlines()
        last = sh("git", "log", "-1", "--format=%h %ad", "--date=short", "--", path).strip()
        history[n] = (first[-1] if first else "—", last)
    gen_mentions = {}
    tracked = sh("git", "ls-files", "*.py", "*.php", "*.sh").split()
    for n in zip_names:
        hits = [f for f in tracked if n in (ROOT / f).read_text(encoding="utf-8", errors="ignore")]
        gen_mentions[n] = hits
    nested = ROOT / "products-storage/04-polnyy-komplekt-pto/01-30-bazovye-pakety/03-dogovor-podryada"
    nested_hits = []
    if nested.is_dir():
        for p in sorted(nested.rglob("*")):
            if p.is_file() and p.name in disk_hash:
                nested_hits.append((rel(p), sha256(p.read_bytes()) == disk_hash[p.name]))

    ls_all = sh("git", "ls-files").splitlines()
    ref_like = [f for f in ls_all if re.search(r"(?i)(referens|reference|skan|scan|original|obrazec).*(dogovor|protokol)|"
                                                 r"(dogovor|protokol).*(referens|reference|skan|scan|original)", f)]
    other_repo = ROOT.parent / "ai-business-os"
    other_hits = []
    if (other_repo / ".git").exists():
        other_hits = [f for f in subprocess.run(["git", "ls-files"], cwd=other_repo, capture_output=True,
                                                text=True).stdout.splitlines()
                      if re.search(r"(?i)(dogovor[-_]?pod|subpodr|protokol[-_]?razn|contract_ref|referens)", f)]

    # ── START-HERE §4 против архива
    sh_text = extracted["00-START-HERE.txt"]["text"]
    sh_sec4 = sh_text.split("4. ЧТО В АРХИВЕ", 1)[1].split("Файлы пронумерованы", 1)[0]
    sh_list = re.findall(r"^\s{3}(\d\d-[\w.-]+\.\w+)", sh_sec4, flags=re.M)
    sh_count = re.search(r"файлов в архиве: (\d+)", sh_text)
    manifest_list = re.findall(r"`([^`]+)`", (KIT / "MANIFEST.md").read_text(encoding="utf-8"))
    kit_files = [n for n, _, _ in mod.KIT_FILES]
    page_counts = sorted(set(re.findall(r"(\d+) (?:готовых )?файл", html_text)))

    shutil.rmtree(tmp)

    # ─────────────────────────────── отчёт ───────────────────────────────
    R = []
    add = R.append
    add("# MB001 · Р-026 · P7 — проверочный пакет для нормативного перечита")
    add("")
    add("Дата сборки пакета: 29.09.2026. Пакет собран скриптом "
        f"`{PACKET_REL}/_tools/build_packet.py`, проверен `verify_packet.py` рядом с ним.")
    add("")
    add("## 1. Режим и границы")
    add("")
    add("- Только подготовка материалов для независимого перечита по Р-026. Юридических выводов, "
        "оценок «верно / неверно», предложений правок в пакете нет.")
    add("- Продукт, договор, протокол, инструкция, страница товара, генератор, цены, SKU, #315, `main` "
        "и другие ветки не менялись. Коммит пакета добавляет только "
        "`tools/candidates/MB001_R026_P7_REVIEW_PACKET.md` и каталог "
        f"`{PACKET_REL}/`.")
    add("- В таблицах записаны факты: текст, место, источник, связь либо отсутствие источника. "
        "Слова «совпал / не совпал» в картах означают побайтовое или построчное сравнение двух "
        "текстов, а не оценку содержания.")
    add("- Полные извлечения лежат в каталоге доказательств; здесь — индекс и сводки.")
    add("")
    add(f"## 2. Preflight и версия #{args.pr}")
    add("")
    add("| Что | Команда | Результат |")
    add("|---|---|---|")
    add(f"| PR | GitHub API `pull_request_read get #{args.pr}` | Draft, open; ветка `{args.pr_head_ref}`; "
        f"голова `{args.pr_head_sha}`; база `main` `{args.pr_base_sha}` |")
    add(f"| Ветка пакета | `git checkout -B <ветка пакета> {args.pr_head_sha}` | пакет стоит на голове #{args.pr}: "
        f"`git merge-base --is-ancestor {args.pr_head_sha[:7]} HEAD` → {'да' if is_anc else 'нет'} |")
    add("| Рабочее дерево до создания файлов | `git status --porcelain \\| wc -l` | 0 (измерено 29.09.2026 до первого файла пакета) |")
    add(f"| Файлы продукта не изменены относительно #{args.pr} | `git diff --quiet {args.pr_head_sha[:7]} -- products-storage products {CONFIG_REL} {IMPL_REL}` | "
        f"{'изменений нет' if kit_same else 'ЕСТЬ ИЗМЕНЕНИЯ'} |")
    add(f"| PHP | `php -r 'echo PHP_VERSION;'` | {php_version} |")
    add(f"| ZipArchive | `php -r \"echo class_exists('ZipArchive');\"` | {zip_class} |")
    add(f"| Реальная выдача | `php {PACKET_REL}/_tools/build_p7.php . <tmp>` → `mvb_build_product_zip('p7')` "
        f"(`{CONFIG_REL}:{sh('grep', '-n', 'function mvb_build_product_zip', CONFIG_REL).split(':')[0]}`) | архив собран, "
        f"{len(entries)} записей |")
    add(f"| Каталог продукта | `{CONFIG_REL}`: `'p7' => dir 03-dogovor-podryada, zip 03-dogovor-podryada.zip` | "
        f"`{KIT_REL}/` |")
    add(f"| Страница P7 | — | `{PAGE_REL}` |")
    add(f"| START-HERE | — | `{KIT_REL}/00-START-HERE.txt` (в ZIP) |")
    add(f"| MANIFEST | — | `{KIT_REL}/MANIFEST.md` (служебный, в ZIP не идёт) |")
    add(f"| Генератор | `ls products-storage/build_paid_03.py` | `{GEN_REL}` — пишет 01, 10, 00-INSTRUKCIYA.docx/.pdf, 00-START-HERE.txt |")
    add(f"| Отчёт реализации #{args.pr} | — | `{IMPL_REL}` (в дереве головы #{args.pr}) |")
    add(f"| Gap-отчёт (#307) | `git fetch origin {GAP_REF}`; `git cat-file -e {GAP_SHA[:7]}:{GAP_REL}` | "
        f"{'найден' if gap_present else 'НЕ НАЙДЕН'}: ветка `{GAP_REF}`, коммит `{GAP_SHA}`; в дереве #{args.pr} файла нет |")
    add("| Исходный договор / референс владельца | поиск по `git ls-files` обоих репозиториев (раздел 7) | "
        "в репозиториях не найден; отчёт реализации #315 §0: «В репозиторий файл не кладётся» |")
    add("")
    add("## 3. Состав реального ZIP с SHA-256")
    add("")
    add(f"Архив: `03-dogovor-podryada.zip`, {len(zbytes)} байт, SHA-256 `{zsha}`.")
    add(f"Повторная сборка в чистый каталог: SHA-256 `{zsha2}` — "
        f"{'байты архива совпали' if zsha == zsha2 else 'байты архива различаются'}; "
        f"имена и SHA-256 записей {'совпали' if same_entries else 'различаются'}.")
    add("")
    add("| # | Файл в ZIP | Тип | Байт | SHA-256 записи | SHA-256 мастера на диске |")
    add("|---|---|---|---|---|---|")
    for i, e in enumerate(sorted(entries, key=lambda e: e["name"]), 1):
        kind = KIND.get(Path(e["name"]).suffix.lower(), "другое")
        dh = disk_hash.get(e["name"], "—")
        add(f"| {i} | `{e['name']}` | {kind} | {e['size']} | `{e['sha256']}` | "
            f"{'= запись' if dh == e['sha256'] else f'`{dh}`'} |")
    add("")
    by_kind = {}
    for e in entries:
        by_kind.setdefault(KIND.get(Path(e["name"]).suffix.lower(), "другое"), []).append(e["name"])
    add("По типам: " + "; ".join(f"{k} — {len(v)} ({', '.join(sorted(v))})" for k, v in sorted(by_kind.items())) + ".")
    add("")
    junk_re = re.compile(r"^(\.|__MACOSX|~\$|Thumbs\.db)")
    junk = sum(bool(junk_re.match(Path(e["name"]).name)) for e in entries)
    add("Служебные и лишние файлы:")
    add(f"- Каталог `{KIT_REL}/` на диске: {len(disk)} файлов. В ZIP не вошли: "
        + ", ".join(f"`{n}`" for n in disk if n not in zip_names)
        + " — это список `$service` функции выдачи (`.htaccess`, `00-PISMO-POSLE-POKUPKI.txt`, `MANIFEST.md`).")
    add(f"- Записей-каталогов в ZIP: {sum(e['is_dir'] for e in entries)}; вложенных путей: "
        f"{sum('/' in e['name'] for e in entries)}; скрытых и системных имён (`.`, `__MACOSX`, `~$`, `Thumbs.db`): "
        f"{junk}.")
    add(f"- Состав ZIP = файлы каталога минус `$service`: {'да' if zip_names == disk_payload else 'нет'}.")
    add(f"- Перечень START-HERE §4 ({len(sh_list)} имён) = состав ZIP: {'да' if sorted(sh_list) == zip_names else 'нет'}; "
        f"«файлов в архиве: {sh_count.group(1) if sh_count else '—'}».")
    add(f"- Таблица состава инструкции (`KIT_FILES`, {len(kit_files)} имён) = состав ZIP: "
        f"{'да' if sorted(kit_files) == zip_names else 'нет'}.")
    add(f"- Перечень `MANIFEST.md` ({len(manifest_list)} имён) = состав ZIP: "
        f"{'да' if sorted(manifest_list) == zip_names else 'нет'}.")
    add(f"- Числа файлов, названные на странице (шаблон «N файл…/N готовых файл…»): {', '.join(page_counts) or 'нет'}.")
    add("")
    add("## 4. Полный индекс текстов и мест в документах")
    add("")
    add("Каждой записи ZIP — полное извлечение. Метки мест: DOCX — `B<блок> P` или "
        "`B<блок> T<таблица> R<строка> C<колонка> p<абзац>`, плюс части колонтитулов, сносок, примечаний; "
        "XLSX — `Лист!Ячейка` с формулой и кэшем; PDF — страницы; TXT — `L<строка>`.")
    add("")
    add("| Файл в ZIP | Тип | Сводка извлечения | Полный текст |")
    add("|---|---|---|---|")
    for row in index:
        add(f"| `{row['name']}` | {row['kind']} | {row['summary']} | `{row['evidence']}` |")
    add("")
    add("Способ извлечения: DOCX — разбор `word/document.xml`, `word/header*.xml`, `word/footer*.xml`, "
        "`word/footnotes.xml`, `word/endnotes.xml`, `word/comments.xml` (lxml), включая таблицы и вложенные "
        "таблицы; XLSX — openpyxl дважды (формулы и кэшированные значения), ошибкой считается тип `e` или "
        "значение из набора `#REF!`, `#DIV/0!`, `#VALUE!`, `#NAME?`, `#N/A`, `#NUM!`, `#NULL!`; PDF — "
        "`pdfinfo`, `pdftotext -layout -enc UTF-8`.")
    add("")
    add("## 5. Карта происхождения разделов договора")
    add("")
    add(f"Договор `01-dogovor-subpodryada.docx` пишет `build_contract()` (`{GEN_REL}:{gl['defs']['build_contract']}`) "
        f"из списка `CONTRACT` (`{GEN_REL}:{gl['defs']['CONTRACT']}`). Порядок блоков тела смоделирован по коду "
        f"`build_contract` и сверен с извлечением: блоков {total_blocks}, модель "
        f"{'совпала' if sim_ok else 'НЕ совпала'} (первый абзац каждого пункта и каждый заголовок раздела "
        "сравнены с текстом генератора). Пересборка генератором во временный каталог "
        f"(`build(out, pdf=False)`): " + "; ".join(
            f"`{n}` — {'байты совпали' if a == b else 'байты различаются'}" for n, (a, b) in regen_cmp.items()) + ".")
    add("")
    add("Референс владельца, по которому написан текст (отчёт реализации §0, §2.1), в репозиториях отсутствует: "
        "столбец «Источник» называет файлы, которые в дереве есть, и строку отчёта реализации, где записано, "
        "что взято из референса и что обобщено полями.")
    add("")
    add("| Фрагмент / раздел | Выдаваемый файл и место | Источник в репозитории | Генератор / ручной файл | Связанная страница или инструкция |")
    add("|---|---|---|---|---|")
    heads = [s for s in seq if s["kind"] == "заголовок раздела"]
    contract_lines = {k: (a, b) for k, a, b in gl["contract"]}
    g_heads = [(k, a) for k, a, b in gl["contract"] if k == "h"]

    def section_rows(n_idx):
        h = heads[n_idx]
        nxt = heads[n_idx + 1]["first"] if n_idx + 1 < len(heads) else total_blocks + 1
        clauses = [s for s in seq if s["kind"] == "пункт" and h["first"] < s["first"] < nxt]
        tables = [s for s in seq if s["kind"].startswith("таблица") and h["first"] < s["first"] < nxt]
        return h, nxt, clauses, tables

    for fname, pat in (("01-dogovor-subpodryada.docx",
                        r"\(01\)|01-dogovor|файла 01|шаблон 01|шаблона 01|текст 01|^Договор субподряда$"),
                       ("10-protokol-raznoglasiy.docx",
                        r"\(10\)|10-protokol|протокол 10|протокола 10|^Протокол разногласий")):
        rx = re.compile(pat)
        pg = [f"L{b['line']}" for b in page["blocks"] if rx.search(b["text"])]
        st = [l for l, t in extracted["00-START-HERE.txt"]["lines"] if rx.search(t)]
        ins = list(dict.fromkeys(l.split(" p")[0] for l, t in extracted["00-INSTRUKCIYA.docx"]["lines"]
                                 if rx.search(t)))
        add(f"| Файл `{fname}` целиком | `{fname}` (весь файл) | `{GEN_REL}` | генератор | "
            f"страница `{PAGE_REL}`: {', '.join(pg) or '—'}; START-HERE: {', '.join(st) or '—'}; "
            f"`00-INSTRUKCIYA.docx`: {', '.join(ins) or '—'} |")
    add(f"| Служебные строки, заголовок, место и дата, преамбула | `01` B0001–B0007 "
        "(B0004 — таблица T1 место/дата) | "
        f"`{GEN_REL}:{gl['defs']['build_contract']}` (`service_lines`, `place_and_date` "
        f"`:{gl['defs']['place_and_date']}`, `parties_preamble` `:{gl['defs']['parties_preamble']}`); "
        f"отчёт реализации `{IMPL_REL}:{impl_line.get('pre', '—')}` | генератор | "
        f"инструкция «Как заполнять» (служебные строки, имена полей сторон): `{PACKET_REL}/zip/00-INSTRUKCIYA.docx.txt` |")
    for i in range(len(heads)):
        h, nxt, clauses, tables = section_rows(i)
        num = h["key"].split(".")[0]
        gh_line = g_heads[i][1]
        g_end = (g_heads[i + 1][1] - 1) if i + 1 < len(g_heads) else gl["contract"][-1][2]
        protos = [clause_proto[c["key"]] for c in clauses if c["key"] in clause_proto]
        protos += [r for r in proto_rows if r["kind"].startswith("новый пункт") and r["key"].split(".")[0] == num]
        pk = {}
        for r in protos:
            k = r["kind"].split(" (")[0].split(",")[0]
            pk[k] = pk.get(k, 0) + 1
        proto_txt = (f"; протокол `10`: {len(protos)} поз. (" + ", ".join(f"{k} {v}" for k, v in pk.items()) +
                     "), строки " + ", ".join(f"R{r['row']}" for r in sorted(protos, key=lambda r: r['row']))) if protos else "; протокол `10`: позиций нет"
        clause_range = (f"пункты {clauses[0]['key']}–{clauses[-1]['key']} ({len(clauses)} блоков "
                        f"`CONTRACT`)") if clauses else "нумерованных пунктов нет"
        tbl = "".join(f"; {t['kind']} B{t['first']:04d}" for t in tables)
        refs = refs_for_section(num)
        add(f"| {h['key']} | `01` B{h['first']:04d}–B{nxt - 1:04d}, {clause_range}{tbl}{proto_txt} | `{GEN_REL}:{gh_line}-{g_end}`; отчёт реализации "
            f"`{IMPL_REL}:{impl_line.get(num, '—')}` | генератор (`CONTRACT`) | "
            f"{'; '.join(refs) if refs else 'явных ссылок на номер раздела или пункта нет'} |")
    add("")
    add(f"Полная карта по каждому пункту — `{PACKET_REL}/maps/contract_clauses.md`; по каждой позиции "
        f"протокола — `{PACKET_REL}/maps/protocol_rows.md`.")
    add("")
    add(f"Протокол разногласий `10-protokol-raznoglasiy.docx`: `build_protocol()` "
        f"(`{GEN_REL}:{gl['defs']['build_protocol']}`), таблица `{ptable_id}`, позиций {len(proto_rows)}: "
        + ", ".join(f"{k} — {v}" for k, v in kinds.items()) +
        f". Левая колонка всех позиций совпала с текстом пункта из генератора: "
        f"{'да' if all(r['left_ok'] for r in proto_rows) else 'нет'}; правая колонка совпала с `PROTOCOL`: "
        f"{'да' if all(r['right_ok'] for r in proto_rows) else 'нет'}.")
    add("")
    add("Остальные выдаваемые файлы и их источник:")
    add("")
    add("| Файл | Генератор / ручной файл | Первое появление в git | Последнее изменение | Упоминания имени в *.py/*.php/*.sh |")
    add("|---|---|---|---|---|")
    generated = {"01-dogovor-subpodryada.docx", "10-protokol-raznoglasiy.docx", "00-INSTRUKCIYA.docx",
                 "00-START-HERE.txt", "00-INSTRUKCIYA.pdf"}
    for n in zip_names:
        src = (f"генератор `{GEN_REL}`" + (" (PDF печатает LibreOffice из 00-INSTRUKCIYA.docx, `build_pdf`)"
                                             if n.endswith(".pdf") else "")) if n in generated else \
            "генератор в репозитории не найден: файл лежит готовым"
        add(f"| `{n}` | {src} | {history[n][0]} | {history[n][1]} | "
            f"{', '.join(f'`{m}`' for m in gen_mentions[n]) or '—'} |")
    add("")
    add("## 6. Тексты страницы, START-HERE и MANIFEST")
    add("")
    add(f"Страница `{PAGE_REL}` (SHA-256 `{page_sha}`):")
    add(f"- видимый текст по блокам с номерами строк исходника — `{PACKET_REL}/site/p7-dogovor-podryada.visible.txt` "
        f"({len(page['blocks'])} блоков, из них с пометкой [скрыт] {sum(b['hidden'] for b in page['blocks'])});")
    add(f"- FAQ: видимый блок `section.faq` и `FAQPage` из JSON-LD — `{PACKET_REL}/site/p7-dogovor-podryada.faq.txt` "
        f"(вопросов в JSON-LD: {len(ldq)});")
    add(f"- JSON-LD целиком — `{PACKET_REL}/site/p7-dogovor-podryada.jsonld.json` ({len(page['jsonld'])} блок(а)).")
    add(f"- Блок «Как это выглядит внутри» генерируется `{PREVIEW_REL}` из файлов комплекта "
        f"(описание PR #{args.pr}, коммит `{args.pr_head_sha[:7]}`).")
    add("")
    add(f"START-HERE (`{KIT_REL}/00-START-HERE.txt`, SHA-256 `{disk_hash['00-START-HERE.txt']}`), полный текст:")
    add("")
    add("```text")
    add(sh_text.rstrip())
    add("```")
    add("")
    mf = (KIT / "MANIFEST.md").read_text(encoding="utf-8")
    add(f"MANIFEST (`{KIT_REL}/MANIFEST.md`, SHA-256 `{disk_hash['MANIFEST.md']}`, в ZIP не выдаётся), полный текст:")
    add("")
    add("```text")
    add(mf.rstrip())
    add("```")
    add("")
    add(f"Второй служебный файл каталога — `00-PISMO-POSLE-POKUPKI.txt` (в ZIP не выдаётся): "
        f"`{PACKET_REL}/kit-service/00-PISMO-POSLE-POKUPKI.txt.txt`; копия MANIFEST с номерами строк — "
        f"`{PACKET_REL}/kit-service/MANIFEST.md.txt`.")
    add("")
    add("## 7. Найденные исходники и отчёты")
    add("")
    add("| Что | Путь | Коммит / ветка | SHA-256 | Копия в пакете |")
    add("|---|---|---|---|---|")
    add(f"| Генератор P7 | `{GEN_REL}` | #{args.pr}, `{args.pr_head_sha[:7]}` | `{sha256((ROOT / GEN_REL).read_bytes())}` | нет (в дереве) |")
    add(f"| Сборка выдачи | `{CONFIG_REL}` (`mvb_products`, `mvb_build_product_zip`) | `{args.pr_head_sha[:7]}` | "
        f"`{sha256((ROOT / CONFIG_REL).read_bytes())}` | нет (в дереве) |")
    add(f"| Тест комплекта | `{TEST_REL}` | #{args.pr} | `{sha256((ROOT / TEST_REL).read_bytes())}` | нет (в дереве) |")
    add(f"| Страница P7 | `{PAGE_REL}` | #{args.pr} | `{page_sha}` | извлечения в `site/` |")
    add(f"| Отчёт реализации #{args.pr} | `{IMPL_REL}` | `{args.pr_head_ref}` | `{sha256(impl_bytes)}` | "
        f"`{PACKET_REL}/source-reports/MB001_R2_P7_IMPLEMENTATION_REPORT.md` (байт в байт) |")
    if gap_present:
        add(f"| Gap-отчёт (спецификация пробелов, #307) | `{GAP_REL}` | `{GAP_REF}` `{GAP_SHA}` | `{sha256(gap_bytes)}` | "
            f"`{PACKET_REL}/source-reports/MB001_R2_P7_GAP_SPEC.md` (байт в байт из `git show`) |")
    else:
        add(f"| Gap-отчёт (#307) | `{GAP_REL}` | `{GAP_REF}` | — | не найден в локальном git |")
    add("| Исходный договор / референс владельца | не найден | — | — | — |")
    add("")
    add("Поиск референса: `git ls-files` сайта по шаблону имён «referens|reference|skan|scan|original|obrazec» рядом "
        f"с «dogovor|protokol» — совпадений: {len(ref_like)}{(' (' + ', '.join(ref_like) + ')') if ref_like else ''}; "
        "`git ls-files` ai-business-os по шаблону «dogovor-pod|subpodr|protokol-razn|contract_ref|referens» — "
        f"совпадений: {len(other_hits)}{(' (' + ', '.join(other_hits[:10]) + ')') if other_hits else ''}.")
    add("")
    if nested_hits:
        add(f"Копии файлов P7 вне каталога продукта: в `{rel(nested)}/` найдено "
            f"{len(nested_hits)} файлов с именами из P7; SHA-256 совпадает с текущим мастером у "
            f"{sum(ok for _, ok in nested_hits)} ({', '.join(Path(n).name for n, ok in nested_hits if ok)}), "
            f"отличается у {sum(not ok for _, ok in nested_hits)} "
            f"({', '.join(Path(n).name for n, ok in nested_hits if not ok)}). "
            "Этот каталог не адресуется sku `p7` и в ZIP p7 не входит.")
        add("")
    add("## 8. Ограничения пакета")
    add("")
    add("- **Референса владельца нет в репозиториях.** Отчёт реализации #315 называет его сканом на 23 страницы, "
        "переданным 29.09.2026, и прямо говорит, что файл не кладётся в репозиторий. Поэтому дословное "
        "происхождение каждого пункта от референса этим пакетом не проверяется: карта доходит до генератора "
        "и до строки отчёта реализации, а не до страницы референса.")
    add("- Файлы 02–09 генератора в репозитории не имеют: их текст извлечён, происхождение — только история git.")
    add("- SHA-256 архива зависит от времени правки файлов в рабочей копии и от порядка обхода каталога "
        "(`ZipArchive` пишет mtime, `RecursiveDirectoryIterator` не сортирует). В другой рабочей копии хеш "
        "архива будет иным; стабильны имена и SHA-256 записей.")
    add("- Вёрстка не проверялась: извлечён текст, а не вид страниц. Скрытость блоков страницы определена по "
        "атрибутам `hidden` / `aria-hidden`, CSS не вычислялся. Страница не открывалась в браузере.")
    add("- PDF извлечён `pdftotext`: переносы и колонки восстановлены по раскладке, а не по структуре документа.")
    add("- В XLSX формул нет; значения, которые покупатель введёт сам, пакет не моделирует.")
    add("- Связь «раздел договора → инструкция / страница / другие файлы» установлена только по явным ссылкам "
        "на номер пункта или раздела. Тематические совпадения без номера не сопоставлялись.")
    add("- Нормативные акты, на которые ссылаются тексты, в пакет не входят: пакет не сверяет текст с нормой, "
        "это работа перечита.")
    add("")
    add("## 9. Команды фактической проверки и их результаты")
    add("")
    add("```bash")
    add(f"# голова #{args.pr} и ветка пакета")
    add(f"git fetch origin {args.pr_head_ref} && git rev-parse FETCH_HEAD    # {args.pr_head_sha}")
    add(f"git merge-base --is-ancestor {args.pr_head_sha[:7]} HEAD && echo ok     # {'ok' if is_anc else 'нет'}")
    add("# реальная выдача")
    add(f"php {PACKET_REL}/_tools/build_p7.php \"$PWD\" \"$(mktemp -d)\"")
    add(f"unzip -Z1 <архив> | sort                                   # {len(entries)} имён, см. раздел 3")
    add("# пересборка пакета и проверка готовности")
    add(f"python3 {PACKET_REL}/_tools/build_packet.py --pr-head-sha {args.pr_head_sha} "
        f"--pr-head-ref {args.pr_head_ref} --pr-base-sha {args.pr_base_sha}")
    add(f"python3 {PACKET_REL}/_tools/verify_packet.py --pr-head-sha {args.pr_head_sha}")
    add("```")
    add("")
    add("Результаты сборки (этот прогон):")
    add("")
    add(f"- `mvb_build_product_zip('p7')` дважды: {len(entries)} записей; архив `{zsha[:16]}…` / `{zsha2[:16]}…`; "
        f"записи {'совпали' if same_entries else 'различаются'}.")
    add(f"- Состав ZIP = каталог минус `$service`: {'да' if zip_names == disk_payload else 'нет'}; "
        f"= START-HERE §4: {'да' if sorted(sh_list) == zip_names else 'нет'}; = `KIT_FILES`: "
        f"{'да' if sorted(kit_files) == zip_names else 'нет'}.")
    add(f"- Пересборка генератором: " + "; ".join(
        f"`{n}` {'=' if a == b else '≠'}" for n, (a, b) in regen_cmp.items()) + ".")
    add(f"- Модель блоков договора сверена с извлечением: {'да' if sim_ok else 'нет'}; протокол: левая колонка "
        f"{sum(r['left_ok'] for r in proto_rows)}/{len(proto_rows)}, правая {sum(r['right_ok'] for r in proto_rows)}/{len(proto_rows)}.")
    add("")
    add("Результаты проверки готовности (DONE) — вывод `verify_packet.py` на момент коммита:")
    add("")
    add("<!-- VERIFY:BEGIN -->")
    add("(заполняется verify_packet.py --write)")
    add("<!-- VERIFY:END -->")
    add("")
    REPORT.write_text("\n".join(R) + "\n", encoding="utf-8")

    facts = {"zip_sha256": zsha, "zip_bytes": len(zbytes),
             "entries": [{"name": e["name"], "size": e["size"], "sha256": e["sha256"]}
                         for e in sorted(entries, key=lambda e: e["name"])]}
    (PACKET / "zip_sha256.json").write_text(json.dumps(facts, ensure_ascii=False, indent=2) + "\n",
                                            encoding="utf-8")
    print(f"ZIP {zsha} {len(entries)} записей; отчёт {rel(REPORT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
