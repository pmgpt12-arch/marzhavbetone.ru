#!/usr/bin/env python3
"""Read-only evidence extraction for one owner-approved Excel reference."""
import collections, copy, hashlib, json, shutil, sys, zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'P1_04_Uchet_Business_Design_v1.xlsx'
if not SOURCE.exists():
    SOURCE = ROOT.parent / 'excel-design-reference-20261006/P1_04_Uchet_Business_Design_v1.xlsx'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def sha(data): return hashlib.sha256(data).hexdigest()
def xml(obj): return ET.tostring(obj.to_tree(), encoding='unicode')
def dump(name, obj):
    (ROOT/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str)+'\n')

def invariants(w):
    """Preserve semantics, status formulas, input contracts and actual print settings."""
    result = {'sheetnames': w.sheetnames, 'calculation': xml(w.calculation),
              'defined_names': [xml(v) for v in w.defined_names.values()], 'sheets': {}}
    for s in w:
        cells = {c.coordinate: {'value': c.value, 'data_type': c.data_type,
                  'number_format': c.number_format, 'protection': xml(c.protection)}
                 for c in s._cells.values() if not isinstance(c, openpyxl.cell.cell.MergedCell)}
        result['sheets'][s.title] = {
            'cells': cells, 'defined_names': [xml(v) for v in s.defined_names.values()],
            'data_validations': xml(s.data_validations), 'protection': xml(s.protection),
            'merged_ranges': sorted(map(str, s.merged_cells.ranges)),
            'conditional_formatting': [xml(r) for rules in s.conditional_formatting._cf_rules.values() for r in rules],
            'conditional_formatting_targets': [str(k.sqref) for k in s.conditional_formatting._cf_rules],
            'auto_filter': xml(s.auto_filter), 'print_area': str(s.print_area),
            'print_title_rows': s.print_title_rows, 'print_title_cols': s.print_title_cols,
            'page_setup': xml(s.page_setup), 'page_margins': xml(s.page_margins),
            'print_options': xml(s.print_options), 'header_footer': xml(s.HeaderFooter),
            'row_breaks': xml(s.row_breaks), 'col_breaks': xml(s.col_breaks),
            'sheet_properties': xml(s.sheet_properties), 'freeze_panes': s.freeze_panes,
            'sheet_state': s.sheet_state, 'tables': [xml(t) for t in s.tables.values()]}
    return result

def main():
    raw = SOURCE.read_bytes()
    w = openpyxl.load_workbook(SOURCE, data_only=False)
    with zipfile.ZipFile(SOURCE) as z:
        assert z.testzip() is None
        parts = [{'path': n, 'bytes': len(z.read(n)), 'sha256': sha(z.read(n))} for n in z.namelist()]
        assert len(parts) == len({v['path'] for v in parts})
        styles_xml = z.read('xl/styles.xml')
        relationships = {n.attrib['Id']: n.attrib['Target'] for n in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        native_styles = {}
        for sheet in ET.fromstring(z.read('xl/workbook.xml')).find('s:sheets', NS):
            target = relationships[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            part = target.lstrip('/') if target.startswith('/') else 'xl/'+target
            native_styles[sheet.attrib['name']] = {n.attrib['r']:int(n.attrib.get('s','0'))
                for n in ET.fromstring(z.read(part)).findall('s:sheetData/s:row/s:c', NS)}
        styles_root = ET.fromstring(styles_xml)
        styles = {k: [ET.tostring(c, encoding='unicode') for c in node]
                  for k in ['fonts','fills','borders','cellXfs','cellStyleXfs','numFmts','dxfs','cellStyles']
                  if (node := styles_root.find('s:'+k, NS)) is not None}
    sheets=[]
    for s in w:
        styles_by_id=collections.defaultdict(list)
        for coord, style_id in native_styles[s.title].items(): styles_by_id[str(style_id)].append(coord)
        assert all(0 <= int(i) < len(styles['cellXfs']) for i in styles_by_id)
        sheets.append({'name': s.title, 'dimensions': s.calculate_dimension(),
          'max_row': s.max_row, 'max_column': s.max_column,
          'populated_cells': sum(c.value is not None for c in s._cells.values()),
          'formula_count': sum(c.data_type == 'f' for c in s._cells.values()),
          'formula_styles': dict(collections.Counter(c.style_id for c in s._cells.values() if c.data_type == 'f')),
          'blank_styles': dict(collections.Counter(c.style_id for c in s._cells.values() if c.value is None)),
          'styles_by_id_coordinates': dict(styles_by_id),
          'column_dimensions': {k: xml(v) for k,v in s.column_dimensions.items()},
          'row_dimensions': {str(k): dict(v) for k,v in s.row_dimensions.items()},
          'sheet_format': xml(s.sheet_format), 'sheet_view': xml(s.sheet_view),
          'data_validations': xml(s.data_validations), 'sheet_protection': xml(s.protection),
          'print_area': str(s.print_area), 'page_setup': xml(s.page_setup),
          'page_margins': xml(s.page_margins), 'print_options': xml(s.print_options),
          'merged_ranges': sorted(map(str,s.merged_cells.ranges))})
    samples=[('Взаиморасчёты','A1','title'),('Взаиморасчёты','A2','instructions'),
             ('Взаиморасчёты','B4','table_header'),('Взаиморасчёты','B5','date_input_candidate'),
             ('Взаиморасчёты','C5','operation_input_validation'),('Взаиморасчёты','E5','amount_input_candidate'),
             ('Взаиморасчёты','G5','calculation_formula'),('Взаиморасчёты','I5','status_formula'),
             ('Сводка','B2','brand'),('Сводка','B3','summary_heading'),
             ('Акт сверки','A1','print_form_heading'),('Акт сверки','C3','print_status_formula')]
    roles=[]
    for sn,coord,role in samples:
        c=w[sn][coord]
        roles.append({'observed_role': role, 'sheet': sn, 'cell': coord,
          'value': c.value, 'style_id': c.style_id, 'font': xml(c.font), 'fill': xml(c.fill),
          'border': xml(c.border), 'alignment': xml(c.alignment), 'number_format':c.number_format,
          'protection': xml(c.protection),
          'semantic_basis': 'Formula/value/header and native validation, never color alone; input_candidate is not definitive target mapping.'})
    baseline=invariants(w)
    checks=[]
    checks.append({'case':'unmodified_reference','passed':invariants(w)==baseline})
    checks.append({'case':'native_style_references_resolve','passed':all(0 <= i < len(styles['cellXfs']) for sm in native_styles.values() for i in sm.values())})
    mutations=[('formula_loss',lambda q:setattr(q['Взаиморасчёты']['G5'],'value',0)),
       ('validation_loss',lambda q:setattr(q['Взаиморасчёты'].data_validations,'dataValidation',[])),
       ('sheet_protection_changed',lambda q:setattr(q['Взаиморасчёты'].protection,'sheet',True)),
       ('cell_protection_changed',lambda q:setattr(q['Взаиморасчёты']['B5'],'protection',openpyxl.styles.Protection(locked=False))),
       ('status_formula_changed',lambda q:setattr(q['Акт сверки']['C3'],'value','=1')),
       ('print_area_changed',lambda q:setattr(q['Акт сверки'],'print_area','A1:B2')),
       ('input_number_format_changed',lambda q:setattr(q['Взаиморасчёты']['B5'],'number_format','General'))]
    for label,mutate in mutations:
        q=openpyxl.load_workbook(SOURCE, data_only=False); mutate(q)
        checks.append({'case':label,'passed':invariants(q)!=baseline})
    assert all(t['passed'] for t in checks)
    assert SOURCE.read_bytes()==raw
    ref={'library_file_id':'libfile_b61f0d4293648191b05c951f04497af3','filename':SOURCE.name,
         'sha256':sha(raw),'size_bytes':len(raw),
         'owner_decision':'Этот вариант намного лучше, применяем такое оформление',
         'owner_decision_date':'2026-10-06',
         'owner_decision_scope':'Business visual design reference; not independent legal, recalculation or full release acceptance.'}
    contract={'schema_version':1,'status':'EXTRACTED_REFERENCE_CONTRACT_NOT_TARGET_APPLICATION',
        'reference':ref,'extraction':{'tool':'Python openpyxl + ZIP/XML','openpyxl':openpyxl.__version__,
         'read_only_reference':True,'zip_crc':'PASS','duplicate_zip_members':False,'model_calls':0},
        'styles_xml_sha256':sha(styles_xml),'style_tables':styles,'sheets':sheets,'grounded_role_samples':roles,
        'coordinate_style_origin':'Raw native worksheet c/@s; default 0. Openpyxl may synthesize merged-cell border styles; those must never be presented as native style IDs.',
        'rules':{'semantic_roles':'Require exact target sheet/ranges/formulas/validations from target evidence; do not infer input/output by fill or protection.',
         'style_ids':'Local to this reference; copy actual style components into target, not numeric style IDs.',
         'colors':'Retain exact RGB/theme/index/tint objects; do not replace theme-based colors with guessed RGB.',
         'fonts_fills_borders_alignment':'Select only from actual role samples/style tables; no decorative charts or macros.',
         'preserve':'Formulas, constants, text including warning/status instructions, validation native XML, names, conditional logic, cell/sheet protection, number formats and actual print contract remain target-owned.',
         'layout':'Reference widths/heights describe reference only. Target mapping must justify fit and retain readable input/status and print output; never copy dimensions blindly.',
         'summary':'Owner-approved compact summary is observed on Сводка; new summary formulas/content require separately scoped task and validation.'},
        'negative_gates':checks,'boundaries':{'native_Excel_render':'NOT_PERFORMED_THIS_TASK','target_modified':False,
        'recalculation':'NOT_PERFORMED','legal_acceptance':'NOT_PERFORMED','sale_ready':False,
        'target_style_mapping':'REQUIRES_NEXT_ATOMIC_TASK'}}
    dump('contract.json',contract); dump('preservation-snapshot.json',baseline)
    dump('zip-parts.json',parts); dump('verification.json',{'checks':checks,'source_sha256_after':sha(SOURCE.read_bytes()),'reference_unchanged':True})
    if SOURCE.resolve() != (ROOT/SOURCE.name).resolve():
        shutil.copyfile(SOURCE,ROOT/SOURCE.name)
    print(json.dumps({'reference':ref,'sheet_counts':[{k:s[k] for k in ['name','max_row','max_column','formula_count']} for s in sheets],
       'style_tables':{k:len(v) for k,v in styles.items()},'negative_gates':checks},ensure_ascii=False))
if __name__=='__main__': main()
