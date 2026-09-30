import openpyxl,sys,glob,os
B='<scratch>/calc'
def v(x): 
    return x.strftime('%d.%m.%Y') if hasattr(x,'strftime') else x
for s in ['C1','C3','C4']:
    print('#####',s)
    wb=openpyxl.load_workbook(f'{B}/{s}/02-proverka-i-kontrol-otveta.xlsx',data_only=True)
    m=wb['Маршрут']; print('02 Маршрут C22..C28:',[v(m[f'C{r}'].value) for r in (22,23,24,25,27,28)])
    ws=wb['Контроль ответа']
    for r in range(5,9):
        if ws[f'B{r}'].value: print(f'02 КО r{r}:',v(ws[f'B{r}'].value),'|',v(ws[f'D{r}'].value),'|E=',ws[f'E{r}'].value,'|',ws[f'F{r}'].value,'| G=',ws[f'G{r}'].value)
    wb=openpyxl.load_workbook(f'{B}/{s}/04-uchet-raschetov-i-otpravok.xlsx',data_only=True)
    ws=wb['Взаиморасчёты']
    print('04 Взаим I5..I7:',[ws[f'I{r}'].value for r in (5,6,7)],'M4..M13:',[ws[f'M{r}'].value for r in range(4,14)])
    ws=wb['Долг по актам']
    for r in (5,6,7): print('04 Долг r',r,[v(ws[f'{c}{r}'].value) for c in 'BDEFGHIJKL'])
    print('04 Долг O4..O8',[ws[f'O{r}'].value for r in range(4,9)])
    ws=wb['Реестр передачи']
    for r in range(5,9):
        if ws[f'B{r}'].value: print('04 РП',r,ws[f'B{r}'].value,'| K=',ws[f'K{r}'].value)
    ws=wb['Реестр приложений']; print('04 РПрил I56 (нужно к иску, но нет)=',ws['I56'].value)
    # errors anywhere
    errs=[(w.title,c.coordinate,c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('#')]
    print('04 ячеек с ошибкой #:',len(errs),errs[:3])
    for f in sorted(glob.glob(f'{B}/{s}/08-*.xlsx')):
        wb=openpyxl.load_workbook(f,data_only=True)
        r=wb['Расчёт']; a=wb['Акты']; p=wb['Проверки']
        print(os.path.basename(f),'статус',r['D2'].value,'проценты',r['D3'].value,'долг',r['D4'].value,'| Акты H:',[a[f'H{i}'].value for i in range(5,8) if a[f'B{i}'].value],
              '| per-act P7..P9:',[r[f'P{i}'].value for i in range(7,10)])
        bl=[(p[f'B{i}'].value,p[f'C{i}'].value) for i in range(2,18) if str(p[f'C{i}'].value).startswith('БЛОК')]
        if bl: print('  БЛОК:',bl)
