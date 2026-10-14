#!/usr/bin/env python3
import hashlib
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
INPUT = ROOT.parents[1] / "portfolio-owner-editorial-v3-20261013/05-ks-bez-vozvrata/07-sroki-hraneniya.docx"
OUTDIR = ROOT / "candidate"
OUTPUT = OUTDIR / "07-sroki-hraneniya.docx"
EXPECTED_INPUT_SHA = "98be8b956c7068e58e66872ba0222f9454d3b8e09ba8afb6e67a75708e996311"

GRAPHITE = "2F343A"
GOLD = "B89A5A"
PALE_GOLD = "F2EAD8"
LIGHT = "F4F5F6"
WHITE = "FFFFFF"
TEXT = RGBColor(47, 52, 58)

def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def atomic_json(path, payload):
    path = Path(path)
    tmp = path.with_suffix(path.suffix + ".tmp")
    data = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    with open(tmp, "wb") as f:
        f.write(data)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)
    fd = os.open(path.parent, os.O_DIRECTORY)
    os.fsync(fd)
    os.close(fd)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_width(cell, mm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(mm * 56.7)))

def set_font(run, size, bold=False, color=TEXT):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def replace_paragraph(paragraph, text, size=9.5, bold=False, color=TEXT):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return paragraph

def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def apply_table_text(cell, text, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_font(run, 7.7 if not header else 8.0, bold=header, color=RGBColor(255,255,255) if header else TEXT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def main():
    actual_input_sha = sha256(INPUT)
    if actual_input_sha != EXPECTED_INPUT_SHA:
        raise SystemExit(f"input sha mismatch: {actual_input_sha}")

    doc = Document(INPUT)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(17)
    section.right_margin = Mm(17)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = TEXT

    title = doc.paragraphs[0]
    title.style = doc.styles["Normal"]
    replace_paragraph(title, "ПАМЯТКА ПО ОПРЕДЕЛЕНИЮ СРОКОВ ХРАНЕНИЯ ДОКУМЕНТОВ КОМПЛЕКТА КС", 15, True, TEXT)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.paragraph_format.space_after = Pt(5)

    replace_paragraph(doc.paragraphs[2], "1. Классифицируйте документ до назначения срока", 11.5, True, TEXT)
    doc.paragraphs[2].paragraph_format.space_after = Pt(5)

    rows = [
        ["Документ", "Базовая квалификация", "Что проверить", "Рабочее действие"],
        ["КС-2, КС-3", "Первичные учетные документы, если ими оформлен факт хозяйственной жизни", "Ст. 29 Закона № 402-ФЗ: не менее 5 лет после отчетного года. Сверьте также архивный, налоговый, договорный и спорный сроки.", "Зафиксируйте отчетный год и примените наиболее длительный действующий срок."],
        ["Акты освидетельствования скрытых работ", "Исполнительная документация по фактической функции документа", "Договор, состав исполнительной документации, требования передачи заказчику, номенклатура дел и точная статья перечня Росархива.", "Не уничтожайте до классификации и завершения передачи, гарантии, претензий и споров."],
        ["Исполнительные схемы", "Исполнительная документация", "Требования проекта и договора, комплект передачи заказчику, номенклатура дел, точная статья перечня.", "Храните в составе идентифицированного дела; срок считайте от указанного для этого дела события."],
        ["Журналы работ", "Исполнительная документация и журналы конкретного вида", "Какой именно журнал, кто обязан его вести, кому передается, точная статья перечня.", "Не переносите срок с другого вида журнала. Назначьте ответственного и событие начала срока."],
        ["Сертификаты и документы качества материалов", "Подтверждающие документы к материалам и работам", "Связь с партией, работой и исполнительным делом; договор; номенклатура; точная статья перечня.", "Храните с реестром и ссылкой на партию/работу до закрытия применимых обязательств."],
        ["Паспорта на изделия", "Документы изготовителя и эксплуатации", "Назначение паспорта, срок службы изделия, договор и требования передачи/эксплуатации.", "Не назначайте единый срок всем паспортам; привяжите решение к изделию и делу."],
        ["Переписка с заказчиком", "Переписка по конкретному вопросу", "Предмет переписки: цена, приемка, недостатки, претензия, ЭП и т.д. Пункт 572 Приказа Росархива № 236 не является общей нормой для всей переписки.", "Классифицируйте по предмету. При споре или претензии сохраните полный контекст и вложения."],
        ["Сметная документация", "Сметные, договорные и/или учетные документы по фактической роли", "Связь со сметой, договором, КС-2/КС-3, учетом и архивным делом.", "Разделите версии и утвержденный экземпляр; срок назначьте по каждой фактической функции."]
    ]
    table = doc.tables[0]
    widths = [34, 43, 57, 42]
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        if r_idx == 0:
            repeat_header(row)
        for c_idx, cell in enumerate(row.cells):
            apply_table_text(cell, rows[r_idx][c_idx], header=(r_idx == 0))
            set_cell_width(cell, widths[c_idx])
            if r_idx == 0:
                set_cell_shading(cell, GRAPHITE)
            elif c_idx == 0:
                set_cell_shading(cell, PALE_GOLD)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            else:
                set_cell_shading(cell, WHITE)

    replace_paragraph(doc.paragraphs[4], "2. Алгоритм назначения срока", 11.5, True, TEXT)
    first = doc.paragraphs[5]
    replace_paragraph(first, "1) Зафиксируйте точное наименование, версию, дату, роль документа и связанный договор/объект.", 9.5)
    first.paragraph_format.space_after = Pt(2)
    for text in [
        "2) Определите фактическую функцию: учетная, исполнительная, договорная, эксплуатационная или переписка по конкретному вопросу.",
        "3) Сопоставьте точную статью действующего перечня, правила учета, договор, проект и утвержденную номенклатуру дел. Не переносите срок с похожего документа.",
        "4) Запишите срок, событие начала его исчисления, место хранения, ответственного и основание в реестре.",
        "5) До уничтожения проверьте передачу заказчику, гарантию, аудит, налоговую проверку, претензии и споры. При открытом основании уничтожение остановите."
    ]:
        p = doc.add_paragraph()
        replace_paragraph(p, text, 9.5)
        p.paragraph_format.space_after = Pt(2)

    h = doc.add_paragraph()
    replace_paragraph(h, "3. Нормативные опоры и границы", 11.5, True, TEXT)
    h.paragraph_format.space_before = Pt(5)
    sources = [
        "Федеральный закон от 06.12.2011 № 402-ФЗ, ст. 29: первичные учетные документы хранятся по правилам архивного дела, но не менее пяти лет после отчетного года. https://minfin.gov.ru/ru/document?id_4=15014-federalnyi_zakon_ot_06.12.2011__402-fz_o_bukhgalterskom_uchete",
        "Приказ Росархива от 20.12.2019 № 236: срок выбирают по точной категории документа. Ранее указанные пункты 424 и 572 не подтверждают универсальные сроки для комплекта КС. https://publication.pravo.gov.ru/Document/View/0001202002070036",
        "Приказ Минстроя России от 16.05.2023 № 344/пр используется для состава и ведения исполнительной документации, но сам по себе не задает единый срок всем перечисленным здесь документам. https://publication.pravo.gov.ru/document/0001202306010019",
        "Эта памятка не заменяет номенклатуру дел, договор, решение архива или правовую оценку конкретного объекта."
    ]
    for i, text in enumerate(sources):
        p = doc.add_paragraph()
        replace_paragraph(p, text, 8.3, bold=(i == 3), color=TEXT)
        p.paragraph_format.space_after = Pt(2)
        if i == 3:
            p.paragraph_format.space_before = Pt(3)
            for run in p.runs:
                run.font.color.rgb = RGBColor(125, 94, 32)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.keep_together = True

    OUTDIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    out_doc = Document(OUTPUT)
    full_text = "\n".join(p.text for p in out_doc.paragraphs)
    full_text += "\n" + "\n".join(cell.text for t in out_doc.tables for row in t.rows for cell in row.cells)
    checks = {
        "input_sha_exact": actual_input_sha == EXPECTED_INPUT_SHA,
        "table_9x4": len(out_doc.tables) == 1 and len(out_doc.tables[0].rows) == 9 and len(out_doc.tables[0].columns) == 4,
        "a4": int(out_doc.sections[0].page_width.mm) == 210 and int(out_doc.sections[0].page_height.mm) == 297,
        "wrong_item_424_removed": "п. 424" not in full_text and "пункт 424" not in full_text,
        "wrong_universal_572_removed": "Перечень 2019 (п. 572)" not in full_text,
        "permanent_claim_removed": "Постоянно" not in full_text,
        "402fz_minimum_present": "не менее 5 лет после отчетного года" in full_text,
        "classification_required": "точную статью" in full_text and "фактическую функцию" in full_text,
        "hold_conditions_present": "претензии и споры" in full_text and "уничтожение остановите" in full_text,
        "official_sources_present": "minfin.gov.ru" in full_text and "publication.pravo.gov.ru" in full_text
    }
    failed = [name for name, ok in checks.items() if not ok]
    payload = {
        "schema_version": 1,
        "unit_id": "P3-RETENTION07-REMEDIATION-20261014",
        "status": "AUTHOR_PASS_RETENTION07_SCOPE_ONLY" if not failed else "FAIL",
        "source_path": str(INPUT.relative_to(ROOT.parents[3])),
        "source_sha256": actual_input_sha,
        "output_path": str(OUTPUT.relative_to(ROOT.parents[3])),
        "output_sha256": sha256(OUTPUT),
        "checks": checks,
        "passed": sum(checks.values()),
        "total": len(checks),
        "failed": failed,
        "limits": {
            "legal_acceptance": False,
            "native_microsoft_word": False,
            "independent_review": False,
            "product_release": False,
            "canonical_replacement": False,
            "model_calls": 0,
            "paid_calls": 0
        },
        "next_step": "Render and inspect every page, then run independent exact-SHA subject/buyer review and native Word."
    }
    atomic_json(ROOT / "author-result.json", payload)
    if failed:
        raise SystemExit("checks failed: " + ", ".join(failed))
    print(json.dumps(payload, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
