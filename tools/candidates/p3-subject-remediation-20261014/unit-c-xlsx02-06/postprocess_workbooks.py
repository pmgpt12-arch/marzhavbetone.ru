"""Apply deterministic A4/print metadata and build non-buyer scenario fixtures."""
import json,os,shutil
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.workbook.properties import CalcProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font
ROOT=Path(__file__).resolve().parent
CAND=ROOT/"candidate"
STATE=Path(os.environ.get("P3_UNITC_STATE", ROOT/"acceptance-state"))
def atomic_json(path,data):
    tmp=path.with_suffix(path.suffix+".tmp")
    with open(tmp,"w",encoding="utf-8") as f:
        json.dump(data,f,ensure_ascii=False,indent=2,default=str);f.write("\n");f.flush();os.fsync(f.fileno())
    os.replace(tmp,path);fd=os.open(path.parent,os.O_DIRECTORY);os.fsync(fd);os.close(fd)
def formula02(r):
    allowed_b=f'OR(B{r}="КС-2 — акт о приёмке выполненных работ",B{r}="КС-3 — справка о стоимости выполненных работ и затрат",B{r}="АОСР — акт освидетельствования скрытых работ",B{r}="Исполнительная схема",B{r}="Паспорт / сертификат",B{r}="Иное приложение")'
    return f'=IF(COUNTA(B{r}:H{r})=0,"",IF(OR(NOT({allowed_b}),C{r}="",NOT(ISNUMBER(D{r})),D{r}<DATE(2000,1,1),D{r}>TODAY(),NOT(AND(ISNUMBER(E{r}),E{r}=INT(E{r}),E{r}>=1,E{r}<=10000)),NOT(OR(F{r}="Приложено",F{r}="Не требуется",F{r}="Отсутствует")),NOT(OR(H{r}="Да",H{r}="Нет")),AND(F{r}="Не требуется",G{r}="")),"ЗАПОЛНИТЬ",IF(AND(F{r}="Приложено",H{r}="Да"),"ГОТОВО",IF(F{r}="Не требуется","ПРОВЕРИТЬ ОСНОВАНИЕ","НЕ ГОТОВО"))))'
def formula06(r):
    return f'=IF(COUNTA(B{r}:K{r})=0,"",IF(OR(NOT(ISNUMBER(B{r})),B{r}<DATE(2000,1,1),B{r}>TODAY(),C{r}="",D{r}="",E{r}="",F{r}="",NOT(AND(ISNUMBER(G{r}),G{r}=INT(G{r}),G{r}>=1,G{r}<=10000)),NOT(OR(H{r}="Подпись на экземпляре",H{r}="ЭДО",H{r}="Письмо / курьер",H{r}="Иное подтверждение")),NOT(OR(I{r}="На проверке",I{r}="Принято",I{r}="Возврат"))),"ЗАПОЛНИТЬ",IF(I{r}="Принято","ЗАКРЫТО",IF(I{r}="Возврат",IF(OR(NOT(ISNUMBER(J{r})),J{r}<B{r},J{r}>TODAY(),K{r}=""),"УКАЗАТЬ ВОЗВРАТ","ВОЗВРАТ ЗАФИКСИРОВАН"),"В РАБОТЕ"))))'
def enforce_controls(ws,status_col,last_row,formula_factory):
    quantity_col="E" if status_col=="I" else "G"
    for r in range(5,last_row+1):
        ws[f"{status_col}{r}"]=formula_factory(r)
        ws[f"{quantity_col}{r}"].number_format="0.##"
    for dv in ws.data_validations.dataValidation:
        dv.showErrorMessage=True;dv.errorStyle="stop";dv.errorTitle="Недопустимое значение";dv.error="Исправьте значение по правилу поля."
        if dv.type=="date": dv.formula1="DATE(2000,1,1)";dv.formula2="TODAY()"
    ws.conditional_formatting._cf_rules.clear()
    styles={
      "ГОТОВО":("C6EFCE","215E21"),"ЗАКРЫТО":("C6EFCE","215E21"),
      "НЕ ГОТОВО":("FFC7CE","9C0006"),"ЗАПОЛНИТЬ":("FFC7CE","9C0006"),
      "УКАЗАТЬ ВОЗВРАТ":("FFC7CE","9C0006"),"ПРОВЕРИТЬ ОСНОВАНИЕ":("FFEB9C","7F6000"),
      "ВОЗВРАТ ЗАФИКСИРОВАН":("FFEB9C","7F6000"),"ПРОВЕРИТЬ":("FFEB9C","7F6000"),
      "В РАБОТЕ":("DDEBF7","1F4E78")}
    for status,(fill,font) in styles.items():
        ws.conditional_formatting.add(f"{status_col}5:{status_col}{last_row}",FormulaRule(formula=[f'${status_col}5="{status}"'],stopIfTrue=True,fill=PatternFill("solid",fgColor=fill),font=Font(bold=True,color=font)))
def setup(name,main,last_col,last_row,data_row_height=None):
    p=CAND/name;wb=openpyxl.load_workbook(p)
    ws=wb[main]
    ws.sheet_view.showGridLines=False
    ws.auto_filter.ref=f"A4:{last_col}{last_row}"
    ws.print_area=f"A1:{last_col}{last_row}"
    ws.print_title_rows="4:4"
    if data_row_height is not None:
        for row in range(5,last_row+1):
            ws.row_dimensions[row].height=data_row_height
    ws.page_setup.orientation="landscape"
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_margins.left=0.25;ws.page_margins.right=0.25;ws.page_margins.top=0.35;ws.page_margins.bottom=0.35
    ins=wb["Инструкция"]
    ins.print_area="A1:B7";ins.page_setup.paperSize=ins.PAPERSIZE_A4
    ins.page_setup.orientation="portrait";ins.page_setup.fitToWidth=1;ins.page_setup.fitToHeight=1
    ins.sheet_properties.pageSetUpPr.fitToPage=True
    if wb.calculation is None:
        wb.calculation=CalcProperties()
    wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode="auto"
    if name.startswith("02-"):
        enforce_controls(ws,"I",last_row,formula02)
    else:
        enforce_controls(ws,"L",last_row,formula06)
    wb.save(p)
setup("02-reestr-prilozheniy.xlsx","Реестр приложений","I",54,21)
setup("06-zhurnal-peredachi.xlsx","Журнал передачи комплекта КС","L",54)
STATE.mkdir(parents=True,exist_ok=True)
fx=STATE/"scenario-fixtures";fx.mkdir(exist_ok=True)
for p in fx.glob("*.xlsx"): p.unlink()
# 02 fixture
p=CAND/"02-reestr-prilozheniy.xlsx";wb=openpyxl.load_workbook(p);ws=wb["Реестр приложений"]
cases02=[
 ("blank",{},""),
 ("blank_date",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-1","E":3,"F":"Приложено","H":"Да"},"ЗАПОЛНИТЬ"),
 ("text_date",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-1","D":"не дата","E":3,"F":"Приложено","H":"Да"},"ЗАПОЛНИТЬ"),
 ("complete",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-1","D":datetime(2026,10,1),"E":3,"F":"Приложено","H":"Да"},"ГОТОВО"),
 ("not_required",{"B":"Иное приложение","C":"N/A","D":datetime(2026,10,1),"E":1,"F":"Не требуется","G":"По договору не применяется","H":"Да"},"ПРОВЕРИТЬ ОСНОВАНИЕ"),
 ("not_checked",{"B":"КС-2 — акт о приёмке выполненных работ","C":"КС-2-1","D":datetime(2026,10,1),"E":2,"F":"Приложено","H":"Нет"},"НЕ ГОТОВО"),
 ("invalid_doc_type",{"B":"Неизвестный документ","C":"X-1","D":datetime(2026,10,1),"E":1,"F":"Приложено","H":"Да"},"ЗАПОЛНИТЬ"),
 ("future_date",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-2","D":datetime(2027,1,1),"E":1,"F":"Приложено","H":"Да"},"ЗАПОЛНИТЬ"),
 ("fractional_pages",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-3","D":datetime(2026,10,1),"E":1.5,"F":"Приложено","H":"Да"},"ЗАПОЛНИТЬ"),
 ("invalid_presence",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-4","D":datetime(2026,10,1),"E":1,"F":"Возможно","H":"Да"},"ЗАПОЛНИТЬ"),
 ("invalid_checked",{"B":"АОСР — акт освидетельствования скрытых работ","C":"АОСР-5","D":datetime(2026,10,1),"E":1,"F":"Приложено","H":"Возможно"},"ЗАПОЛНИТЬ"),
 ("not_required_no_reason",{"B":"Иное приложение","C":"N/A","D":datetime(2026,10,1),"E":1,"F":"Не требуется","H":"Да"},"ЗАПОЛНИТЬ"),
]
for i,(case,vals,expected) in enumerate(cases02,5):
    for col,v in vals.items(): ws[f"{col}{i}"]=v
out=fx/"02-scenarios.xlsx";wb.save(out)
# 06 fixture
p=CAND/"06-zhurnal-peredachi.xlsx";wb=openpyxl.load_workbook(p);ws=wb["Журнал передачи комплекта КС"]
base={"B":datetime(2026,10,1),"C":"КС-2-1","D":"КС-3-1","E":"Иванов И.И.","F":"Инженер","G":8,"H":"ЭДО"}
cases06=[
 ("blank",{},""),
 ("blank_date",{"C":"КС-2-1","D":"КС-3-1","E":"Иванов И.И.","F":"Инженер","G":8,"H":"ЭДО","I":"Принято"},"ЗАПОЛНИТЬ"),
 ("return_missing",dict(base,I="Возврат"),"УКАЗАТЬ ВОЗВРАТ"),
 ("return_before",dict(base,I="Возврат",J=datetime(2026,9,30),K="Исправить"),"УКАЗАТЬ ВОЗВРАТ"),
 ("return_complete",dict(base,I="Возврат",J=datetime(2026,10,2),K="Исправить"),"ВОЗВРАТ ЗАФИКСИРОВАН"),
 ("accepted",dict(base,I="Принято"),"ЗАКРЫТО"),
 ("review",dict(base,I="На проверке"),"В РАБОТЕ"),
 ("missing_recipient",dict(base,E="",I="Принято"),"ЗАПОЛНИТЬ"),
 ("future_transfer",dict(base,B=datetime(2027,1,1),I="Принято"),"ЗАПОЛНИТЬ"),
 ("invalid_method",dict(base,H="Телепатия",I="Принято"),"ЗАПОЛНИТЬ"),
 ("invalid_status",dict(base,I="Закрыть"),"ЗАПОЛНИТЬ"),
 ("fractional_count",dict(base,G=1.5,I="Принято"),"ЗАПОЛНИТЬ"),
 ("return_future",dict(base,I="Возврат",J=datetime(2027,1,1),K="Исправить"),"УКАЗАТЬ ВОЗВРАТ"),
]
for i,(case,vals,expected) in enumerate(cases06,5):
    for col,v in vals.items(): ws[f"{col}{i}"]=v
out=fx/"06-scenarios.xlsx";wb.save(out)
atomic_json(STATE/"scenario-map.json",{"02":[{"row":i,"case":c,"expected":e} for i,(c,v,e) in enumerate(cases02,5)],"06":[{"row":i,"case":c,"expected":e} for i,(c,v,e) in enumerate(cases06,5)]})
print("POSTPROCESS_AND_FIXTURES_READY")
