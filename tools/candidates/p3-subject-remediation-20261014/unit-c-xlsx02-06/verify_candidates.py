"""Verify final candidate structure and cached scenario results after LibreOffice recalc."""
import hashlib,json,os
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parent
WT=ROOT.parents[3]
STATE=Path(os.environ.get("P3_UNITC_STATE", ROOT/"acceptance-state"))
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def atomic_json(path,data):
    tmp=path.with_suffix(path.suffix+".tmp")
    with open(tmp,"w",encoding="utf-8") as f:
        json.dump(data,f,ensure_ascii=False,indent=2);f.write("\n");f.flush();os.fsync(f.fileno())
    os.replace(tmp,path);fd=os.open(path.parent,os.O_DIRECTORY);os.fsync(fd);os.close(fd)
expected_sources={
 "products-storage/05-ks-bez-vozvrata/02-reestr-prilozheniy.xlsx":"ccc422fe4d9f76580d204091b9206243410f5b709600fc314863f91a0f0542f8",
 "products-storage/05-ks-bez-vozvrata/06-zhurnal-peredachi.xlsx":"8da8ec7d25eca4d5e3c729d029d506317420eccc82cb6776260312ee23c210de",
}
checks={}
for rel,h in expected_sources.items():
    checks["source_unchanged_"+Path(rel).name]=sha(WT/rel)==h
spec=[
 ("02-reestr-prilozheniy.xlsx","Реестр приложений","I",5,54,5,["КС-11","Ведомость ресурсов"]),
 ("06-zhurnal-peredachi.xlsx","Журнал передачи комплекта КС","L",5,54,5,[]),
]
artifacts=[]
for name,sheet,col,r1,r2,min_dv,banned in spec:
    p=ROOT/"candidate"/name
    wb=openpyxl.load_workbook(p,data_only=False);ws=wb[sheet]
    formulas=[ws[f"{col}{r}"].value for r in range(r1,r2+1)]
    checks[name+"_formula_coverage_50"]=len(formulas)==50 and all(isinstance(x,str) and x.startswith("=") for x in formulas)
    dvs=ws.data_validations.dataValidation
    checks[name+"_data_validation_ranges"]=len(dvs)>=min_dv
    checks[name+"_data_validation_error_alerts"]=all(dv.showErrorMessage is True and dv.errorStyle=="stop" for dv in dvs)
    date_dvs=[dv for dv in dvs if dv.type=="date"]
    checks[name+"_date_bounds_explicit"]=bool(date_dvs) and all(str(dv.formula1).lstrip("=")=="DATE(2000,1,1)" and str(dv.formula2).lstrip("=")=="TODAY()" for dv in date_dvs)
    cf_rules=[rule for values in ws.conditional_formatting._cf_rules.values() for rule in values]
    cf_formulas=[f for rule in cf_rules for f in (rule.formula or [])]
    expected_statuses=["ГОТОВО","ЗАКРЫТО","НЕ ГОТОВО","ЗАПОЛНИТЬ","УКАЗАТЬ ВОЗВРАТ","ПРОВЕРИТЬ ОСНОВАНИЕ","ВОЗВРАТ ЗАФИКСИРОВАН","ПРОВЕРИТЬ","В РАБОТЕ"]
    expected_cf={f'${col}5="{status}"' for status in expected_statuses}
    rules=len(cf_rules)
    checks[name+"_conditional_status_rules"]=rules==9 and all(rule.type=="expression" for rule in cf_rules)
    checks[name+"_conditional_exact_equality"]=set(cf_formulas)==expected_cf and all("SEARCH(" not in f.upper() for f in cf_formulas)
    checks[name+"_a4_landscape_fit_width"]=str(ws.page_setup.paperSize)==str(ws.PAPERSIZE_A4) and ws.page_setup.orientation=="landscape" and ws.page_setup.fitToWidth==1
    checks[name+"_print_title_row_4"]=ws.print_title_rows=="$4:$4"
    if name.startswith("02-"):
        checks[name+"_data_row_height_21"]=all(ws.row_dimensions[r].height==21 for r in range(r1,r2+1))
    checks[name+"_frozen_header"]=str(ws.freeze_panes)=="A5"
    text="\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    checks[name+"_banned_labels_absent"]=all(x not in text for x in banned)
    artifacts.append({"path":"candidate/"+name,"sha256":sha(p),"bytes":p.stat().st_size,"formula_rows":50,"data_validations":len(ws.data_validations.dataValidation),"conditional_rules":rules})
    wb.close()
scenario=json.loads((STATE/"scenario-map.json").read_text())
checks["scenario_count_02_exact_12"]=len(scenario.get("02",[]))==12
checks["scenario_count_06_exact_13"]=len(scenario.get("06",[]))==13
for key,name,sheet,col in [
 ("02","02-scenarios.xlsx","Реестр приложений","I"),
 ("06","06-scenarios.xlsx","Журнал передачи комплекта КС","L"),
]:
    p=STATE/"scenario-recalc"/name
    wb=openpyxl.load_workbook(p,data_only=True);ws=wb[sheet]
    got=[]
    for case in scenario[key]:
        raw=ws[f"{col}{case['row']}"].value
        actual="" if raw is None else raw
        got.append({"case":case["case"],"row":case["row"],"expected":case["expected"],"actual":actual,"cached_raw":raw,"pass":actual==case["expected"]})
    checks[name+"_scenario_expected_values"]=all(x["pass"] for x in got)
    scenario[key]=got;wb.close()
for name in ["02-reestr-prilozheniy.xlsx","06-zhurnal-peredachi.xlsx"]:
    wb=openpyxl.load_workbook(ROOT/"candidate"/name,data_only=True)
    errs=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.startswith("#"): errs.append({"sheet":ws.title,"cell":c.coordinate,"value":c.value})
    checks[name+"_cached_formula_errors_zero"]=not errs
    wb.close()
status="PASS_AUTHOR_FUNCTIONAL_AND_LIBREOFFICE_RECALC" if all(checks.values()) else "FAIL"
report={"id":"P3-XLSX02-06-AUTHOR-ACCEPTANCE-20261014","status":status,"checks":checks,"passed":sum(checks.values()),"total":len(checks),"scenarios":scenario,"artifacts":artifacts,"renderer":"LibreOffice Calc headless + Poppler; artifact authoring by @oai/artifact-tool","native_microsoft_excel":"NOT_RUN","independent_buyer_review":"PENDING","canonical_replacement":False,"release":"HOLD","model_calls":0,"paid_calls":0}
atomic_json(ROOT/"author-acceptance.json",report)
print(json.dumps({"status":status,"passed":report["passed"],"total":report["total"],"artifacts":artifacts,"scenarios":scenario},ensure_ascii=False,indent=2))
if status!="PASS_AUTHOR_FUNCTIONAL_AND_LIBREOFFICE_RECALC": raise SystemExit(1)
