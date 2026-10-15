#!/usr/bin/env python3
from pathlib import Path
import json,hashlib,re
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
SOURCE=HERE.parent/'p3-full-11-preflight-20261015/candidate/03-akt-skrytyh-rabot.docx'
OUTPUT=HERE/'03-akt-skrytyh-rabot.docx'
SHA='cd8e5d4b1533bbbfebb3d85a06784a1b861ef6a8b54bcfb027d557fcb64bfdb5'
def hashfile(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def text(par,s,size=8.2,bold=False,color='263039',italic=False,underline=False):
 r=par.add_run(s);r.bold=bold;r.italic=italic;r.underline=underline;r.font.name='Times New Roman'
 r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string(color)
def style(par,before=0,after=0,space=1.03,keep=False):
 x=par.paragraph_format;x.space_before=Pt(before);x.space_after=Pt(after);x.line_spacing=space;x.keep_with_next=keep
def border(par, shade='C6AA70',sz='5'):
 pr=par._p.get_or_add_pPr();b=OxmlElement('w:pBdr');e=OxmlElement('w:bottom')
 for n,v in [('val','single'),('sz',sz),('color',shade)]:e.set(qn('w:'+n),v)
 b.append(e);pr.append(b)
def cellborder(c):
 pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom')
 for k,v in [('val','single'),('sz','4'),('color','B5B5B5')]:e.set(qn('w:'+k),v)
 b.append(e);pr.append(b)
def margins(c):
 pr=c._tc.get_or_add_tcPr();el=OxmlElement('w:tcMar')
 for k in ('top','bottom','start','end'):
  v=OxmlElement('w:'+k);v.set(qn('w:w'),'70');v.set(qn('w:type'),'dxa');el.append(v)
 pr.append(el)
def p(d,s,size=8.4,before=0,after=0,bold=False,color='29313B',center=False):
 a=d.add_paragraph();style(a,before,after);a.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT;text(a,s,size,bold,color);return a
def h(d,s):
 a=d.add_paragraph();style(a,5,2,keep=True);text(a,s.upper(),9.1,True,'40484E');border(a,'D2D5D6','4');return a
def split(s):
 match=re.search(r'\{\{[^{}]+\}\}',s)
 if match:return s[:match.start()].strip(),s[match.start():].strip()
 return s,''
def table(d,rows,left=71,right=105):
 t=d.add_table(rows=0,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for label,field,note in rows:
  row=t.add_row();mark=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(mark)
  a,b=row.cells;a.width=Mm(left);b.width=Mm(right)
  for c in (a,b):c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;margins(c)
  pa=a.paragraphs[0];style(pa);text(pa,label,8.9,True,'303840')
  pb=b.paragraphs[0];style(pb,after=0)
  for frag in re.split(r'(\{\{[^{}]+\}\})',field):
   if frag:text(pb,frag,9.1 if '{{' in frag else 8.7,color='303840',underline=frag.startswith('{{'))
  if note:
   pn=b.add_paragraph();style(pn,space=1)
   text(pn,note,7.9,color='65696D',italic=True)
  cellborder(b)
 return t
def rows(src,pairs):
 ans=[]
 for a,b in pairs:
  title,field=split(src[a]);ans.append((title,field,src[b] if b is not None else ''))
 return ans
doc=Document(SOURCE)
assert hashfile(SOURCE)==SHA
raw=[q.text for q in doc.paragraphs]
assert len(raw)==281
assert all(raw[16+i*53].startswith('АКТ ОСВИДЕТЕЛЬСТВОВАНИЯ') for i in range(5))
new=Document()
sec=new.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(14);sec.bottom_margin=Mm(14);sec.left_margin=Mm(18);sec.right_margin=Mm(16)
sec.header_distance=Mm(6);sec.footer_distance=Mm(7)
normal=new.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(8.4)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1.03
hdr=sec.header.paragraphs[0];style(hdr);hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
text(hdr,'МАРЖА В БЕТОНЕ  /  КОНТРОЛЬ СКРЫТЫХ РАБОТ',7.1,color='646A70');border(hdr)
ftr=sec.footer.paragraphs[0];style(ftr);ftr.alignment=WD_ALIGN_PARAGRAPH.RIGHT;text(ftr,'АОСР  •  рабочая форма   |   стр. ',7.1,color='646A70')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');ftr._p.append(f)
cover=p(new,raw[0],15,23,8,True,'222831');border(cover,sz='8')
p(new,raw[1],10,0,7,True)
p(new,raw[2],9.2,0,7)
for s in raw[3:8]:p(new,s,9.2,1,1)
h(new,raw[8])
for s in raw[9:16]:p(new,s,8.6,2,2)
for k in range(5):
 a=raw[16+53*k:16+53*(k+1)]
 t=p(new,a[0],11.15,7,4,True,'202830',True)
 t.paragraph_format.page_break_before=True;t.paragraph_format.keep_with_next=True;border(t,sz='7')
 p(new,a[1],7.15,0,4,color='666666',center=True)
 h(new,'Сведения об объекте и организациях')
 table(new,rows(a,[(2,3),(4,5),(6,7),(8,9)]),62,114)
 p(new,a[10],10,5,4,True,center=True)
 h(new,a[11])
 table(new,rows(a,[(12,13),(14,15),(16,17),(18,19),(20,21)]),73,103)
 p(new,a[22],8.9,5,4)
 h(new,'Содержание освидетельствования').paragraph_format.page_break_before=True
 table(new,rows(a,[(23,24),(25,26),(27,28),(29,30),(31,32),(33,34),(35,36)]),68,108)
 h(new,'Дополнительные сведения и приложения')
 table(new,rows(a,[(37,None),(38,39),(40,41)]),55,121)
 h(new,a[42])
 table(new,[(a[i],a[i+1],'') for i in [43,45,47,49,51]],95,81)
new.save(OUTPUT)
from collections import Counter
def tokens(s):return Counter(re.findall(r'\{\{[^{}]+\}\}',s))
src=tokens('\n'.join(raw))
outdoc=Document(OUTPUT)
outtext=''.join(t.text or '' for t in outdoc.element.body.iter(qn('w:t')))
dst=tokens(outtext)
assert src==dst, {'missing':list((src-dst).elements()),'extra':list((dst-src).elements())}
assert sum(1 for s in raw if s.startswith('АКТ ОСВИДЕТЕЛЬСТВОВАНИЯ'))==5
report={'source_sha256':hashfile(SOURCE),'candidate_sha256':hashfile(OUTPUT),'forms':5,
'source_paragraphs':len(raw),'placeholder_tokens_before':sum(src.values()),
'placeholder_tokens_after':sum(dst.values()),'all_placeholder_identities_preserved':True,
'status':'PASS_STRUCTURE_PENDING_VISUAL_NATIVE','canonical_changes':False,'release':'HOLD'}
(HERE/'qa-content.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
