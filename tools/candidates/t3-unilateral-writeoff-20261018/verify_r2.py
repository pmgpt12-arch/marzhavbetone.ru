#!/usr/bin/env python3
"""T3 current edition: complete terminology coverage, exact files and all T3 prices."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from docx import Document
from bs4 import BeautifulSoup
import hashlib,re,subprocess,json,sys
root=Path(__file__).resolve().parents[3]
folder=Path(__file__).resolve().parent
archive=Path(sys.argv[1]) if len(sys.argv)>1 else folder/'T3-odnostoronnee-spisanie-5-files-R2.zip'
title='Система подготовки возражений на односторонний акт или списание в одностороннем порядке'
expected=['00-START-HERE.pdf','01-poryadok-deystviy.docx','02-vozrazheniya-na-akt.docx','03-vozrazheniya-na-odnostoronnee-spisanie.docx','04-slovar-poley.docx']
no_old=lambda t: re.search(r'зач[её]т',t,re.I) is None
with ZipFile(archive) as z:
 assert z.testzip() is None and len(z.namelist())==5 and set(z.namelist())==set(expected)
 texts={}
 for n in expected:
  if n.endswith('.pdf'):
   texts[n]=subprocess.run(['pdftotext','-','-'],input=z.read(n),capture_output=True,check=True).stdout.decode()
  else:
   d=Document(BytesIO(z.read(n)))
   texts[n]=' '.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
  assert no_old(texts[n]), 'old buyer terminology '+n
 assert all(n in texts[expected[0]] for n in expected[1:])
 for n in expected[2:4]:
  d=Document(BytesIO(z.read(n)));assert d.paragraphs[0].text.startswith('Кому:')
  fields=set(re.findall(r'\{\{([А-ЯЁ_0-9]+)\}\}',texts[n]))
  assert fields and all(f in texts[expected[4]] for f in fields), 'dictionary gaps '+n
 src=root/'products-storage/17-odnostoronniy-akt-i-zachet'
 assert sorted(f.name for f in src.iterdir())==sorted(expected)
 assert all((src/n).read_bytes()==z.read(n) for n in expected), 'source byte mismatch'
 # The act objection has no occurrence to replace and stays byte-identical.
 assert z.read(expected[2])==(root/'tools/candidates/t3-buyer-review-20261017/R1'/expected[2]).read_bytes()
page=BeautifulSoup((root/'products/t3-odnostoronniy-akt-i-zachet.html').read_text(),'html.parser')
assert no_old(page.get_text()) and no_old(str(page.select_one('head'))), 'old T3 page terminology'
assert page.select_one('h1').get_text()==title
assert all(n in page.get_text() for n in expected)
assert page.select_one('[data-sku=t3]')['data-price']=='1490000'
cat=BeautifulSoup((root/'katalog.html').read_text(),'html.parser')
card=cat.select_one('[data-sku=t3]').find_parent('article')
assert title in card.get_text() and no_old(card.get_text()) and '14 900 ₽' in card.get_text()
def visit(x):
 if isinstance(x,dict):
  if x.get('@type')=='Product' and 't3-odnostoronniy' in str(x.get('offers',{})):
   assert x['name']==title and str(x['offers']['price'])=='14900'
  for v in x.values():visit(v)
 elif isinstance(x,list):
  for v in x:visit(v)
for node in cat.select('script[type="application/ld+json"]'):visit(json.loads(node.string))
for path in ['articles/garantiynoe-uderzhanie-v-ks-2-i-ks-3.html','articles/raboty-ne-prinyaty.html','materialy/uderzhaniya.html']:
 soup=BeautifulSoup((root/path).read_text(),'html.parser')
 for a in soup.select('a[href*="t3-odnostoronniy"]'):
  block=a.find_parent('article') or a.find_parent('p')
  assert no_old(block.get_text()) and '14 900 ₽' in block.get_text(), 'stale T3 CTA '+path
diag=(root/'diagnostika.html').read_text()
for m in re.finditer(r'"url":"/products/t3-odnostoronniy-akt-i-zachet.html","price":(\d+)',diag):assert m.group(1)=='14900'
print('T3_R2_GATE_PASS buyer_files=5 source_sha=5/5 no_old_term=PASS dictionary=PASS all_T3_prices=14900 act02=UNCHANGED')
