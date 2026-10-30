from pathlib import Path
from collections import Counter
from lxml import etree as E
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader,PdfWriter
import hashlib,json,re,zipfile
R=Path(__file__).resolve().parent;B=R/'candidate';I=R/'internal';O=R/'output';O.mkdir(exist_ok=True)
H=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();checks=[]
def check(n,v):
 if not v:raise AssertionError(n)
 checks.append({'name':n,'status':'PASS'})
norm=lambda s:re.sub(r'\s+',' ',s).strip()
files=sorted(B.iterdir());check('26 buyer files',len(files)==26)
check('format composition',Counter(p.suffix for p in files)=={'.docx':13,'.xlsx':6,'.pdf':7})
check('one initial instruction',[p.name for p in files if p.name.startswith('00-')]==['00-INSTRUKCIYA.pdf'])
guide=norm(' '.join(p.extract_text() for p in PdfReader(B/'00-INSTRUKCIYA.pdf').pages))
check('instruction exact composition',all(p.name in guide for p in files))
check('instruction distinguishes working forms','отличается от рекомендуемого образца' in guide and 'не заменяют установленный общий журнал' in guide)
for p in sorted((R/'source').iterdir()):
 if p.suffix=='.pdf' or (p.suffix=='.xlsx' and not p.name.startswith('38-')):check('unchanged '+p.name,H(p)==H(B/p.name))
render=json.loads((I/'RENDER_RECEIPT.json').read_text());tables=[]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};q=lambda n:'{'+ns['w']+'}'+n
for p in sorted(B.glob('*.docx')):
 d=Document(p);with_source=(R/'source'/p.name)
 text=norm(' '.join(d.element.body.xpath('.//w:t/text()')))
 check('editable '+p.name,not d.settings.element.xpath('./w:documentProtection'))
 check('render exact input '+p.name,any(x['file']==p.name and x['exit_code']==0 and x['sha256']==H(p) for x in render))
 if with_source.exists():
  old=Document(with_source)
  check('source body preserved '+p.name,all(norm(x.text) in text for x in old.paragraphs if x.text.strip() and x.text!='{{ДОЛЖНОСТЬ, ПОДПИСЬ, ФИО}}'))
 for ti,t in enumerate(d.tables,1):
  widths=[int(x.get(q('w')))for x in t._tbl.tblGrid.gridCol_lst]
  check('width fixed '+p.name+' '+str(ti),sum(widths)==9866 and t._tbl.tblPr.find(q('tblLayout')).get(q('type'))=='fixed')
  tables.append({'file':p.name,'table':ti,'width_twips':sum(widths)})
 pdf=R/'render-final'/p.stem/(p.stem+'.pdf');n=len(PdfReader(pdf).pages)
 check('single page '+p.name,n==1)
 check('page PNG exists '+p.name,(pdf.parent/'page-1.png').exists())
check('accepted editable AOSR',H(B/'31A-rabochiy-akt-skrytyh-rabot.docx')=='a0b0e8f48f7f0b8310a96358c226038a2066a17c159021d6f34f742b7969d7d7')
a=load_workbook(R/'source/38-reestr-id.xlsx');b=load_workbook(B/'38-reestr-id.xlsx')
check('38 sheets preserved',a.sheetnames==b.sheetnames)
check('38 original data preserved',all(c.value==b['Реестр'][c.coordinate].value for row in a['Реестр'] for c in row))
check('38 six explicit fields',[b['Реестр'].cell(3,n).value for n in range(12,18)]==['Версия документа','Опись номер и дата','Адресат и представитель','Дата получения','Подтверждение получения','Ответ номер и дата'])
check('38 print includes all fields',str(b['Реестр'].print_area)=="'Реестр'!$A$3:$Q$103" and b['Реестр'].print_title_cols=='$A:$C' and b['Реестр'].page_setup.fitToWidth==2)
spreadsheet_pages={}
for p in sorted(B.glob('*.xlsx')):
 pdf=R/'render-excel'/(p.stem+'.pdf');pages=PdfReader(pdf).pages
 spreadsheet_pages[p.name]=len(pages)
 check('spreadsheet PDF '+p.name,len(pages)>0)
 if p.name.startswith('38-'):
  t=norm(' '.join(page.extract_text()for page in pages))
  check('38 all six fields printed',all(norm(b['Реестр'].cell(3,n).value)in t for n in range(12,18)))
check('guide four pages',len(PdfReader(B/'00-INSTRUKCIYA.pdf').pages)==4)
archive=O/'P4_PTO_bez_zamechaniy_v1.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for p in files:
  info=zipfile.ZipInfo(p.name,(2026,10,30,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,p.read_bytes())
with zipfile.ZipFile(archive)as z:check('archive CRC and all hashes',z.testzip()is None and all(hashlib.sha256(z.read(p.name)).hexdigest()==H(p) for p in files))
preview=O/'P4_predprosmotr_38str.pdf';w=PdfWriter();w.append(B/'00-INSTRUKCIYA.pdf',outline_item='Начальная инструкция')
for p in sorted(B.glob('*.docx')):w.append(R/'render-final'/p.stem/(p.stem+'.pdf'),outline_item=p.name)
for p in sorted(B.glob('*.xlsx')):w.append(R/'render-excel'/(p.stem+'.pdf'),outline_item=p.name)
w.add_metadata({'/Title':'P4 ПТО без замечаний Предпросмотр Word и Excel','/Author':'На стороне подрядчика'})
with preview.open('wb')as f:w.write(f)
check('preview 38 pages',len(PdfReader(preview).pages)==38)
receipt={'schema':'NSP_QA16_P4_OWNER_REVIEW_v1','sku':'p4','edition':'P4-layout-v1','status':'READY_FOR_OWNER_REVIEW','owner_approval':'PENDING','release_allowed':False,'production_changed':False,'full_qa_pass':False,'scope':'Composition, editable working AOSR, Word table layout, 38 explicit workflow fields, LibreOffice visual print review','source_commit':'5361c80125107d9bb47d355bd5730f8b6e2da2c1','source_archive_sha256':H(R/'source.zip'),'source_archive_path':'tools/nsp-next-corrections-20261022/candidate-archives/08-pto-bez-zamechaniy.zip','file_count':26,'formats':{'docx':13,'xlsx':6,'pdf':7},'table_count':len(tables),'table_geometry':tables,'excel_06_07_changed':False,'new_native_office_test_performed':False,'render_engine':'LibreOffice; not native Microsoft Office acceptance','visual_pages':38,'visual_scope':'4 guide, 13 Word, 21 spreadsheet export pages; unchanged normative PDFs preserved by SHA','spreadsheet_pages':spreadsheet_pages,'checks':checks,'check_count':len(checks),'files':[{'name':p.name,'bytes':p.stat().st_size,'sha256':H(p)}for p in files],'archive':{'name':archive.name,'sha256':H(archive),'bytes':archive.stat().st_size},'preview':{'name':preview.name,'sha256':H(preview),'bytes':preview.stat().st_size},'remaining':'Owner review; no publication permission inferred from P3 approval. No new full native product pass claimed.'}
(I/'REVIEW_MANIFEST.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:receipt[k]for k in ['status','file_count','table_count','visual_pages','check_count','archive','preview']},ensure_ascii=False,indent=2))
