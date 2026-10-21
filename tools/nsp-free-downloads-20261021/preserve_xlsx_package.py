#!/usr/bin/env python3
"""Apply only Artifact Tool-authored CTA values/merges to the original package.
Reject wholesale round-trip drift (None becoming empty strings in input cells).
"""
from pathlib import Path
from lxml import etree as E
from openpyxl import load_workbook
import json,zipfile,hashlib
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda n:'{'+NS+'}'+n
spec=[('00-free-bankrotstvo-proverka','02-reestr-priznakov-po-obektu.xlsx',63,'K'),('00-free-uderzhaniya-do-podpisi','02-reestr-uderzhaniy-po-obektu.xlsx',40,'N')];proof=[]
for folder,file,start,last in spec:
 original=Path('/tmp/nsp-free-original')/(folder+'--'+file);candidate=Path('/tmp/nsp-free-xlsx')/file;target=ROOT/'products-storage'/folder/file
 source=load_workbook(original);authored=load_workbook(candidate);a=source.active;b=authored.active
 drift=sum(c.value is None and b[c.coordinate].value=='' for row in a for c in row)
 values={f'A{r}':b[f'A{r}'].value for r in range(start+1,start+4)};assert all(isinstance(v,str) for v in values.values())
 with zipfile.ZipFile(original) as z:entries={n:z.read(n) for n in z.namelist()}
 tree=E.fromstring(entries['xl/worksheets/sheet1.xml']);sd=tree.find(q('sheetData'))
 for r in range(start,start+4):
  row=sd.find(q('row')+f'[@r="{r}"]');assert row is not None
  row.set('ht',str(36 if r==start+1 else 28));row.set('customHeight','1')
  key='A'+str(r)
  if key in values:
   cell=row.find(q('c')+f'[@r="{key}"]');assert cell is not None
   for child in list(cell):cell.remove(child)
   cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,q('is')),q('t')).text=values[key]
  for cell in list(row):
   if cell.tag==q('c') and cell.get('r')!=key:
    assert not ''.join(cell.itertext()).strip(),'MERGE_NONEMPTY_CELL';row.remove(cell)
 mc=tree.find(q('mergeCells'));assert mc is None
 mc=E.Element(q('mergeCells'),count='4')
 for r in range(start,start+4):E.SubElement(mc,q('mergeCell'),ref=f'A{r}:{last}{r}')
 index=next((i for i,c in enumerate(tree) if c.tag==q('pageMargins')),len(tree));tree.insert(index,mc)
 entries['xl/worksheets/sheet1.xml']=E.tostring(tree,encoding='utf-8')
 with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
  for name,raw in entries.items():z.writestr(name,raw)
 final=load_workbook(target);assert source.sheetnames==final.sheetnames
 for s,t in zip(source,final):
  for row in s:
   for cell in row:
    key=cell.coordinate;new=t[key];expected=values.get(key,cell.value)
    assert new.value==expected,('UNINTENDED_VALUE_CHANGE',key,cell.value,new.value)
    assert new.data_type==cell.data_type or key in values or cell.value is None,('UNINTENDED_TYPE_CHANGE',key)
    assert new.number_format==cell.number_format,('FORMAT_CHANGE',key)
    if not(start<=cell.row<start+4):assert new._style==cell._style,('STYLE_CHANGE',key)
  assert s.freeze_panes==t.freeze_panes and s.auto_filter.ref==t.auto_filter.ref
  assert len(s.data_validations.dataValidation)==len(t.data_validations.dataValidation)
  assert [(c.coordinate,c.value) for row in s for c in row if c.data_type=='f']==[(c.coordinate,c.value) for row in t for c in row if c.data_type=='f']
 with zipfile.ZipFile(original) as a,zipfile.ZipFile(target) as b:
  assert set(a.namelist())==set(b.namelist());changed=[n for n in a.namelist() if a.read(n)!=b.read(n)];assert changed==['xl/worksheets/sheet1.xml'],changed
 proof.append({'path':str(target.relative_to(ROOT)),'before':hashlib.sha256(original.read_bytes()).hexdigest(),'after':hashlib.sha256(target.read_bytes()).hexdigest(),'authored_with':'Artifact Tool','rejected_blank_value_drift':drift,'changed_package_parts':changed,'changed_values':values,'all_other_values_formulas_formats':'UNCHANGED'})
(HERE/'xlsx-preservation.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n');print('STRICT_XLSX_PRESERVATION=PASS')
