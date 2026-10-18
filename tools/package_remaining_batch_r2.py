#!/usr/bin/env python3
"""Reproducible R2 buyer-review packaging and checks against immutable R1."""
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from hashlib import sha256
from io import BytesIO
from docx import Document
from openpyxl import load_workbook
import json, subprocess, sys, re
from refine_remaining_batch_r2 import BASE, OUT, SKUS, source
from office_core_metadata_gate import check_bytes as check_office_metadata

COUNTS={'P1':19,'P2':15,'P4':23,'P7':13,'P8':11,'P9':10,'P10':8}
ALLOWED={
 'P1':{'14-raschet-procentov-395-gk.xlsx':{'Входные данные!A9'}},
 'P4':{name:{'Инструкция!A3'} for name in ('35-obshiy-zhurnal-rabot.xlsx','36-vhodnoy-kontrol.xlsx','37-reestr-pasportov.xlsx','38-reestr-id.xlsx')},
 'P7':{name:{'Инструкция!A3'} for name in ('02-perechen-rabot.xlsx','03-kalendarnyy-plan.xlsx','05-grafik-platezhey.xlsx')},
 'P8':{'02-kalkulyator-usloviy-avansa.xlsx':{'Условия!A34','Условия!A35','Сравнение!A19','Сравнение!A20','Как читать!A17','Как читать!A18','Пример заполнения!A26'}},
}
def digest(data):return sha256(data).hexdigest()
def cells(book):
 return {f'{s.title}!{c.coordinate}':c.value for s in book for row in s for c in row if c.value is not None}
def workbook_check(before,after,sku,name):
 a=load_workbook(BytesIO(before));b=load_workbook(BytesIO(after))
 assert a.sheetnames==b.sheetnames,(sku,name,'sheet names')
 ac,bc=cells(a),cells(b)
 assert {k:v for k,v in ac.items() if isinstance(v,str) and v.startswith('=')}=={k:v for k,v in bc.items() if isinstance(v,str) and v.startswith('=')},(sku,name,'formulas')
 delta={k for k in ac.keys()|bc.keys() if ac.get(k)!=bc.get(k)}
 assert delta<=ALLOWED.get(sku,{}).get(name,set()),(sku,name,'unexpected cell edits',delta)
 for x,y in zip(a,b):
  assert set(str(m) for m in x.merged_cells.ranges)==set(str(m) for m in y.merged_cells.ranges),(sku,name,'merges')
  assert len(x.data_validations.dataValidation)==len(y.data_validations.dataValidation),(sku,name,'validations')
  assert len(x.conditional_formatting)==len(y.conditional_formatting),(sku,name,'conditional formats')
  assert x.freeze_panes==y.freeze_panes,(sku,name,'freeze panes')
 return len([v for v in bc.values() if isinstance(v,str) and v.startswith('=')]),sorted(delta)
def docx_check(data,sku,name):
 d=Document(BytesIO(data))
 for s in d.sections:
  w,h=s.page_width.cm,s.page_height.cm
  assert abs(sorted((w,h))[0]-21)<0.05 and abs(sorted((w,h))[1]-29.7)<0.05,(sku,name,'paper',w,h)
 if sku in ('P1','P2','P4'):
  for style in d.styles:
   if style.name=='Normal':assert style.font.name=='Arial',(sku,name,'Normal font')
  for p in d.paragraphs:
   for r in p.runs:
    if r.text:assert r.font.name in (None,'Arial'),(sku,name,'run font',r.font.name)
 text=' '.join(p.text for p in d.paragraphs)
 assert not re.search(r'(зачем этот документ|чего (?:комплект|файл|документ) не делает|не юридическая консультация|до подписи)',text,re.I),(sku,name,'editorial')
 return len(d.sections)
def pdf_check(sku,name):
 p=(OUT/sku/name) if name.endswith('.pdf') else (Path('/tmp/qa-r2')/sku/(Path(name).stem+'.pdf'))
 assert p.exists(),(sku,name,'missing rendered PDF')
 text=subprocess.check_output(['pdftotext','-layout',str(p),'-'],text=True)
 pages=text.split('\f')[:-1]
 assert pages and all(len(re.sub(r'\s+','',x))>40 for x in pages),(sku,name,'blank/orphan page',list(map(len,pages)))
 if sku=='P1' and name=='02-ks-3.docx':assert len(pages)==1,(sku,name,'orphan signature')
 return len(pages)
summary=[]
for sku in SKUS:
 with ZipFile(source(sku)) as old:
  prior={n:old.read(n) for n in old.namelist()}
 files=sorted(x for x in (OUT/sku).iterdir() if x.is_file())
 assert len(files)==COUNTS[sku],(sku,len(files))
 assert [x.name for x in files]==sorted(prior),(sku,'names')
 intro=[x for x in files if x.name.startswith('00-START-HERE')]
 if intro and intro[0].suffix=='.txt':
  body=intro[0].read_text()
  assert all(x.name in body for x in files if x!=intro[0]),(sku,'inventory')
 changed=[]; formulas=0; details=[]; printpages={}
 for f in files:
  data=f.read_bytes();name=f.name
  if f.suffix in {'.docx','.xlsx'}:check_office_metadata(data,sku+'/'+name)
  if prior[name]!=data:changed.append(name)
  if name.endswith('.xlsx'):
   n,edits=workbook_check(prior[name],data,sku,name);formulas+=n
  if name.endswith('.docx'):docx_check(data,sku,name)
  if name.endswith(('.docx','.xlsx')):
   with ZipFile(BytesIO(data)) as office:assert office.testzip() is None,(sku,name,'inner CRC')
   printpages[name]=pdf_check(sku,name)
  if name.endswith('.pdf'):
   assert data.startswith(b'%PDF'),(sku,name)
   printpages[name]=pdf_check(sku,name)
  details.append({'name':name,'sha256':digest(data),'size_bytes':len(data)})
 target=BASE/f'{sku}-buyer-review-20261018-R2.zip'
 with ZipFile(target,'w',compression=ZIP_DEFLATED,compresslevel=9) as z:
  for f in files:
   info=ZipInfo(f.name,date_time=(2026,10,18,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o644<<16
   z.writestr(info,f.read_bytes(),compress_type=ZIP_DEFLATED,compresslevel=9)
 with ZipFile(target) as z:assert z.testzip() is None and z.namelist()==[x.name for x in files]
 g=subprocess.run([sys.executable,str(BASE.parents[1]/'editorial_business_style_gate.py'),str(target)],capture_output=True,text=True)
 assert g.returncode==0,(sku,g.stdout,g.stderr)
 manifest={'sku':sku,'status':'BUYER_REVIEW_ONLY','archive':target.name,'sha256':digest(target.read_bytes()),'size_bytes':target.stat().st_size,'file_count':len(files),'changed_from_R1':changed,'members':details,'excel_formulas':formulas,'print_pages':printpages,'checks':{'CRC':'PASS','inner_office_CRC':'PASS','editorial':g.stdout.strip(),'A4_fonts':'PASS','signatures_and_page_flow':'PASS','XLSX_formulas_and_features':'PASS','no_blank_or_orphan_pages':'PASS'}}
 (BASE/f'{sku}-manifest-R2.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
 summary.append(manifest)
 print(sku,len(files),len(changed),formulas,manifest['sha256'])
(BASE/'BATCH-MANIFEST-R2.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
