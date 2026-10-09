#!/usr/bin/env python3
"""Synchronize the P1 quick-start instruction with its delivered files."""
from pathlib import Path
import subprocess
import tempfile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "products-storage/01-zakrytie-rabot"
MIRROR = ROOT / "products-storage/04-polnyy-komplekt-pto/01-30-bazovye-pakety/01-zakrytie-rabot"
path = BASE / "00-INSTRUKCIYA.docx"
doc = Document(path)
table = doc.tables[0]
rows = [
    ("11-slovar-poley.docx", "Для заполнения форм: значения полей и источники данных", "Проверены исходные данные для заполнения"),
    ("12-pretenziya-na-neoplatu-po-ks-2.docx", "После проверки приёмки, срока оплаты и претензионного порядка", "Адаптированная претензия с приложениями"),
    ("13-uvedomlenie-o-prosrochke-oplaty.docx", "При наступлении просрочки оплаты по договору", "Адаптированное уведомление и доказательство отправки"),
    ("14-raschet-procentov-395-gk.xlsx", "Для расчёта процентов по ст. 395 ГК РФ по периодам", "Проверенный расчёт с исходными данными"),
    ("15-algoritm-pri-zaderzhke-oplaty.docx", "Для выбора следующего шага при задержке оплаты", "План действий и контроль сроков"),
    ("16-iskovoe-zayavlenie-o-vzyskanii.docx", "Перед подготовкой обращения в арбитражный суд", "Адаптированный проект иска и список приложений"),
]
existing = {r.cells[0].text for r in table.rows}
for filename, when, result in rows:
    if filename not in existing:
        row = table.add_row()
        for cell, text in zip(row.cells, (filename, when, result)):
            cell.text = text
short_rows = {
    "01-ks-2.docx": ("При оформлении объёмов работ", "Проект КС-2 с проверенными объёмами"),
    "02-ks-3.docx": ("При своде стоимости по КС-2", "Проект КС-3 с проверенной суммой"),
    "03-akt-vypolnennyh-rabot.docx": ("Для фиксации выполненных работ", "Акт с датой и подписями"),
    "04-akt-priemki.docx": ("При согласовании результата работ", "Акт приёмки с замечаниями или подписями"),
    "05-peredatochnyy-akt.docx": ("При передаче пакета заказчику", "Состав переданного и подтверждение вручения"),
    "06-zhurnal-obemov.xlsx": ("Для фиксации выполненных объёмов", "Журнал и ссылки на доказательства"),
    "07-reestr-zamechaniy.xlsx": ("Для учёта и закрытия замечаний", "Статусы, сроки и доказательства закрытия"),
    "08-reestr-peredachi.xlsx": ("Для фиксации состава и передачи", "Дата, канал, получатель и подтверждение"),
    "09-checklist-peredachi.pdf": ("Перед передачей комплекта", "Отмеченные проверки и пробелы"),
}
for row in table.rows:
    if row.cells[0].text == "11-slovar-poley.docx":
        row.cells[1].text = "Для заполнения форм: значения полей и источники данных"
    if row.cells[0].text in short_rows:
        row.cells[1].text, row.cells[2].text = short_rows[row.cells[0].text]
    if row.cells[0].text == "10-sroki-hraneniya.pdf":
        row.cells[1].text = "Для организации хранения и подтверждений передачи"
        row.cells[2].text = "Определены правила хранения по виду документов"
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
for i, row in enumerate(table.rows):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if i == 0 and trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))
    for cell in row.cells:
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)
            for run in para.runs:
                run.font.size = Pt(9)
old_note = "Сверьте правовые нормы на дату использования: исходная проверка расширения проведена 16.08.2026; ставки, сроки и реквизиты уточняйте по делу."
note = "Сверьте правовые нормы и реквизиты на дату использования; ставки, сроки и порядок берите из источников, указанных в соответствующем файле."
criteria = ("Пакет внедрён, когда по каждому объёму зафиксированы основание, факт выполнения и доказательства в журнале 06, "
            "суммы согласованы в формах 01–03, состав и передача отмечены в 05 и 08, замечания отражены в 07. "
            "Срок оплаты берите из договора и проверяйте по словарю 11. Частичные оплаты сверяйте с банковскими выписками перед расчётом в 14.")
old_referral = "Спорные юридические формулировки согласуйте с профильным специалистом."
referral_action = "Сопоставьте формулировки с условиями договора и подтверждающими документами объекта."
criteria_found = 0
for paragraph in doc.paragraphs:
    replacement = note if paragraph.text == old_note else criteria if paragraph.text.startswith("Пакет внедрён, когда") else None
    if paragraph.text == "Не отправляйте документ с квадратными скобками и подсказками.":
        replacement = ("Заполните поля в квадратных и двойных фигурных скобках, затем удалите обозначения полей и подсказки. "
                       "Проверьте заполненную форму перед отправкой.")
    if paragraph.text.startswith("Заполните поля в квадратных и двойных фигурных скобках"):
        replacement = ("Заполните поля в квадратных и двойных фигурных скобках, затем удалите обозначения полей и подсказки. "
                       "Проверьте заполненную форму перед отправкой.")
    if paragraph.text == old_referral:
        replacement = referral_action
    if replacement is not None:
        if paragraph.text.startswith("Пакет внедрён, когда"):
            criteria_found += 1
        if paragraph.runs:
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(replacement)
assert criteria_found == 1, "Inspect the instruction before replacing its implementation criterion"

if not any(p.text == note for p in doc.paragraphs):
    for p in doc.paragraphs:
        if p.text == "Структура папки объекта":
            p.insert_paragraph_before(note, style="List Bullet")
            break
for section in doc.sections:
    for p in section.header.paragraphs:
        for run in p.runs:
            run.text = run.text.replace("РАБОЧИЙ ШАБЛОН", "ФОРМЫ И ПОРЯДОК")
    for p in section.footer.paragraphs:
        if "версия 21.07.2026" in p.text:
            for run in p.runs:
                run.text = run.text.replace("версия 21.07.2026", "редакция 08.10.2026")
# Stage both navigation formats before replacing delivered copies.
with tempfile.TemporaryDirectory(prefix="p1-instruction-") as temporary:
    stage = Path(temporary)
    staged_docx = stage / path.name
    doc.save(staged_docx)
    subprocess.run([
        "libreoffice", "-env:UserInstallation=" + (stage / "profile").as_uri(),
        "--headless", "--convert-to", "pdf", "--outdir", str(stage), str(staged_docx),
    ], check=True, timeout=30, capture_output=True)
    staged_pdf = stage / path.with_suffix(".pdf").name
    assert staged_pdf.is_file(), "Instruction PDF export did not produce an artifact"
    for destination in (BASE, MIRROR):
        (destination / path.name).write_bytes(staged_docx.read_bytes())
        (destination / staged_pdf.name).write_bytes(staged_pdf.read_bytes())
print("P1 instruction:",len(table.rows)-1,"listed deliverables; DOCX/PDF synchronized")
