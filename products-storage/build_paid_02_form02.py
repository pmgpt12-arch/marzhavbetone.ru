"""Собрать форму 02 для двух маршрутов комплекта дополнительных работ."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FILE_NAME = "02-soglasovanie-obema.docx"
GRAPHITE = RGBColor(38, 47, 55)
GOLD = RGBColor(142, 111, 50)
LIGHT = "F3F0E8"


def shade(cell, color: str = LIGHT) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def fill(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.2)
    run.font.color.rgb = GRAPHITE
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def two_column(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(13.1)
    for key, value in rows:
        cells = table.add_row().cells
        no_split(table.rows[-1])
        fill(cells[0], key, bold=True)
        shade(cells[0])
        fill(cells[1], value)


def heading(doc: Document, title: str, *, new_page: bool = False) -> None:
    p = doc.add_heading(title, level=1)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = new_page


def body(doc: Document, text: str, *, bold_start: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)


def position_table(doc: Document) -> None:
    headings = ["№", "Работа, место и границы", "Отличие от договора / проекта", "Доп. объём и ед.", "Основание / подтверждение"]
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    table.autofit = False
    widths = [0.8, 4.1, 4.0, 2.8, 5.4]
    for idx, (name, width) in enumerate(zip(headings, widths)):
        table.columns[idx].width = Cm(width)
        fill(table.rows[0].cells[idx], name, bold=True)
        shade(table.rows[0].cells[idx])
    no_split(table.rows[0])
    for n in range(1, 4):
        cells = table.add_row().cells
        no_split(table.rows[-1])
        for cell, value in zip(cells, [str(n), "[___]", "[___]", "[___]", "[___]"]):
            fill(cell, value)
    body(doc, "Если позиций больше трёх, добавьте строки. К каждой позиции укажите документ, по которому проверяются границы и количество.")


def attachment_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, ["Приложение", "Номер и дата", "Какие позиции подтверждает"]):
        fill(cell, title, bold=True)
        shade(cell)
    no_split(table.rows[0])
    for label in ["Схема, чертёж или дефектная ведомость", "Расчёт стоимости (файл 07)", "Журнал, фото, переписка или акт"]:
        cells = table.add_row().cells
        no_split(table.rows[-1])
        for cell, value in zip(cells, [label, "[___]", "[___]"]):
            fill(cell, value)
    body(doc, "Оставьте в перечне только приложенные документы и укажите их фактические реквизиты.")


def signature_table(doc: Document) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    titles = table.add_row().cells
    no_split(table.rows[-1])
    for cell, title in zip(titles, ["Подрядчик", "Заказчик"]):
        fill(cell, title, bold=True)
        shade(cell)
    for left, right in [
        ("Должность, Ф. И. О. [___]", "Должность, Ф. И. О. [___]"),
        ("Основание полномочий [___]", "Основание полномочий [___]"),
        ("Подпись [___]  Дата [___]", "Подпись [___]  Дата [___]"),
    ]:
        cells = table.add_row().cells
        no_split(table.rows[-1])
        fill(cells[0], left)
        fill(cells[1], right)


def make_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = GRAPHITE
    normal.paragraph_format.space_after = Pt(5)
    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = "Times New Roman"
    heading_style.font.color.rgb = GOLD
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.paragraph_format.space_before = Pt(9)
    heading_style.paragraph_format.space_after = Pt(5)

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(17)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(9)
    border = title_style.element.get_or_add_pPr().find(qn("w:pBdr"))
    if border is not None:
        title_style.element.get_or_add_pPr().remove(border)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(9)
    title_run = title.add_run("Согласование объёма дополнительных работ")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(17)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body(doc, "Заполните данные объекта и выберите один маршрут. Для направления заказчику сохраните экземпляр и подтверждение доставки; после ответа зафиксируйте решение по позициям.")

    heading(doc, "Реквизиты")
    two_column(doc, [
        ("Номер и дата", "№ [___] от [___]; место составления [___]"),
        ("Договор", "Договор подряда № [___] от [___], приложение / раздел исходного объёма [___]"),
        ("Объект", "Название, адрес, участок или захватка [___]"),
        ("Подрядчик", "Полное наименование, ИНН, адрес [___]"),
        ("Заказчик", "Полное наименование, ИНН, адрес [___]"),
    ])

    heading(doc, "Основание для дополнительного объёма")
    two_column(doc, [
        ("Источник задания", "Поручение, письмо, протокол или выявленное обстоятельство: кто, когда, № документа [___]"),
        ("Исходный объём", "Пункт договора, сметы, проекта или ведомости и соответствующая позиция [___]"),
        ("Причина отличия", "Что обнаружено или изменено на объекте и почему нужна отдельная позиция [___]"),
        ("Уведомление", "Файл 03: исходящий № [___], дата [___], канал и подтверждение направления [___]"),
    ])

    heading(doc, "Позиции дополнительного объёма")
    position_table(doc)

    heading(doc, "Выберите маршрут")
    body(doc, "□ А. Работы ещё не начаты. Планируемая дата начала [___]. Письменное поручение заказчика (файл 01): № [___] от [___]. Решение об объёме и приложениях получите до начала соответствующих работ.")
    body(doc, "□ Б. Работы уже выполнены. Фактический период [___] — [___]. Просьба заказчика подтверждается [документ, дата, отправитель/подписант: ___]. Выполнение подтверждается [журнал, акт, фото с датами, переписка: ___]. Оформляйте согласование текущей датой.")
    body(doc, "Для смешанного случая оформите отдельные строки по планируемому и выполненному объёму. Перед направлением оставьте нужный маршрут и конкретные подтверждения.")

    heading(doc, "Приложения и расчёт", new_page=True)
    attachment_table(doc)
    body(doc, "Цена и порядок оплаты: расчёт 07 № [___] от [___]; предложение о согласовании цены — допсоглашение 11 № [___] от [___]. Влияние на срок: документ 05 № [___] от [___]; при изменении срока — допсоглашение 12 № [___] от [___]. Заполняйте ссылки по фактически подготовленным документам.")

    heading(doc, "Решение по объёму")
    body(doc, "□ Объём по позициям № [___] согласован сторонами в указанных границах и единицах измерения.")
    body(doc, "□ По позициям № [___] нужны уточнения: [___]. Срок ответа [___].")
    body(doc, "□ Позиции № [___] заказчик не согласовал; ответ № [___] от [___].")
    body(doc, "Подписанные обеими сторонами позиции с отметкой «согласован» фиксируют решение об объёме. Цена, порядок оплаты и изменение сроков отражаются в названных документах 11 и 12 после их оформления.")

    heading(doc, "Направление и подписи")
    two_column(doc, [
        ("Передача", "Канал [___], дата [___], адресат [___], подтверждение доставки [___]"),
        ("Ответ", "Дата [___], входящий № / письмо [___], связь с позициями [___]"),
    ])
    signature_table(doc)
    body(doc, "Копию подписанного документа и приложения сохраните вместе с договором, уведомлением 03 и записью в журнале 08.")
    return doc


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, help="Путь для проверки; по умолчанию обновляется активная форма P2")
    args = parser.parse_args()
    if args.output:
        outputs = [args.output]
    else:
        root = Path(__file__).resolve().parent
        outputs = [root / "02-dopraboty-bez-poter" / FILE_NAME]
    document = make_document()
    for output in outputs:
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        print(output)


if __name__ == "__main__":
    main()
