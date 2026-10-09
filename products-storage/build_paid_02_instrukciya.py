"""Инструкция покупателя P2: `00-INSTRUKCIYA.docx` и `00-INSTRUKCIYA.pdf`.

Зачем скрипт существует. `00-START-HERE.txt` велит открыть инструкцию
первой и обещает в ней два случая: работы ещё не начаты и работы уже
выполнены без подписанных бумаг. В архиве лежал общий шаблон линейки под
прежним именем «Допработы без потерь», без допсоглашений 11 и 12 и без
обоих случаев (дефект P2-01, отчёт PR #300, issue #305). Генератора у того
файла в репозитории не было, поэтому и исправить его было нечем.

Одно содержание — два файла. DOCX и PDF собираются из одной модели
`СОДЕРЖАНИЕ` ниже: PDF, отставший от DOCX, ловит
`tools/test_p2_kit_instruction.py`.

Границы текста. Инструкция описывает порядок работы с файлами комплекта и
его пределы. Ссылок на нормы и обещаний результата в ней нет: нормативная
часть живёт в самих формах и проходит сверку отдельно.

Воспроизводимость. Даты свойств DOCX и метки времени внутри его zip
фиксированы, PDF пишется reportlab с `invariant=1`: повторный прогон даёт
те же байты. Пишется только папка отдельного комплекта — вложенная копия в
`04-polnyy-komplekt-pto` — исторический архив, не адресуется ни одним sku
и не меняется.

Запуск: python3 products-storage/build_paid_02_instrukciya.py
"""
import datetime
import io
import os
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.join(BASE_DIR, "02-dopraboty-bez-poter")
DOCX = os.path.join(KIT, "00-INSTRUKCIYA.docx")
PDF = os.path.join(KIT, "00-INSTRUKCIYA.pdf")

# Версия текста. Она же — дата в свойствах файлов и метка времени в zip:
# меняется только вместе с текстом.
ВЕРСИЯ = datetime.datetime(2026, 10, 5, 0, 0, 0)

ТОВАР = "Дополнительные работы: как получить оплату"
ЗАГОЛОВОК = f"Инструкция к комплекту «{ТОВАР}»"

# Названия случаев дословно те, что стоят в START-HERE §2 в кавычках:
# тест ищет их заголовками здесь.
СЛУЧАЙ_А = "Работы ещё не начаты"
СЛУЧАЙ_Б = "Работы уже выполнены без подписанных бумаг"

# Модель содержания: (вид, данные). Виды: lead — вводный абзац, h — раздел,
# p — абзац, steps — нумерованные шаги, bullets — список, table — таблица
# (первая строка — шапка).
СОДЕРЖАНИЕ = [
    ("lead",
     "Комплект помогает оформить дополнительный объём так, чтобы у вас "
     "были письменные основания: поручение, согласованные объём, цена и "
     "срок, — и включить этот объём в закрывающие документы. Порядок "
     "работы зависит от того, начаты ли работы. Сначала определите свой "
     "случай, потом идите по шагам своего раздела."),

    ("h", "1. Какой у вас случай"),
    ("bullets", [
        f"«{СЛУЧАЙ_А}» — заказчик просит сделать сверх договора, а вы ещё "
        "не приступили. Идите по разделу 2.",
        f"«{СЛУЧАЙ_Б}» — объём сделан по устной просьбе или по ходу работ, "
        "письменного согласования до начала нет. Идите по разделу 3. Этот "
        "маршрут помогает собрать подтверждения поручения и фактического объёма.",
        "Часть уже сделана, часть ещё нет — сделанную часть ведите по "
        "разделу 3, остальное по разделу 2, отдельными строками журнала.",
    ]),
    ("p",
     "Общий порядок на одной странице — 10-algoritm-doprabot.pdf. Разделы "
     "2 и 3 раскладывают его по файлам комплекта."),

    ("h", f"2. Случай А. {СЛУЧАЙ_А}"),
    ("steps", [
        "Зафиксируйте просьбу письменно: 03-uvedomlenie-o-doprabotah.docx — "
        "уведомление заказчику о выявленной работе и запросе решения. "
        "Выберите маршрут А, заполните фактические реквизиты, состав "
        "объёма, срок ответа, статус участка и приложения. Маршрут Б "
        "перед отправкой удалите.",
        "Отделите дополнительный объём от договорного и опишите его: "
        "02-soglasovanie-obema.docx — причина, отличие от исходного объёма, "
        "границы, единица, количество и подтверждающий документ. "
        "Выберите маршрут А и удалите маршрут Б в отправляемом экземпляре.",
        "Посчитайте цену: 07-raschet-stoimosti.xlsx — одна позиция в "
        "строке, в колонке «Источник цены» укажите, откуда взята цена.",
        "Оцените влияние на срок: 05-izmenenie-srokov.docx — затронутые "
        "операции, дополнительная продолжительность, новая контрольная дата "
        "и обновлённый график.",
        "Получите письменное поручение до начала работ: "
        "01-prikaz-na-dopobem.docx — кто поручает, что, где, в каком "
        "объёме, где зафиксированы цена, срок и приёмка. Проверьте, "
        "что подписывает лицо с полномочиями по вашему договору.",
        "Закрепите объём и цену допсоглашением: "
        "11-dopsoglashenie-obem-i-cena.docx, приложения — описание объёма "
        "из 02-soglasovanie-obema.docx и расчёт из "
        "07-raschet-stoimosti.xlsx. Если меняется срок — "
        "12-dopsoglashenie-sroki-doprabot.docx.",
        "Если ответа нет по применимому сроку и соответствующие работы "
        "приостановлены — 06-pismo-o-priostanovke.docx. Укажите точные "
        "позиции, дату, меры сохранности и подтверждение доставки уведомления 03.",
        "По ходу работ ведите 08-zhurnal-doprabot.xlsx (каждая позиция — "
        "строка со статусом) и снимайте по 09-checklist-fotofiksacii.pdf. "
        "Работы, которые закроются следующими, оформите до закрытия: "
        "04-akt-skrytyh-rabot.docx — форму и подписантов сверьте с проектом "
        "и требованиями заказчика.",
        "Включите согласованный объём в закрывающие документы и отметьте в "
        "журнале колонки «Включено в КС-2» и «Оплачено».",
    ]),
    ("p",
     "Случай А закрыт, когда дополнительный объём согласован письменно до "
     "начала работ либо оформлен допсоглашением с ценой и сроком."),

    ("h", f"3. Случай Б. {СЛУЧАЙ_Б}"),
    ("p",
     "Соберите документы, подтверждающие просьбу заказчика и выполненный объём. "
     "Зафиксируйте их в журнале, рассчитайте стоимость и направьте предложение "
     "оформить дополнительный объём текущей датой."),
    ("steps", [
        "Оформляйте документы текущей датой. Письменное поручение 01 относится "
        "к работам до их начала; для выполненного объёма переходите к журналу, "
        "подтверждениям и уведомлению.",
        "Соберите, чем подтверждается просьба заказчика и выполнение: "
        "переписку, записи в журналах работ, протоколы совещаний, "
        "исполнительную документацию, фото с датами. Внесите каждую позицию "
        "в 08-zhurnal-doprabot.xlsx, в колонке «Основание» — чем она "
        "подтверждена.",
        "Для каждой позиции укажите в журнале конкретное подтверждение поручения "
        "и выполнения. Отдельно отметьте позиции, по которым нужно собрать "
        "дополнительные документы или переписку.",
        "Разложите имеющиеся фото по 09-checklist-fotofiksacii.pdf и "
        "отметьте, каких кадров нет. Если выполненное ещё доступно для "
        "осмотра, а следующие работы его закроют, — оформите "
        "04-akt-skrytyh-rabot.docx до закрытия.",
        "Опишите объём в 02-soglasovanie-obema.docx (маршрут Б: фактический "
        "период, просьба заказчика и подтверждения по позициям) и посчитайте цену в "
        "07-raschet-stoimosti.xlsx.",
        "Направьте заказчику 03-uvedomlenie-o-doprabotah.docx: когда и по "
        "чьей просьбе выполнен объём, приложения — 02 и расчёт из 07, "
        "просьба согласовать объём и цену. Выберите маршрут Б с фактическими "
        "датами и документами; удалите маршрут А и относящиеся к нему поля. "
        "В приложениях укажите 02, расчёт 07 и подтверждения выполнения. "
        "Не меняйте дату уже выполненных работ и дату текущего уведомления.",
        "Предложите оформить объём допсоглашением "
        "11-dopsoglashenie-obem-i-cena.docx — в разделе «Основание» "
        "перечислите, чем зафиксировано поручение. Если выполненные работы "
        "сдвинули сроки — 12-dopsoglashenie-sroki-doprabot.docx.",
        "После подписания оформите приёмочные документы по договору и "
        "отметьте в журнале включение в КС-2 и фактическое поступление оплаты отдельно.",
        "При отказе или отсутствии ответа сохраните уведомление, доказательства "
        "доставки, переписку, журнал, акты и расчёт стоимости единым пакетом "
        "для следующего этапа работы с задолженностью.",
    ]),

    ("h", "4. Что лежит в архиве"),
    ("table", [
        ["Файл", "Что это", "Случай"],
        ["00-START-HERE.txt", "С чего начать: ситуация, состав, границы", "—"],
        ["00-INSTRUKCIYA.docx", "Эта инструкция, Word", "—"],
        ["00-INSTRUKCIYA.pdf", "Эта инструкция, PDF", "—"],
        ["01-prikaz-na-dopobem.docx",
         "Письменное поручение на дополнительный объём", "А"],
        ["02-soglasovanie-obema.docx",
         "Описание и согласование дополнительного объёма", "А, Б"],
        ["03-uvedomlenie-o-doprabotah.docx",
         "Уведомление заказчику о дополнительных работах", "А, Б"],
        ["04-akt-skrytyh-rabot.docx",
         "Акт освидетельствования скрываемых работ", "А, Б"],
        ["05-izmenenie-srokov.docx",
         "Расчёт влияния на срок и новая контрольная дата", "А"],
        ["06-pismo-o-priostanovke.docx",
         "Письмо о приостановке до решения заказчика", "А"],
        ["07-raschet-stoimosti.xlsx",
         "Расчёт стоимости дополнительных работ", "А, Б"],
        ["08-zhurnal-doprabot.xlsx",
         "Журнал: основание, объём, статус, КС-2, оплата", "А, Б"],
        ["09-checklist-fotofiksacii.pdf", "Чек-лист фотофиксации", "А, Б"],
        ["10-algoritm-doprabot.pdf",
        "Два маршрута и контрольный порядок", "А, Б"],
        ["11-dopsoglashenie-obem-i-cena.docx",
         "Допсоглашение об изменении объёма и цены", "А, Б"],
        ["12-dopsoglashenie-sroki-doprabot.docx",
         "Допсоглашение об изменении сроков", "А, Б"],
    ]),
    ("p",
     "Word и Excel редактируются, PDF — чек-лист и алгоритм для работы, их "
     "заказчику не отправляют."),

    ("h", "5. Как закрепить результат работы"),
    ("bullets", [
        "Сохраните поручение, описание дополнительного объёма и согласованную цену вместе с исходным договором.",
        "Свяжите каждую позицию журнала с письмом, актом, фотографиями и расчётом стоимости.",
        "После подписания включите объём в закрывающие документы и отслеживайте его оплату.",
        "При отказе или отсутствии ответа сохраните подтверждение направления документов и собранные доказательства.",
    ]),

    ("h", "6. Перед отправкой любого документа"),
    ("bullets", [
        "Проверьте заполненные реквизиты сторон, договора, объекта, даты, "
        "подписи, полномочия подписантов и приложения в формах 01–06, 11 и 12.",
        "Замените поля в квадратных скобках фактическими данными.",
        "В формах 02, 03 и 11 оставьте только выбранный маршрут и относящиеся "
        "к нему заполненные сведения. В направляемых документах не должно "
        "остаться квадратных скобок и второго варианта на выбор.",
        "Перед отправкой удалите рабочие подсказки, оставьте заполненные реквизиты, "
        "основной текст и приложения, которые относятся к вашему объекту.",
        "Проверьте, что на экземпляре для заказчика указаны стороны, объект, дата, "
        "подписи и перечень приложений.",
        "Сохраните отправленную версию и подтверждение доставки рядом с "
        "записью в журнале 08-zhurnal-doprabot.xlsx.",
    ]),

    ("h", "7. Если ситуация изменилась"),
    ("p",
     "Допы согласованы и выполнены, работы закрыты, а оплаты нет — спор "
     "перешёл с доказанности поручения на сами деньги, и это другой набор "
     "документов. Куда идти дальше, сказано в 00-START-HERE.txt, раздел 7."),
    ("p", "Вопросы по составу комплекта: marzhavbetone@yandex.ru"),
]


# --- DOCX ------------------------------------------------------------------

def _docx_bytes():
    doc = Document()
    стиль = doc.styles["Normal"]
    стиль.font.name = "Arial"
    стиль.font.size = Pt(11)
    стиль.font.color.rgb = RGBColor(45, 48, 51)
    for имя in ("Title", "Heading 1"):
        doc.styles[имя].font.name = "Arial"
        doc.styles[имя].font.color.rgb = RGBColor(45, 48, 51)
    doc.styles["Title"].font.size = Pt(20)
    граница = doc.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
    if граница is not None:
        doc.styles["Title"].element.get_or_add_pPr().remove(граница)
    for раздел in doc.sections:
        раздел.left_margin = раздел.right_margin = Cm(2)
        раздел.top_margin = раздел.bottom_margin = Cm(2)

    doc.add_heading(ЗАГОЛОВОК, level=0)
    for вид, данные in СОДЕРЖАНИЕ:
        if вид == "h":
            doc.add_heading(данные, level=1)
        elif вид in ("p", "lead"):
            doc.add_paragraph(данные)
        elif вид == "steps":
            # Номера пишутся текстом: у встроенного «List Number» одна
            # нумерация на документ, и шаги раздела 3 продолжили бы раздел 2.
            for n, шаг in enumerate(данные, 1):
                абзац = doc.add_paragraph(f"{n}. {шаг}")
                абзац.paragraph_format.left_indent = Cm(0.8)
                абзац.paragraph_format.first_line_indent = Cm(-0.8)
        elif вид == "bullets":
            for пункт in данные:
                doc.add_paragraph(пункт, style="List Bullet")
        elif вид == "table":
            таблица = doc.add_table(rows=0, cols=len(данные[0]))
            таблица.style = "Table Grid"
            for i, строка in enumerate(данные):
                ячейки = таблица.add_row().cells
                for ячейка, значение in zip(ячейки, строка):
                    ячейка.text = значение
                    if i == 0:
                        for r in ячейка.paragraphs[0].runs:
                            r.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                        заливка = OxmlElement("w:shd")
                        заливка.set(qn("w:fill"), "303336")
                        ячейка._tc.get_or_add_tcPr().append(заливка)
                ячейки[-1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            raise ValueError(вид)

    свойства = doc.core_properties
    свойства.title = ЗАГОЛОВОК
    свойства.author = "Маржа в бетоне"
    свойства.last_modified_by = "Маржа в бетоне"
    свойства.created = свойства.modified = ВЕРСИЯ
    свойства.revision = 1

    сырой = io.BytesIO()
    doc.save(сырой)
    # python-docx ставит каждой части zip текущее время. Пересобираем архив
    # с меткой версии и тем же порядком частей — повторный прогон даёт те
    # же байты.
    вход = zipfile.ZipFile(io.BytesIO(сырой.getvalue()))
    выход_буфер = io.BytesIO()
    with zipfile.ZipFile(выход_буфер, "w", zipfile.ZIP_DEFLATED) as выход:
        for часть in вход.infolist():
            инфо = zipfile.ZipInfo(часть.filename, ВЕРСИЯ.timetuple()[:6])
            инфо.compress_type = zipfile.ZIP_DEFLATED
            инфо.external_attr = 0o644 << 16
            выход.writestr(инфо, вход.read(часть.filename))
    return выход_буфер.getvalue()


# --- PDF -------------------------------------------------------------------

_ШРИФТЫ = (
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
)


def _шрифты():
    """Кириллический TTF обязателен: стандартные шрифты reportlab кириллицы
    не содержат и дают пустую страницу без единой ошибки."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for обычный, жирный in _ШРИФТЫ:
        if os.path.exists(обычный) and os.path.exists(жирный):
            if "MVBSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("MVBSans", обычный))
                pdfmetrics.registerFont(TTFont("MVBSans-Bold", жирный))
            return "MVBSans", "MVBSans-Bold"
    raise SystemExit("Нет кириллического TTF (Liberation Sans или DejaVu Sans)")


def _pdf_bytes():
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                    SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    обычный, жирный = _шрифты()
    графит = colors.HexColor("#303336")
    золото = colors.HexColor("#A98A54")
    текст = ParagraphStyle("t", fontName=обычный, fontSize=10.5, leading=14,
                           textColor=графит,
                           spaceAfter=4)
    вводный = ParagraphStyle("lead", parent=текст, spaceAfter=8)
    заголовок = ParagraphStyle("title", fontName=жирный, fontSize=16,
                               textColor=графит,
                               leading=20, spaceAfter=10)
    раздел = ParagraphStyle("h", fontName=жирный, fontSize=12.5, leading=16,
                            textColor=графит,
                            spaceBefore=10, spaceAfter=5)
    ячейка = ParagraphStyle("cell", parent=текст, fontSize=9, leading=11.5,
                            spaceAfter=0)
    шапка = ParagraphStyle("head", parent=ячейка, fontName=жирный,
                           textColor=colors.white)

    поток = [Paragraph(escape(ЗАГОЛОВОК), заголовок)]
    for вид, данные in СОДЕРЖАНИЕ:
        if вид == "h":
            поток.append(Paragraph(escape(данные), раздел))
        elif вид == "lead":
            поток.append(Paragraph(escape(данные), вводный))
        elif вид == "p":
            поток.append(Paragraph(escape(данные), текст))
        elif вид in ("steps", "bullets"):
            пункты = [ListItem(Paragraph(escape(x), текст)) for x in данные]
            if вид == "steps":
                поток.append(ListFlowable(пункты, bulletType="1",
                                          bulletFontName=обычный,
                                          bulletFontSize=10.5,
                                          leftIndent=18, bulletFormat="%s."))
            else:
                поток.append(ListFlowable(пункты, bulletType="bullet",
                                          start="•", bulletFontName=обычный,
                                          leftIndent=14))
            поток.append(Spacer(1, 4))
        elif вид == "table":
            строки = [[Paragraph(escape(v), шапка if i == 0 else ячейка)
                       for v in строка] for i, строка in enumerate(данные)]
            таблица = Table(строки, colWidths=[6.4 * cm, 8.6 * cm, 2.0 * cm],
                            repeatRows=1)
            таблица.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C6C7C5")),
                ("BACKGROUND", (0, 0), (-1, 0), графит),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (2, 0), (2, -1), "CENTER"),
            ]))
            поток.append(таблица)
            поток.append(Spacer(1, 6))
        else:
            raise ValueError(вид)

    def колонтитул(canvas, документ):
        canvas.saveState()
        canvas.setFont(обычный, 8)
        canvas.setFillColor(графит)
        canvas.setStrokeColor(золото)
        canvas.line(2 * cm, 1.55 * cm, A4[0] - 2 * cm, 1.55 * cm)
        canvas.drawString(2 * cm, 1.2 * cm, f"Комплект «{ТОВАР}»")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm,
                               f"стр. {документ.page}")
        canvas.restoreState()

    буфер = io.BytesIO()
    документ = SimpleDocTemplate(
        буфер, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm, title=ЗАГОЛОВОК,
        author="Маржа в бетоне", creator="build_paid_02_instrukciya.py",
        producer="reportlab", invariant=1)
    документ.build(поток, onFirstPage=колонтитул, onLaterPages=колонтитул)
    return буфер.getvalue()


def main():
    for путь, данные in ((DOCX, _docx_bytes()), (PDF, _pdf_bytes())):
        with open(путь, "wb") as f:
            f.write(данные)
        print(f"  {os.path.relpath(путь, BASE_DIR)}  {len(данные)} байт")


if __name__ == "__main__":
    main()

