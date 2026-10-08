import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Название, цена и адрес платного комплекта — из products-config.php, ключ p1.
# Здесь они повторены текстом; при расхождении верным считается config.
# До 16.09.2026 здесь стояла цена снятой линейки и строка-заглушка вместо
# адреса: файл уходил читателю с ценой, которой в кассе нет, и без ссылки,
# по которой её можно проверить.
PAID_KEY = "p1"
PAID_NAME = "Акты выполненных работ, КС-2 и КС-3: закрытие и взыскание оплаты"
PAID_PRICE = "29 900 ₽"
PAID_URL = "https://marzhavbetone.ru/products/p1-oplata-po-ks2.html"

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)


def create_excel(filepath, sheet_name="Лист1", headers=None, rows=None, col_widths=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    if headers:
        ws.append(headers)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
    
    if rows:
        for r in rows:
            ws.append(r)
            row_num = ws.max_row
            for col_idx in range(1, len(r) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    if col_widths:
        for idx, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = w
    
    if sheet_name == "Калькулятор срока оплаты":
        ws["B2"].number_format = "DD.MM.YYYY"
        ws["B3"].number_format = "DD.MM.YYYY"
        ws["B6"].number_format = "DD.MM.YYYY"
        ws["B7"].number_format = "DD.MM.YYYY"
        ws["B10"].number_format = "0.00%"
        ws["B11"].number_format = '#,##0.00 "₽"'
        ws["B12"].number_format = '#,##0.00 "₽"'
        kind = DataValidation(type="list", formula1='"рабочие,календарные"', allow_blank=True)
        ws.add_data_validation(kind)
        kind.add(ws["B5"])
    wb.save(filepath)
    print(f"  [Excel] {filepath}")


def create_word_doc(filepath, title, sections):
    doc = Document()
    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    
    for section in sections:
        if "heading" in section:
            doc.add_heading(section["heading"], level=1)
        
        if "text" in section:
            doc.add_paragraph(section["text"])
        
        if "bullet" in section:
            for item in section["bullet"]:
                doc.add_paragraph(item, style='List Bullet')
        
        if "table" in section:
            table_data = section["table"]
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for i, row_data in enumerate(table_data):
                    for j, cell_text in enumerate(row_data):
                        table.cell(i, j).text = str(cell_text)
                doc.add_paragraph()
        
        if "numbered" in section:
            for item in section["numbered"]:
                doc.add_paragraph(item, style='List Number')
    
    doc.save(filepath)
    print(f"  [Word]  {filepath}")


def create_pdf_simple(filepath, title, lines):
    c = canvas.Canvas(filepath, pagesize=A4)
    width, height = A4
    
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(width/2, height - 2*cm, title)
    
    c.setFont("Helvetica", 11)
    y = height - 3.5*cm
    for line in lines:
        if y < 2*cm:
            c.showPage()
            y = height - 2*cm
            c.setFont("Helvetica", 11)
        c.drawString(2*cm, y, line)
        y -= 0.6*cm
    
    c.save()
    print(f"  [PDF]   {filepath}")


# ==================== FREE 01 ====================
def build_free_01():
    folder = os.path.join(BASE_DIR, "00-free-ks-podpisany-deneg-net")
    print(f"\n[BUILD] {folder}")
    
    create_word_doc(
        os.path.join(folder, "01-algoritm-pervichnoy-proverki.docx"),
        "АЛГОРИТМ ПЕРВИЧНОЙ ПРОВЕРКИ: КС подписаны, денег нет",
        [
            {"heading": "Шаг 1. Проверьте срок оплаты по договору",
             "numbered": [
                 "Откройте раздел 'Условия оплаты' в договоре подряда.",
                 "Определите: аванс / поэтапно / по факту.",
                 "Посчитайте крайний срок оплаты (рабочие или календарные дни - важно!)."
             ]},
            {"heading": "Шаг 2. Проверьте подписание КС-2 / КС-3",
             "numbered": [
                 "Есть ли подписи обеих сторон на КС-2?",
                 "Есть ли подписи на КС-3?",
                 "Совпадают ли даты подписания с фактом выполнения работ?",
                 "Если КС не подписана, проверьте порядок сдачи и мотивы отказа; маршрут требует отдельной оценки."
             ]},
            {"heading": "Шаг 3. Проверьте срок исковой давности",
             "text": "Общий срок исковой давности — три года. Начало течения срока и обстоятельства его изменения проверяйте отдельно по обязательству и датам."},
            {"heading": "Шаг 4. Соберите доказательства выполнения работ",
             "bullet": [
                 "Акты освидетельствования скрытых работ (АОСР), если применимо",
                 "Журналы работ",
                 "Фото/видео фиксация объемов",
                 "Переписка с заказчиком (email, Telegram, WhatsApp)",
                 "Протоколы совещаний"
             ]},
            {"heading": "Шаг 5. Проверьте наличие претензионного порядка",
             "text": "Для денежного требования из договора в арбитражном процессе проверьте ч. 5 ст. 4 АПК РФ и договор: как направить претензию и сколько ждать до иска. Иной порядок может быть установлен законом или договором."},
            {"heading": "Шаг 6. Рассчитайте неустойку",
             "text": "Проверьте договорную неустойку. Если её нет, отдельно оцените проценты по ст. 395 ГК РФ. Калькулятор ниже считает только пеню, прямо предусмотренную договором."},
            {"heading": "Шаг 7. Подготовьте письмо-напоминание",
             "text": "Используйте шаблон письма из этого комплекта. Фиксируйте все коммуникации в письменном виде."},
            {"heading": "⚠️ Что этот материал НЕ закрывает",
             "bullet": [
                 "Составление полноценной претензии с юридическим обоснованием",
                 "Подготовка искового заявления и представление интересов в суде",
                 "Работа с исполнительным производством",
                 "Анализ сложных случаев: банкротство заказчика, субподрядные цепочки, залоги"
             ]},
            {"heading": "🔗 Переход к полному комплекту",
             "text": f"Комплект «{PAID_NAME}» включает: уведомление, претензию, расчёт процентов по ст. 395 ГК РФ, исковое заявление и алгоритм действий при задержке оплаты.\n\nСтоимость: {PAID_PRICE}\nСтраница комплекта: {PAID_URL}"},
        ]
    )
    
    create_excel(
        os.path.join(folder, "02-kalkulyator-sroka-oplaty.xlsx"),
        "Калькулятор срока оплаты",
        headers=["Параметр", "Значение", "Примечание"],
        rows=[
            ["Дата подписания КС-2", "", "Введите дату"],
            ["Дата подписания КС-3", "", "Введите дату"],
            ["Срок оплаты по договору (дней)", "", "Например: 15 рабочих дней"],
            ["Тип срока", "", "рабочие / календарные"],
            ["Дата начала срока по договору", "", "Укажите событие из договора, от которого считается срок."],
            ["Ориентир срока оплаты", "=IF(OR(NOT(ISNUMBER(B6)),NOT(ISNUMBER(B4)),B4<0,AND(B5<>\"рабочие\",B5<>\"календарные\")),\"\",IF(B5=\"рабочие\",WORKDAY(B6,B4),B6+B4))", "Проверьте праздники и перенос срока по ст. 193 ГК РФ."],
            ["Просрочка (дней)", "=IF(NOT(ISNUMBER(B7)),\"\",MAX(0,TODAY()-B7))", "От ориентировочной даты; проверьте её по договору."],
            ["", "", ""],
            ["Ставка пени по договору (%/день)", "", "Введите договорную ставку как процент, например 0,1%."],
            ["Сумма по КС-2 (руб.)", "", "Введите сумму"],
            ["Начисленные пени", "=IF(OR(NOT(ISNUMBER(B8)),NOT(ISNUMBER(B10)),NOT(ISNUMBER(B11))),\"\",B8*B10*B11)", "Только если договор предусматривает пеню."],
        ],
        col_widths=[35, 25, 40]
    )
    
    create_word_doc(
        os.path.join(folder, "03-perechen-dokumentov.docx"),
        "ПЕРЕЧЕНЬ НЕОБХОДИМЫХ ДОКУМЕНТОВ для претензионной работы",
        [
            {"heading": "Базовые документы для проверки",
             "numbered": [
                 "Договор подряда (все приложения, допсоглашения)",
                 "КС-2 или согласованная договором форма акта; подписи либо доказательства направления и отказа",
                 "КС-3 или согласованная договором справка, если она предусмотрена",
                 "Журнал учёта работ и исполнительная документация, если применимо",
                 "Акты освидетельствования скрытых работ (АОСР), если применимо",
                 "Ведомость ресурсов (КС-4) - если применимо"
             ]},
            {"heading": "Дополнительные документы (усиливают позицию)",
             "bullet": [
                 "Календарный план производства работ",
                 "Журнал производства работ",
                 "Фото/видео материалы выполнения работ",
                 "Протоколы совещаний с заказчиком",
                 "Переписка (email, мессенджеры) - экспорт с метаданными",
                 "Письма заказчику с подтверждением доставки (заказные, email-уведомления)",
                 "Служебные записки внутри организации"
             ]},
            {"heading": "Если работы выполнены, но КС не подписана",
             "text": "Проверьте уведомление о готовности к сдаче, направление акта, доказательства получения, отметку об отказе и мотивы возражений. При споре о качестве или объёме может потребоваться отдельная оценка."},
            {"heading": "Совет",
             "text": "Собирайте документы ПОСТЕПЕННО, в процессе работы. Не откладывайте на 'потом' - когда конфликт начался, заказчик может перестать отвечать и предоставлять документы."},
        ]
    )
    
    create_word_doc(
        os.path.join(folder, "04-shablon-pisma.docx"),
        "ШАБЛОН ПИСЬМА ЗАКАЗЧИКУ о задержке оплаты",
        [
            {"text": "[На бланке организации]"},
            {"text": ""},
            {"text": "Исх. № _____ от '___' ___________ 20___ г."},
            {"text": ""},
            {"text": "Генеральному директору"},
            {"text": "[Наименование заказчика]"},
            {"text": "[ФИО руководителя]"},
            {"text": ""},
            {"text": "От [Наименование подрядчика]"},
            {"text": ""},
            {"text": "УВЕДОМЛЕНИЕ"},
            {"text": "о нарушении срока оплаты по договору подряда № [номер] от [дата]"},
            {"text": ""},
            {"text": "Уважаемый [ФИО]!"},
            {"text": ""},
            {"text": "Настоящим сообщаем, что в соответствии с п. [номер] договора подряда № [номер] от [дата] (далее - Договор) срок оплаты выполненных работ составляет [указать срок] после [договорное событие: подписание акта, передача комплекта или иное]."},
            {"text": ""},
            {"text": "Форма КС-2 № [номер] от [дата] подписана обеими сторонами [дата подписания заказчиком]. Крайний срок оплаты наступил [дата]."},
            {"text": ""},
            {"text": "По состоянию на '___' ___________ 20___ г. оплата не поступила. Просрочка составляет [количество] дней."},
            {"text": ""},
            {"text": "Если п. [номер] Договора предусматривает неустойку за просрочку оплаты, заказчик уплачивает пени в размере [процент]% от суммы просроченной оплаты за каждый день просрочки. На текущую дату начисленные пени составляют [сумма] руб."},
            {"text": ""},
            {"text": "Просим в срок до [дата] оплатить подтверждённую сумму долга [сумма] руб. При наличии договорной неустойки и её расчёта просим также оплатить [сумма] руб.; итог требования [итог] руб."},
            {"text": ""},
            {"text": "Если оплата не поступит, оставляем за собой право после соблюдения досудебного порядка обратиться в суд за взысканием долга, применимой договорной неустойки либо процентов по ст. 395 ГК РФ и судебных расходов."},
            {"text": ""},
            {"text": "Приложения:"},
            {"bullet": [
                "Копия договора подряда № ___ от ___ (л. __)",
                "Копия формы КС-2 № ___ от ___ (л. __)",
                "Копия формы КС-3 № ___ от ___ (л. __)",
                "Расчёт договорной пени (если заявляется, л. __)."
            ]},
            {"text": ""},
            {"text": "Генеральный директор _________________ [ФИО]"},
            {"text": "М.П."},
            {"text": ""},
            {"text": ""},
            {"text": "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━"},
            {"heading": "💡 Как использовать этот шаблон",
             "numbered": [
                 "Заполните все поля в квадратных скобках.",
                 "Проверьте соответствие сроков и сумм с договором и КС-2.",
                 "Направьте по согласованному в договоре каналу и сохраните подтверждение направления и получения.",
                 "Сохраните квитанцию об отправке - это доказательство направления.",
                 "После истечения срока оплаты направьте претензию и проверьте срок ожидания до иска по закону и договору."
             ]},
            {"heading": "🔗 Переход к полному комплекту",
             "text": f"Этот шаблон - только первый шаг. Комплект «{PAID_NAME}» включает: уведомление о просрочке, претензию, расчёт процентов по ст. 395 ГК РФ, исковое заявление и алгоритм действий при задержке оплаты.\n\nСтоимость: {PAID_PRICE}\nСтраница комплекта: {PAID_URL}"},
        ]
    )


if __name__ == "__main__":
    build_free_01()
    print("\nГОТОВО: free-01")
